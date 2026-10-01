from __future__ import annotations
# ONESTOCK_EMBEDDED_LOCAL_IMPORT_BOOTSTRAP
import sys as _onestock_embedded_sys
from pathlib import Path as _OneStockEmbeddedPath
_onestock_embedded_dir = str(_OneStockEmbeddedPath(__file__).resolve().parent)
if _onestock_embedded_dir not in _onestock_embedded_sys.path:
    _onestock_embedded_sys.path.insert(0, _onestock_embedded_dir)

import hashlib
import importlib.util
import json
import copy
import subprocess
import sys
import tempfile
import types
import unittest
from pathlib import Path
from unittest.mock import patch

from docx import Document
from PIL import Image


SCRIPTS = Path(__file__).resolve().parent
SKILL_ROOT = SCRIPTS.parent
if str(SCRIPTS) not in sys.path:
    sys.path.insert(0, str(SCRIPTS))
import live_data_pipeline as PIPELINE

if "pdf2image" not in sys.modules:
    pdf2image_stub = types.ModuleType("pdf2image")
    pdf2image_stub.convert_from_path = lambda *_args, **_kwargs: []
    pdf2image_stub.pdfinfo_from_path = lambda *_args, **_kwargs: {"Pages": 0}
    sys.modules["pdf2image"] = pdf2image_stub
import word_com_render as WORD_RENDER

MODULE_PATH = SCRIPTS / "report_validator.py"
SPEC = importlib.util.spec_from_file_location("leader_report_validator_test", MODULE_PATH)
assert SPEC is not None and SPEC.loader is not None
MODULE = importlib.util.module_from_spec(SPEC)
sys.modules[SPEC.name] = MODULE
SPEC.loader.exec_module(MODULE)
ADAPTER_PATH = SCRIPTS / "canonical_business_adapter.py"
ADAPTER_SPEC = importlib.util.spec_from_file_location(
    "leader_canonical_business_adapter_test",
    ADAPTER_PATH,
)
assert ADAPTER_SPEC is not None and ADAPTER_SPEC.loader is not None
ADAPTER = importlib.util.module_from_spec(ADAPTER_SPEC)
sys.modules[ADAPTER_SPEC.name] = ADAPTER
ADAPTER_SPEC.loader.exec_module(ADAPTER)
LEGACY_PATH = SCRIPTS / "legacy_codex_entry.py"
LEGACY_SPEC = importlib.util.spec_from_file_location(
    "leader_legacy_entry_test",
    LEGACY_PATH,
)
assert LEGACY_SPEC is not None and LEGACY_SPEC.loader is not None
LEGACY = importlib.util.module_from_spec(LEGACY_SPEC)
sys.modules[LEGACY_SPEC.name] = LEGACY
LEGACY_SPEC.loader.exec_module(LEGACY)


class LeaderDeliveryGateTests(unittest.TestCase):
    def test_word_com_paths_drop_windows_extended_prefix(self) -> None:
        self.assertEqual(
            WORD_RENDER.word_com_compatible_path(
                Path(r"\\?\C:\Users\25296\AppData\Local\TieRobot\OneStockAgent\report.docx")
            ),
            r"C:\Users\25296\AppData\Local\TieRobot\OneStockAgent\report.docx",
        )
        self.assertEqual(
            WORD_RENDER.word_com_compatible_path(
                Path(r"\\?\UNC\server\share\report.pdf")
            ),
            r"\\server\share\report.pdf",
        )
        self.assertEqual(
            WORD_RENDER.word_com_compatible_path(Path(r"D:\reports\report.docx")),
            r"D:\reports\report.docx",
        )

    @staticmethod
    def dxx_snapshot(
        codes: list[str],
        *,
        plate_codes: list[str] | None = None,
        fetched_at: str = "2026-08-19T11:00:00+08:00",
    ) -> dict:
        return {
            "generated_at": fetched_at,
            "datasets": {
                "ztpool": {
                    "data": {
                        "list": [[code, f"股票{code}"] for code in codes],
                    }
                },
                "ztplate": {
                    "data": {
                        "list": [
                            {"code": code}
                            for code in (plate_codes if plate_codes is not None else codes)
                        ],
                    }
                },
            },
        }

    @staticmethod
    def eastmoney_snapshot(codes: list[str]) -> dict:
        return {
            "data": {
                "pool": [
                    {"c": code, "n": f"股票{code}"}
                    for code in codes
                ],
            }
        }

    def run_stable_source_fixture(
        self,
        root: Path,
        states: list[tuple[dict, dict]],
        *,
        max_attempts: int,
    ) -> dict:
        dxx_index = 0
        em_index = 0

        def fake_dxx(attempt_dir: Path) -> tuple[dict, Path]:
            nonlocal dxx_index
            payload = states[min(dxx_index, len(states) - 1)][0]
            dxx_index += 1
            attempt_dir.mkdir(parents=True, exist_ok=True)
            output = attempt_dir / "duanxianxia.json"
            output.write_text(json.dumps(payload, ensure_ascii=False), encoding="utf-8")
            return payload, output

        def fake_em(
            trade_date: str,
            attempt_dir: Path,
        ) -> tuple[dict, Path, str]:
            nonlocal em_index
            payload = states[min(em_index, len(states) - 1)][1]
            em_index += 1
            output = attempt_dir / f"eastmoney_ztpool_{trade_date.replace('-', '')}.json"
            output.write_text(json.dumps(payload, ensure_ascii=False), encoding="utf-8")
            return payload, output, "https://example.test/eastmoney"

        with (
            patch.object(PIPELINE, "collect_duanxianxia", side_effect=fake_dxx),
            patch.object(PIPELINE, "collect_eastmoney_pool", side_effect=fake_em),
            patch.object(PIPELINE.time, "sleep"),
        ):
            return PIPELINE.collect_stable_limit_up_sources(
                "2026-08-19",
                root,
                max_attempts=max_attempts,
                retry_delay_seconds=0,
            )

    def test_cross_source_drift_retries_until_full_sets_and_plates_converge(self) -> None:
        states = [
            (
                self.dxx_snapshot(["000001"]),
                self.eastmoney_snapshot(["000001", "000002"]),
            ),
            (
                self.dxx_snapshot(["000001", "000002"]),
                self.eastmoney_snapshot(["000001", "000002"]),
            ),
        ]
        with tempfile.TemporaryDirectory() as temporary:
            result = self.run_stable_source_fixture(
                Path(temporary),
                states,
                max_attempts=3,
            )

            self.assertEqual(result["status"], "CLEAN_PASS")
            self.assertEqual(result["attempt_count"], 2)
            self.assertEqual(result["attempts"][0]["eastmoney_only"], ["000002"])
            self.assertEqual(result["attempts"][0]["plate_missing"], ["000002"])
            self.assertEqual(result["attempts"][1]["eastmoney_only"], [])
            self.assertEqual(result["attempts"][1]["duanxianxia_only"], [])
            self.assertEqual(result["attempts"][1]["plate_missing"], [])
            self.assertNotEqual(
                result["attempts"][0]["duanxianxia_snapshot"],
                result["attempts"][1]["duanxianxia_snapshot"],
            )
            self.assertTrue(Path(result["comparison_log"]).is_file())

    def test_cross_source_drift_failure_reports_all_differences_and_attempts(self) -> None:
        state = (
            self.dxx_snapshot(
                ["000001", "000004"],
                plate_codes=["000001"],
            ),
            self.eastmoney_snapshot(["000001", "000002", "000003"]),
        )
        with tempfile.TemporaryDirectory() as temporary:
            with self.assertRaises(RuntimeError) as caught:
                self.run_stable_source_fixture(
                    Path(temporary),
                    [state],
                    max_attempts=2,
                )

        message = str(caught.exception)
        self.assertIn("cross_source_limit_up_not_stable", message)
        self.assertIn("attempts=2", message)
        self.assertIn("eastmoney_only=['000002', '000003']", message)
        self.assertIn("duanxianxia_only=['000004']", message)
        self.assertIn("plate_missing=['000002', '000003']", message)

    def test_cross_source_overlap_is_not_accepted_as_full_reconciliation(self) -> None:
        state = (
            self.dxx_snapshot(
                ["000001", "000002"],
                plate_codes=["000001", "000002", "000003"],
            ),
            self.eastmoney_snapshot(["000001", "000003"]),
        )
        with tempfile.TemporaryDirectory() as temporary:
            with self.assertRaises(RuntimeError) as caught:
                self.run_stable_source_fixture(
                    Path(temporary),
                    [state],
                    max_attempts=2,
                )

        message = str(caught.exception)
        self.assertIn("eastmoney_only=['000003']", message)
        self.assertIn("duanxianxia_only=['000002']", message)
        self.assertIn("plate_missing=[]", message)

    def content_identity_contract(self) -> dict:
        return json.loads(
            (SKILL_ROOT / "references" / "report_template_contract.json").read_text(
                encoding="utf-8"
            )
        )

    @staticmethod
    def stock_payload(count: int = 40) -> dict:
        return {
            "stocks": [
                {"stock_code": f"{index:06d}"}
                for index in range(1, count + 1)
            ]
        }

    def test_default_canonical_action_is_real_run(self) -> None:
        with tempfile.TemporaryDirectory() as temporary:
            command = ADAPTER.business_command(Path(temporary), [])

        self.assertIn("run", command)
        self.assertNotIn("info", command)
        self.assertIn("--run-dir", command)
        self.assertNotIn("--output", command)

    def test_default_deliverable_uses_chinese_name(self) -> None:
        output = LEGACY.default_output_path(Path("C:/tmp"), "2026-08-07")

        self.assertEqual(output.name, "龙头深度研究_2026-08-07.docx")
        self.assertNotIn("canary", output.name.casefold())
        self.assertNotIn("leader_report", output.name.casefold())

    def test_execution_scripts_do_not_reference_retired_renderer(self) -> None:
        forbidden = (
            "libre" + "office",
            "sof" + "fice",
            "render_" + "docx",
        )
        execution_files = (
            SCRIPTS / "legacy_codex_entry.py",
            SCRIPTS / "canonical_business_adapter.py",
            SCRIPTS / "word_com_render.py",
        )
        for path in execution_files:
            source = path.read_text(encoding="utf-8").lower()
            for token in forbidden:
                self.assertNotIn(token, source, f"{token} found in {path}")

    def test_report_without_current_run_markers_is_blocked(self) -> None:
        with tempfile.TemporaryDirectory() as temporary:
            root = Path(temporary)
            payload_path = root / "payload.json"
            payload_path.write_text(
                json.dumps(
                    {
                        "trade_date": "2026-08-08",
                        "concept_summary": [],
                        "report_claims": [],
                        "morning_limit_ups": [],
                        "lhb_net_buy_ge_100m": [],
                    },
                    ensure_ascii=False,
                ),
                encoding="utf-8",
            )
            contract_path = (
                SKILL_ROOT / "references" / "report_template_contract.json"
            )
            manifest_path = root / "manifest.json"
            manifest_path.write_text(
                json.dumps(
                    {
                        "result_sha256": hashlib.sha256(
                            payload_path.read_bytes()
                        ).hexdigest(),
                        "template_contract_sha256": hashlib.sha256(
                            contract_path.read_bytes()
                        ).hexdigest(),
                    }
                ),
                encoding="utf-8",
            )
            docx_path = root / "report.docx"
            document = Document()
            document.add_paragraph("龙头深度研究")
            document.save(docx_path)

            result = MODULE.validate(
                docx_path,
                payload_path,
                manifest_path,
                contract_path,
                None,
            )

        self.assertEqual(result["status"], "BLOCKED")
        self.assertIn(
            "machine_marker_missing_or_wrong:data_evidence_sha256",
            result["errors"],
        )
        self.assertIn(
            "machine_marker_missing_or_wrong:template_contract_sha256",
            result["errors"],
        )
        self.assertIn(
            "machine_marker_missing_or_wrong:STOCK_DATA_TRADE_DATE",
            result["errors"],
        )
        self.assertIn("business_cover_missing", result["errors"])
        self.assertIn("render_dir_required", result["errors"])

    def test_missing_render_pages_are_blocked(self) -> None:
        contract = self.content_identity_contract()
        with tempfile.TemporaryDirectory() as temporary:
            errors, detail = MODULE.check_render(Path(temporary), contract)

        self.assertEqual(errors, ["render_pages_missing"])
        self.assertEqual(detail["page_count"], 0)

    def test_sparse_final_orphan_page_is_blocked(self) -> None:
        contract = self.content_identity_contract()
        with tempfile.TemporaryDirectory() as temporary:
            root = Path(temporary)
            for page_number, ink_width in ((1, 500), (2, 70)):
                image = Image.new("L", (1200, 1600), 255)
                for x in range(ink_width):
                    for y in range(300):
                        image.putpixel((x, y), 0)
                image.save(root / f"page-{page_number}.png")
            errors, _ = MODULE.check_render(root, contract)

        self.assertTrue(
            any(error.startswith("render_final_page_orphaned:") for error in errors)
        )

    def test_explanatory_audit_word_is_blocked_by_content_identity(self) -> None:
        document = Document()
        document.core_properties.title = "龙头深度研究 数据真实性核验报告"
        document.core_properties.subject = "工作流说明与结论生成机制"
        document.add_paragraph(
            "本报告用于解释数据真实性核验、审计步骤和交付机制。"
            "不包含正式涨停板龙头分类业务内容。" * 220
        )

        errors, detail = MODULE.check_content_identity(
            document,
            self.stock_payload(),
            self.content_identity_contract(),
        )

        self.assertEqual(detail["status"], "BLOCKED")
        self.assertIn(
            "content_identity_title_term_missing:A股涨停板龙头分类报告",
            errors,
        )
        self.assertIn("content_identity_core_metadata_forbidden:核验", errors)
        self.assertIn(
            "content_identity_visible_substitute:数据真实性核验",
            errors,
        )
        self.assertIn(
            "content_identity_stock_code_coverage_low:0<12",
            errors,
        )

    def test_content_light_low_stock_coverage_report_is_blocked(self) -> None:
        document = Document()
        document.core_properties.title = (
            "2026-08-10 龙头深度研究 A股涨停板龙头分类报告"
        )
        document.core_properties.subject = (
            "按一级投资板块、龙头性质、上午封板时间和龙虎榜单日净额分类"
        )
        document.add_paragraph(
            "投资风险提示 A股每日涨停板 当日结论 000001 000002"
        )

        errors, detail = MODULE.check_content_identity(
            document,
            self.stock_payload(),
            self.content_identity_contract(),
        )

        self.assertEqual(detail["status"], "BLOCKED")
        self.assertTrue(
            any(error.startswith("content_identity_text_too_short:") for error in errors)
        )
        self.assertIn(
            "content_identity_stock_code_coverage_low:2<12",
            errors,
        )

    def test_formal_substantive_report_passes_content_identity(self) -> None:
        document = Document()
        document.core_properties.title = (
            "2026-08-10 龙头深度研究 A股涨停板龙头分类报告"
        )
        document.core_properties.subject = (
            "按一级投资板块、细分方向、当日推动因素、龙头性质、"
            "板块地位、上午封板时间和龙虎榜单日净额分类"
        )
        codes = " ".join(f"{index:06d}" for index in range(1, 41))
        document.add_paragraph(
            codes
            + (
                " 当日核心判断覆盖市场高度、主导板块、连板梯队、封板强度、"
                "成交额、换手率、龙虎榜席位净额和板块内龙头地位。"
            )
            * 100
        )

        errors, detail = MODULE.check_content_identity(
            document,
            self.stock_payload(),
            self.content_identity_contract(),
        )

        self.assertEqual(errors, [])
        self.assertEqual(detail["status"], "PASS")
        self.assertGreaterEqual(detail["visible_character_count"], 4500)
        self.assertEqual(detail["matched_payload_stock_code_count"], 40)

    def test_visible_process_pollution_is_blocked(self) -> None:
        document = Document()
        document.core_properties.title = (
            "2026-08-13 龙头深度研究 A股涨停板龙头分类报告"
        )
        document.core_properties.subject = (
            "按一级投资板块、细分方向、当日推动因素、龙头性质、"
            "板块地位、上午封板时间和龙虎榜单日净额分类"
        )
        codes = " ".join(f"{index:06d}" for index in range(1, 41))
        document.add_paragraph(
            codes
            + (
                " 市场高度、主导板块、连板梯队、封板强度和龙虎榜资金。"
            )
            * 120
            + "反证：若任一来源代码集合发生差异，本轮报告停止生成。"
            + "数据来源与差异处理：东方财富与短线侠逐只对账。"
        )

        errors, detail = MODULE.check_content_identity(
            document,
            self.stock_payload(),
            self.content_identity_contract(),
        )

        self.assertEqual(detail["status"], "BLOCKED")
        self.assertIn(
            "content_identity_visible_process_pollution:反证：",
            errors,
        )
        self.assertIn(
            "content_identity_visible_process_pollution:停止生成",
            errors,
        )
        self.assertIn(
            "content_identity_visible_process_pollution:逐只对账",
            errors,
        )

    def test_payload_claim_with_basis_or_counterevidence_is_blocked(self) -> None:
        sample = LEGACY.synthetic_fixture()
        polluted = copy.deepcopy(sample)
        polluted["report_claims"][0]["basis_text"] = "内部数据对账过程"
        polluted["report_claims"][0]["counterevidence_text"] = "若失败则停止生成"

        with self.assertRaisesRegex(
            Exception,
            "禁止携带依据/反证过程字段",
        ):
            LEGACY.audit_payload(polluted)

    def test_statistics_only_board_logic_is_blocked(self) -> None:
        sample = LEGACY.synthetic_fixture()
        polluted = copy.deepcopy(sample)
        polluted["concept_summary"][0]["logic_summary"] = (
            "大科技共有1只涨停，其中连板1只、10:00前首次封板1只、"
            "零开板1只；最高近期高度为1天1板。"
        )

        with self.assertRaisesRegex(
            Exception,
            "不能仅罗列统计数据|可见结论与结构化逻辑模型不一致",
        ):
            LEGACY.audit_payload(polluted)

    def test_copied_position_reasons_without_relative_logic_are_blocked(self) -> None:
        sample = LEGACY.synthetic_fixture()
        polluted = copy.deepcopy(sample)
        rows = polluted["concept_summary"][0]["position_top"]
        for row in rows:
            row["reason"] = (
                "属于覆盖面最大的共同机制；其后90分钟有成员封板；"
                "首次09:30、最终09:30，开板0次。"
            )
            for field in (
                "rank_driver",
                "relative_comparison",
                "weakness",
                "invalidation",
            ):
                row.pop(field, None)

        with self.assertRaisesRegex(
            Exception,
            "缺少胜出依据|缺少相对比较|地位依据重复",
        ):
            LEGACY.audit_payload(polluted)

    def test_theme_cooccurrence_cannot_be_called_catalyst(self) -> None:
        sample = LEGACY.synthetic_fixture()
        polluted = copy.deepcopy(sample)
        polluted["concept_summary"][0]["logic_summary"] = (
            "主催化机器人覆盖1/1只，并形成扩散。"
        )

        with self.assertRaisesRegex(
            Exception,
            "含模板措辞: 主催化|可见结论与结构化逻辑模型不一致",
        ):
            LEGACY.audit_payload(polluted)

    def test_ambiguous_evidence_completeness_language_is_blocked(self) -> None:
        sample = LEGACY.synthetic_fixture()
        polluted = copy.deepcopy(sample)
        polluted["concept_summary"][0]["validation_constraints"] = (
            "盘面验证：1/1只在10:00前首次封板。"
            "结构约束：高度依赖单一样本。逻辑证据完整度为高。"
        )

        with self.assertRaisesRegex(
            Exception,
            "含模板措辞: 证据完整度|可见结论与结构化逻辑模型不一致",
        ):
            LEGACY.audit_payload(polluted)

    def test_missing_delivery_inputs_exit_nonzero(self) -> None:
        validator = SCRIPTS / "report_validator.py"
        contract = SKILL_ROOT / "references" / "report_template_contract.json"
        with tempfile.TemporaryDirectory() as temporary:
            root = Path(temporary)
            payload = root / "payload.json"
            manifest = root / "manifest.json"
            docx = root / "report.docx"
            render_dir = root / "render"
            render_dir.mkdir()
            payload.write_text(
                json.dumps(
                    {
                        "trade_date": "2026-08-08",
                        "concept_summary": [],
                        "report_claims": [],
                        "morning_limit_ups": [],
                        "lhb_net_buy_ge_100m": [],
                    }
                ),
                encoding="utf-8",
            )
            manifest.write_text("{}", encoding="utf-8")
            document = Document()
            document.add_paragraph("delivery canary")
            document.save(docx)

            cases = {
                "missing_data": (root / "missing-payload.json", manifest, docx, contract),
                "missing_template": (payload, manifest, docx, root / "missing-contract.json"),
                "missing_docx": (payload, manifest, root / "missing.docx", contract),
                "missing_pages": (payload, manifest, docx, contract),
            }
            for name, (case_payload, case_manifest, case_docx, case_contract) in cases.items():
                with self.subTest(name=name):
                    completed = subprocess.run(
                        [
                            sys.executable,
                            str(validator),
                            "--docx",
                            str(case_docx),
                            "--input",
                            str(case_payload),
                            "--manifest",
                            str(case_manifest),
                            "--template-contract",
                            str(case_contract),
                            "--render-dir",
                            str(render_dir),
                        ],
                        capture_output=True,
                        text=True,
                        timeout=30,
                    )
                    self.assertNotEqual(completed.returncode, 0)


if __name__ == "__main__":
    unittest.main()
