from __future__ import annotations
# ONESTOCK_EMBEDDED_LOCAL_IMPORT_BOOTSTRAP
import sys as _onestock_embedded_sys
from pathlib import Path as _OneStockEmbeddedPath
_onestock_embedded_dir = str(_OneStockEmbeddedPath(__file__).resolve().parent)
if _onestock_embedded_dir not in _onestock_embedded_sys.path:
    _onestock_embedded_sys.path.insert(0, _onestock_embedded_dir)

import argparse
import hashlib
import json
import re
from collections import Counter
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_ALIGN_VERTICAL, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

from report_conclusion_blocks import build_conclusion_rows


PARSER = argparse.ArgumentParser(description="龙头深度研究正式 Word 报告生成器")
PARSER.add_argument("--input", required=True, type=Path)
PARSER.add_argument("--manifest", required=True, type=Path)
PARSER.add_argument("--output", required=True, type=Path)
PARSER.add_argument("--template-contract", required=True, type=Path)
ARGS = PARSER.parse_args()

DATA_PATH = ARGS.input.resolve()
MANIFEST_PATH = ARGS.manifest.resolve()
FINAL_PATH = ARGS.output.resolve()
TEMPLATE_CONTRACT_PATH = ARGS.template_contract.resolve()
if not DATA_PATH.is_file() or not MANIFEST_PATH.is_file() or not TEMPLATE_CONTRACT_PATH.is_file():
    raise RuntimeError("本轮结果或清单不存在")
FINAL_PATH.parent.mkdir(parents=True, exist_ok=True)

data = json.loads(DATA_PATH.read_text(encoding="utf-8"))
manifest = json.loads(MANIFEST_PATH.read_text(encoding="utf-8"))
if hashlib.sha256(DATA_PATH.read_bytes()).hexdigest() != manifest.get("result_sha256"):
    raise RuntimeError("结构化结果哈希与清单不一致")
if manifest.get("trade_date") != data.get("trade_date"):
    raise RuntimeError("结构化结果交易日与清单不一致")
template_contract_sha256 = hashlib.sha256(TEMPLATE_CONTRACT_PATH.read_bytes()).hexdigest()
if manifest.get("template_contract_sha256") != template_contract_sha256:
    raise RuntimeError("模板合同哈希与清单不一致")

for stock in data["stocks"]:
    stock.setdefault("final_limit_time", stock.get("latest_seal_time_at_cutoff"))
    stock.setdefault("consecutive_limit_ups", stock.get("consecutive_limit_count", 1))
    stock.setdefault("recent_limit_up_hits", stock.get("recent_limit_count", stock["consecutive_limit_ups"]))
    stock.setdefault(
        "recent_limit_up_label",
        f"{stock.get('recent_limit_days', stock['recent_limit_up_hits'])}天"
        f"{stock['recent_limit_up_hits']}板",
    )
    stock.setdefault("turnover_amount", stock.get("amount", 0))

stocks = {item["stock_code"]: item for item in data["stocks"]}
for item in data["concept_summary"]:
    members = [stock for stock in data["stocks"] if stock["primary_concept"] == item["concept"]]
    item.setdefault(
        "earliest_first_limit_time",
        min((stock["first_limit_time"] for stock in members), default="--:--:--"),
    )
    for nature in item.get("nature_top") or []:
        stock = stocks[nature["stock_code"]]
        nature.setdefault(
            "nature_names",
            [role["name"] for role in stock.get("leader_roles") or []]
            or ["板块前排"],
        )

morning_rows = []
for rank, item in enumerate(data["morning_limit_ups"], start=1):
    stock = stocks[item["stock_code"]]
    morning_rows.append(
        {
            **item,
            "rank": rank,
            "stock_name": stock["stock_name"],
            "first_limit_time": stock["first_limit_time"],
            "final_limit_time": stock["final_limit_time"],
            "primary_concept": stock["primary_concept"],
            "open_count": stock.get("open_count", 0),
        }
    )
data["morning_limit_ups"] = morning_rows

lhb_rows = []
for rank, item in enumerate(data["lhb_net_buy_ge_100m"], start=1):
    stock = stocks[item["stock_code"]]
    lhb_rows.append(
        {
            **item,
            "rank": rank,
            "stock_name": stock["stock_name"],
            "primary_concept": stock["primary_concept"],
            "first_limit_time": stock["first_limit_time"],
            "final_limit_time": stock["final_limit_time"],
        }
    )
data["lhb_net_buy_ge_100m"] = lhb_rows
data.setdefault("business_data_sha256", manifest["result_sha256"])
data.setdefault("business_data_trade_date", data["trade_date"])

FONT = "微软雅黑"
NAVY = "102A43"
DEEP_BLUE = "173E5C"
RED = "BC1823"
DARK_RED = "8B0E19"
GOLD = "D6A43E"
PALE_GOLD = "FAF2DC"
PALE_RED = "FCeded".upper()
PALE_BLUE = "ECF3F8"
ICE_BLUE = "F6F9FC"
LIGHT_GRAY = "F4F6F8"
MID_GRAY = "5E6872"
LIGHT_BORDER = "CFD7DF"
WHITE = "FFFFFF"
BLACK = "1F242A"

CONTENT_DXA = 10150
TABLE_INDENT_DXA = 120

DISPLAY_REPLACEMENTS = (
    ("AM5", "五型康复训练系统"),
    ("Kimi", "月之暗面智能助手"),
    ("OLED", "有机发光屏"),
    ("PEEK", "聚醚醚酮"),
    ("AI", "人工智能"),
)


def display_text(value):
    text = str(value)
    for source, target in DISPLAY_REPLACEMENTS:
        text = text.replace(source, target)
    return text


def set_run_font(run, size=11, bold=False, color=BLACK, italic=False):
    run.font.name = FONT
    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), FONT)
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), FONT)
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), FONT)
    run.font.size = Pt(size)
    run.bold = bold
    run.italic = italic
    run.font.color.rgb = RGBColor.from_string(color)


def shade_paragraph(paragraph, fill):
    p_pr = paragraph._p.get_or_add_pPr()
    shd = p_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        p_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_paragraph_border(paragraph, edge="bottom", color=GOLD, size=12):
    p_pr = paragraph._p.get_or_add_pPr()
    p_bdr = p_pr.find(qn("w:pBdr"))
    if p_bdr is None:
        p_bdr = OxmlElement("w:pBdr")
        p_pr.append(p_bdr)
    node = OxmlElement(f"w:{edge}")
    node.set(qn("w:val"), "single")
    node.set(qn("w:sz"), str(size))
    node.set(qn("w:space"), "1")
    node.set(qn("w:color"), color)
    p_bdr.append(node)


def add_para(
    text="",
    *,
    size=11,
    bold=False,
    color=BLACK,
    align=WD_ALIGN_PARAGRAPH.LEFT,
    before=0,
    after=6,
    line=1.1,
    keep=False,
    page_break_before=False,
    fill=None,
    left=0,
    right=0,
):
    p = doc.add_paragraph()
    p.alignment = align
    pf = p.paragraph_format
    pf.space_before = Pt(before)
    pf.space_after = Pt(after)
    pf.line_spacing = line
    pf.keep_with_next = keep
    pf.page_break_before = page_break_before
    pf.left_indent = Cm(left)
    pf.right_indent = Cm(right)
    if fill:
        shade_paragraph(p, fill)
    if text:
        run = p.add_run(display_text(text))
        set_run_font(run, size=size, bold=bold, color=color)
    return p


def add_multiline_para(lines, **kwargs):
    p = add_para("", **kwargs)
    for index, line in enumerate(lines):
        if index:
            p.add_run().add_break()
        run = p.add_run(display_text(line[0]))
        set_run_font(run, size=line[1], bold=line[2], color=line[3])
    return p


def add_band(text, fill, color, size, *, before=0, after=4, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER, pad_left=0.25, pad_right=0.25):
    return add_para(
        text,
        size=size,
        bold=bold,
        color=color,
        align=align,
        before=before,
        after=after,
        line=1.0,
        fill=fill,
        left=pad_left,
        right=pad_right,
    )


def set_cell_margins(cell, top=90, start=120, bottom=90, end=120):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for margin, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{margin}"))
        if node is None:
            node = OxmlElement(f"w:{margin}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_table_borders(table, color=LIGHT_BORDER, size=4):
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.first_child_found_in("w:tblBorders")
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        node = borders.find(qn(f"w:{edge}"))
        if node is None:
            node = OxmlElement(f"w:{edge}")
            borders.append(node)
        node.set(qn("w:val"), "single")
        node.set(qn("w:sz"), str(size))
        node.set(qn("w:color"), color)


def set_table_geometry(table, widths_dxa, indent_dxa=TABLE_INDENT_DXA):
    if sum(widths_dxa) != CONTENT_DXA:
        raise ValueError(f"表格宽度必须合计{CONTENT_DXA}，实际{sum(widths_dxa)}")
    table.autofit = False
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    tbl_pr = table._tbl.tblPr
    tbl_w = tbl_pr.first_child_found_in("w:tblW")
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), str(CONTENT_DXA))
    tbl_w.set(qn("w:type"), "dxa")
    tbl_ind = tbl_pr.first_child_found_in("w:tblInd")
    if tbl_ind is None:
        tbl_ind = OxmlElement("w:tblInd")
        tbl_pr.append(tbl_ind)
    tbl_ind.set(qn("w:w"), str(indent_dxa))
    tbl_ind.set(qn("w:type"), "dxa")
    grid = table._tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for width in widths_dxa:
        col = OxmlElement("w:gridCol")
        col.set(qn("w:w"), str(width))
        grid.append(col)
    for row in table.rows:
        for index, cell in enumerate(row.cells):
            tc_w = cell._tc.get_or_add_tcPr().first_child_found_in("w:tcW")
            if tc_w is None:
                tc_w = OxmlElement("w:tcW")
                cell._tc.get_or_add_tcPr().append(tc_w)
            tc_w.set(qn("w:w"), str(widths_dxa[index]))
            tc_w.set(qn("w:type"), "dxa")
            cell.width = Cm(widths_dxa[index] / 1440 * 2.54)


def set_row_no_split(row, repeat=False):
    tr_pr = row._tr.get_or_add_trPr()
    cant = OxmlElement("w:cantSplit")
    tr_pr.append(cant)
    if repeat:
        header = OxmlElement("w:tblHeader")
        header.set(qn("w:val"), "true")
        tr_pr.append(header)


def set_cell(cell, text, *, size=9.8, bold=False, color=BLACK, fill=WHITE, align=WD_ALIGN_PARAGRAPH.LEFT, margins=(90, 120, 90, 120)):
    cell.text = ""
    cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)
    set_cell_margins(cell, *margins)
    p = cell.paragraphs[0]
    p.alignment = align
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.line_spacing = 1.05
    pieces = display_text(text).split("\n")
    for index, piece in enumerate(pieces):
        if index:
            p.add_run().add_break()
        run = p.add_run(piece)
        set_run_font(run, size=size, bold=bold, color=color)


def add_table(headers, rows, widths, aligns=None, *, font_size=9.8, accent_first=False, header_fill=NAVY, after=7, body_margins=(90, 120, 90, 120)):
    has_header = any(str(header).strip() for header in headers)
    table = doc.add_table(rows=1 if has_header else 0, cols=len(headers))
    set_table_geometry(table, widths)
    set_table_borders(table)
    aligns = aligns or [WD_ALIGN_PARAGRAPH.LEFT] * len(headers)
    if has_header:
        for index, header in enumerate(headers):
            set_cell(table.rows[0].cells[index], header, size=max(10, font_size), bold=True, color=WHITE, fill=header_fill, align=aligns[index], margins=(100, 120, 100, 120))
            table.rows[0].cells[index].paragraphs[0].paragraph_format.keep_with_next = True
        set_row_no_split(table.rows[0], repeat=True)
    for row_index, values in enumerate(rows, start=1):
        cells = table.add_row().cells
        fill = ICE_BLUE if row_index % 2 else WHITE
        for col_index, value in enumerate(values):
            cell_fill = PALE_BLUE if accent_first and col_index == 0 else fill
            set_cell(
                cells[col_index],
                value,
                size=font_size,
                bold=accent_first and col_index == 0,
                color=RED if accent_first and col_index == 0 else BLACK,
                fill=cell_fill,
                align=aligns[col_index],
                margins=body_margins,
            )
        set_row_no_split(table.rows[-1])
    add_para("", after=after)
    return table


def add_callout(
    label,
    text,
    *,
    accent=NAVY,
    fill=PALE_BLUE,
    body_size=11,
    body_line=1.15,
    page_break_before=False,
):
    p1 = add_para(label, size=12.5, bold=True, color=WHITE, before=3, after=0, keep=True, fill=accent, page_break_before=page_break_before, left=0.25, right=0.25)
    p2 = add_para(text, size=body_size, color=BLACK, before=0, after=8, line=body_line, fill=fill, left=0.25, right=0.25)
    return p1, p2


def add_section_title(number, title, subtitle, *, page_break_before=False):
    p = add_para(f"{number}  {title}", size=21, bold=True, color=WHITE, before=0, after=2, keep=True, page_break_before=page_break_before, fill=NAVY, left=0.3, right=0.3)
    if subtitle:
        add_para(subtitle, size=10.5, color=GOLD, before=0, after=12, keep=True, fill=NAVY, left=0.3, right=0.3)
    return p


def add_heading(text, level=2, *, page_break_before=False, color=None):
    if level == 1:
        return add_para(text, size=21, bold=True, color=color or NAVY, before=14, after=8, keep=True, page_break_before=page_break_before)
    if level == 2:
        return add_para(text, size=16, bold=True, color=color or DEEP_BLUE, before=11, after=6, keep=True, page_break_before=page_break_before)
    return add_para(text, size=13, bold=True, color=color or RED, before=8, after=4, keep=True, page_break_before=page_break_before)


def display_name(name):
    return re.sub(r"-U$", "", str(name))


def money_yi(value):
    return f"{value / 100_000_000:.2f}亿元"


def compact_reason(text):
    pieces = [part.strip("。 ") for part in re.split(r"[；。]", str(text or "")) if part.strip("。 ")]
    result = []
    seen = set()
    for piece in pieces:
        if piece not in seen:
            seen.add(piece)
            result.append(piece)
    return "；".join(result) + ("。" if result else "")


def add_page_field(paragraph):
    run = paragraph.add_run()
    fld_begin = OxmlElement("w:fldChar")
    fld_begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = " PAGE "
    fld_end = OxmlElement("w:fldChar")
    fld_end.set(qn("w:fldCharType"), "end")
    run._r.extend([fld_begin, instr, fld_end])
    set_run_font(run, size=9.5, color=MID_GRAY)


doc = Document()
section = doc.sections[0]
section.page_width = Cm(21.0)
section.page_height = Cm(29.7)
section.top_margin = Cm(1.55)
section.bottom_margin = Cm(1.45)
section.left_margin = Cm(1.55)
section.right_margin = Cm(1.55)
section.header_distance = Cm(0.65)
section.footer_distance = Cm(0.7)
section.different_first_page_header_footer = True

normal = doc.styles["Normal"]
normal.font.name = FONT
normal._element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
normal.font.size = Pt(11)
normal.font.color.rgb = RGBColor.from_string(BLACK)
normal.paragraph_format.space_after = Pt(6)
normal.paragraph_format.line_spacing = 1.1

for style_name, size, color, before, after in (
    ("Title", 30, NAVY, 0, 8),
    ("Heading 1", 21, NAVY, 16, 8),
    ("Heading 2", 16, DEEP_BLUE, 12, 6),
    ("Heading 3", 13, RED, 8, 4),
):
    style = doc.styles[style_name]
    style.font.name = FONT
    style._element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
    style.font.size = Pt(size)
    style.font.bold = True
    style.font.color.rgb = RGBColor.from_string(color)
    style.paragraph_format.space_before = Pt(before)
    style.paragraph_format.space_after = Pt(after)
    style.paragraph_format.keep_with_next = True

header = section.header
hp = header.paragraphs[0]
hp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
set_run_font(hp.add_run(f"{data['trade_date']}｜龙头深度研究｜A股涨停板龙头分类"), size=9.5, color=MID_GRAY)
footer = section.footer
fp = footer.paragraphs[0]
fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
add_page_field(fp)

# 第一页：业务封面与核心判断
add_para("A股每日涨停板", size=30, bold=True, color=NAVY, before=6, after=0, keep=True)
title2 = add_para("龙头分类报告", size=30, bold=True, color=RED, before=0, after=7, keep=True)
set_paragraph_border(title2, color=GOLD, size=14)
add_para(f"{data['trade_date']}收盘复盘", size=15, bold=True, color=NAVY, before=5, after=2)
add_para(f"数据截止时间：{data['generated_at'][:16].replace('T', ' ')}｜普通A股涨停口径", size=10.5, color=MID_GRAY, after=10)

metric_rows = [["普通A股涨停", "投资板块", "上午最终封板", "龙虎榜净买过亿"], [f"{len(data['stocks'])}只", f"{len(data['concept_summary'])}类", f"{len(data['morning_limit_ups'])}只", f"{len(data['lhb_net_buy_ge_100m'])}只"]]
metric_table = doc.add_table(rows=2, cols=4)
set_table_geometry(metric_table, [2538, 2538, 2537, 2537])
set_table_borders(metric_table, color=WHITE, size=5)
for col, text in enumerate(metric_rows[0]):
    set_cell(metric_table.rows[0].cells[col], text, size=10.3, bold=True, color=WHITE, fill=NAVY, align=WD_ALIGN_PARAGRAPH.CENTER, margins=(100, 100, 100, 100))
for col, text in enumerate(metric_rows[1]):
    set_cell(metric_table.rows[1].cells[col], text, size=22, bold=True, color=RED, fill=PALE_GOLD, align=WD_ALIGN_PARAGRAPH.CENTER, margins=(140, 100, 140, 100))
for row in metric_table.rows:
    set_row_no_split(row)
add_heading("当日核心判断", level=2)
conclusion_rows = build_conclusion_rows(data)
add_table(
    ["", ""],
    conclusion_rows,
    [1550, 8600],
    [WD_ALIGN_PARAGRAPH.CENTER, WD_ALIGN_PARAGRAPH.LEFT],
    font_size=10.8,
    accent_first=True,
    header_fill=NAVY,
    after=4,
    body_margins=(60, 120, 60, 120),
)

# 一页看懂
add_section_title("01", "一页看懂当天结构", "")
top_two = sum(item["verified_limit_up_count"] for item in data["concept_summary"][:2])
morning_count = len(data["morning_limit_ups"])
top_lhb = data["lhb_net_buy_ge_100m"][0] if data["lhb_net_buy_ge_100m"] else None
dash = doc.add_table(rows=2, cols=4)
set_table_geometry(dash, [2538, 2538, 2537, 2537])
set_table_borders(dash, color=WHITE, size=5)
dash_labels = ["前两大板块合计", "占全部涨停", "上午最终封板占比", "龙虎榜最高净买"]
dash_values = [f"{top_two}只", f"{top_two / len(data['stocks']) * 100:.1f}%", f"{morning_count / len(data['stocks']) * 100:.1f}%", money_yi(top_lhb["net_amount"]) if top_lhb else "无符合项"]
for col in range(4):
    set_cell(dash.rows[0].cells[col], dash_labels[col], size=10.2, bold=True, color=WHITE, fill=NAVY, align=WD_ALIGN_PARAGRAPH.CENTER)
    set_cell(dash.rows[1].cells[col], dash_values[col], size=18, bold=True, color=RED, fill=PALE_GOLD, align=WD_ALIGN_PARAGRAPH.CENTER, margins=(130, 100, 130, 100))
    set_row_no_split(dash.rows[0]); set_row_no_split(dash.rows[1])
top_names = "、".join(f"{item['concept']}{item['verified_limit_up_count']}只" for item in data['concept_summary'][:2])
add_callout("当天最集中的方向", f"前两大投资板块为{top_names}，合计{top_two}只，占{len(data['stocks'])}只普通A股涨停的{top_two / len(data['stocks']) * 100:.1f}%。", accent=GOLD, fill=PALE_GOLD, body_size=11.2)

# 市场高度
height_rows = sorted(data["stocks"], key=lambda row: (-row["recent_limit_up_hits"], -row["consecutive_limit_ups"], row["first_limit_time"], row["stock_code"]))[:2]
market_leader = next((row for row in data["stocks"] if any(role["name"] == "市场总龙头" for role in row["leader_roles"])), None)
height_subtitle = f"{market_leader['stock_name']}形成唯一最高近期涨停高度" if market_leader else "最高近期涨停高度存在并列，不强行指定市场总龙头"
add_section_title("02", "市场高度与龙头性质", height_subtitle, page_break_before=True)
leaders = [
    ["最高近期高度", "第二高度参照"],
    [
        f"{height_rows[0]['stock_name']}\n{height_rows[0]['stock_code']}｜{height_rows[0]['recent_limit_up_label']}\n标准连板{height_rows[0]['consecutive_limit_ups']}板\n首次{height_rows[0]['first_limit_time'][:5]}，最终{height_rows[0]['final_limit_time'][:5]}",
        f"{height_rows[1]['stock_name']}\n{height_rows[1]['stock_code']}｜{height_rows[1]['recent_limit_up_label']}\n标准连板{height_rows[1]['consecutive_limit_ups']}板\n首次{height_rows[1]['first_limit_time'][:5]}，最终{height_rows[1]['final_limit_time'][:5]}",
    ],
]
leader_table = doc.add_table(rows=2, cols=2)
set_table_geometry(leader_table, [5075, 5075])
set_table_borders(leader_table)
set_cell(leader_table.rows[0].cells[0], leaders[0][0], size=11, bold=True, color=WHITE, fill=RED, align=WD_ALIGN_PARAGRAPH.CENTER)
set_cell(leader_table.rows[0].cells[1], leaders[0][1], size=11, bold=True, color=WHITE, fill=GOLD, align=WD_ALIGN_PARAGRAPH.CENTER)
set_cell(leader_table.rows[1].cells[0], leaders[1][0], size=12, bold=True, color=DARK_RED, fill=PALE_RED, align=WD_ALIGN_PARAGRAPH.CENTER, margins=(150, 130, 150, 130))
set_cell(leader_table.rows[1].cells[1], leaders[1][1], size=12, bold=True, color=DARK_RED, fill=PALE_GOLD, align=WD_ALIGN_PARAGRAPH.CENTER, margins=(150, 130, 150, 130))
for row in leader_table.rows: set_row_no_split(row)
if market_leader:
    add_callout("市场总龙头判断", f"{market_leader['stock_name']}以{market_leader['recent_limit_up_label']}形成全市场唯一最高近期涨停高度；标准连板为{market_leader['consecutive_limit_ups']}板。", accent=DEEP_BLUE, fill=ICE_BLUE, body_size=11.2)
else:
    add_callout("市场总龙头判断", f"最高近期涨停高度存在并列，当日{len(data['stocks'])}只涨停股中没有形成唯一市场总龙头。", accent=DEEP_BLUE, fill=ICE_BLUE, body_size=11.2)

# 投资板块全景
add_section_title("03", f"{len(data['concept_summary'])}类投资板块全景", "")
overview_rows = []
for i, item in enumerate(data["concept_summary"], start=1):
    overview_rows.append([item["concept"], f"{item['verified_limit_up_count']}只", item["earliest_first_limit_time"][:5]])
add_table(["投资板块", "涨停家数", "最早首次封板"], overview_rows, [4700, 2200, 3250], [WD_ALIGN_PARAGRAPH.LEFT, WD_ALIGN_PARAGRAPH.CENTER, WD_ALIGN_PARAGRAPH.CENTER], font_size=11, accent_first=True)

# 仅深析涨停家数前三板块
add_section_title("04", "涨停家数前三板块深度分析", "")
for concept_index, item in enumerate(data["concept_summary"][:3], start=1):
    add_heading(f"{concept_index:02d}｜{item['concept']}  ·  {item['verified_limit_up_count']}只涨停", level=1)
    model = item["logic_model"]
    add_para(
        f"逻辑类型：{model['logic_type']}｜"
        f"盘面结构证据强度：{model['structural_evidence_strength']}",
        size=10.5,
        bold=True,
        color=MID_GRAY,
        before=0,
        after=5,
        keep=True,
    )
    add_callout("核心逻辑", item["logic_summary"], accent=DEEP_BLUE, fill=ICE_BLUE, body_size=11.2)
    add_callout("机制、事件与时序", item["mechanism_evidence"], accent=GOLD, fill=PALE_GOLD, body_size=11.2)
    add_callout("盘面验证与约束", item["validation_constraints"], accent=RED, fill=PALE_RED, body_size=11.2)
    add_callout("板块内部地位结论", model["position_conclusion"], accent=DEEP_BLUE, fill=ICE_BLUE, body_size=11.2)
    add_heading("性质榜", level=3)
    if item["nature_top"]:
        rows = []
        for nature in item["nature_top"]:
            stock = stocks[nature["stock_code"]]
            rows.append([str(nature["rank"]), f"{display_name(stock['stock_name'])}\n{stock['stock_code']}", "、".join(nature["nature_names"]), compact_reason(nature["reason"])])
        add_table(["序", "个股", "当天承担的作用", "事实依据"], rows, [550, 1700, 2400, 5500], [WD_ALIGN_PARAGRAPH.CENTER, WD_ALIGN_PARAGRAPH.CENTER, WD_ALIGN_PARAGRAPH.LEFT, WD_ALIGN_PARAGRAPH.LEFT], font_size=9.8, accent_first=True)
    add_heading("地位榜", level=3)
    rows = []
    for position in item["position_top"]:
        stock = stocks[position["stock_code"]]
        rows.append([
            f"龙{position['rank']}",
            f"{display_name(stock['stock_name'])}\n{stock['stock_code']}",
            (
                f"{stock['recent_limit_up_label']}｜连板{stock['consecutive_limit_count']}｜"
                f"开板{stock['open_count']}｜最终{stock['final_limit_time'][:5]}"
            ),
            compact_reason(position["relative_comparison"]),
            compact_reason(position["weakness"] + "；" + position["invalidation"]),
        ])
    add_table(
        ["地位", "个股", "硬指标", "为何在此位", "短板与失效"],
        rows,
        [600, 1500, 2250, 3100, 2700],
        [
            WD_ALIGN_PARAGRAPH.CENTER,
            WD_ALIGN_PARAGRAPH.CENTER,
            WD_ALIGN_PARAGRAPH.LEFT,
            WD_ALIGN_PARAGRAPH.LEFT,
            WD_ALIGN_PARAGRAPH.LEFT,
        ],
        font_size=9.5,
        accent_first=True,
    )

# 上午封板
add_section_title("05", "上午最终封板时间分类", "")
def bucket(value):
    hhmm = value[:5]
    if hhmm <= "09:30": return "09:25—09:30"
    if hhmm <= "10:00": return "09:31—10:00"
    if hhmm <= "10:30": return "10:01—10:30"
    if hhmm <= "11:00": return "10:31—11:00"
    return "11:01—11:30"

bucket_order = ["09:25—09:30", "09:31—10:00", "10:01—10:30", "10:31—11:00", "11:01—11:30"]
bucket_counts = Counter(bucket(item["final_limit_time"]) for item in data["morning_limit_ups"])
bar_rows = []
max_bucket = max(bucket_counts.values())
for label in bucket_order:
    blocks = max(1, round(bucket_counts[label] / max_bucket * 22))
    bar_rows.append([label, "█" * blocks, f"{bucket_counts[label]}只"])
add_heading(f"{len(data['morning_limit_ups'])}只股票的最终封板时段", level=2)
add_table(["时段", "数量对比", "家数"], bar_rows, [2500, 6000, 1650], [WD_ALIGN_PARAGRAPH.LEFT, WD_ALIGN_PARAGRAPH.LEFT, WD_ALIGN_PARAGRAPH.RIGHT], font_size=10.8, accent_first=True)
morning_rows = [[str(item["rank"]), item["first_limit_time"][:5], item["final_limit_time"][:5], f"{display_name(item['stock_name'])}\n{item['stock_code']}", item["primary_concept"], str(item["open_count"])] for item in data["morning_limit_ups"]]
add_table(["序", "首次封板", "最终封板", "个股", "投资板块", "开板次数"], morning_rows, [650, 1250, 1250, 1900, 3950, 1150], [WD_ALIGN_PARAGRAPH.CENTER, WD_ALIGN_PARAGRAPH.CENTER, WD_ALIGN_PARAGRAPH.CENTER, WD_ALIGN_PARAGRAPH.CENTER, WD_ALIGN_PARAGRAPH.LEFT, WD_ALIGN_PARAGRAPH.CENTER], font_size=10, accent_first=True)

# 龙虎榜
add_section_title("06", "龙虎榜单日净买入过亿元", "")
lhb_rows = [[str(item["rank"]), f"{display_name(item['stock_name'])}\n{item['stock_code']}", item["primary_concept"], money_yi(item["buy_amount"]), money_yi(item["sell_amount"]), money_yi(item["net_amount"]), f"{item['first_limit_time'][:5]}\n{item['final_limit_time'][:5]}"] for item in data["lhb_net_buy_ge_100m"]]
add_table(["序", "个股", "投资板块", "买入", "卖出", "净买入", "首次\n最终"], lhb_rows, [500, 1550, 1200, 1500, 1500, 1600, 2300], [WD_ALIGN_PARAGRAPH.CENTER, WD_ALIGN_PARAGRAPH.CENTER, WD_ALIGN_PARAGRAPH.LEFT, WD_ALIGN_PARAGRAPH.RIGHT, WD_ALIGN_PARAGRAPH.RIGHT, WD_ALIGN_PARAGRAPH.RIGHT, WD_ALIGN_PARAGRAPH.CENTER], font_size=10, accent_first=True)
for item_index, item in enumerate(data["lhb_net_buy_ge_100m"]):
    seat_rows = [[seat["seat_name"], seat["seat_identity"], money_yi(seat["buy_amount"]), money_yi(seat["sell_amount"]), money_yi(seat["net_amount"])] for seat in item["seats"]]
    add_heading(
        f"{display_name(item['stock_name'])} {item['stock_code']}｜公开榜单席位",
        level=2,
        page_break_before=item_index > 0 and len(seat_rows) >= 5,
    )
    add_table(["席位名称", "身份", "买入", "卖出", "净额"], seat_rows, [4100, 1700, 1450, 1450, 1450], [WD_ALIGN_PARAGRAPH.LEFT, WD_ALIGN_PARAGRAPH.CENTER, WD_ALIGN_PARAGRAPH.RIGHT, WD_ALIGN_PARAGRAPH.RIGHT, WD_ALIGN_PARAGRAPH.RIGHT], font_size=9.5, accent_first=True)

doc.core_properties.title = f"{data['trade_date']} 龙头深度研究 A股涨停板龙头分类报告"
doc.core_properties.subject = "按一级投资板块、细分方向、当日推动因素、龙头性质、板块地位、上午封板时间和龙虎榜单日净额分类"
doc.core_properties.author = "小牛"
doc.core_properties.comments = (
    f"data_evidence_sha256={manifest['result_sha256']} | "
    f"template_contract_sha256={template_contract_sha256} | "
    f"STOCK_DATA_TRADE_DATE={data['trade_date']}"
)

temporary_path = FINAL_PATH.with_name(FINAL_PATH.stem + ".building.docx")
doc.save(temporary_path)
temporary_path.replace(FINAL_PATH)
print(json.dumps({
    "status": "PASS",
    "docx": str(FINAL_PATH),
    "bytes": FINAL_PATH.stat().st_size,
    "stocks": len(data["stocks"]),
    "concepts": len(data["concept_summary"]),
    "morning": len(data["morning_limit_ups"]),
    "lhb": len(data["lhb_net_buy_ge_100m"]),
    "result_sha256": manifest["result_sha256"],
    "template_contract_sha256": template_contract_sha256,
}, ensure_ascii=False, indent=2))
