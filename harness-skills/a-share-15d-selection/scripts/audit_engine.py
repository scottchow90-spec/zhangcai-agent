#!/usr/bin/env python3
from __future__ import annotations
# ONESTOCK_EMBEDDED_LOCAL_IMPORT_BOOTSTRAP
import sys as _onestock_embedded_sys
from pathlib import Path as _OneStockEmbeddedPath
_onestock_embedded_dir = str(_OneStockEmbeddedPath(__file__).resolve().parent)
if _onestock_embedded_dir not in _onestock_embedded_sys.path:
    _onestock_embedded_sys.path.insert(0, _onestock_embedded_dir)

import argparse
import csv
import hashlib
import json
import math
import os
import zipfile
from collections import Counter
from pathlib import Path
from typing import Any

from docx import Document


SKILL_ROOT = Path(__file__).resolve().parents[1]
SKILLS_ROOT = SKILL_ROOT.parent
GLOBAL_SCORE_CONTRACT = SKILLS_ROOT / "stock-unified" / "references" / "short_term_strong_stock_scoring_contract.json"
GLOBAL_SCORE_CONTRACT_PAYLOAD = json.loads(GLOBAL_SCORE_CONTRACT.read_text(encoding="utf-8"))
GLOBAL_SCORE_CONTRACT_SHA256 = hashlib.sha256(GLOBAL_SCORE_CONTRACT.read_bytes()).hexdigest()


def audit_reports_root() -> Path:
    onestock_root = os.environ.get("ONESTOCK_STOCK_DATA_ROOT", "").strip()
    if onestock_root:
        return Path(onestock_root).resolve() / "stock-custom-audit" / SKILL_ROOT.name
    return SKILL_ROOT / "reports" / "audit"


TEMPLATE = SKILL_ROOT / "assets" / "15维选股_教学案例精美Word模板.docx"
TEMPLATE_SHA256 = "EC71F04176CB13FBF56AC73D60EE6701CCC71AB877C457D365E130A1795993F5"
REQUIRED_FORMULAS = (
    "大牛线4.0",
    "飞龙在天",
    "游资资金监控",
    "机构资金监控",
)
REQUIRED_FORMULA_SUMMARY = (
    "大牛线",
    "飞龙在天",
    "游资资金",
    "机构资金",
    "庄家资金",
)
REQUIRED_VALIDATIONS = (
    "complete_scan",
    "result_count_matches",
    "result_count_lte_target",
    "scarcity_disclosed",
    "all_top5_hard_gates",
    "all_top5_26_factors",
    "all_top5_15_risks",
    "all_top5_7_hard_exclusions",
    "all_top5_required_formulas",
    "all_top5_segment_le_80",
    "all_top5_score_math",
    "ranking_consistent",
    "score_contract_v6",
    "ranking_count_matches",
    "all_ranked_hard_gates",
    "all_ranked_26_factors",
    "all_ranked_15_risks",
    "all_ranked_7_hard_exclusions",
    "all_ranked_required_formulas",
    "all_ranked_segment_le_80",
    "all_ranked_score_math",
    "full_ranking_consistent",
    "top5_is_ranking_prefix",
    "pool_coverage_complete",
    "market_scope_complete",
    "current_trade_date",
)
CSV_HEADERS = (
    "排名",
    "代码",
    "名称",
    "现价",
    "涨幅",
    "正向得分",
    "风险扣分",
    "最终得分",
    "飞龙波",
    "飞龙段",
    "主趋势线",
    "失效条件",
)
POSITIVE_DIMENSIONS = (
    ("强势收盘强度", 4.0),
    ("前高突破强度", 4.0),
    ("短周期收益加速度", 4.0),
    ("健康放量质量", 4.0),
    ("成交额放大强度", 3.0),
    ("大牛线攻击信号", 4.0),
    ("飞龙启动信号", 5.0),
    ("游资点火强度", 4.0),
    ("机构净流入强度", 3.0),
    ("多周期相对强度", 5.0),
    ("均线多头结构", 4.0),
    ("趋势斜率", 4.0),
    ("趋势效率", 4.0),
    ("飞龙波段阶段", 4.0),
    ("趋势乖离健康度", 3.0),
    ("量价一致性", 3.0),
    ("资金合力持续性", 4.0),
    ("催化延续证据", 4.0),
    ("市场广度", 4.0),
    ("涨停生态", 4.0),
    ("主线题材共振", 7.0),
    ("全市场相对强度", 5.0),
    ("流动性分位", 3.0),
    ("日内承接强度", 3.0),
    ("流动性稳定度", 2.0),
    ("波动可控性", 2.0),
)
SCORE_CONTRACT_VERSION = "A-SHARE-STRONG-26F-100-V6.1"
FACTOR_AXES = {
    **{name: "爆发力" for name, _weight in POSITIVE_DIMENSIONS[:9]},
    **{name: "持续性" for name, _weight in POSITIVE_DIMENSIONS[9:18]},
    **{name: "市场协同" for name, _weight in POSITIVE_DIMENSIONS[18:22]},
    **{name: "可交易性" for name, _weight in POSITIVE_DIMENSIONS[22:]},
}
FACTOR_ROLES = {
    **{name: "candidate_ranking" for name, _weight in POSITIVE_DIMENSIONS},
    "强势收盘强度": "pool_confirmation",
    "催化延续证据": "candidate_evidence",
    "市场广度": "market_regime",
    "涨停生态": "market_regime",
}
RANKING_FACTOR_ROLES = {"candidate_ranking", "candidate_evidence"}
AXIS_MAX_SCORES = {"爆发力": 35.0, "持续性": 35.0, "市场协同": 20.0, "可交易性": 10.0}
RISK_DEDUCTIONS = (
    ("大牛线庄出", 2),
    ("飞龙高位背离", 2),
    ("飞龙波段转弱", 2),
    ("机构大单出", 2),
    ("游资买方意向转负", 2),
    ("跌破主趋势线", 3),
    ("极端放量与大幅分歧", 2),
    ("重大公告或财务风险", 3),
    ("波段过高且无资金承接", 3),
    ("大额解禁或减持", 2),
    ("失控加速且承接不足", 2),
    ("高位钝化区", 2),
    ("强压力位压制", 1),
    ("五浪末端与指标背离", 2),
    ("反弹陷阱", 1),
)
HARD_NAMES = (
    "风险警示或退市整理",
    "立案调查或财务造假",
    "当日一字板无法参与",
    "波大于90且无主升启动信号",
    "近3日极端加速且无启动承接",
    "近5日或10日极端加速且无资金承接",
    "飞龙在天段大于80或字段缺失",
)


def sha256_file(path: Path) -> str:
    return hashlib.sha256(path.read_bytes()).hexdigest()


def file_artifact(path: Path) -> dict[str, Any]:
    return {
        "path": str(path.resolve()),
        "size": path.stat().st_size,
        "sha256": sha256_file(path),
    }


def _finite_number(value: Any) -> float | None:
    if isinstance(value, bool) or not isinstance(value, (int, float)):
        return None
    number = float(value)
    return number if math.isfinite(number) else None


def _same_number(left: Any, right: Any, tolerance: float = 0.011) -> bool:
    left_number = _finite_number(left)
    right_number = _finite_number(right)
    return left_number is not None and right_number is not None and abs(left_number - right_number) <= tolerance


def _same_csv_number(value: str | None, expected: Any, tolerance: float = 0.011) -> bool:
    try:
        return _same_number(float(value or ""), expected, tolerance)
    except (TypeError, ValueError):
        return False


def _coerce_number(value: Any) -> float | None:
    if isinstance(value, list):
        value = value[-1] if value else None
    if value is None or value == "" or isinstance(value, bool):
        return None
    try:
        number = float(str(value).replace("%", "").strip())
        return number if math.isfinite(number) else None
    except (TypeError, ValueError):
        return None


def _signal_active(value: Any, positive: bool = True) -> bool:
    number = _coerce_number(value)
    if number is not None:
        return number > 0
    text = str(value or "").strip()
    tokens = ("启动", "打板", "秘进", "龙头", "点火", "金叉", "庄进") if positive else ("庄出", "卖出", "离场", "转弱")
    return bool(text) and any(token in text for token in tokens)


def _ma(rows: list[dict[str, Any]], days: int) -> float | None:
    if len(rows) < days:
        return None
    values = [_coerce_number(row.get("close")) for row in rows[-days:]]
    return sum(values) / days if all(value is not None for value in values) else None


def _ema(values: list[float], period: int) -> list[float]:
    if not values:
        return []
    alpha = 2.0 / (period + 1)
    result = [values[0]]
    for value in values[1:]:
        result.append(alpha * value + (1 - alpha) * result[-1])
    return result


def _macd_bearish_divergence(closes: list[float]) -> bool:
    if len(closes) < 10:
        return False
    e12, e26 = _ema(closes, 12), _ema(closes, 26)
    dif = [left - right for left, right in zip(e12, e26)]
    dea = _ema(dif, 9)
    hist = [(left - right) * 2 for left, right in zip(dif, dea)]
    prior_positive_peak = max(hist[-10:-1])
    return closes[-1] >= max(closes[-10:]) and prior_positive_peak > 0 and hist[-1] >= 0 and hist[-1] < prior_positive_peak * 0.75


def _cross_sectional_stars(value: float, cohort: list[float]) -> float:
    if value <= 0:
        return 0.0
    valid = [float(peer) for peer in cohort if math.isfinite(float(peer)) and float(peer) > 0]
    if len(valid) < 5:
        return 5.0
    stronger = sum(peer > value + 1e-12 for peer in valid)
    percentile = stronger / len(valid)
    if percentile < 0.10:
        return 5.0
    if percentile < 0.30:
        return 4.0
    if percentile < 0.60:
        return 3.0
    return 1.5


def _percentile_rank(value: float | None, cohort: list[float]) -> float:
    if value is None or not math.isfinite(float(value)):
        return 0.0
    valid = [float(peer) for peer in cohort if math.isfinite(float(peer))]
    if not valid:
        return 0.0
    numeric = float(value)
    below = sum(peer < numeric - 1e-12 for peer in valid)
    equal = sum(abs(peer - numeric) <= 1e-12 for peer in valid)
    return (below + 0.5 * equal) / len(valid)


def _percentile_stars(percentile: float) -> float:
    return 5.0 if percentile >= 0.90 else 4.0 if percentile >= 0.70 else 3.0 if percentile >= 0.40 else 1.5 if percentile >= 0.20 else 0.0


def _ratio_to_prior_average(rows: list[dict[str, Any]], field: str, days: int = 20) -> float:
    history = rows[-days - 1:-1] if len(rows) >= days + 1 else rows[:-1]
    values = [_coerce_number(row.get(field)) for row in history]
    usable = [value for value in values if value is not None and value > 0]
    latest = _coerce_number(rows[-1].get(field)) if rows else None
    return latest / (sum(usable) / len(usable)) if latest is not None and usable else 0.0


def _return_acceleration(state: dict[str, Any]) -> float:
    gain3 = _coerce_number(state.get("recent_gain_3d_pct"))
    gain10 = _coerce_number(state.get("recent_gain_10d_pct"))
    return gain3 / 3.0 - (gain10 - gain3) / 7.0 if gain3 is not None and gain10 is not None else 0.0


def _relative_strength(state: dict[str, Any]) -> float:
    gains = [_coerce_number(state.get(key)) for key in ("recent_gain_3d_pct", "recent_gain_5d_pct", "recent_gain_10d_pct")]
    return 0.25 * gains[0] + 0.35 * gains[1] + 0.40 * gains[2] if all(value is not None for value in gains) else 0.0


def _volume_quality(rows: list[dict[str, Any]]) -> float:
    ratio = _ratio_to_prior_average(rows, "volume")
    return math.exp(-abs(math.log(ratio / 1.8))) if 0.4 <= ratio <= 6.0 else 0.0


def _trend_slope(rows: list[dict[str, Any]]) -> float:
    if len(rows) < 10:
        return 0.0
    current = sum(float(row["close"]) for row in rows[-5:]) / 5
    prior = sum(float(row["close"]) for row in rows[-10:-5]) / 5
    return (current / prior - 1) * 100 if prior > 0 else 0.0


def _trend_efficiency(rows: list[dict[str, Any]], days: int = 10) -> float:
    if len(rows) < days + 1:
        return 0.0
    values = [float(row["close"]) for row in rows[-days - 1:]]
    path = sum(abs(right - left) for left, right in zip(values, values[1:]))
    return max(0.0, (values[-1] - values[0]) / path) if path > 0 else 0.0


def _price_volume_consistency(rows: list[dict[str, Any]], days: int = 10) -> float:
    sample = rows[-days - 1:]
    matches = 0
    observations = 0
    for previous, current in zip(sample, sample[1:]):
        price_change = float(current["close"]) - float(previous["close"])
        volume_change = float(current["volume"]) - float(previous["volume"])
        if price_change == 0:
            continue
        observations += 1
        matches += int((price_change > 0 and volume_change >= 0) or (price_change < 0 and volume_change <= 0))
    return matches / observations if observations else 0.0


def _institution_strength(item: dict[str, Any]) -> float:
    inputs = item.get("audit_inputs") if isinstance(item, dict) else None
    formulas = inputs.get("formula") if isinstance(inputs, dict) else None
    inst = formulas.get("institution") if isinstance(formulas, dict) and isinstance(formulas.get("institution"), dict) else {}
    values = [_coerce_number(inst.get(key)) for key in ("机构大单进", "机构大单出", "大单动向", "大户大单进", "散户资金进")]
    if any(value is None for value in values):
        return 0.0
    inst_in, inst_out, big_order, large_in, retail_in = values
    normalized_net = max(0.0, inst_in - inst_out) / (abs(inst_in) + abs(inst_out) + 1e-9)
    return normalized_net + sum((big_order > 0, large_in > 0, retail_in < 0)) / 3.0


def _liquidity_stability(rows: list[dict[str, Any]], days: int = 20) -> float:
    values = [float(row["amount"]) for row in rows[-days:] if float(row["amount"]) > 0]
    if len(values) < 5:
        return 0.0
    mean = sum(values) / len(values)
    variance = sum((value - mean) ** 2 for value in values) / len(values)
    return 1.0 / (1.0 + math.sqrt(variance) / mean) if mean > 0 else 0.0


def _atr_percent(rows: list[dict[str, Any]], days: int = 10) -> float:
    if len(rows) < days + 1:
        return 0.0
    ranges = []
    for previous, current in zip(rows[-days - 1:-1], rows[-days:]):
        high, low, prev_close = float(current["high"]), float(current["low"]), float(previous["close"])
        ranges.append(max(high - low, abs(high - prev_close), abs(low - prev_close)))
    close = float(rows[-1]["close"])
    return sum(ranges) / len(ranges) / close * 100 if close > 0 else 0.0


def _volatility_quality(rows: list[dict[str, Any]]) -> float:
    atr_value = _atr_percent(rows)
    if atr_value <= 0:
        return 0.0
    return math.exp(-abs(math.log(atr_value / 5.0)))


def _market_breadth_stars(macro: dict[str, Any]) -> float:
    advancers = int(macro.get("advancers") or 0)
    decliners = int(macro.get("decliners") or 0)
    breadth = advancers / (advancers + decliners) if advancers + decliners else 0.0
    return 5.0 if breadth >= 0.60 else 4.0 if breadth >= 0.52 else 3.0 if breadth >= 0.45 else 1.5 if breadth >= 0.38 else 0.0


def _limit_up_ecology_stars(macro: dict[str, Any]) -> float:
    market_count = int(macro.get("market_count") or 0)
    density = int(macro.get("limit_up_count") or 0) / market_count if market_count else 0.0
    return 5.0 if density >= 0.015 else 4.0 if density >= 0.010 else 3.0 if density >= 0.005 else 1.5 if density >= 0.002 else 0.0


def _short_fund_strength(item: dict[str, Any]) -> float:
    inputs = item.get("audit_inputs")
    if not isinstance(inputs, dict):
        return 0.0
    rows = inputs.get("rows")
    formulas = inputs.get("formula")
    if not isinstance(rows, list) or not rows or not isinstance(formulas, dict):
        return 0.0
    hot = formulas.get("hot") if isinstance(formulas.get("hot"), dict) else {}
    close = _coerce_number(rows[-1].get("close"))
    buyer = _coerce_number(hot.get("买方意向"))
    aaa = _coerce_number(hot.get("AAA"))
    ddd = _coerce_number(hot.get("DDD"))
    if None in (close, buyer, aaa, ddd) or close <= 0 or buyer <= 0 or aaa <= ddd or aaa <= 0:
        return 0.0
    return buyer / close


def _theme_strength(
    concepts: list[Any],
    concept_frequency: Counter[str],
    population: int,
) -> float:
    peer_counts = sorted(
        (max(0, int(concept_frequency[str(token)]) - 1) for token in set(concepts)),
        reverse=True,
    )
    return sum(peer_counts[:3]) / max(1, population)


def _build_scoring_context(
    ranking: list[dict[str, Any]],
    concept_frequency: Counter[str],
    macro: dict[str, Any] | None = None,
) -> dict[str, Any]:
    population = max(1, len(ranking))
    concepts_by_item = []
    states = []
    rows_by_item = []
    for item in ranking:
        inputs = item.get("audit_inputs") if isinstance(item, dict) else None
        concepts = inputs.get("concepts") if isinstance(inputs, dict) else None
        concepts_by_item.append(concepts if isinstance(concepts, list) else [])
        state = inputs.get("state") if isinstance(inputs, dict) else None
        states.append(state if isinstance(state, dict) else {})
        rows = inputs.get("rows") if isinstance(inputs, dict) else None
        rows_by_item.append(rows if isinstance(rows, list) else [])
    resolved_macro = dict(macro or {})
    if not resolved_macro and states and isinstance(states[0].get("market_context"), dict):
        resolved_macro = dict(states[0]["market_context"])
    return {
        "short_fund_strengths": [_short_fund_strength(item) for item in ranking],
        "institution_strengths": [_institution_strength(item) for item in ranking],
        "theme_strengths": [
            _theme_strength(concepts, concept_frequency, population)
            for concepts in concepts_by_item
        ],
        "return_accelerations": [_return_acceleration(state) for state in states],
        "relative_strengths": [_relative_strength(state) for state in states],
        "volume_qualities": [_volume_quality(rows) for rows in rows_by_item],
        "amount_ratios": [_ratio_to_prior_average(rows, "amount") for rows in rows_by_item],
        "trend_slopes": [_trend_slope(rows) for rows in rows_by_item],
        "trend_efficiencies": [_trend_efficiency(rows) for rows in rows_by_item],
        "amounts": [float(rows[-1]["amount"]) if rows else 0.0 for rows in rows_by_item],
        "liquidity_stabilities": [_liquidity_stability(rows) for rows in rows_by_item],
        "volatility_qualities": [_volatility_quality(rows) for rows in rows_by_item],
        "macro": resolved_macro,
    }


def _expected_candidate_semantics(
    item: dict[str, Any],
    concept_frequency: Counter[str] | None = None,
    scoring_context: dict[str, list[float]] | None = None,
) -> tuple[list[float], list[int]] | None:
    inputs = item.get("audit_inputs")
    if not isinstance(inputs, dict):
        return None
    rows = inputs.get("rows")
    formulas = inputs.get("formula")
    evidence = inputs.get("evidence")
    state = inputs.get("state")
    concepts = inputs.get("concepts")
    if not isinstance(rows, list) or not rows or not isinstance(formulas, dict) or not isinstance(evidence, dict) or not isinstance(state, dict) or not isinstance(concepts, list):
        return None
    big = formulas.get("big") if isinstance(formulas.get("big"), dict) else {}
    fly = formulas.get("fly") if isinstance(formulas.get("fly"), dict) else {}
    hot = formulas.get("hot") if isinstance(formulas.get("hot"), dict) else {}
    inst = formulas.get("institution") if isinstance(formulas.get("institution"), dict) else {}
    dealer = formulas.get("dealer") if isinstance(formulas.get("dealer"), dict) else {}
    last = rows[-1]
    close = _coerce_number(last.get("close"))
    high = _coerce_number(last.get("high"))
    low = _coerce_number(last.get("low"))
    amount = _coerce_number(last.get("amount"))
    volume = _coerce_number(last.get("volume"))
    if None in (close, high, low, amount, volume):
        return None
    trend = _coerce_number(big.get("主趋势线"))
    ema9 = _coerce_number(big.get("EMA9"))
    ema10 = _coerce_number(big.get("EMA10"))
    ema11 = _coerce_number(big.get("EMA11"))
    wave = _coerce_number(fly.get("波"))
    segment = _coerce_number(fly.get("段"))
    buyer = _coerce_number(hot.get("买方意向"))
    aaa = _coerce_number(hot.get("AAA"))
    ddd = _coerce_number(hot.get("DDD"))
    inst_in = _coerce_number(inst.get("机构大单进"))
    inst_out = _coerce_number(inst.get("机构大单出"))
    big_order = _coerce_number(inst.get("大单动向"))
    large_in = _coerce_number(inst.get("大户大单进"))
    retail_in = _coerce_number(inst.get("散户资金进"))
    ma5, ma10, ma20, ma60 = (_ma(rows, days) for days in (5, 10, 20, 60))
    history20 = rows[-21:-1] if len(rows) >= 21 else rows[:-1]
    prior_volumes = [_coerce_number(row.get("volume")) for row in history20]
    usable_volumes = [value for value in prior_volumes if value is not None and value > 0]
    vol_ratio = volume / (sum(usable_volumes) / len(usable_volumes)) if usable_volumes else 0.0
    amount_ratio = _ratio_to_prior_average(rows, "amount")
    support20 = min((_coerce_number(row.get("low")) for row in history20), default=low)
    pressure20 = max((_coerce_number(row.get("high")) for row in history20), default=0.0)
    if support20 is None or pressure20 is None:
        return None
    big_signal = any(_signal_active(big.get(field)) for field in ("OUTPUT3", "OUTPUT4", "OUTPUT5", "OUTPUT9"))
    ignition = any(_signal_active(fly.get(field)) for field in ("OUTPUT3", "OUTPUT4", "OUTPUT5", "OUTPUT6"))
    institution_complete = None not in (inst_in, inst_out, big_order, large_in, retail_in)
    trend_mas_complete = None not in (ma5, ma10, ma20)
    wave_valid = wave is not None and segment is not None and 0 <= wave <= 100 and 0 <= segment <= 80
    downside = (close / trend - 1) * 100 if trend is not None and trend > 0 else None
    breakout = pressure20 > 0 and close > pressure20 * 1.005
    population = max(1, int(_coerce_number(state.get("universe_size")) or 1))
    frequencies = concept_frequency or Counter(str(token) for token in concepts)
    hot_strength = _short_fund_strength(item)
    inst_strength = _institution_strength(item)
    theme_metric = _theme_strength(concepts, frequencies, population)
    context = scoring_context or {
        "short_fund_strengths": [hot_strength],
        "institution_strengths": [inst_strength],
        "theme_strengths": [theme_metric],
        "return_accelerations": [_return_acceleration(state)],
        "relative_strengths": [_relative_strength(state)],
        "volume_qualities": [_volume_quality(rows)],
        "amount_ratios": [amount_ratio],
        "trend_slopes": [_trend_slope(rows)],
        "trend_efficiencies": [_trend_efficiency(rows)],
        "amounts": [amount],
        "liquidity_stabilities": [_liquidity_stability(rows)],
        "volatility_qualities": [_volatility_quality(rows)],
        "macro": state.get("market_context") if isinstance(state.get("market_context"), dict) else {},
    }
    macro = context.get("macro") if isinstance(context.get("macro"), dict) else {}
    market = str(state.get("market") or "SZ")
    code = str(state.get("code") or "000001")
    limit_rate = 0.30 if market == "BJ" else 0.20 if code.startswith(("300", "301", "688", "689")) else 0.10
    pct = _coerce_number(state.get("pct")) or 0.0
    strength_ratio = pct / (limit_rate * 100)
    limit_price = _coerce_number(state.get("limit_price"))
    at_limit = limit_price is not None and close >= limit_price - 0.011
    breakout_pct = (close / pressure20 - 1) * 100 if pressure20 > 0 else None
    acceleration = _return_acceleration(state)
    volume_quality = _volume_quality(rows)
    relative_strength = _relative_strength(state)
    slope = _trend_slope(rows)
    efficiency = _trend_efficiency(rows)
    consistency = _price_volume_consistency(rows)
    capital_checks = [
        hot_strength > 0,
        inst_strength > 0,
        institution_complete and big_order > 0,
        institution_complete and large_in > 0,
        institution_complete and retail_in < 0,
        evidence.get("lhb") is True,
    ]
    cohorts = macro.get("market_return_cohorts") if isinstance(macro.get("market_return_cohorts"), dict) else {}
    market_values = {
        "1d": _coerce_number(state.get("pct")),
        "3d": _coerce_number(state.get("recent_gain_3d_pct")),
        "5d": _coerce_number(state.get("recent_gain_5d_pct")),
        "10d": _coerce_number(state.get("recent_gain_10d_pct")),
    }
    persisted_percentiles = state.get("market_percentiles") if isinstance(state.get("market_percentiles"), dict) else {}
    percentiles = {
        horizon: _percentile_rank(value, cohorts.get(horizon, []))
        if isinstance(cohorts.get(horizon), list) and cohorts[horizon]
        else float(_coerce_number(persisted_percentiles.get(horizon)) or 0.0)
        for horizon, value in market_values.items()
    }
    market_relative = 0.15 * percentiles["1d"] + 0.35 * percentiles["3d"] + 0.30 * percentiles["5d"] + 0.20 * percentiles["10d"]
    market_relative_available = all(
        (isinstance(cohorts.get(horizon), list) and cohorts[horizon])
        or _coerce_number(persisted_percentiles.get(horizon)) is not None
        for horizon in market_values
    )
    previous_close = _coerce_number(rows[-2].get("close")) if len(rows) >= 2 else None
    recovery_pct = (close - low) / previous_close * 100 if previous_close else 0.0
    gap_pct = (float(last.get("open")) / previous_close - 1) * 100 if previous_close else 0.0
    liquidity_stability = _liquidity_stability(rows)
    atr10 = _atr_percent(rows)
    volatility_quality = _volatility_quality(rows)
    stars = [
        5.0 if at_limit else 4.0 if strength_ratio >= 0.85 else 3.0 if strength_ratio >= 0.60 else 1.5 if strength_ratio > 0 else 0.0,
        5.0 if breakout_pct is not None and breakout_pct >= 3 else 4.0 if breakout_pct is not None and breakout_pct >= 1 else 3.0 if breakout_pct is not None and breakout_pct > 0 else 1.5 if breakout_pct is not None and breakout_pct >= -2 else 0.0,
        _cross_sectional_stars(acceleration, context.get("return_accelerations", [])),
        _cross_sectional_stars(volume_quality, context.get("volume_qualities", [])),
        _cross_sectional_stars(amount_ratio, context.get("amount_ratios", [])),
        1.25 * sum((trend is not None and close >= trend, None not in (ema9, ema10, ema11) and ema9 >= ema10 >= ema11, ema9 is not None and close >= ema9, big_signal)),
        5.0 if ignition else 0.0,
        _cross_sectional_stars(hot_strength, context.get("short_fund_strengths", [])),
        _cross_sectional_stars(inst_strength, context.get("institution_strengths", [])),
        _cross_sectional_stars(relative_strength, context.get("relative_strengths", [])),
        5.0 if trend_mas_complete and close >= ma5 >= ma10 >= ma20 else 3.0 if trend_mas_complete and close >= ma20 and ma5 >= ma10 else 1.5 if ma20 is not None and close >= ma20 else 0.0,
        _cross_sectional_stars(slope, context.get("trend_slopes", [])),
        _cross_sectional_stars(efficiency, context.get("trend_efficiencies", [])),
        5.0 if wave_valid and 20 <= segment <= 65 and 20 <= wave <= 85 and wave >= segment else 3.5 if wave_valid and 10 <= segment <= 70 and 10 <= wave <= 90 else 1.5 if wave_valid else 0.0,
        5.0 if downside is not None and 0 <= downside <= 8 else 4.0 if downside is not None and 0 <= downside <= 12 else 2.0 if downside is not None and 0 <= downside <= 18 else 0.0,
        5.0 if consistency >= 0.70 else 4.0 if consistency >= 0.55 else 3.0 if consistency >= 0.40 else 1.5 if consistency >= 0.25 else 0.0,
        sum(capital_checks) / len(capital_checks) * 5.0,
        5.0 if evidence.get("catalyst") is True else 0.0,
        _market_breadth_stars(macro),
        _limit_up_ecology_stars(macro),
        _cross_sectional_stars(theme_metric, context.get("theme_strengths", [])),
        _percentile_stars(market_relative) if market_relative_available else 0.0,
        _cross_sectional_stars(amount, context.get("amounts", [])),
        5.0 if not state.get("is_one_price") and 1 <= recovery_pct <= 10 and gap_pct <= 7 else 4.0 if not state.get("is_one_price") and 0.3 <= recovery_pct <= 14 and gap_pct <= 9 else 2.0 if not state.get("is_one_price") and recovery_pct <= 16 else 0.0,
        _cross_sectional_stars(liquidity_stability, context.get("liquidity_stabilities", [])),
        _cross_sectional_stars(volatility_quality, context.get("volatility_qualities", [])),
    ]
    closes = [_coerce_number(row.get("close")) for row in rows]
    if any(value is None for value in closes):
        return None
    high_wave_divergence = wave is not None and segment is not None and wave >= 85 and abs(wave - segment) <= 5
    fund_support = hot_strength > 0 or inst_strength > 0
    prev_close = _coerce_number(rows[-2].get("close")) if len(rows) >= 2 else None
    amplitude = (high - low) / (prev_close or 1.0) * 100
    risks = [
        2 if _signal_active(big.get("OUTPUT6"), positive=False) else 0,
        2 if high_wave_divergence else 0,
        2 if wave is not None and segment is not None and wave < segment and not high_wave_divergence else 0,
        2 if inst_in is not None and inst_out is not None and inst_out > inst_in and inst_out > 0 else 0,
        2 if buyer is not None and buyer < 0 else 0,
        3 if trend is not None and close < trend else 0,
        2 if vol_ratio > 3.5 and amplitude >= 8 else 0,
        3 if evidence.get("major_risk") is True else 0,
        3 if wave is not None and wave >= 80 and not fund_support else 0,
        2 if evidence.get("reduction") is True else 0,
        2 if ((_coerce_number(state.get("recent_gain_3d_pct")) or 0) > 25 or (_coerce_number(state.get("recent_gain_5d_pct")) or 0) > 35) and not (ignition or fund_support) else 0,
        2 if segment is not None and segment >= 70 else 0,
        1 if pressure20 > 0 and pressure20 * 0.98 <= close <= pressure20 * 1.005 else 0,
        2 if _macd_bearish_divergence(closes) else 0,
        1 if ma60 is not None and close < ma60 and (_coerce_number(state.get("pct")) or 0) >= 9.5 else 0,
    ]
    return stars, risks


def _ranking_key(item: Any) -> tuple[float, float, float, str]:
    if not isinstance(item, dict):
        return (float("inf"), float("inf"), float("inf"), "")
    final_score = _finite_number(item.get("final_score"))
    positive_score = _finite_number(item.get("positive_score"))
    amount_yi = _finite_number(item.get("amount_yi"))
    return (
        -final_score if final_score is not None else float("inf"),
        -positive_score if positive_score is not None else float("inf"),
        -amount_yi if amount_yi is not None else float("inf"),
        str(item.get("code", "")),
    )


def validate_candidate_score(
    item: dict[str, Any],
    index: int,
    concept_frequency: Counter[str] | None = None,
    scoring_context: dict[str, Any] | None = None,
) -> list[str]:
    prefix = f"candidate_{index}"
    errors: list[str] = []
    dimensions = item.get("positive_dimensions")
    risks = item.get("risks")
    dimension_total = 0.0
    risk_total = 0.0
    semantics = _expected_candidate_semantics(item, concept_frequency, scoring_context)
    expected_stars, expected_risks = semantics if semantics is not None else ([], [])
    if semantics is None:
        errors.append(f"{prefix}_audit_inputs_invalid")
    if not isinstance(dimensions, list) or len(dimensions) != len(POSITIVE_DIMENSIONS):
        errors.append(f"{prefix}_positive_dimension_count")
    else:
        for position, (row, (expected_name, expected_weight)) in enumerate(zip(dimensions, POSITIVE_DIMENSIONS), 1):
            if not isinstance(row, dict):
                errors.append(f"{prefix}_dimension_{position}_not_object")
                continue
            if row.get("name") != expected_name:
                errors.append(f"{prefix}_dimension_{position}_name_mismatch")
            if row.get("axis") != FACTOR_AXES[expected_name]:
                errors.append(f"{prefix}_dimension_{position}_axis_mismatch")
            if row.get("role") != FACTOR_ROLES[expected_name]:
                errors.append(f"{prefix}_dimension_{position}_role_mismatch")
            if row.get("ranking_relevant") is not (FACTOR_ROLES[expected_name] in RANKING_FACTOR_ROLES):
                errors.append(f"{prefix}_dimension_{position}_ranking_relevance_mismatch")
            if not _same_number(row.get("weight"), expected_weight):
                errors.append(f"{prefix}_dimension_{position}_weight_mismatch")
            star_value = _finite_number(row.get("stars"))
            score_value = _finite_number(row.get("score"))
            if star_value is None or not 0 <= star_value <= 5:
                errors.append(f"{prefix}_dimension_{position}_stars_invalid")
            if position <= len(expected_stars) and not _same_number(star_value, expected_stars[position - 1]):
                errors.append(f"{prefix}_dimension_{position}_semantic_mismatch")
            if score_value is None:
                errors.append(f"{prefix}_dimension_{position}_score_invalid")
                continue
            dimension_total += score_value
            if star_value is not None and not _same_number(score_value, round(star_value / 5 * expected_weight, 2)):
                errors.append(f"{prefix}_dimension_{position}_score_mismatch")
    if not isinstance(risks, list) or len(risks) != len(RISK_DEDUCTIONS):
        errors.append(f"{prefix}_risk_count")
    else:
        for position, (row, (expected_name, maximum)) in enumerate(zip(risks, RISK_DEDUCTIONS), 1):
            if not isinstance(row, dict):
                errors.append(f"{prefix}_risk_{position}_not_object")
                continue
            if row.get("name") != expected_name:
                errors.append(f"{prefix}_risk_{position}_name_mismatch")
            deduction = _finite_number(row.get("deduction"))
            if deduction is None or deduction not in (0.0, float(maximum)):
                errors.append(f"{prefix}_risk_{position}_deduction_invalid")
                continue
            if position <= len(expected_risks) and not _same_number(deduction, expected_risks[position - 1]):
                errors.append(f"{prefix}_risk_{position}_semantic_mismatch")
            risk_total += deduction
    expected_positive = round(dimension_total, 2)
    expected_risk = round(risk_total, 2)
    expected_final = round(max(0.0, expected_positive - expected_risk), 2)
    if not _same_number(item.get("positive_score"), expected_positive):
        errors.append(f"{prefix}_positive_score_mismatch")
    if not _same_number(item.get("risk_deduction"), expected_risk):
        errors.append(f"{prefix}_risk_deduction_mismatch")
    if not _same_number(item.get("final_score"), expected_final):
        errors.append(f"{prefix}_final_score_mismatch")
    axis_scores = item.get("axis_scores")
    if not isinstance(axis_scores, dict):
        errors.append(f"{prefix}_axis_scores_invalid")
    else:
        for axis, maximum in AXIS_MAX_SCORES.items():
            expected_raw = round(sum(
                float(row.get("score") or 0.0)
                for row in dimensions or []
                if isinstance(row, dict) and row.get("axis") == axis
            ), 2)
            axis_row = axis_scores.get(axis)
            if (
                not isinstance(axis_row, dict)
                or not _same_number(axis_row.get("raw"), expected_raw)
                or not _same_number(axis_row.get("max"), maximum)
                or not _same_number(axis_row.get("normalized"), round(expected_raw / maximum * 100, 2))
            ):
                errors.append(f"{prefix}_axis_{axis}_mismatch")
        explosion = axis_scores.get("爆发力", {}).get("normalized") if isinstance(axis_scores.get("爆发力"), dict) else None
        persistence = axis_scores.get("持续性", {}).get("normalized") if isinstance(axis_scores.get("持续性"), dict) else None
        if explosion is not None and persistence is not None:
            expected_style = (
                "爆发延续型" if explosion >= 75 and persistence >= 70
                else "爆发脉冲型" if explosion >= 75
                else "趋势接力型" if persistence >= 70
                else "观察型"
            )
            if item.get("style_profile") != expected_style:
                errors.append(f"{prefix}_style_profile_mismatch")
    if isinstance(dimensions, list):
        expected_coverage = round(sum(
            float(row.get("weight") or 0.0)
            for row in dimensions
            if isinstance(row, dict) and row.get("available") is True
        ), 2)
        if not _same_number(item.get("score_coverage"), expected_coverage):
            errors.append(f"{prefix}_score_coverage_mismatch")
    contract = item.get("score_contract")
    if (
        not isinstance(contract, dict)
        or contract.get("version") != SCORE_CONTRACT_VERSION
        or contract.get("global_contract_path") != str(GLOBAL_SCORE_CONTRACT)
        or contract.get("global_contract_sha256") != GLOBAL_SCORE_CONTRACT_SHA256
        or contract.get("scale") != 100
        or contract.get("factor_count") != len(POSITIVE_DIMENSIONS)
        or not _same_number(contract.get("positive_weight_total"), 100.0)
        or not _same_number(contract.get("fundamental_positive_weight"), 0.0)
        or contract.get("missing_evidence_policy") != "zero_without_renormalization"
        or contract.get("factor_role_policy") != "pool_and_regime_factors_calibrate_absolute_context_but_do_not_claim_cross_sectional_ranking_power"
        or contract.get("tie_break") != "final_score_desc,positive_score_desc,amount_yi_desc,code_asc"
    ):
        errors.append(f"{prefix}_score_contract_mismatch")
    return errors


def validate_candidate_hard_gates(item: dict[str, Any], index: int) -> list[str]:
    prefix = f"candidate_{index}"
    errors: list[str] = []
    inputs = item.get("audit_inputs")
    if not isinstance(inputs, dict):
        return [f"{prefix}_hard_gate_audit_inputs_invalid"]
    formulas = inputs.get("formula")
    evidence = inputs.get("evidence")
    state = inputs.get("state")
    if not isinstance(formulas, dict) or not isinstance(evidence, dict) or not isinstance(state, dict):
        return [f"{prefix}_hard_gate_audit_inputs_invalid"]
    big = formulas.get("big") if isinstance(formulas.get("big"), dict) else {}
    fly = formulas.get("fly") if isinstance(formulas.get("fly"), dict) else {}
    hot = formulas.get("hot") if isinstance(formulas.get("hot"), dict) else {}
    inst = formulas.get("institution") if isinstance(formulas.get("institution"), dict) else {}
    wave = _coerce_number(fly.get("波"))
    segment = _coerce_number(fly.get("段"))
    ignition = any(_signal_active(fly.get(field)) for field in ("OUTPUT3", "OUTPUT4", "OUTPUT5", "OUTPUT6"))
    gain3 = _coerce_number(state.get("recent_gain_3d_pct"))
    gain5 = _coerce_number(state.get("recent_gain_5d_pct"))
    gain10 = _coerce_number(state.get("recent_gain_10d_pct"))
    buyer = _coerce_number(hot.get("买方意向"))
    aaa = _coerce_number(hot.get("AAA"))
    ddd = _coerce_number(hot.get("DDD"))
    inst_in = _coerce_number(inst.get("机构大单进"))
    inst_out = _coerce_number(inst.get("机构大单出"))
    fund_support = (
        buyer is not None and aaa is not None and ddd is not None and buyer > 0 and aaa > ddd
    ) or (
        inst_in is not None and inst_out is not None and inst_in > inst_out
    )
    expected = {
        HARD_NAMES[0]: not any(token in str(state.get("name") or "").upper() for token in ("ST", "退市")),
        HARD_NAMES[1]: int(_coerce_number(evidence.get("candidate_text_source_count")) or 0) > 0 and evidence.get("risk") is not True,
        HARD_NAMES[2]: int(_coerce_number(state.get("one_price_streak")) or 0) == 0,
        HARD_NAMES[3]: wave is not None and not (wave > 90 and not ignition),
        HARD_NAMES[4]: gain3 is not None and (gain3 <= 40 or ignition or fund_support),
        HARD_NAMES[5]: gain5 is not None and gain10 is not None and ((gain5 <= 60 and gain10 <= 90) or fund_support),
        HARD_NAMES[6]: segment is not None and segment <= 80,
    }
    rows = item.get("hard_exclusions")
    actual = {row.get("name"): row.get("passed") for row in rows if isinstance(row, dict)} if isinstance(rows, list) else {}
    for position, name in enumerate(HARD_NAMES, 1):
        if actual.get(name) is not expected[name]:
            errors.append(f"{prefix}_hard_gate_{position}_semantic_mismatch")
    formula_expected = {
        "大牛线4.0": all(_coerce_number(big.get(field)) is not None for field in ("主趋势线", "EMA9", "EMA10", "EMA11")),
        "飞龙在天": all(_coerce_number(fly.get(field)) is not None for field in ("波", "段")),
        "游资资金监控": all(_coerce_number(hot.get(field)) is not None for field in ("买方意向", "AAA", "DDD")),
        "机构资金监控": all(_coerce_number(inst.get(field)) is not None for field in ("机构大单进", "机构大单出", "大单动向", "大户大单进", "散户资金进")),
    }
    formula_actual = item.get("formula_required_ok")
    if not isinstance(formula_actual, dict) or any(formula_actual.get(name) is not passed for name, passed in formula_expected.items()):
        errors.append(f"{prefix}_required_formula_semantic_mismatch")
    expected_passed = all(expected.values()) and all(formula_expected.values())
    if item.get("passed_hard_gate") is not expected_passed:
        errors.append(f"{prefix}_passed_hard_gate_semantic_mismatch")
    return errors


def classify_engine_result(result: dict) -> str:
    top5 = result.get("top5")
    ranking = result.get("ranking")
    result_count = result.get("result_count")
    ranking_count = result.get("ranking_count")
    pool_candidate_count = result.get("pool_candidate_count")
    excluded = result.get("excluded")
    candidate_pool = result.get("candidate_pool")
    validation = result.get("validation")
    macro = result.get("macro")
    scan = macro.get("scan_coverage") if isinstance(macro, dict) else None
    scan_complete = (
        isinstance(scan, dict)
        and scan.get("complete") is True
        and scan.get("read_error_count") == 0
        and isinstance(scan.get("eligible_file_count"), int)
        and isinstance(scan.get("current_row_count"), int)
        and isinstance(scan.get("stale_or_short_count"), int)
        and scan.get("eligible_file_count") == scan.get("current_row_count") + scan.get("stale_or_short_count")
        and set(scan.get("markets", [])) == {"SH", "SZ", "BJ"}
    )
    expected_selection_status = (
        "TOP5" if result_count == 5 else "NO_CANDIDATES" if result_count == 0 else "PARTIAL"
    )
    valid = (
        result.get("status") == "CLEAN_PASS"
        and isinstance(top5, list)
        and isinstance(ranking, list)
        and isinstance(result_count, int)
        and isinstance(ranking_count, int)
        and isinstance(pool_candidate_count, int)
        and isinstance(excluded, list)
        and isinstance(candidate_pool, dict)
        and result_count == len(top5)
        and ranking_count == len(ranking)
        and 0 <= result_count <= 5
        and top5 == ranking[:5]
        and len(ranking) + len(excluded) == pool_candidate_count
        and candidate_pool.get("member_count") == pool_candidate_count
        and result.get("selection_status") == expected_selection_status
        and isinstance(validation, dict)
        and scan_complete
        and all(
            validation.get(key) is True
            for key in (
                "complete_scan",
                "result_count_matches",
                "result_count_lte_target",
                "scarcity_disclosed",
                "market_scope_complete",
            )
        )
    )
    return "CLEAN_PASS" if valid else "BLOCKED"


def _docx_text(path: Path) -> str:
    with zipfile.ZipFile(path) as package:
        required_parts = {"[Content_Types].xml", "word/document.xml"}
        if not required_parts.issubset(package.namelist()):
            raise ValueError("docx_required_ooxml_part_missing")
        corrupt_member = package.testzip()
        if corrupt_member is not None:
            raise ValueError(f"docx_corrupt_member:{corrupt_member}")
    document = Document(path)
    values = [paragraph.text for paragraph in document.paragraphs]
    for table in document.tables:
        for row in table.rows:
            for cell in row.cells:
                values.extend(paragraph.text for paragraph in cell.paragraphs)
    return "\n".join(values)


def validate_artifacts(
    json_path: Path,
    csv_path: Path,
    docx_path: Path,
    output_path: Path,
) -> dict[str, Any]:
    json_path = Path(json_path).resolve()
    csv_path = Path(csv_path).resolve()
    docx_path = Path(docx_path).resolve()
    output_path = Path(output_path).resolve()
    errors: list[str] = []
    inputs = {"json": json_path, "csv": csv_path, "docx": docx_path}
    for name, path in inputs.items():
        if not path.is_file():
            errors.append(f"{name}_missing:{path}")

    result: dict[str, Any] = {}
    if json_path.is_file():
        try:
            loaded = json.loads(json_path.read_text(encoding="utf-8"))
            if not isinstance(loaded, dict):
                raise ValueError("json_root_not_object")
            result = loaded
        except Exception as exc:
            errors.append(f"json_invalid:{type(exc).__name__}:{exc}")

    top5 = result.get("top5") if result else None
    ranking = result.get("ranking") if result else None
    result_count = result.get("result_count") if result else None
    if result and classify_engine_result(result) != "CLEAN_PASS":
        errors.append("engine_result_contract_blocked")
    validations = result.get("validation") if result else None
    if result and (
        not isinstance(validations, dict)
        or not all(validations.get(key) is True for key in REQUIRED_VALIDATIONS)
    ):
        errors.append("engine_detailed_validation_blocked")
    source = result.get("source") if result else None
    if result and (
        not isinstance(source, dict)
        or str(source.get("template_sha256", "")).upper() != TEMPLATE_SHA256
    ):
        errors.append("json_template_hash_mismatch")
    if not TEMPLATE.is_file() or sha256_file(TEMPLATE).upper() != TEMPLATE_SHA256:
        errors.append("locked_template_hash_mismatch")
    factor_registry = result.get("factor_registry") if result else None
    enabled_registry = factor_registry.get("enabled") if isinstance(factor_registry, dict) else None
    global_dimensions = [
        (str(row.get("name")), float(row.get("weight", 0)))
        for row in GLOBAL_SCORE_CONTRACT_PAYLOAD.get("factors", [])
        if isinstance(row, dict)
    ]
    global_axes = {
        str(row.get("name")): str(row.get("axis"))
        for row in GLOBAL_SCORE_CONTRACT_PAYLOAD.get("factors", [])
        if isinstance(row, dict)
    }
    global_roles = {
        str(row.get("name")): str(row.get("role"))
        for row in GLOBAL_SCORE_CONTRACT_PAYLOAD.get("factors", [])
        if isinstance(row, dict)
    }
    if (
        GLOBAL_SCORE_CONTRACT_PAYLOAD.get("version") != SCORE_CONTRACT_VERSION
        or GLOBAL_SCORE_CONTRACT_PAYLOAD.get("fundamental_policy", {}).get("positive_weight") != 0
        or global_dimensions != list(POSITIVE_DIMENSIONS)
        or global_axes != FACTOR_AXES
        or global_roles != FACTOR_ROLES
        or not isinstance(enabled_registry, list)
        or len(enabled_registry) != len(POSITIVE_DIMENSIONS)
        or [row.get("name") for row in enabled_registry if isinstance(row, dict)] != [name for name, _weight in POSITIVE_DIMENSIONS]
        or any(
            not isinstance(row, dict)
            or row.get("axis") != FACTOR_AXES[name]
            or row.get("role") != FACTOR_ROLES[name]
            or row.get("ranking_relevant") is not (FACTOR_ROLES[name] in RANKING_FACTOR_ROLES)
            or not _same_number(row.get("weight"), weight)
            or row.get("status") != "ENABLED_REPRODUCIBLE"
            for row, (name, weight) in zip(enabled_registry or [], POSITIVE_DIMENSIONS)
        )
        or factor_registry.get("enabled_count") != len(POSITIVE_DIMENSIONS)
        or int(factor_registry.get("pending_count") or 0) < 10
        or int(factor_registry.get("rejected_count") or 0) < 5
    ):
        errors.append("factor_registry_contract_mismatch")
    evaluation_contract = result.get("evaluation_contract") if result else None
    if (
        not isinstance(evaluation_contract, dict)
        or evaluation_contract.get("decision_time") != "trade_date_after_close"
        or evaluation_contract.get("entry_assumption") != "next_trade_day_open"
        or evaluation_contract.get("forward_horizons") != [1, 3, 5]
        or evaluation_contract.get("walk_forward_required") is not True
        or evaluation_contract.get("lookahead_policy") != "features_and_evidence_timestamp_at_or_before_decision_time"
    ):
        errors.append("forward_evaluation_contract_mismatch")
    predictive_validation = result.get("predictive_validation") if result else None
    if (
        not isinstance(predictive_validation, dict)
        or predictive_validation.get("status") != "UNVERIFIED_REQUIRES_POINT_IN_TIME_FORWARD_LABELS"
    ):
        errors.append("predictive_validation_status_mismatch")

    if isinstance(ranking, list):
        concept_frequency: Counter[str] = Counter(
            str(token)
            for ranked_item in ranking
            if isinstance(ranked_item, dict)
            for token in (
                ranked_item.get("audit_inputs", {}).get("concepts", [])
                if isinstance(ranked_item.get("audit_inputs"), dict)
                else []
            )
        )
        scoring_context = _build_scoring_context(
            ranking,
            concept_frequency,
            result.get("macro") if isinstance(result.get("macro"), dict) else None,
        )
        for index, item in enumerate(ranking, 1):
            prefix = f"candidate_{index}"
            if not isinstance(item, dict):
                errors.append(f"{prefix}_not_object")
                continue
            if item.get("rank") != index:
                errors.append(f"{prefix}_rank_mismatch")
            if item.get("passed_hard_gate") is not True:
                errors.append(f"{prefix}_hard_gate_not_passed")
            if len(item.get("positive_dimensions", [])) != len(POSITIVE_DIMENSIONS):
                errors.append(f"{prefix}_positive_factor_count")
            if len(item.get("risks", [])) != 15:
                errors.append(f"{prefix}_risk_count")
            hard_exclusions = item.get("hard_exclusions", [])
            if len(hard_exclusions) != 7 or not all(
                isinstance(row, dict) and row.get("passed") is True
                for row in hard_exclusions
            ):
                errors.append(f"{prefix}_hard_exclusion_contract")
            formula_status = item.get("formula_required_ok")
            if not isinstance(formula_status, dict) or not all(
                formula_status.get(name) is True for name in REQUIRED_FORMULAS
            ):
                errors.append(f"{prefix}_required_formula_contract")
            formula_summary = item.get("formula_summary")
            if not isinstance(formula_summary, dict) or not all(
                name in formula_summary for name in REQUIRED_FORMULA_SUMMARY
            ):
                errors.append(f"{prefix}_formula_summary_contract")
            segment = item.get("segment")
            if not isinstance(segment, (int, float)) or segment > 80:
                errors.append(f"{prefix}_segment_contract")
            errors.extend(validate_candidate_score(item, index, concept_frequency, scoring_context))
            errors.extend(validate_candidate_hard_gates(item, index))
        expected_order = sorted(ranking, key=_ranking_key)
        if ranking != expected_order:
            errors.append("full_ranking_mismatch")
        if isinstance(top5, list) and top5 != ranking[:5]:
            errors.append("top5_ranking_prefix_mismatch")
        diagnostics = result.get("factor_diagnostics")
        expected_dead = []
        expected_saturated = []
        expected_effective = 0
        expected_ranking_dead = []
        expected_ranking_saturated = []
        expected_ranking_effective = 0
        for name, _weight in POSITIVE_DIMENSIONS:
            values = []
            for item in ranking:
                dimensions = item.get("positive_dimensions") if isinstance(item, dict) else None
                by_name = {
                    row.get("name"): row
                    for row in dimensions or []
                    if isinstance(row, dict)
                }
                value = _finite_number(by_name.get(name, {}).get("stars"))
                values.append(value if value is not None else 0.0)
            distinct = len({round(value, 4) for value in values})
            expected_effective += int(distinct > 1)
            ranking_relevant = FACTOR_ROLES[name] in RANKING_FACTOR_ROLES
            expected_ranking_effective += int(ranking_relevant and distinct > 1)
            if values and all(value == 0 for value in values):
                expected_dead.append(name)
                if ranking_relevant:
                    expected_ranking_dead.append(name)
            if values and sum(value == 5 for value in values) / len(values) >= 0.80:
                expected_saturated.append(name)
                if ranking_relevant:
                    expected_ranking_saturated.append(name)
        if (
            not isinstance(diagnostics, dict)
            or diagnostics.get("sample_count") != len(ranking)
            or diagnostics.get("effective_factor_count") != expected_effective
            or diagnostics.get("ranking_eligible_factor_count") != sum(role in RANKING_FACTOR_ROLES for role in FACTOR_ROLES.values())
            or diagnostics.get("ranking_effective_factor_count") != expected_ranking_effective
            or diagnostics.get("context_factor_count") != sum(role not in RANKING_FACTOR_ROLES for role in FACTOR_ROLES.values())
            or diagnostics.get("dead_factors") != expected_dead
            or diagnostics.get("saturated_factors") != expected_saturated
            or diagnostics.get("ranking_dead_factors") != expected_ranking_dead
            or diagnostics.get("ranking_saturated_factors") != expected_ranking_saturated
        ):
            errors.append("factor_diagnostics_mismatch")

    csv_rows: list[dict[str, str]] = []
    if csv_path.is_file():
        try:
            with csv_path.open("r", newline="", encoding="utf-8-sig") as handle:
                reader = csv.DictReader(handle)
                if tuple(reader.fieldnames or ()) != CSV_HEADERS:
                    errors.append("csv_header_mismatch")
                csv_rows = list(reader)
        except Exception as exc:
            errors.append(f"csv_invalid:{type(exc).__name__}:{exc}")
    if isinstance(ranking, list):
        if len(csv_rows) != len(ranking):
            errors.append("csv_result_count_mismatch")
        for index, (row, item) in enumerate(zip(csv_rows, ranking), 1):
            if not isinstance(item, dict):
                continue
            if row.get("排名") != str(item.get("rank", "")):
                errors.append(f"csv_candidate_{index}_rank_mismatch")
            if row.get("代码") != str(item.get("code", "")):
                errors.append(f"csv_candidate_{index}_code_mismatch")
            if row.get("名称") != str(item.get("name", "")):
                errors.append(f"csv_candidate_{index}_name_mismatch")
            numeric_fields = {
                "现价": "latest",
                "涨幅": "pct",
                "正向得分": "positive_score",
                "风险扣分": "risk_deduction",
                "最终得分": "final_score",
                "飞龙波": "wave",
                "飞龙段": "segment",
                "主趋势线": "trend_line",
            }
            for csv_name, json_name in numeric_fields.items():
                if not _same_csv_number(row.get(csv_name), item.get(json_name)):
                    errors.append(f"csv_candidate_{index}_{json_name}_mismatch")
            if row.get("失效条件") != str(item.get("invalidation", "")):
                errors.append(f"csv_candidate_{index}_invalidation_mismatch")

    if docx_path.is_file() and result:
        try:
            visible_text = _docx_text(docx_path)
            trade_date = str(result.get("trade_date", ""))
            formatted_date = (
                f"{trade_date[:4]}-{trade_date[4:6]}-{trade_date[6:8]}"
                if len(trade_date) == 8 and trade_date.isdigit()
                else trade_date
            )
            if not formatted_date or formatted_date not in visible_text:
                errors.append("docx_trade_date_missing")
            if isinstance(ranking, list) and ranking:
                for index, item in enumerate(ranking, 1):
                    if not isinstance(item, dict):
                        continue
                    if str(item.get("name", "")) not in visible_text:
                        errors.append(f"docx_candidate_{index}_name_missing")
                    if str(item.get("code", "")) not in visible_text:
                        errors.append(f"docx_candidate_{index}_code_missing")
            elif "本轮无候选" not in visible_text:
                errors.append("docx_no_candidate_disclosure_missing")
        except Exception as exc:
            errors.append(f"docx_invalid:{type(exc).__name__}:{exc}")

    errors = list(dict.fromkeys(errors))
    payload = {
        "schema": "A_SHARE_15D_ARTIFACT_VALIDATION_V1",
        "skill_id": SKILL_ROOT.name,
        "status": "CLEAN_PASS" if not errors else "BLOCKED",
        "errors": errors,
        "trade_date": result.get("trade_date") if result else None,
        "selection_status": result.get("selection_status") if result else None,
        "result_count": result_count,
        "top5_count": len(top5) if isinstance(top5, list) else None,
        "ranking_count": len(ranking) if isinstance(ranking, list) else None,
        "template": file_artifact(TEMPLATE) if TEMPLATE.is_file() else {
            "path": str(TEMPLATE),
            "exists": False,
        },
        "artifacts": {
            name: file_artifact(path) if path.is_file() else {
                "path": str(path),
                "exists": False,
            }
            for name, path in inputs.items()
        },
    }
    output_path.parent.mkdir(parents=True, exist_ok=True)
    output_path.write_text(
        json.dumps(payload, ensure_ascii=False, indent=2) + "\n",
        encoding="utf-8",
    )
    return payload


def main() -> int:
    parser = argparse.ArgumentParser(
        description="Validate the persisted A-share 15-dimension JSON/CSV/DOCX artifacts."
    )
    parser.add_argument("--json", required=True)
    parser.add_argument("--csv", required=True)
    parser.add_argument("--docx", required=True)
    parser.add_argument("--output", required=True)
    args = parser.parse_args()
    payload = validate_artifacts(
        Path(args.json),
        Path(args.csv),
        Path(args.docx),
        Path(args.output),
    )
    print(json.dumps(payload, ensure_ascii=False, indent=2))
    return 0 if payload["status"] == "CLEAN_PASS" else 2


if __name__ == "__main__" and os.environ.get("ONESTOCK_STOCK_CANONICAL_CHILD") != "1":
    print("canonical_stock_legacy_entry_direct_execution_blocked", file=_onestock_embedded_sys.stderr)
    raise SystemExit(2)


if __name__ == "__main__":
    raise SystemExit(main())
