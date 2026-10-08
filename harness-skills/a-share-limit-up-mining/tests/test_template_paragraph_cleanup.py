from docx import Document

from scripts.build_docx import remove_unused_section_paragraphs


def test_unused_template_paragraphs_are_removed_from_section():
    doc = Document()
    keep = doc.add_paragraph("current run")
    stale = doc.add_paragraph("stale template stock")

    remove_unused_section_paragraphs([keep, stale], keep)

    assert [paragraph.text for paragraph in doc.paragraphs] == ["current run"]
