#!/usr/bin/env python3
# -*- coding: utf-8 -*-
from __future__ import annotations
# ONESTOCK_EMBEDDED_LOCAL_IMPORT_BOOTSTRAP
import sys as _onestock_embedded_sys
from pathlib import Path as _OneStockEmbeddedPath
_onestock_embedded_dir = str(_OneStockEmbeddedPath(__file__).resolve().parent)
if _onestock_embedded_dir not in _onestock_embedded_sys.path:
    _onestock_embedded_sys.path.insert(0, _onestock_embedded_dir)

import json
import sys
from pathlib import Path
from typing import Iterable

_APP_SCRIPTS = Path(__file__).resolve().parents[3] / "scripts"
if str(_APP_SCRIPTS) not in sys.path:
    sys.path.insert(0, str(_APP_SCRIPTS))
from tdx_path_config import resolve_data_root

try:
    sys.stdout.reconfigure(encoding="utf-8", errors="replace")
    sys.stderr.reconfigure(encoding="utf-8", errors="replace")
except Exception:
    pass

TEMPLATE = Path(__file__).resolve().parents[1] / "assets" / "连板挖掘模板.docx"


def set_cell(cell, text: str) -> None:
    cell.text = str(text)


def add_table(doc, headers: Iterable[str], rows: int) -> None:
    headers = list(headers)
    table = doc.add_table(rows=max(rows, 1), cols=len(headers))
    table.style = "Table Grid"
    for idx, header in enumerate(headers):
        set_cell(table.rows[0].cells[idx], header)
    for row in table.rows[1:]:
        for idx, cell in enumerate(row.cells):
            set_cell(cell, f"模板字段{idx + 1}")


def build_template(path: Path = TEMPLATE) -> Path:
    try:
        from docx import Document
        from docx.enum.text import WD_ALIGN_PARAGRAPH
        from docx.shared import Pt
    except Exception as exc:
        raise SystemExit(f"python-docx unavailable: {type(exc).__name__}: {exc}")

    path.parent.mkdir(parents=True, exist_ok=True)
    doc = Document()

    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = title.add_run("连板挖掘全流程选股报告")
    run.bold = True
    run.font.size = Pt(24)

    subtitle = doc.add_paragraph("涨停池 · 连板股 · 首板 · 一进二候选 · 精品")
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER

    doc.add_heading("最终结果", level=1)
    doc.add_paragraph("涨停池、连板股、首板、一进二候选和精品层结论必须在执行后由脚本写入。")

    step_names = ["第一步", "第二步", "第三步", "第四步", "第五步", "第六步", "第七步", "第八步", "第九步", "第十步"]
    for step_name in step_names:
        doc.add_heading(step_name, level=1)
        doc.add_paragraph("本段为流程锚点，业务数据由唯一入口脚本自动写入表格。")

    add_table(doc, ["涨停池", "连板股", "首板", "一进二候选", "精品"], 2)
    add_table(doc, ["排名", "代码名称", "行业", "评分", "龙虎榜净买", "核心理由"], 6)
    add_table(doc, ["涨停池", "全池", "首板", "连板", "风险"], 2)
    add_table(doc, ["代码名称", "主线", "行业", "连板", "涨停统计", "首次封板", "最后封板", "炸板", "封板资金", "换手率"], 2)
    add_table(doc, ["代码名称", "主线", "行业", "连板", "净买", "买入", "卖出", "炸板", "纯度", "上榜原因"], 2)
    add_table(doc, ["主线", "数量", "首板", "连板", "最高连板", "封板资金", "龙虎榜净买", "平均炸板", "评分"], 2)
    add_table(doc, ["代码", "名称", "报告名称", "评级", "机构", "行业", "日期"], 2)
    add_table(doc, ["代码", "名称", "新闻标题", "新闻内容", "发布时间", "来源"], 2)
    add_table(doc, ["代码名称", "主线", "行业", "连板", "涨停统计", "首次封板", "炸板", "净买", "纯度", "诊断"], 2)
    add_table(doc, ["主线", "主题", "首板数", "零炸", "龙虎榜数", "封板资金"], 2)
    add_table(doc, ["代码名称", "主线", "行业", "首次封板", "炸板", "净买", "纯度", "备注", "评分"], 2)
    add_table(doc, ["序号", "代码名称", "主线", "行业", "首次封板", "炸板", "封板资金", "净买", "评级", "评分"], 2)
    for _ in range(5):
        add_table(doc, ["项目", "内容"], 7)
    add_table(doc, ["代码名称", "主线", "行业", "连板", "炸板", "净买", "卖出", "风险原因"], 2)
    add_table(doc, ["步骤", "结果"], 11)

    doc.save(str(path))
    return path


def validate_template(path: Path = TEMPLATE) -> dict:
    from docx import Document

    if not path.exists():
        return {"ok": False, "reason": "missing", "path": str(path)}
    try:
        doc = Document(str(path))
    except Exception as exc:
        return {"ok": False, "reason": f"unreadable:{type(exc).__name__}", "path": str(path)}
    table_shapes = [[len(t.rows), len(t.columns)] for t in doc.tables]
    required_cols = [5, 6, 5, 10, 10, 9, 7, 6, 10, 6, 9, 10, 2, 2, 2, 2, 2, 8, 2]
    headings_text = "\n".join(p.text for p in doc.paragraphs)
    required_steps = ["第一步", "第二步", "第三步", "第四步", "第五步", "第十步"]
    steps_ok = all(step in headings_text for step in required_steps)
    ok = steps_ok and len(doc.tables) >= len(required_cols) and all(
        len(doc.tables[idx].columns) >= cols for idx, cols in enumerate(required_cols)
    )
    return {
        "ok": ok,
        "path": str(path),
        "size": path.stat().st_size,
        "tables": len(doc.tables),
        "steps_ok": steps_ok,
        "table_shapes": table_shapes[:25],
    }


def main() -> int:
    template = TEMPLATE
    status = validate_template(template)
    created = False
    if not status.get("ok"):
        template = resolve_data_root() / "templates" / "a-share-limit-up-mining" / "连板挖掘模板.docx"
        build_template(template)
        created = True
        status = validate_template(template)
    result = {"status": "CLEAN_PASS" if status.get("ok") else "BLOCKED", "created": created, **status}
    print(json.dumps(result, ensure_ascii=False, indent=2))
    return 0 if status.get("ok") else 1


if __name__ == "__main__":
    raise SystemExit(main())
