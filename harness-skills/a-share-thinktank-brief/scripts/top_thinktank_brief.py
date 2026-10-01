#!/usr/bin/env python3
# -*- coding: utf-8 -*-
from __future__ import annotations
# ONESTOCK_EMBEDDED_LOCAL_IMPORT_BOOTSTRAP
import sys as _onestock_embedded_sys
from pathlib import Path as _OneStockEmbeddedPath
_onestock_embedded_dir = str(_OneStockEmbeddedPath(__file__).resolve().parent)
if _onestock_embedded_dir not in _onestock_embedded_sys.path:
    _onestock_embedded_sys.path.insert(0, _onestock_embedded_dir)

import argparse
import datetime as dt
import hashlib
import html
import json
import os
import re
import shutil
import subprocess
import sys
import tempfile
from pathlib import Path
from typing import Any
from zipfile import ZipFile

os.environ.setdefault("PYTHONIOENCODING", "utf-8")
os.environ.setdefault("PYTHONUTF8", "1")
try:
    sys.stdout.reconfigure(encoding="utf-8")
    sys.stderr.reconfigure(encoding="utf-8")
except Exception:
    pass

try:
    from docx import Document
    from docx.enum.section import WD_SECTION
    from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.opc.constants import RELATIONSHIP_TYPE
    from docx.shared import Inches, Pt, RGBColor
except Exception as exc:  # pragma: no cover - reported by doctor/selftest
    Document = None
    WD_SECTION = None
    WD_CELL_VERTICAL_ALIGNMENT = None
    WD_TABLE_ALIGNMENT = None
    WD_ALIGN_PARAGRAPH = None
    OxmlElement = None
    qn = None
    RELATIONSHIP_TYPE = None
    Inches = None
    Pt = None
    RGBColor = None
    DOCX_IMPORT_ERROR = f"{type(exc).__name__}: {exc}"
else:
    DOCX_IMPORT_ERROR = ""


SKILL_DIR = Path(__file__).resolve().parents[1]
ENTRY = SKILL_DIR / "scripts" / "top_thinktank_brief.py"
CANARY_BRIEF = SKILL_DIR / "references" / "canary_brief.md"
LAYOUT_SPEC = SKILL_DIR / "references" / "docx-layout-template.md"

FONT_LATIN = "Calibri"
FONT_CN = "Microsoft YaHei"
CONTENT_DXA = 9360
NAVY = "1F3556"
BLUE = "2F5597"
TEAL = "0F766E"
GOLD = "9C6500"
RED = "9B1C1C"
INK = "111827"
MUTED = "5B677A"
WHITE = "FFFFFF"
BORDER = "D9E2F3"
FILL_BLUE = "EAF1F8"
FILL_TEAL = "E7F4F1"
FILL_GOLD = "FFF4D6"
FILL_GRAY = "F3F6FA"


KNOWN_SOURCES: list[dict[str, Any]] = [
    {
        "name": "布鲁金斯学会",
        "aliases": ["布鲁金斯", "Brookings", "Brookings Institution"],
        "url": "https://www.brookings.edu/",
        "focus": "综合政策、产业治理、全球经济",
        "signal": "观察美国政策中间派对产业政策、竞争战略和对华议题的主流表达。",
    },
    {
        "name": "兰德公司",
        "aliases": ["兰德", "RAND", "RAND Corporation"],
        "url": "https://www.rand.org/",
        "focus": "国防安全、科技风险、供应链韧性",
        "signal": "观察安全部门对先进技术、关键基础设施和长期风险建模的判断。",
    },
    {
        "name": "美国企业研究所",
        "aliases": ["美国企业研究所", "American Enterprise Institute", "AEI"],
        "url": "https://www.aei.org/",
        "focus": "财政、产业监管、对外政策",
        "signal": "观察保守派政策网络对税制、监管和产业竞争的约束变量。",
    },
    {
        "name": "传统基金会",
        "aliases": ["传统基金会", "Heritage", "Heritage Foundation"],
        "url": "https://www.heritage.org/",
        "focus": "保守派政策、国家安全、行政议程",
        "signal": "观察行政政策转向、出口管制和国家安全叙事的强弱。",
    },
    {
        "name": "卡内基国际和平基金会",
        "aliases": ["卡内基", "Carnegie", "Carnegie Endowment"],
        "url": "https://carnegieendowment.org/",
        "focus": "外交政策、地区安全、国际秩序",
        "signal": "观察地缘风险、盟友协调和大国关系缓冲空间。",
    },
    {
        "name": "胡佛研究所",
        "aliases": ["胡佛", "Hoover", "Hoover Institution"],
        "url": "https://www.hoover.org/",
        "focus": "安全战略、技术竞争、保守派学术网络",
        "signal": "观察技术竞争叙事是否从政策倡议进入长期制度安排。",
    },
    {
        "name": "战略与国际问题研究中心",
        "aliases": ["战略与国际问题研究中心", "CSIS", "Center for Strategic and International Studies"],
        "url": "https://www.csis.org/",
        "focus": "国际安全、产业供应链、技术政策",
        "signal": "观察出口管制、供应链安全和国会政策议程的联动。",
    },
    {
        "name": "外交关系协会",
        "aliases": ["外交关系协会", "CFR", "Council on Foreign Relations"],
        "url": "https://www.cfr.org/",
        "focus": "外交政策、全球贸易、战略竞争",
        "signal": "观察对外政策共识是否影响贸易、金融和科技监管节奏。",
    },
    {
        "name": "新美国安全中心",
        "aliases": ["新美国安全中心", "CNAS", "Center for a New American Security"],
        "url": "https://www.cnas.org/",
        "focus": "防务科技、人工智能安全、产业安全",
        "signal": "观察防务和科技政策是否向算力、软件、无人系统扩散。",
    },
    {
        "name": "大西洋理事会",
        "aliases": ["大西洋理事会", "Atlantic Council"],
        "url": "https://www.atlanticcouncil.org/",
        "focus": "跨大西洋关系、制裁、网络安全",
        "signal": "观察盟友协同、制裁工具和网络安全投资线索。",
    },
    {
        "name": "美国进步中心",
        "aliases": ["美国进步中心", "Center for American Progress", "CAP"],
        "url": "https://www.americanprogress.org/",
        "focus": "进步派产业政策、劳工、绿色转型",
        "signal": "观察清洁能源、制造回流和产业补贴的民主党政策基础。",
    },
    {
        "name": "彼得森国际经济研究所",
        "aliases": ["彼得森", "PIIE", "Peterson Institute"],
        "url": "https://www.piie.com/",
        "focus": "国际贸易、宏观经济、全球化",
        "signal": "观察关税、贸易摩擦、汇率和全球需求对产业链的影响。",
    },
    {
        "name": "卡托研究所",
        "aliases": ["卡托", "Cato", "Cato Institute"],
        "url": "https://www.cato.org/",
        "focus": "自由市场、贸易、监管",
        "signal": "观察产业政策约束、监管反弹和财政纪律争论。",
    },
    {
        "name": "曼哈顿研究所",
        "aliases": ["曼哈顿研究所", "Manhattan Institute"],
        "url": "https://manhattan.institute/",
        "focus": "城市治理、公共安全、教育与科技",
        "signal": "观察城市基础设施、安全治理和公共服务数字化方向。",
    },
    {
        "name": "信息技术与创新基金会",
        "aliases": ["信息技术与创新基金会", "ITIF", "Information Technology and Innovation Foundation"],
        "url": "https://itif.org/",
        "focus": "创新政策、半导体、人工智能、数字经济",
        "signal": "观察先进制造、标准竞争和创新补贴的产业细节。",
    },
    {
        "name": "阿斯彭研究所",
        "aliases": ["阿斯彭", "Aspen", "Aspen Institute"],
        "url": "https://www.aspeninstitute.org/",
        "focus": "科技伦理、社会治理、跨界对话",
        "signal": "观察人工智能治理、数据隐私和社会约束变量。",
    },
    {
        "name": "新美国基金会",
        "aliases": ["新美国基金会", "新美国", "New America"],
        "url": "https://www.newamerica.org/",
        "focus": "数字政策、教育、数据治理",
        "signal": "观察数据治理、平台监管和公共技术基础设施。",
    },
    {
        "name": "威尔逊中心",
        "aliases": ["威尔逊", "Wilson Center"],
        "url": "https://www.wilsoncenter.org/",
        "focus": "地区研究、国会关联、政策交流",
        "signal": "观察区域地缘风险和国会政策渠道。",
    },
    {
        "name": "贝尔弗中心",
        "aliases": ["贝尔弗", "Belfer Center"],
        "url": "https://www.belfercenter.org/",
        "focus": "国家安全、科技风险、能源与核安全",
        "signal": "观察国家安全人才网络对科技风险和能源安全的判断。",
    },
]


BANNED_TERMS = [
    "UTF-8",
    "本页把",
    "来源文件",
    "原始智库清单",
    "原生可编辑",
    "紧凑精美终版",
    "链接保留",
    "INTELLIGENCE",
    "THINK-TANK",
    "SOURCE INTAKE",
    "A-SHARE",
    "SCENARIOS",
    "SOURCES &",
    "TECH × GEO",
    "TECH x GEO",
    "US THINK",
    "Evidence-based",
    "Research Brief",
    "Brookings",
    "CSIS",
    "RAND",
    "CNAS",
    "Carnegie",
    "CFR",
    "PIIE",
    "ITIF",
    "Hoover",
    "American Enterprise",
    "Heritage",
    "Cato",
    "Manhattan Institute",
    "Atlantic Council",
    "New America",
    "Wilson Center",
    "Belfer",
    "Choice",
    "截图",
    "脚本",
    "路径",
    "浏览器",
    "登录",
    "token",
    "cookie",
    "TODO",
    "TBD",
    "待补",
    "公众号写作骨架",
    "对公众号文章",
    "新版文章",
    "不能独占全文",
    "全文",
    "写成",
    "标题方向",
    "导语",
    "第一节",
    "第二节",
    "第三节",
    "第四节",
    "第五节",
    "第六节",
    "模板",
    "预设",
    "兜底",
]


TRADING_DIRECTIVE_TERMS = [
    "买入",
    "卖出",
    "加仓",
    "减仓",
    "满仓",
    "清仓",
    "止损",
    "止盈",
]


def now_cn_date() -> str:
    return dt.datetime.now().astimezone().strftime("%Y-%m-%d")


def json_out(payload: dict[str, Any], code: int = 0) -> int:
    print(json.dumps(payload, ensure_ascii=False, indent=2, default=str))
    return code


def require_docx() -> None:
    if Document is None:
        raise RuntimeError(f"python-docx unavailable: {DOCX_IMPORT_ERROR}")


def read_text(path: Path) -> str:
    for enc in ("utf-8-sig", "utf-8", "gb18030"):
        try:
            return path.read_text(encoding=enc)
        except UnicodeDecodeError:
            continue
    return path.read_text(encoding="utf-8", errors="ignore")


def rgb(hex_color: str) -> RGBColor:
    h = hex_color.replace("#", "")
    return RGBColor(int(h[0:2], 16), int(h[2:4], 16), int(h[4:6], 16))


def set_run_font(run, size=None, color=None, bold=None, italic=None, name=FONT_LATIN):
    run.font.name = name
    rpr = run._element.get_or_add_rPr()
    rfonts = rpr.rFonts
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.append(rfonts)
    rfonts.set(qn("w:ascii"), name)
    rfonts.set(qn("w:hAnsi"), name)
    rfonts.set(qn("w:eastAsia"), FONT_CN)
    if size is not None:
        run.font.size = Pt(size)
    if color is not None:
        run.font.color.rgb = rgb(color)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic


def style_font(style, size=None, color=None, bold=None, name=FONT_LATIN):
    style.font.name = name
    if size is not None:
        style.font.size = Pt(size)
    if color is not None:
        style.font.color.rgb = rgb(color)
    if bold is not None:
        style.font.bold = bold
    rpr = style._element.get_or_add_rPr()
    rfonts = rpr.rFonts
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.append(rfonts)
    rfonts.set(qn("w:ascii"), name)
    rfonts.set(qn("w:hAnsi"), name)
    rfonts.set(qn("w:eastAsia"), FONT_CN)


def set_cell_shading(cell, fill: str):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=40, start=70, bottom=40, end=70):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = tc_pr.find(qn("w:tcMar"))
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for key, value in {"top": top, "start": start, "bottom": bottom, "end": end}.items():
        node = tc_mar.find(qn(f"w:{key}"))
        if node is None:
            node = OxmlElement(f"w:{key}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_cell_border(cell, color=BORDER, size=6):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_borders = tc_pr.find(qn("w:tcBorders"))
    if tc_borders is None:
        tc_borders = OxmlElement("w:tcBorders")
        tc_pr.append(tc_borders)
    for edge in ("top", "bottom", "start", "end"):
        node = tc_borders.find(qn(f"w:{edge}"))
        if node is None:
            node = OxmlElement(f"w:{edge}")
            tc_borders.append(node)
        node.set(qn("w:val"), "single")
        node.set(qn("w:sz"), str(size))
        node.set(qn("w:space"), "0")
        node.set(qn("w:color"), color)


def set_table_geometry(table, widths: list[int], indent: int = 0):
    table.autofit = False
    table.allow_autofit = False
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    tbl = table._tbl
    tbl_pr = tbl.tblPr
    tbl_w = tbl_pr.find(qn("w:tblW"))
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), str(sum(widths)))
    tbl_w.set(qn("w:type"), "dxa")
    tbl_ind = tbl_pr.find(qn("w:tblInd"))
    if tbl_ind is None:
        tbl_ind = OxmlElement("w:tblInd")
        tbl_pr.append(tbl_ind)
    tbl_ind.set(qn("w:w"), str(indent))
    tbl_ind.set(qn("w:type"), "dxa")
    old_grid = tbl.find(qn("w:tblGrid"))
    if old_grid is not None:
        tbl.remove(old_grid)
    grid = OxmlElement("w:tblGrid")
    for width in widths:
        col = OxmlElement("w:gridCol")
        col.set(qn("w:w"), str(width))
        grid.append(col)
    tbl.insert(0, grid)
    for row in table.rows:
        for idx, cell in enumerate(row.cells):
            tc_pr = cell._tc.get_or_add_tcPr()
            tc_w = tc_pr.find(qn("w:tcW"))
            if tc_w is None:
                tc_w = OxmlElement("w:tcW")
                tc_pr.append(tc_w)
            tc_w.set(qn("w:w"), str(widths[idx]))
            tc_w.set(qn("w:type"), "dxa")
            cell.width = Inches(widths[idx] / 1440)


def clear_cell(cell):
    cell.text = ""
    for p in cell.paragraphs:
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(0)


def add_para(target, text="", style=None, size=9.4, color=INK, bold=False, italic=False, before=0, after=3, align=None):
    p = target.add_paragraph(style=style)
    p.paragraph_format.space_before = Pt(before)
    p.paragraph_format.space_after = Pt(after)
    p.paragraph_format.line_spacing = 0.96
    if align is not None:
        p.alignment = align
    if text:
        r = p.add_run(text)
        set_run_font(r, size=size, color=color, bold=bold, italic=italic)
    return p


def add_hyperlink(paragraph, text: str, url: str, color: str = BLUE):
    part = paragraph.part
    r_id = part.relate_to(url, RELATIONSHIP_TYPE.HYPERLINK, is_external=True)
    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("r:id"), r_id)
    new_run = OxmlElement("w:r")
    r_pr = OxmlElement("w:rPr")
    r_fonts = OxmlElement("w:rFonts")
    r_fonts.set(qn("w:ascii"), FONT_LATIN)
    r_fonts.set(qn("w:hAnsi"), FONT_LATIN)
    r_fonts.set(qn("w:eastAsia"), FONT_CN)
    r_pr.append(r_fonts)
    c = OxmlElement("w:color")
    c.set(qn("w:val"), color)
    r_pr.append(c)
    u = OxmlElement("w:u")
    u.set(qn("w:val"), "single")
    r_pr.append(u)
    new_run.append(r_pr)
    text_node = OxmlElement("w:t")
    text_node.text = text
    new_run.append(text_node)
    hyperlink.append(new_run)
    paragraph._p.append(hyperlink)
    return hyperlink


def style_doc(doc):
    sec = doc.sections[0]
    sec.page_width = Inches(8.27)
    sec.page_height = Inches(11.69)
    sec.top_margin = Inches(0.34)
    sec.bottom_margin = Inches(0.34)
    sec.left_margin = Inches(0.48)
    sec.right_margin = Inches(0.48)
    style_font(doc.styles["Normal"], size=8.4, color=INK)
    for name, size, color, bold in [
        ("Title", 24, NAVY, True),
        ("Heading 1", 15, NAVY, True),
        ("Heading 2", 11.5, BLUE, True),
    ]:
        if name in doc.styles:
            style_font(doc.styles[name], size=size, color=color, bold=bold)


def add_section_title(doc, title: str, subtitle: str = ""):
    p = add_para(doc, title, size=12.8, color=NAVY, bold=True, before=0, after=1)
    p.paragraph_format.keep_with_next = True
    if subtitle:
        add_para(doc, subtitle, size=7.7, color=MUTED, after=3)


def add_table(doc, headers: list[str], rows: list[list[str]], widths: list[int], header_fill: str = NAVY):
    table = doc.add_table(rows=1, cols=len(headers))
    set_table_geometry(table, widths)
    table.style = "Table Grid"
    header_cells = table.rows[0].cells
    for i, h in enumerate(headers):
        clear_cell(header_cells[i])
        set_cell_shading(header_cells[i], header_fill)
        set_cell_border(header_cells[i], color=header_fill)
        set_cell_margins(header_cells[i], top=70, bottom=70)
        header_cells[i].vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        add_para(header_cells[i], h, size=7.7, color=WHITE, bold=True, after=0, align=WD_ALIGN_PARAGRAPH.CENTER)
    for r_idx, row_data in enumerate(rows):
        row = table.add_row()
        for c_idx, text in enumerate(row_data):
            cell = row.cells[c_idx]
            clear_cell(cell)
            set_cell_shading(cell, WHITE if r_idx % 2 == 0 else FILL_GRAY)
            set_cell_border(cell)
            set_cell_margins(cell)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            add_para(cell, fit_text(text, 86), size=7.25, color=INK, after=0)
    add_para(doc, "", size=2, after=1)
    return table


def add_metric_grid(doc, items: list[tuple[str, str, str]]):
    table = doc.add_table(rows=1, cols=3)
    set_table_geometry(table, [3120, 3120, 3120])
    fills = [FILL_BLUE, FILL_TEAL, FILL_GOLD]
    accents = [BLUE, TEAL, GOLD]
    for i, (label, value, note) in enumerate(items[:3]):
        cell = table.rows[0].cells[i]
        clear_cell(cell)
        set_cell_shading(cell, fills[i])
        set_cell_border(cell, color=accents[i])
        set_cell_margins(cell, top=140, start=150, bottom=140, end=150)
        add_para(cell, label, size=8.2, color=accents[i], bold=True, after=1, align=WD_ALIGN_PARAGRAPH.CENTER)
        add_para(cell, value, size=15, color=NAVY, bold=True, after=1, align=WD_ALIGN_PARAGRAPH.CENTER)
        add_para(cell, note, size=8.1, color=MUTED, after=0, align=WD_ALIGN_PARAGRAPH.CENTER)
    add_para(doc, "", size=2, after=4)
    return table


def fit_text(text: str, max_chars: int = 130) -> str:
    text = sanitize_visible_text(text)
    text = re.sub(r"\s+", " ", text).strip(" |")
    if len(text) <= max_chars:
        return text
    return text[: max_chars - 1].rstrip("，。；、 ") + "。"


def sanitize_visible_text(text: str) -> str:
    if not text:
        return ""
    out = str(text)
    out = re.sub(r"https?://\S+|www\.\S+", "", out, flags=re.I)
    out = out.replace("A-share", "A股").replace("A-SHARE", "A股").replace("A Share", "A股")
    out = out.replace("AI", "人工智能").replace("EDA", "电子设计自动化")
    out = out.replace("东方财富Choice", "东方财富数据").replace("Choice", "数据")
    replacements: list[tuple[str, str]] = []
    for src in KNOWN_SOURCES:
        for alias in src["aliases"]:
            if alias != src["name"]:
                replacements.append((alias, src["name"]))
    replacements.extend(
        [
            ("Research Brief", "研究简报"),
            ("Evidence-based", "基于证据"),
            ("THINK-TANK", "智库"),
            ("SOURCE INTAKE", "来源"),
            ("SCENARIOS", "情景"),
            ("TECH × GEO", "科技与地缘"),
            ("TECH x GEO", "科技与地缘"),
            ("US THINK", "美国智库"),
            ("UTF-8", ""),
            ("原生可编辑", ""),
            ("紧凑精美终版", ""),
            ("链接保留", ""),
            ("本页把", ""),
        ]
    )
    for old, new in sorted(replacements, key=lambda x: len(x[0]), reverse=True):
        if re.search(r"[A-Za-z]", old):
            out = re.sub(rf"\b{re.escape(old)}\b", new, out, flags=re.I)
        elif old not in new:
            out = out.replace(old, new)
    out = re.sub(r"[`*_>#]", "", out)
    out = re.sub(r"\s+", " ", out)
    return out.strip()


META_SECTION_MARKERS = [
    "公众号写作",
    "新版文章",
    "自测",
    "反偷懒",
    "深度要求",
    "表达标准",
]


META_LINE_MARKERS = [
    "对公众号文章",
    "标题方向",
    "导语",
    "第一节",
    "第二节",
    "第三节",
    "第四节",
    "第五节",
    "第六节",
    "不要写成",
    "写成“",
    "不能独占全文",
    "不得把",
    "不得只",
    "必须避免",
]


AXIS_DOMAIN_TERMS = {
    "fiscal": ["财政", "债务", "利率", "金融条件", "赤字"],
    "trade": ["贸易", "关税", "出口", "全球失衡", "供应链"],
    "security": ["外交", "安全", "能源", "航运", "冲突", "伊朗", "中东"],
    "resources": ["关键矿产", "稀土", "永磁", "资源", "矿产"],
    "technology": ["科技", "人工智能", "半导体", "算力", "芯片", "工业软件"],
}


def iter_markdown_sections(text: str) -> list[tuple[str, str]]:
    headings = list(re.finditer(r"(?m)^\s{0,3}#{1,4}\s*(.+?)\s*$", text))
    sections: list[tuple[str, str]] = []
    if not headings:
        return [("", text)]
    prefix = text[: headings[0].start()].strip()
    if prefix:
        sections.append(("", prefix))
    for idx, match in enumerate(headings):
        title = match.group(1).strip()
        start = match.end()
        end = headings[idx + 1].start() if idx + 1 < len(headings) else len(text)
        sections.append((title, text[start:end].strip()))
    return sections


def clean_research_text(text: str) -> str:
    cleaned: list[str] = []
    for title, body in iter_markdown_sections(text):
        if any(marker in title for marker in META_SECTION_MARKERS):
            continue
        body_lines: list[str] = []
        for raw in body.splitlines():
            if any(marker in raw for marker in META_LINE_MARKERS):
                continue
            body_lines.append(raw)
        body_clean = "\n".join(body_lines).strip()
        if not body_clean:
            continue
        if title:
            cleaned.append(f"## {title}\n\n{body_clean}")
        else:
            cleaned.append(body_clean)
    return "\n\n".join(cleaned).strip()


def extract_section(text: str, keywords: list[str]) -> str:
    headings = list(re.finditer(r"(?m)^\s{0,3}#{1,4}\s*(.+?)\s*$", text))
    for idx, match in enumerate(headings):
        title = match.group(1)
        if any(k in title for k in keywords):
            start = match.end()
            end = headings[idx + 1].start() if idx + 1 < len(headings) else len(text)
            return text[start:end].strip()
    return ""


def extract_items(text: str, keywords: list[str] | None = None, max_items: int = 6, min_chars: int = 8) -> list[str]:
    segment = extract_section(text, keywords or []) if keywords else text
    if not segment:
        return []
    items: list[str] = []
    for raw in segment.splitlines():
        line = raw.strip()
        if not line or line.startswith("| ---") or set(line) <= {"|", "-", " "}:
            continue
        line = re.sub(r"^\s*[-*+]\s*", "", line)
        line = re.sub(r"^\s*\d+[\.、]\s*", "", line)
        line = line.strip("| ")
        if len(line) < min_chars:
            continue
        clean = fit_text(line, 150)
        if clean and clean not in items:
            items.append(clean)
        if len(items) >= max_items:
            break
    return items[:max_items]


def require_items(name: str, items: list[str], minimum: int) -> list[str]:
    if len(items) < minimum:
        raise ValueError(f"{name} insufficient: {len(items)} < {minimum}")
    return items


def split_prefixed_item(line: str, prefixes: list[str]) -> str:
    for prefix in prefixes:
        for sep in ("：", ":"):
            marker = prefix + sep
            if line.startswith(marker):
                return line[len(marker):].strip()
    return ""


def split_label_body(line: str) -> tuple[str, str]:
    clean = sanitize_visible_text(line)
    for sep in ("：", ":"):
        if sep in clean:
            label, body = clean.split(sep, 1)
            label = label.strip()
            body = body.strip()
            if label and body:
                return fit_text(label, 28), fit_text(body, 180)
    return "", fit_text(clean, 180)


def extract_axes(text: str) -> list[dict[str, str]]:
    axes: list[dict[str, str]] = []
    for title, body in iter_markdown_sections(text):
        if "主线" not in title:
            continue
        items = extract_items(f"## {title}\n{body}", [title], max_items=12, min_chars=10)
        if not items:
            continue
        signal = ""
        mapping = ""
        verify = ""
        logic = ""
        for item in items:
            mapping = mapping or split_prefixed_item(item, ["A股映射"])
            verify = verify or split_prefixed_item(item, ["验证指标"])
            logic = logic or split_prefixed_item(item, ["深层逻辑"])
            if not signal and not any(item.startswith(p) for p in ["A股映射", "验证指标", "深层逻辑"]):
                signal = item
        if signal and mapping and verify and logic:
            axes.append(
                {
                    "title": sanitize_visible_text(title),
                    "signal": signal,
                    "mapping": mapping,
                    "verify": verify,
                    "logic": logic,
                }
            )
    return axes


def axis_domain_coverage(axes: list[dict[str, str]]) -> dict[str, bool]:
    joined = "\n".join(" ".join(axis.values()) for axis in axes)
    return {name: any(term in joined for term in terms) for name, terms in AXIS_DOMAIN_TERMS.items()}


def extract_market_facts(items: list[str]) -> list[str]:
    terms = ["上证指数", "深证成指", "创业板指", "成交额", "上涨", "下跌", "涨跌家数", "涨停", "个股", "板块", "概念"]
    facts = [item for item in items if any(term in item for term in terms)]
    return facts[:8]


def market_fact_checks(text: str) -> dict[str, bool]:
    return {
        "index_move": bool(re.search(r"(上证指数|深证成指|创业板指).{0,18}(涨|跌)\s*\d+(?:\.\d+)?\s*%", text)),
        "turnover": bool(re.search(r"成交额.{0,18}\d+(?:\.\d+)?\s*万?亿元", text)),
        "breadth": bool(
            re.search(r"(上涨|下跌|涨跌家数).{0,18}\d+.{0,8}(只|家|个)", text)
            or re.search(r"\d+.{0,8}(只|家|个).{0,12}(上涨|下跌)", text)
        ),
        "sector_or_stock": bool(
            re.search(r"(涨停|板块|概念|个股).{0,80}(有研硅|有研新材|华天科技|半导体|游戏|中芯)", text)
            or re.search(r"(有研硅|有研新材|华天科技|半导体|游戏|中芯).{0,80}(涨停|板块|概念|个股)", text)
        ),
    }


def market_variable_label(fact: str) -> str:
    if any(k in fact for k in ["上证指数", "深证成指", "创业板指"]):
        return "指数方向"
    if "成交额" in fact:
        return "成交额"
    if any(k in fact for k in ["上涨", "下跌", "涨跌家数"]):
        return "市场宽度"
    if any(k in fact for k in ["涨停", "个股", "板块", "概念"]):
        return "强弱结构"
    return "事实变量"


def derive_market_badge(facts: list[str]) -> str:
    joined = " ".join(facts)
    index_facts = [f for f in facts if any(k in f for k in ["上证指数", "深证成指", "创业板指"])]
    if index_facts and all("跌" in f for f in index_facts[:3]):
        return "三指下跌"
    if index_facts and all("涨" in f for f in index_facts[:3]):
        return "三指上涨"
    if "成交额" in joined and any(k in joined for k in ["上涨", "下跌"]):
        return "成交分化"
    return fit_text(facts[0], 8) if facts else "盘面待验"


def extract_markdown_sources(text: str) -> list[dict[str, str]]:
    sources: list[dict[str, str]] = []
    seen: set[str] = set()
    for label, url in re.findall(r"\[([^\]]+)\]\((https?://[^)\s]+)\)", text, flags=re.I):
        clean_label = sanitize_visible_text(label)
        if not clean_label or url in seen:
            continue
        seen.add(url)
        matched = match_source(label + " " + url)
        focus = str(matched["focus"]) if matched else "当次事实来源"
        sources.append(
            {
                "name": fit_text(clean_label, 42),
                "url": url,
                "focus": focus,
                "signal": fit_text(f"引用材料：{clean_label}", 78),
            }
        )
    return sources


def parse_scenario_rows(items: list[str]) -> list[list[str]]:
    rows: list[list[str]] = []
    for item in items:
        label, body = split_label_body(item)
        if label and body:
            rows.append([label.replace("情景", ""), body])
        else:
            rows.append([f"情景{len(rows) + 1}", body])
    return rows


def parse_monitor_rows(items: list[str]) -> list[list[str]]:
    rows: list[list[str]] = []
    for item in items:
        label, body = split_label_body(item)
        rows.append([label or f"变量{len(rows) + 1}", body])
    return rows


def parse_risk_rows(items: list[str]) -> list[list[str]]:
    rows: list[list[str]] = []
    for idx, item in enumerate(items, 1):
        label, body = split_label_body(item)
        rows.append([label or f"边界{idx}", body])
    return rows


def pick_evidence(judgement: str, facts: list[str], used: set[int]) -> str:
    tokens = [t for t in re.findall(r"[\u4e00-\u9fff]{2,}", judgement) if len(t) >= 2]
    best_idx = None
    best_score = -1
    for idx, fact in enumerate(facts):
        if idx in used:
            continue
        score = sum(1 for token in tokens if token in fact)
        if re.search(r"\d", fact):
            score += 1
        if score > best_score:
            best_idx = idx
            best_score = score
    if best_idx is None:
        raise ValueError("not enough fact evidence for core judgements")
    used.add(best_idx)
    return facts[best_idx]


def validate_research_material(text: str) -> dict[str, Any]:
    text = clean_research_text(text)
    compact = re.sub(r"\s+", "", text)
    links = re.findall(r"\[[^\]]+\]\(https?://[^)\s]+\)", text, flags=re.I)
    numbers = re.findall(r"\d+(?:\.\d+)?\s*(?:%|亿元|万亿元|亿美元|万亿美元|家|只|个|年|月|日|nm|英寸|层|点|页|条|项)", text, flags=re.I)
    market_checks = market_fact_checks(text)
    axes = extract_axes(text)
    axis_checks = axis_domain_coverage(axes)
    fact_items = extract_items(text, ["数字事实池"], max_items=80, min_chars=6)
    core_items = extract_items(text, ["核心判断"], max_items=8, min_chars=10)
    scenario_items = extract_items(text, ["情景推演", "情景"], max_items=8, min_chars=10)
    risk_items = extract_items(text, ["风险边界", "风险"], max_items=8, min_chars=10)
    monitor_items = extract_items(text, ["交叉主线监测表", "监测"], max_items=10, min_chars=10)
    checks = {
        "min_chars": len(compact) >= 6500,
        "min_links": len(set(links)) >= 12,
        "min_numeric_facts": len(set(numbers)) >= 18,
        "core_judgements": len(core_items) >= 3,
        "fact_pool": len(fact_items) >= 18,
        "scenario_rows": len(scenario_items) >= 4,
        "risk_rows": len(risk_items) >= 4,
        "monitor_rows": len(monitor_items) >= 5,
        "axis_count": len(axes) >= 5,
        "axis_domain_coverage": all(axis_checks.values()),
        "a_share_index_move": market_checks["index_move"],
        "a_share_turnover": market_checks["turnover"],
        "a_share_breadth": market_checks["breadth"],
        "a_share_sector_or_stock": market_checks["sector_or_stock"],
        "thinktank_policy": any(k in text for k in ["智库", "布鲁金斯", "兰德", "战略与国际问题研究中心", "出口管制", "美国商务部", "联邦公报"]),
        "macro_or_trade": any(k in text for k in ["财政", "债务", "利率", "通胀", "关税", "贸易", "税收", "赤字", "金融条件"]),
        "geo_policy": any(k in text for k in ["地缘", "出口管制", "供应链", "盟友", "制裁", "国防", "关键矿产", "海外子公司"]),
        "industry_or_resources": any(k in text for k in ["半导体设备", "硅片", "先进封装", "算力", "工业软件", "网络安全", "汽车芯片", "能源", "稀土", "关键矿产", "电力", "储能"]),
        "a_share_market": any(k in text for k in ["上证指数", "深证成指", "创业板指", "成交额", "涨停", "A股"]),
        "company_or_sector": any(k in text for k in [
            "有研硅", "华天科技", "有研新材", "TCL中环", "沪硅产业", "长电科技", "盛美上海", "中芯",
            "中国稀土", "北方稀土", "金力永磁", "中国海油", "中国石油", "中国核电", "中航沈飞",
            "中航西飞", "中国船舶", "紫金矿业", "工商银行", "招商银行", "中信证券", "创新药", "军工", "稀土"
        ]),
        "risk_boundary": any(k in text for k in ["风险", "降级", "失效", "兑现", "毛利", "现金流", "应收"]),
        "no_meta_residue": not any(term in text for term in ["公众号写作骨架", "对公众号文章", "新版文章必须避免", "不能独占全文", "标题方向"]),
    }
    missing = [name for name, ok in checks.items() if not ok]
    return {
        "ok": not missing,
        "missing": missing,
        "metrics": {
            "chars_compact": len(compact),
            "markdown_links": len(set(links)),
            "numeric_facts": len(set(numbers)),
        },
        "checks": checks,
    }


def detect_sources(source_text: str) -> list[dict[str, str]]:
    found: dict[str, dict[str, str]] = {}
    hay = source_text.casefold()
    for src in KNOWN_SOURCES:
        matched = False
        for alias in src["aliases"]:
            if alias.casefold() in hay:
                matched = True
                break
        if src["url"].replace("https://", "").replace("http://", "").strip("/").casefold() in hay:
            matched = True
        if matched:
            found[src["name"]] = {k: str(src[k]) for k in ("name", "url", "focus", "signal")}
    for label, url in re.findall(r"\[([^\]]+)\]\((https?://[^)\s]+)\)", source_text, flags=re.I):
        matched_src = match_source(label + " " + url)
        if matched_src:
            found[matched_src["name"]] = {k: str(matched_src[k]) for k in ("name", "url", "focus", "signal")}
        elif len(found) < 16:
            name = f"中文来源{len(found) + 1}"
            found[name] = {
                "name": name,
                "url": url,
                "focus": "补充资料",
                "signal": "用于交叉验证当次研究材料。",
            }
    return list(found.values())


def match_source(text: str) -> dict[str, Any] | None:
    hay = text.casefold()
    for src in KNOWN_SOURCES:
        if any(alias.casefold() in hay for alias in src["aliases"]):
            return src
        domain = src["url"].replace("https://", "").replace("http://", "").strip("/").casefold()
        if domain in hay:
            return src
    return None


def safe_filename(text: str) -> str:
    text = re.sub(r'[\\/:*?"<>|]+', "_", text)
    text = re.sub(r"\s+", "_", text).strip("_")
    return text[:90] or "顶级智库观点简报"


def configure_core_props(doc, title: str):
    props = doc.core_properties
    props.title = title
    props.subject = "科技产业、地缘政治、A股映射"
    props.author = "Codex 顶级智库观点"
    props.keywords = "智库,科技产业,地缘政治,A股"


def build_docx(source_path: Path, brief_path: Path, out_dir: Path, title: str | None = None, as_of: str | None = None) -> Path:
    require_docx()
    if not LAYOUT_SPEC.is_file():
        raise FileNotFoundError(f"layout specification missing: {LAYOUT_SPEC}")
    if not source_path.exists():
        raise FileNotFoundError(f"thinktank-md not found: {source_path}")
    if not brief_path.exists():
        raise FileNotFoundError(f"brief-md not found: {brief_path}")
    source_text = read_text(source_path)
    brief_text = clean_research_text(read_text(brief_path))
    material_gate = validate_research_material(brief_text)
    if not material_gate["ok"]:
        raise ValueError(f"research material depth gate failed: {material_gate}")
    sources = extract_markdown_sources(brief_text)
    known_sources = detect_sources(source_text + "\n" + brief_text)
    known_seen = {s["name"] for s in sources}
    for src in known_sources:
        if src["name"] not in known_seen:
            sources.append(src)
            known_seen.add(src["name"])
    if len(sources) < 10:
        for src in KNOWN_SOURCES:
            if src["name"] not in {s["name"] for s in sources}:
                sources.append({k: str(src[k]) for k in ("name", "url", "focus", "signal")})
            if len(sources) >= 10:
                break
    if len(sources) < 10:
        raise ValueError("not enough source anchors; at least 10 Chinese source labels are required")

    as_of = as_of or now_cn_date()
    title = title or "美国智库五条压力链：A股研究简报"
    title = sanitize_visible_text(title)

    core = require_items("core judgements", extract_items(brief_text, ["核心判断"], max_items=5, min_chars=10), 3)
    fact_items = require_items("numeric fact pool", extract_items(brief_text, ["数字事实池"], max_items=80, min_chars=6), 18)
    axes = extract_axes(brief_text)
    if len(axes) < 5:
        raise ValueError(f"axis extraction failed: {len(axes)} < 5")
    market_facts = require_items("A-share market facts", extract_market_facts(fact_items + core), 4)
    scenarios = require_items("scenario rows", extract_items(brief_text, ["情景推演", "情景"], max_items=6, min_chars=10), 4)
    risks = require_items("risk rows", extract_items(brief_text, ["风险边界", "风险"], max_items=8, min_chars=10), 4)
    monitors = require_items("monitor rows", extract_items(brief_text, ["交叉主线监测表", "监测"], max_items=8, min_chars=10), 5)
    numbers = re.findall(r"\d+(?:\.\d+)?\s*(?:%|亿元|万亿元|亿美元|万亿美元|家|只|个|年|月|日|条|项)", brief_text)

    doc = Document()
    style_doc(doc)
    configure_core_props(doc, title)

    # Page 1
    add_para(doc, "美国智库信号", size=8.5, color=TEAL, bold=True, after=2)
    p = add_para(doc, title, size=22, color=NAVY, bold=True, after=2)
    p.paragraph_format.line_spacing = 0.96
    add_para(doc, f"数据日：{as_of}    来源锚点：{len(sources)}个    主线：{len(axes)}条", size=9.0, color=MUTED, after=7)
    add_metric_grid(
        doc,
        [
            ("来源锚点", f"{len(sources)}个", "中文可见标签"),
            ("数字事实", f"{len(set(numbers))}条", "含盘面与来源事实"),
            ("盘面基准", derive_market_badge(market_facts), "指数、成交、宽度"),
        ],
    )
    add_section_title(doc, "核心判断", "每条判断绑定材料中的事实依据")
    evidence_used: set[int] = set()
    core_rows = [[core[i], pick_evidence(core[i], fact_items, evidence_used), market_facts[min(i, len(market_facts) - 1)]] for i in range(3)]
    add_table(
        doc,
        ["判断", "事实依据", "盘面依据"],
        core_rows,
        [3700, 3100, 2560],
        header_fill=NAVY,
    )
    add_section_title(doc, "A股盘面事实", "只列当次材料中可复核的市场变量")
    add_table(
        doc,
        ["变量", "盘面事实"],
        [[market_variable_label(fact), fact] for fact in market_facts[:5]],
        [1800, 7560],
        header_fill=TEAL,
    )

    doc.add_page_break()

    # Page 2
    add_section_title(doc, "美国智库信号", "机构只显示中文名称，超链接绑定在名称背后")
    source_rows = [[s["name"], s["focus"], s["signal"]] for s in sources[:5]]
    add_table(doc, ["机构", "观察重心", "可跟踪信号"], source_rows, [2100, 2600, 4660], header_fill=NAVY)
    add_section_title(doc, "五条事实链", "主线、智库信号和A股映射均来自研究材料")
    axis_rows = [[axis["title"], axis["signal"], axis["mapping"]] for axis in axes[:5]]
    add_table(doc, ["主线", "事实信号", "A股映射"], axis_rows, [1800, 3900, 3660], header_fill=BLUE)

    doc.add_page_break()

    # Page 3
    add_section_title(doc, "传导链条", "从美国政策事实推导到A股验证变量")
    transmission_rows = [[axis["title"], axis["logic"], axis["verify"]] for axis in axes[:5]]
    add_table(doc, ["主线", "传导逻辑", "验证指标"], transmission_rows, [1700, 5000, 2660], header_fill=NAVY)
    add_section_title(doc, "数字事实池", "保留材料中的硬事实骨架")
    fact_rows = [[str(i), fact] for i, fact in enumerate(fact_items[:6], 1)]
    add_table(doc, ["序号", "事实"], fact_rows, [700, 8660], header_fill=TEAL)

    doc.add_page_break()

    # Page 4
    add_section_title(doc, "情景推演", "所有情景来自材料原文")
    scenario_rows = parse_scenario_rows(scenarios)
    add_table(doc, ["情景", "触发和市场含义"], scenario_rows, [1200, 8160], header_fill=BLUE)
    add_section_title(doc, "监测清单", "把观点转成可复核的跟踪信号")
    monitor_rows = parse_monitor_rows(monitors)
    add_table(doc, ["主线", "跟踪变量与判断口径"], monitor_rows[:5], [1500, 7860], header_fill=TEAL)

    doc.add_page_break()

    # Page 5
    add_section_title(doc, "中文来源", "可见标签为中文，链接绑定在来源名称中")
    source_table = doc.add_table(rows=1, cols=3)
    set_table_geometry(source_table, [900, 2500, 5960])
    for idx, header in enumerate(["序号", "来源", "用途"]):
        cell = source_table.rows[0].cells[idx]
        clear_cell(cell)
        set_cell_shading(cell, NAVY)
        set_cell_border(cell, color=NAVY)
        set_cell_margins(cell)
        add_para(cell, header, size=8.8, color=WHITE, bold=True, after=0, align=WD_ALIGN_PARAGRAPH.CENTER)
    for i, src in enumerate(sources[:8], 1):
        row = source_table.add_row()
        vals = [str(i), src["name"], src["focus"]]
        for c_idx, value in enumerate(vals):
            cell = row.cells[c_idx]
            clear_cell(cell)
            set_cell_shading(cell, WHITE if i % 2 else FILL_GRAY)
            set_cell_border(cell)
            set_cell_margins(cell)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            if c_idx == 1:
                p = add_para(cell, "", size=8.0, after=0)
                add_hyperlink(p, sanitize_visible_text(value), src["url"])
            else:
                add_para(cell, fit_text(value, 78), size=8.0, color=INK, after=0, align=WD_ALIGN_PARAGRAPH.CENTER if c_idx == 0 else None)
    add_para(doc, "", size=2, after=3)
    add_section_title(doc, "风险边界", "来自材料原文的降级条件")
    add_table(doc, ["边界", "说明"], parse_risk_rows(risks)[:5], [1500, 7860], header_fill=GOLD)

    out_dir.mkdir(parents=True, exist_ok=True)
    out_path = out_dir / f"{safe_filename(title)}_{as_of}.docx"
    doc.save(str(out_path))
    return out_path


def extract_xml_text(xml: str) -> str:
    texts = re.findall(r"<w:t(?:\s[^>]*)?>(.*?)</w:t>", xml, flags=re.S)
    return "".join(html.unescape(re.sub(r"<[^>]+>", "", t)) for t in texts)


def visible_text_from_docx(docx_path: Path) -> str:
    parts: list[str] = []
    with ZipFile(docx_path) as zf:
        for name in zf.namelist():
            if re.match(r"word/(document|header\d*|footer\d*|footnotes|endnotes)\.xml$", name):
                parts.append(extract_xml_text(zf.read(name).decode("utf-8", errors="ignore")))
    return "\n".join(parts)


def page_segments_from_docx(docx_path: Path) -> list[str]:
    with ZipFile(docx_path) as zf:
        xml = zf.read("word/document.xml").decode("utf-8", errors="ignore")
    chunks = re.split(r"<w:br[^>]*w:type=[\"']page[\"'][^>]*/>|<w:lastRenderedPageBreak\s*/>", xml)
    return [extract_xml_text(chunk) for chunk in chunks]


def count_hyperlinks(docx_path: Path) -> int:
    with ZipFile(docx_path) as zf:
        total = 0
        for name in zf.namelist():
            if name.startswith("word/_rels/") and name.endswith(".rels"):
                xml = zf.read(name).decode("utf-8", errors="ignore")
                total += len(re.findall(r'Type="[^"]*/hyperlink"', xml))
        return total


def word_com_page_count_pywin32(docx_path: Path) -> int | None:
    try:
        import win32com.client  # type: ignore
    except Exception:
        return None
    word = None
    doc = None
    try:
        word = win32com.client.DispatchEx("Word.Application")
        word.Visible = False
        doc = word.Documents.Open(str(docx_path), ReadOnly=True, AddToRecentFiles=False)
        return int(doc.ComputeStatistics(2))
    except Exception:
        return None
    finally:
        try:
            if doc is not None:
                doc.Close(False)
        except Exception:
            pass
        try:
            if word is not None:
                word.Quit()
        except Exception:
            pass


def word_com_page_count_powershell(docx_path: Path) -> int | None:
    ps_path = str(docx_path.resolve()).replace("'", "''")
    script = f"""
$ErrorActionPreference = 'Stop'
$word = $null
$doc = $null
try {{
  $word = New-Object -ComObject Word.Application
  $word.Visible = $false
  $word.DisplayAlerts = 0
  $doc = $word.Documents.Open('{ps_path}', $false, $true, $false)
  $doc.Repaginate() | Out-Null
  $pages = [int]$doc.ComputeStatistics(2)
  Write-Output "PAGES:$pages"
  exit 0
}} catch {{
  Write-Output ("ERROR:" + $_.Exception.Message)
  exit 1
}} finally {{
  if ($doc -ne $null) {{ $doc.Close($false) | Out-Null }}
  if ($word -ne $null) {{ $word.Quit() | Out-Null }}
}}
"""
    try:
        proc = subprocess.run(
            ["powershell", "-NoProfile", "-ExecutionPolicy", "Bypass", "-Command", script],
            capture_output=True,
            text=True,
            encoding="utf-8",
            errors="ignore",
            timeout=90,
        )
    except Exception:
        return None
    match = re.search(r"PAGES:(\d+)", proc.stdout or "")
    if proc.returncode == 0 and match:
        pages = int(match.group(1))
        return pages if pages > 0 else None
    return None


def pdf_page_count_with_pdfinfo(pdf_path: Path) -> int | None:
    pdfinfo = shutil.which("pdfinfo")
    if not pdfinfo:
        return None
    try:
        proc = subprocess.run(
            [pdfinfo, str(pdf_path)],
            capture_output=True,
            text=True,
            encoding="utf-8",
            errors="ignore",
            timeout=60,
        )
    except Exception:
        return None
    match = re.search(r"^Pages:\s*(\d+)", proc.stdout or "", flags=re.M)
    if proc.returncode == 0 and match:
        pages = int(match.group(1))
        return pages if pages > 0 else None
    return None


def libreoffice_pdf_page_count(docx_path: Path) -> int | None:
    soffice = shutil.which("soffice") or shutil.which("libreoffice")
    if not soffice:
        return None
    with tempfile.TemporaryDirectory(prefix="top_thinktank_pages_") as td:
        out_dir = Path(td)
        try:
            proc = subprocess.run(
                [soffice, "--headless", "--convert-to", "pdf", "--outdir", str(out_dir), str(docx_path)],
                capture_output=True,
                text=True,
                encoding="utf-8",
                errors="ignore",
                timeout=120,
            )
        except Exception:
            return None
        if proc.returncode != 0:
            return None
        pdfs = list(out_dir.glob("*.pdf"))
        if not pdfs:
            return None
        return pdf_page_count_with_pdfinfo(pdfs[0])


def page_count_probe(docx_path: Path, estimated_pages: int) -> dict[str, Any]:
    pywin32_pages = word_com_page_count_pywin32(docx_path)
    if pywin32_pages:
        return {
            "checked_pages": pywin32_pages,
            "word_pages": pywin32_pages,
            "page_count_method": "word_com_pywin32",
            "warning": "",
        }
    powershell_pages = word_com_page_count_powershell(docx_path)
    if powershell_pages:
        return {
            "checked_pages": powershell_pages,
            "word_pages": powershell_pages,
            "page_count_method": "word_com_powershell",
            "warning": "",
        }
    pdf_pages = libreoffice_pdf_page_count(docx_path)
    if pdf_pages:
        return {
            "checked_pages": pdf_pages,
            "word_pages": None,
            "page_count_method": "pdf_render_pdfinfo",
            "warning": "Word COM unavailable; used PDF render page count",
        }
    return {
        "checked_pages": estimated_pages,
        "word_pages": None,
        "page_count_method": "ooxml_page_break_estimate",
        "warning": "Word COM and PDF page count unavailable; used OOXML page-break estimate",
    }


def scan_docx_file(
    docx_path: Path,
    min_hyperlinks: int = 6,
    min_tables: int = 5,
    min_pages: int = 4,
    max_pages: int = 8,
) -> dict[str, Any]:
    require_docx()
    failures: list[str] = []
    warnings: list[str] = []
    if not docx_path.exists():
        return {"status": "BLOCKED", "failures": [f"file not found: {docx_path}"], "path": str(docx_path)}
    if docx_path.suffix.lower() != ".docx":
        failures.append("not a .docx file")
    with ZipFile(docx_path) as zf:
        names = zf.namelist()
        media_count = sum(1 for n in names if n.startswith("word/media/") and not n.endswith("/"))
    doc = Document(str(docx_path))
    paragraph_count = len([p for p in doc.paragraphs if p.text.strip()])
    table_count = len(doc.tables)
    hyperlink_count = count_hyperlinks(docx_path)
    visible_text = visible_text_from_docx(docx_path)
    visible_norm = visible_text.casefold()
    banned_hits = []
    for term in BANNED_TERMS:
        if term.casefold() in visible_norm:
            banned_hits.append(term)
    directive_hits = []
    for term in TRADING_DIRECTIVE_TERMS:
        for match in re.finditer(re.escape(term), visible_text):
            ctx = visible_text[max(0, match.start() - 6): match.end() + 6]
            if any(ok in ctx for ok in ["净买入", "融资买入", "买入额", "净卖出"]):
                continue
            directive_hits.append(term)
            break
    if banned_hits:
        failures.append(f"banned terms in visible text: {banned_hits}")
    if directive_hits:
        failures.append(f"trading directive terms in visible text: {directive_hits}")
    if re.search(r"https?://|www\.|\.com\b|\.org\b|\.edu\b", visible_text, flags=re.I):
        failures.append("visible URL or domain residue found")
    if media_count != 0:
        failures.append(f"media_count must be 0, got {media_count}")
    if hyperlink_count < min_hyperlinks:
        failures.append(f"hyperlinks too few: {hyperlink_count} < {min_hyperlinks}")
    if table_count < min_tables:
        failures.append(f"tables too few: {table_count} < {min_tables}")
    if paragraph_count < 20:
        failures.append(f"paragraphs too few: {paragraph_count} < 20")
    segments = page_segments_from_docx(docx_path)
    segment_lengths = [len(re.sub(r"\s+", "", s)) for s in segments]
    estimated_pages = len(segments)
    page_probe = page_count_probe(docx_path, estimated_pages)
    word_pages = page_probe["word_pages"]
    checked_pages = page_probe["checked_pages"]
    if checked_pages < min_pages or checked_pages > max_pages:
        failures.append(f"page count outside range: {checked_pages} not in [{min_pages}, {max_pages}]")
    blank_segments = [idx + 1 for idx, ln in enumerate(segment_lengths) if ln < 80]
    if blank_segments:
        failures.append(f"blank or tiny page segments: {blank_segments}")
    if page_probe["warning"]:
        warnings.append(page_probe["warning"])
    return {
        "status": "CLEAN_PASS" if not failures else "BLOCKED",
        "path": str(docx_path),
        "metrics": {
            "paragraphs": paragraph_count,
            "tables": table_count,
            "media_count": media_count,
            "hyperlinks": hyperlink_count,
            "estimated_pages": estimated_pages,
            "word_pages": word_pages,
            "checked_pages": checked_pages,
            "page_count_method": page_probe["page_count_method"],
            "page_text_lengths": segment_lengths,
            "visible_chars": len(visible_text),
        },
        "banned_hits": banned_hits,
        "failures": failures,
        "warnings": warnings,
    }


def write_manifest(out_dir: Path, docx_path: Path, scan: dict[str, Any], command: str) -> Path:
    manifest = {
        "skill": "顶级智库观点",
        "command": command,
        "created_at": dt.datetime.now().astimezone().isoformat(timespec="seconds"),
        "docx": str(docx_path),
        "layout_binding": {
            "path": str(LAYOUT_SPEC.resolve()),
            "sha256": hashlib.sha256(LAYOUT_SPEC.read_bytes()).hexdigest(),
        },
        "validation": scan,
        "status": scan["status"],
    }
    path = out_dir / "顶级智库观点_manifest.json"
    path.write_text(json.dumps(manifest, ensure_ascii=False, indent=2), encoding="utf-8")
    return path


def command_generate(args: argparse.Namespace, command_name: str = "generate") -> int:
    try:
        docx_path = build_docx(Path(args.thinktank_md), Path(args.brief_md), Path(args.out_dir), args.title, args.as_of)
        scan = scan_docx_file(
            docx_path,
            min_hyperlinks=args.min_hyperlinks,
            min_tables=args.min_tables,
            min_pages=args.min_pages,
            max_pages=args.max_pages,
        )
        manifest = write_manifest(Path(args.out_dir), docx_path, scan, command_name)
        result = {
            "skill": "顶级智库观点",
            "status": scan["status"],
            "docx": str(docx_path),
            "manifest": str(manifest),
            "scan": scan,
        }
        return json_out(result, 0 if scan["status"] == "CLEAN_PASS" else 1)
    except Exception as exc:
        return json_out(
            {
                "skill": "顶级智库观点",
                "status": "BLOCKED",
                "error": f"{type(exc).__name__}: {exc}",
            },
            1,
        )


def command_scan(args: argparse.Namespace) -> int:
    try:
        scan = scan_docx_file(
            Path(args.docx_path),
            min_hyperlinks=args.min_hyperlinks,
            min_tables=args.min_tables,
            min_pages=args.min_pages,
            max_pages=args.max_pages,
        )
        return json_out(scan, 0 if scan["status"] == "CLEAN_PASS" else 1)
    except Exception as exc:
        return json_out({"status": "BLOCKED", "error": f"{type(exc).__name__}: {exc}"}, 1)


def command_auto(args: argparse.Namespace) -> int:
    has_inputs = bool(args.thinktank_md or args.brief_md or args.out_dir)
    if not has_inputs:
        return command_doctor(args)
    if not (args.thinktank_md and args.brief_md and args.out_dir):
        return json_out(
            {
                "skill": "顶级智库观点",
                "status": "BLOCKED",
                "reason": "auto generation requires --thinktank-md, --brief-md and --out-dir together",
            },
            1,
        )
    return command_generate(args, command_name="auto")


def command_doctor(_args: argparse.Namespace) -> int:
    checks: list[dict[str, Any]] = []
    checks.append({"name": "python-docx", "ok": Document is not None, "detail": DOCX_IMPORT_ERROR or "available"})
    for p in [SKILL_DIR / "SKILL.md", SKILL_DIR / "references" / "workflow.md", SKILL_DIR / "references" / "checklist.md", LAYOUT_SPEC, CANARY_BRIEF, ENTRY]:
        checks.append({"name": str(p.relative_to(SKILL_DIR)), "ok": p.exists(), "detail": f"size={p.stat().st_size}" if p.exists() else "missing"})
    source_probe = "\n".join(src["name"] for src in KNOWN_SOURCES[:10])
    checks.append({"name": "source_detector", "ok": len(detect_sources(source_probe)) >= 6, "detail": f"detected={len(detect_sources(source_probe))}"})
    ok = all(c["ok"] for c in checks)
    return json_out({"skill": "顶级智库观点", "status": "CLEAN_PASS" if ok else "BLOCKED", "checks": checks}, 0 if ok else 1)


def command_selftest(_args: argparse.Namespace) -> int:
    try:
        require_docx()
        with tempfile.TemporaryDirectory(prefix="top_thinktank_brief_") as td:
            tmp = Path(td)
            source_path = tmp / "美国权威智库.md"
            source_path.write_text(
                "\n".join(
                    [
                        "# 美国权威智库",
                        "布鲁金斯学会、兰德公司、战略与国际问题研究中心、外交关系协会、新美国安全中心、卡内基国际和平基金会、信息技术与创新基金会、彼得森国际经济研究所、胡佛研究所、大西洋理事会。",
                    ]
                ),
                encoding="utf-8",
            )
            out_dir = tmp / "out"
            docx_path = build_docx(source_path, CANARY_BRIEF, out_dir, as_of="2026-07-07")
            scan = scan_docx_file(docx_path)
            result = {
                "skill": "顶级智库观点",
                "status": scan["status"],
                "docx": str(docx_path),
                "scan": scan,
            }
            return json_out(result, 0 if scan["status"] == "CLEAN_PASS" else 1)
    except Exception as exc:
        return json_out({"skill": "顶级智库观点", "status": "BLOCKED", "error": f"{type(exc).__name__}: {exc}"}, 1)


def add_generate_args(parser: argparse.ArgumentParser, required: bool = True):
    parser.add_argument("--thinktank-md", required=required, default="", help="美国权威智库 Markdown 清单")
    parser.add_argument("--brief-md", required=required, default="", help="当次研究材料 Markdown")
    parser.add_argument("--out-dir", required=required, default="", help="输出目录")
    parser.add_argument("--title", default="", help="可选报告标题")
    parser.add_argument("--as-of", default="", help="资料日期，默认当天")
    parser.add_argument("--min-hyperlinks", type=int, default=6)
    parser.add_argument("--min-tables", type=int, default=5)
    parser.add_argument("--min-pages", type=int, default=4)
    parser.add_argument("--max-pages", type=int, default=8)


def add_scan_args(parser: argparse.ArgumentParser):
    parser.add_argument("docx_path")
    parser.add_argument("--min-hyperlinks", type=int, default=6)
    parser.add_argument("--min-tables", type=int, default=5)
    parser.add_argument("--min-pages", type=int, default=4)
    parser.add_argument("--max-pages", type=int, default=8)


def main(argv: list[str] | None = None) -> int:
    parser = argparse.ArgumentParser(description="Codex 顶级智库观点 Word 简报生成与闭环扫描。")
    sub = parser.add_subparsers(dest="command", required=True)
    sub.add_parser("selftest", help="生成临时样例并扫描")
    sub.add_parser("doctor", help="检查技能文件和依赖")
    p_generate = sub.add_parser("generate", help="生成 DOCX 并扫描")
    add_generate_args(p_generate)
    p_auto = sub.add_parser("auto", help="锁定路径：生成 DOCX 并扫描")
    add_generate_args(p_auto, required=False)
    p_scan = sub.add_parser("scan-docx", help="扫描 DOCX 质量与正文清洁度")
    add_scan_args(p_scan)
    args = parser.parse_args(argv)
    if args.command == "selftest":
        return command_selftest(args)
    if args.command == "doctor":
        return command_doctor(args)
    if args.command == "scan-docx":
        return command_scan(args)
    if args.command == "generate":
        return command_generate(args, "generate")
    if args.command == "auto":
        return command_auto(args)
    return json_out({"status": "BLOCKED", "error": f"unknown command: {args.command}"}, 1)


if __name__ == "__main__" and __import__("os").environ.get("ONESTOCK_STOCK_CANONICAL_CHILD") != "1":
    print("canonical_stock_legacy_entry_direct_execution_blocked", file=__import__("sys").stderr)
    raise SystemExit(2)


if __name__ == "__main__":
    raise SystemExit(main())
