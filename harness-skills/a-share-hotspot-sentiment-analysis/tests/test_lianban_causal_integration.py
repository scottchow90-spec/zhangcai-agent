from __future__ import annotations
# ONESTOCK_EMBEDDED_LOCAL_IMPORT_BOOTSTRAP
import sys as _onestock_embedded_sys
from pathlib import Path as _OneStockEmbeddedPath
_onestock_embedded_dir = str(_OneStockEmbeddedPath(__file__).resolve().parent)
if _onestock_embedded_dir not in _onestock_embedded_sys.path:
    _onestock_embedded_sys.path.insert(0, _onestock_embedded_dir)

import importlib.util
import json
from pathlib import Path

from docx import Document


ROOT = Path(r"D:\C盘转移\日志\codex\skills\a-share-hotspot-sentiment-analysis")
SCRIPT = (
    ROOT
    / "components"
    / "a-share-sentiment-workflow"
    / "scripts"
    / "a-share-sentiment-workflow.py"
)


def load_workflow():
    spec = importlib.util.spec_from_file_location("sentiment_workflow_lianban_test", SCRIPT)
    assert spec and spec.loader
    module = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(module)
    return module


def test_sina_feed_labels_are_normalized_to_mainstream_source():
    workflow = load_workflow()

    source, origin = workflow._canonical_chrome_capture_source(
        "primary.json",
        {"site": "A股公开主来源"},
        {"label": "环球市场播报", "source_key": "sina.com.cn"},
    )

    assert source == "新浪财经"
    assert origin == "环球市场播报"
    assert workflow.source_layer(source) == "主流财经"


def test_ai_compute_demand_news_is_a_causal_trigger():
    workflow = load_workflow()
    now = workflow._china_datetime()
    event = {
        "source": "新浪财经",
        "date": workflow.cn_date(now.date()),
        "published_at": now.isoformat(timespec="seconds"),
        "time_precision": "minute",
        "time_evidence": now.isoformat(timespec="minutes"),
        "window_type": workflow.EVENT_WINDOW_TYPE,
        "source_channel": "primary_public_readonly_fallback",
        "url": "https://finance.sina.com.cn/example",
        "event": (
            "CoreWeave公布二季度营收同比翻倍。各大企业持续加码基建建设，"
            "市场对其AI算力资源需求持续攀升。"
        ),
    }

    roles, context, variable, grounding = workflow._causal_roles(
        event,
        "算力",
        [],
    )

    assert roles == ["catalyst"]
    assert "算力资源需求持续攀升" in context
    assert variable == "准入范围与需求预期"
    assert grounding == "direct_theme"


def test_causal_candidates_reject_generic_cross_context_terms_and_normalize_interim_growth():
    workflow = load_workflow()
    now = workflow._china_datetime()
    trading_dates = workflow.recent_trading_dates(now)
    trading_day = trading_dates[-1]
    if (
        trading_day == now.date()
        and now.weekday() < 5
        and (now.hour, now.minute) < (9, 30)
    ):
        trading_day = trading_dates[-2]
    date_text = workflow.cn_date(trading_day)
    target_date = trading_day.isoformat()
    pool = [
        {
            "source": "财联社",
            "date": workflow.cn_date(now.date()),
            "published_at": now.isoformat(timespec="seconds"),
            "time_precision": "minute",
            "time_evidence": now.isoformat(timespec="minutes"),
            "window_type": workflow.EVENT_WINDOW_TYPE,
            "source_channel": workflow.MANDATORY_SENTIMENT_CHANNEL,
            "url": "https://www.cls.cn/example-interim-growth",
            "event": (
                "标的公司发布半年度报告，营业收入同比增长22.36%，"
                "归母净利润同比增长151.25%；一般账户不构成行业主题。"
            ),
        },
        {
            "source": "连板网·市场复盘",
            "date": date_text,
            "market_date": date_text,
            "published_at": now.isoformat(timespec="seconds"),
            "time_precision": "market_close",
            "time_evidence": target_date,
            "window_type": workflow.MARKET_VALIDATION_WINDOW_TYPE,
            "source_channel": "lianban_daily_market_cross_validation",
            "url": f"https://lianban.net/days/{target_date}.html",
            "event": (
                f"{date_text}交易复盘显示，所属概念为中报增长，该方向共有3只涨停，"
                "一般零售股票不是本主题，代表标的出现板块级涨停扩散。"
            ),
        },
        {
            "source": "融资数据",
            "date": date_text,
            "market_date": date_text,
            "published_at": now.isoformat(timespec="seconds"),
            "time_precision": "market_close",
            "time_evidence": target_date,
            "window_type": workflow.MARKET_VALIDATION_WINDOW_TYPE,
            "source_channel": "local_participant_topic_aggregation",
            "url": "local-fund-review",
            "event": (
                "所属概念为中报增长，该方向参与者确认：一般零售股票不是本主题，"
                "代表标的出现机构、融资和主力净流入记录。"
            ),
        },
    ]

    candidates = workflow.causal_candidate_terms(pool)

    assert "中报增长" in candidates
    assert {"一般", "标的", "增长"}.isdisjoint(candidates)
    cluster = workflow._build_causal_cluster("中报增长", pool)
    assert cluster["causal_quality"]["ready"] is True
    assert len(workflow.select_evidence(cluster, 10)) >= 3


def test_explicit_gold_mapping_is_not_discarded_as_a_stock_name_fragment():
    workflow = load_workflow()

    terms = workflow._explicit_theme_terms("A股映射：黄金与贵金属板块")

    assert {"黄金", "贵金属"} <= terms


def test_lianban_topics_join_trade_and_fund_evidence_into_ready_chains(
    tmp_path,
    monkeypatch,
):
    workflow = load_workflow()
    now = workflow._china_datetime()
    trading_day = workflow.recent_trading_dates(now)[-1]
    date_text = workflow.cn_date(trading_day)
    target_date = trading_day.isoformat()
    snapshot = {
        "schema": "LIANBAN_DAILY_SNAPSHOT_V1",
        "status": "CLEAN_PASS",
        "target_date": target_date,
        "source_name": "连板网",
        "source_role": "supplemental_event_theme_sentiment_cross_validation",
        "sources": {
            "page": {"url": f"https://lianban.net/days/{target_date}.html"},
            "open_data": {"url": f"https://lianban.net/opendata/{target_date}.json"},
        },
        "topics": [
            {"name": "算力", "count": 3},
            {"name": "AI应用", "count": 3},
        ],
        "page_details": {
            "stock_items": [
                {"name": "云赛智联", "board": "算力", "theme": "算力租赁"},
                {"name": "鸿博股份", "board": "算力", "theme": "算力"},
                {"name": "城地香江", "board": "算力", "theme": "算力租赁"},
                {"name": "威士顿", "board": "AI应用", "theme": "软件开发"},
                {"name": "格尔软件", "board": "AI应用", "theme": "软件开发"},
                {"name": "延华智能", "board": "AI应用", "theme": "软件开发"},
            ]
        },
    }
    snapshot_path = tmp_path / "lianban-daily.json"
    snapshot_path.write_text(
        json.dumps(snapshot, ensure_ascii=False),
        encoding="utf-8",
    )
    monkeypatch.setenv("CODEX_LIANBAN_STATUS", "CLEAN_PASS")
    monkeypatch.setenv("CODEX_LIANBAN_SNAPSHOT", str(snapshot_path))

    market_base = {
        "date": date_text,
        "published_at": now.isoformat(timespec="seconds"),
        "time_precision": "market_close",
        "time_evidence": target_date,
        "window_type": workflow.MARKET_VALIDATION_WINDOW_TYPE,
        "source_channel": "primary_local_trade_review",
        "url": "local-review",
    }
    local_events = [
        dict(
            market_base,
            source="融资数据",
            event=f"{name}资金流显示，主力净流入1亿元，反映大单资金参与强度。",
        )
        for name in (
            "云赛智联",
            "鸿博股份",
            "城地香江",
            "威士顿",
            "格尔软件",
            "延华智能",
        )
    ]
    lianban_events = workflow.collect_lianban_market_validation_events(
        now,
        local_events,
    )
    trigger_base = {
        "source": "新浪财经",
        "date": date_text,
        "published_at": now.isoformat(timespec="seconds"),
        "time_precision": "minute",
        "time_evidence": now.isoformat(timespec="minutes"),
        "window_type": workflow.EVENT_WINDOW_TYPE,
        "source_channel": "primary_public_readonly_fallback",
        "url": "https://finance.sina.com.cn/example",
    }
    triggers = [
        dict(
            trigger_base,
            event="CoreWeave二季度营收同比翻倍，企业加大AI建设，AI算力资源需求持续攀升。",
        )
    ]
    pool = triggers + local_events + lianban_events

    assert {"AI", "算力"} <= set(workflow.causal_candidate_terms(pool))
    for theme in ("AI", "算力"):
        cluster = workflow._build_causal_cluster(theme, pool)
        assert cluster["score"] >= 60
        assert cluster["causal_quality"]["ready"] is True
        assert cluster["causal_quality"]["direction_quality"]["is_direction_level"] is True


def test_locked_template_renders_two_three_evidence_chains_without_residue(tmp_path):
    workflow = load_workflow()
    now = workflow._china_datetime()
    trading_dates = workflow.recent_trading_dates(now)
    trading_day = trading_dates[-1]
    if (
        trading_day == now.date()
        and now.weekday() < 5
        and (now.hour, now.minute) < (9, 30)
    ):
        trading_day = trading_dates[-2]
    date_text = workflow.cn_date(trading_day)
    target_date = trading_day.isoformat()

    pool = []
    clusters = []
    for theme in ("算力", "机器人"):
        events = [
            {
                "source": "新浪财经",
                "date": date_text,
                "published_at": now.isoformat(timespec="seconds"),
                "time_precision": "minute",
                "time_evidence": now.isoformat(timespec="minutes"),
                "window_type": workflow.EVENT_WINDOW_TYPE,
                "source_channel": "primary_public_readonly_fallback",
                "url": "https://finance.sina.com.cn/example",
                "event": f"{theme}产业需求持续增长，相关企业新增订单并扩大产能。",
            },
            {
                "source": "连板网·市场复盘",
                "date": date_text,
                "market_date": date_text,
                "published_at": now.isoformat(timespec="seconds"),
                "time_precision": "market_close",
                "time_evidence": target_date,
                "window_type": workflow.MARKET_VALIDATION_WINDOW_TYPE,
                "source_channel": "lianban_daily_market_cross_validation",
                "url": f"https://lianban.net/days/{target_date}.html",
                "event": (
                    f"{date_text}交易复盘显示，所属概念为{theme}，"
                    "该方向共有4只涨停，说明该方向出现板块级涨停扩散。"
                ),
            },
            {
                "source": "融资数据",
                "date": date_text,
                "market_date": date_text,
                "published_at": now.isoformat(timespec="seconds"),
                "time_precision": "market_close",
                "time_evidence": target_date,
                "window_type": workflow.MARKET_VALIDATION_WINDOW_TYPE,
                "source_channel": "local_participant_topic_aggregation",
                "url": "local-fund-review",
                "event": (
                    f"所属概念为{theme}，该方向参与者确认："
                    "机构与融资资金记录显示主力净流入，反映资金参与强度。"
                ),
            },
        ]
        pool.extend(events)
        cluster = workflow._build_causal_cluster(theme, events)
        assert cluster["causal_quality"]["ready"] is True
        assert len(workflow.select_evidence(cluster, 10)) == 3
        clusters.append(cluster)

    pool.extend([
        {
            "source": "东方财富全A行情逐股涨跌幅延迟节点·市场复盘",
            "source_key": "eastmoney.com",
            "date": date_text,
            "published_at": now.isoformat(timespec="seconds"),
            "time_precision": "market_close",
            "time_evidence": target_date,
            "window_type": workflow.MARKET_VALIDATION_WINDOW_TYPE,
            "event": f"{date_text}全市场上涨家数3860、下跌家数1217。",
            "url": "https://push2delay.eastmoney.com/example",
        },
        {
            "source": "新浪财经·市场复盘",
            "source_key": "sina.com.cn",
            "date": date_text,
            "published_at": now.isoformat(timespec="seconds"),
            "time_precision": "market_close",
            "time_evidence": target_date,
            "window_type": workflow.MARKET_VALIDATION_WINDOW_TYPE,
            "event": f"{date_text}全市场约1217只个股下跌。",
            "url": "https://finance.sina.com.cn/example",
        },
    ])

    report = tmp_path / "two-ready-chains.docx"
    workflow.build_report(clusters, report, pool)

    document = Document(report)
    assert len(document.paragraphs) == 43
    assert len(document.tables) == 25
    assert [row.cells[0].text for row in document.tables[8].rows] == [
        "明确结论",
        "逻辑分析",
        "情绪与资金",
        "产业链与观察对象",
        "次日观察",
        "失效风险",
    ]
    assert all(row.cells[1].text.strip() for row in document.tables[8].rows)
    assert [row.cells[0].text for row in document.tables[4].rows] == [
        "次日观察",
        "失效风险",
    ]
    report_text = "\n".join(
        paragraph.text for paragraph in document.paragraphs
    ) + "\n" + "\n".join(
        cell.text
        for table in document.tables
        for row in table.rows
        for cell in row.cells
    )
    assert "光学光电" not in report_text
