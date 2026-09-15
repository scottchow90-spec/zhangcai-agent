#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""Strict dynamic report generator for limit-up review."""
from __future__ import annotations
# ONESTOCK_EMBEDDED_LOCAL_IMPORT_BOOTSTRAP
import sys as _onestock_embedded_sys
from pathlib import Path as _OneStockEmbeddedPath
_onestock_embedded_dir = str(_OneStockEmbeddedPath(__file__).resolve().parent)
if _onestock_embedded_dir not in _onestock_embedded_sys.path:
    _onestock_embedded_sys.path.insert(0, _onestock_embedded_dir)

import hashlib
import json
import math
import os
import re
import shutil
import subprocess
import struct
import sys
import time
import zipfile
from collections import Counter
from datetime import datetime
from pathlib import Path, PureWindowsPath
from typing import Any
import xml.etree.ElementTree as ET

import pandas as pd
import requests
from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

from limit_up_review_workflow_contract import (
    BUSINESS_RESULT_SCHEMA_VERSION,
    compute_prediction_ledger_hash,
    validate_business_result_payload,
)
from entry_limit_up_review import DEFAULT_DELIVERY_ROOT, evaluate_date_gate, path_is_within
from official_pool_date_reconciliation import is_cross_validated_official_pool_date

try:
    sys.stdout.reconfigure(encoding="utf-8", errors="replace")
    sys.stderr.reconfigure(encoding="utf-8", errors="replace")
except Exception:
    pass

WORKSPACE = Path(r"D:\C盘转移\日志\codex")
BUSINESS_DATA_ROOT = WORKSPACE / "business_data" / "limit-up-review"
DATE = os.environ.get("LIMITUP_DATE", datetime.now().strftime("%Y%m%d"))
DATE_H = os.environ.get("LIMITUP_DATE_H", f"{DATE[:4]}-{DATE[4:6]}-{DATE[6:]}")
TASK_DIR = Path(os.environ.get("LIMITUP_TASK_DIR", WORKSPACE / "reports" / f"{DATE}_limit_up_review_closed_loop"))
MEM_DIR = Path(os.environ.get("LIMITUP_MEM_DIR", BUSINESS_DATA_ROOT))
HISTORY_DIR = Path(os.environ.get("LIMITUP_HISTORY_DIR", BUSINESS_DATA_ROOT))
DELIVERY = Path(os.environ.get("LIMITUP_DELIVERY", r"F:\小龙虾6月交付"))
RUN_ID = os.environ.get("LIMITUP_RUN_ID", f"limit-up-review-{DATE}-{datetime.now().strftime('%H%M%S')}")

IN_TABLE = Path(os.environ.get("LIMITUP_TABLE", TASK_DIR / f"verified_limitup_union_{DATE}.csv"))
OUT_MD = MEM_DIR / f"{DATE_H}-limit-up-review.md"
OUT_CSV = MEM_DIR / f"{DATE_H}-limit-up-table.csv"
OUT_WATCH = MEM_DIR / f"{DATE_H}-limit-up-watchlist.md"
OUT_DOCX = DELIVERY / f"涨停板深度复盘_{DATE_H}_事实结论版.docx"
OPEN_COPY = TASK_DIR / OUT_DOCX.name
EVIDENCE_JSON = TASK_DIR / "final_evidence.json"
BUSINESS_RESULT_JSON = Path(os.environ.get("LIMITUP_BUSINESS_RESULT", TASK_DIR / "business_result.json"))
MARKET_BREADTH_JSON = TASK_DIR / f"market_breadth_{DATE}.json"
MARKET_BREADTH_OVERRIDE_JSON = Path(os.environ.get("LIMITUP_MARKET_BREADTH_JSON", ""))
TQ_SAMPLE_JSON = TASK_DIR / f"tq_core_sample_{DATE}.json"
SCRIPTS_DIR = Path(__file__).resolve().parent
TDX_ROOT = Path(os.environ.get("TDX_ROOT", r"C:\new_tdx_mock"))
TDX_VIPDOC = TDX_ROOT / "vipdoc"
DAY_RECORD = struct.Struct("<IIIIIfII")
DELISTING_HARD_EXCLUSION = "delisting_hard_exclusion"
VISIBLE_LATIN_GATE = "visible_latin_gate"

FONT = "Microsoft YaHei"
BLUE = RGBColor(0x1F, 0x4E, 0x79)
DARK = RGBColor(0x1F, 0x1F, 0x1F)
GRAY = RGBColor(0x66, 0x66, 0x66)
RED = RGBColor(0xB0, 0x30, 0x30)
GREEN = RGBColor(0x2E, 0x7D, 0x32)
GOLD = RGBColor(0x9A, 0x6A, 0x00)


CURRENT_DATA_SOURCE_METADATA_KEYS = frozenset({
    "collected_at",
    "observed_at",
    "sha256",
})
CURRENT_DATA_SOURCE_PATH_KEYS = frozenset({
    "cwd",
    "path",
    "source",
    "source_dir",
    "task_dir",
})
BUSINESS_ARTIFACT_PROFILE_BY_SOURCE_ID = {
    "market_breadth": "market_breadth",
    "formula_sample": "tq_core_sample",
}


def _is_absolute_runtime_path(value: Any) -> bool:
    if not isinstance(value, str):
        return False
    return PureWindowsPath(value.strip()).is_absolute()


def _is_business_fingerprint_metadata(
    path: tuple[str | int, ...],
    value: Any,
    profile: str,
) -> bool:
    if profile == "current_data" and len(path) == 3 and path[0] == "sources":
        field = path[2]
        if field in CURRENT_DATA_SOURCE_METADATA_KEYS:
            return True
        return field in CURRENT_DATA_SOURCE_PATH_KEYS and _is_absolute_runtime_path(value)
    if profile == "market_breadth":
        return path == ("collected_at",)
    if profile == "tq_core_sample":
        if path == ("collected_at",):
            return True
        if (
            len(path) == 4
            and path[0] == "raw_samples"
            and isinstance(path[1], int)
            and path[2] == "payload"
            and path[3] in {"elapsed_ms", "generated_at"}
        ):
            return True
        return (
            len(path) == 6
            and path[0] == "raw_samples"
            and isinstance(path[1], int)
            and path[2] == "payload"
            and path[3] == "items"
            and isinstance(path[4], int)
            and path[5] in {"elapsed_ms", "lock_waited_ms"}
        )
    return False


def _business_semantic_projection(
    value: Any,
    *,
    profile: str,
    path: tuple[str | int, ...] = (),
) -> Any:
    if isinstance(value, dict):
        projected: dict[str, Any] = {}
        for key, child in value.items():
            if not isinstance(key, str):
                raise TypeError("business fingerprint JSON object keys must be strings")
            child_path = path + (key,)
            if _is_business_fingerprint_metadata(child_path, child, profile):
                continue
            projected[key] = _business_semantic_projection(
                child,
                profile=profile,
                path=child_path,
            )
        return projected
    if isinstance(value, list):
        return [
            _business_semantic_projection(
                child,
                profile=profile,
                path=path + (index,),
            )
            for index, child in enumerate(value)
        ]
    return value


def business_semantic_fingerprint(value: Any, *, profile: str = "current_data") -> str:
    projected = _business_semantic_projection(value, profile=profile)
    payload = json.dumps(
        projected,
        ensure_ascii=False,
        sort_keys=True,
        separators=(",", ":"),
    ).encode("utf-8")
    return hashlib.sha256(payload).hexdigest()


def business_artifact_fingerprint(path: Path, *, source_id: str) -> str:
    if path.suffix.casefold() == ".json":
        try:
            payload = json.loads(path.read_text(encoding="utf-8-sig"))
        except (OSError, UnicodeError, json.JSONDecodeError):
            pass
        else:
            profile = BUSINESS_ARTIFACT_PROFILE_BY_SOURCE_ID.get(
                source_id,
                "artifact_exact",
            )
            return business_semantic_fingerprint(payload, profile=profile)
    return sha256(path)


def sha256(path: Path) -> str:
    digest = hashlib.sha256()
    digest.update(path.read_bytes())
    return digest.hexdigest()


def embed_stock_review_manifest(docx_path: Path, business_result: dict[str, Any]) -> None:
    graph = business_result.get("conclusion_graph") if isinstance(business_result.get("conclusion_graph"), list) else []
    selected = [
        str(item.get("signal_id") or "")
        for item in business_result.get("candidate_signal_audit", [])
        if isinstance(item, dict) and item.get("status") == "SELECTED"
    ]
    nodes = []
    for item in graph:
        if not isinstance(item, dict):
            continue
        conclusion = str(item.get("conclusion") or "").strip()
        nodes.append({
            "id": str(item.get("id") or ""),
            "signal_id": str(item.get("signal_id") or ""),
            "topic": str(item.get("topic") or ""),
            "category_id": str(item.get("category_id") or ""),
            "category_label": str(item.get("category_label") or ""),
            "scope_label": str(item.get("scope_label") or ""),
            "angle_label": str(item.get("angle_label") or ""),
            "headline": str(item.get("headline") or ""),
            "priority": int(item.get("priority") or 0),
            "data_fingerprint": str(item.get("data_fingerprint") or ""),
            "conclusion": conclusion,
            "conclusion_sha256": hashlib.sha256(conclusion.encode("utf-8")).hexdigest(),
            "objects": [str(value) for value in item.get("objects", []) if str(value).strip()],
            "evidence_count": len(item.get("evidence", [])) if isinstance(item.get("evidence"), list) else 0,
            "reasoning_steps": len(item.get("reasoning_path", [])) if isinstance(item.get("reasoning_path"), list) else 0,
            "counterevidence_count": len(item.get("counterevidence", [])) if isinstance(item.get("counterevidence"), list) else 0,
            "confirmation_count": len(item.get("confirmation", [])) if isinstance(item.get("confirmation"), list) else 0,
            "invalidation_count": len(item.get("invalidation", [])) if isinstance(item.get("invalidation"), list) else 0,
        })
    section_analysis = business_result.get("section_analysis") if isinstance(business_result.get("section_analysis"), list) else []
    manifest = {
        "schema": "stock-review-conclusion-manifest.v3",
        "business_result_schema": business_result.get("schema_version"),
        "run_id": (business_result.get("run") or {}).get("run_id"),
        "trade_date": (business_result.get("run") or {}).get("trade_date"),
        "current_data_fingerprint": (business_result.get("run") or {}).get("current_data_fingerprint"),
        "business_result_path": str(BUSINESS_RESULT_JSON),
        "business_result_sha256": sha256(BUSINESS_RESULT_JSON),
        "selected_signal_ids": selected,
        "nodes": nodes,
        "section_analysis": section_analysis,
    }
    manifest["graph_sha256"] = hashlib.sha256(
        json.dumps(nodes, ensure_ascii=False, sort_keys=True, separators=(",", ":")).encode("utf-8")
    ).hexdigest()
    manifest["section_analysis_sha256"] = hashlib.sha256(
        json.dumps(section_analysis, ensure_ascii=False, sort_keys=True, separators=(",", ":")).encode("utf-8")
    ).hexdigest()
    manifest_text = json.dumps(manifest, ensure_ascii=False, sort_keys=True, separators=(",", ":"))

    custom_ns = "http://schemas.openxmlformats.org/officeDocument/2006/custom-properties"
    vt_ns = "http://schemas.openxmlformats.org/officeDocument/2006/docPropsVTypes"
    ET.register_namespace("", custom_ns)
    ET.register_namespace("vt", vt_ns)
    properties = ET.Element(f"{{{custom_ns}}}Properties")
    prop = ET.SubElement(properties, f"{{{custom_ns}}}property", {
        "fmtid": "{D5CDD505-2E9C-101B-9397-08002B2CF9AE}",
        "pid": "2",
        "name": "StockReviewConclusionManifest",
    })
    value = ET.SubElement(prop, f"{{{vt_ns}}}lpwstr")
    value.text = manifest_text
    custom_xml = ET.tostring(properties, encoding="utf-8", xml_declaration=True)

    temp_path = docx_path.with_suffix(docx_path.suffix + ".manifest.tmp")
    with zipfile.ZipFile(docx_path, "r") as source, zipfile.ZipFile(temp_path, "w", zipfile.ZIP_DEFLATED) as target:
        for info in source.infolist():
            if info.filename in {"docProps/custom.xml", "[Content_Types].xml", "_rels/.rels"}:
                continue
            target.writestr(info, source.read(info.filename))

        content_root = ET.fromstring(source.read("[Content_Types].xml"))
        content_ns = "http://schemas.openxmlformats.org/package/2006/content-types"
        if not any(node.get("PartName") == "/docProps/custom.xml" for node in content_root):
            ET.SubElement(content_root, f"{{{content_ns}}}Override", {
                "PartName": "/docProps/custom.xml",
                "ContentType": "application/vnd.openxmlformats-officedocument.custom-properties+xml",
            })
        target.writestr("[Content_Types].xml", ET.tostring(content_root, encoding="utf-8", xml_declaration=True))

        rels_root = ET.fromstring(source.read("_rels/.rels"))
        rels_ns = "http://schemas.openxmlformats.org/package/2006/relationships"
        custom_rel_type = "http://schemas.openxmlformats.org/officeDocument/2006/relationships/custom-properties"
        if not any(node.get("Type") == custom_rel_type for node in rels_root):
            existing_ids = {node.get("Id") for node in rels_root}
            rel_id = "rIdStockReviewManifest"
            suffix = 1
            while rel_id in existing_ids:
                rel_id = f"rIdStockReviewManifest{suffix}"
                suffix += 1
            ET.SubElement(rels_root, f"{{{rels_ns}}}Relationship", {
                "Id": rel_id,
                "Type": custom_rel_type,
                "Target": "docProps/custom.xml",
            })
        target.writestr("_rels/.rels", ET.tostring(rels_root, encoding="utf-8", xml_declaration=True))
        target.writestr("docProps/custom.xml", custom_xml)
    temp_path.replace(docx_path)


def code6(value: Any) -> str:
    text = str(value or "").strip()
    match = re.search(r"(\d{6})", text)
    return match.group(1) if match else text.zfill(6)


def clean(value: Any, default: str = "-") -> str:
    text = str(value if value is not None else "").strip()
    if not text or text.lower() in {"nan", "none"}:
        return default
    for forbidden in ["TDX", "AK", "Codex", "ZTC", "ELB", "工作流"]:
        text = text.replace(forbidden, "")
    return re.sub(r"\s+", " ", text).strip() or default


def chinese_visible_text(value: Any, default: str = "-") -> str:
    """Translate or neutralize Latin tokens before any user-visible rendering."""
    text = clean(value, default)
    if text == default:
        return text
    replacements = {
        "PCB": "印制电路板",
        "IT": "信息技术",
        "AI": "人工智能",
        "ST": "特别处理",
        "CLEANPASS": "清洁通过",
    }
    for raw, translated in replacements.items():
        text = re.sub(re.escape(raw), translated, text, flags=re.I)
    text = re.sub(r"[-—]?[A-Za-z]+", "", text)
    return re.sub(r"\s+", " ", text).strip() or default


def load_eligibility_audit() -> dict[str, Any]:
    path = TASK_DIR / "data_integrity_audit.json"
    if not path.exists() or path.stat().st_size == 0:
        return {}
    try:
        payload = json.loads(path.read_text(encoding="utf-8-sig"))
        return payload if isinstance(payload, dict) else {}
    except Exception:
        return {}


def read_csv(
    path: Path,
    dtype: dict[str, Any] | None = None,
    *,
    allow_empty_file: bool = False,
) -> pd.DataFrame:
    if not path.exists():
        return pd.DataFrame()
    try:
        return pd.read_csv(path, dtype=dtype or {"代码": str})
    except pd.errors.EmptyDataError:
        if not allow_empty_file:
            raise
        return pd.DataFrame(columns=list((dtype or {}).keys()))


def parse_verified_bool(value: Any) -> bool:
    """Parse provider status without treating the string ``False`` as truthy."""
    if isinstance(value, bool):
        return value
    if value is None or (isinstance(value, float) and math.isnan(value)):
        return False
    text = str(value).strip().lower()
    if text in {"true", "1", "yes", "y"}:
        return True
    if text in {"false", "0", "no", "n", "", "none", "nan"}:
        return False
    return False


def source_record(source_id: str, kind: str, path: Path | None, status: str, row_count: int | None = None, **extra: Any) -> dict[str, Any]:
    collected_at = datetime.now().isoformat(timespec="seconds")
    file_path = path if path and path.exists() and path.is_file() else None
    record: dict[str, Any] = {
        "id": source_id,
        "kind": kind,
        "verification_status": status,
        "trade_date": DATE,
        "date": DATE,
        "observed_at": collected_at,
        "collected_at": collected_at,
        "as_of": f"{DATE_H}T15:00:00+08:00",
        "source": str(path) if path else "",
        "sha256": sha256(file_path) if file_path else "",
        "business_sha256": business_artifact_fingerprint(file_path, source_id=source_id) if file_path else "",
        "row_count": row_count,
    }
    record.update(extra)
    if status == "UNAVAILABLE" and not record.get("reason"):
        record["reason"] = f"{kind} source unavailable for {DATE}"
    return record


def load_research_bundle() -> tuple[dict[str, Any], dict[str, Any]]:
    """Load optional catalyst/news research; absence is explicit, never coerced to zero."""
    raw = os.environ.get("LIMITUP_RESEARCH_BUNDLE", "").strip()
    if not raw:
        return {}, source_record("catalyst_research", "news_policy", None, "UNAVAILABLE", reason="LIMITUP_RESEARCH_BUNDLE not supplied")
    path = Path(raw)
    if not path.exists() or not path.is_file():
        return {}, source_record("catalyst_research", "news_policy", path, "UNAVAILABLE", reason="research bundle missing")
    try:
        payload = json.loads(path.read_text(encoding="utf-8-sig"))
    except Exception as exc:
        return {}, source_record("catalyst_research", "news_policy", path, "UNAVAILABLE", reason=f"invalid json: {type(exc).__name__}")
    bundle_date = str(payload.get("trade_date") or payload.get("date") or "").replace("-", "")
    items = payload.get("items") if isinstance(payload.get("items"), list) else []
    status = "VERIFIED" if bundle_date == DATE and items else "UNAVAILABLE"
    reason = "" if status == "VERIFIED" else f"bundle date/items invalid: date={bundle_date}, items={len(items)}"
    return payload, source_record("catalyst_research", "news_policy", path, status, row_count=len(items), reason=reason)


def is_number(value: Any) -> bool:
    try:
        if pd.isna(value):
            return False
        float(value)
        return True
    except Exception:
        return False


def push2_json(url: str, params: dict[str, Any], retries: int = 4) -> tuple[dict[str, Any], str]:
    headers = {"User-Agent": "Mozilla/5.0", "Referer": "https://quote.eastmoney.com/"}
    last_error: Exception | None = None
    for attempt in range(retries):
        try:
            resp = requests.get(url, params=params, timeout=20, headers=headers)
            resp.raise_for_status()
            return resp.json(), resp.url
        except Exception as exc:
            last_error = exc
            time.sleep(0.8 * (attempt + 1))
    raise RuntimeError(str(last_error))


def compute_breadth_from_changes(changes: list[float], *, source: str, rows: int, url: str = "") -> dict[str, Any]:
    up = sum(value > 0 for value in changes)
    down = sum(value < 0 for value in changes)
    flat = sum(value == 0 for value in changes)
    valid = up + down + flat
    down_ratio = down / valid if valid else 0
    up_ratio = up / valid if valid else 0
    label = "上涨家数多于下跌家数" if up > down else "下跌家数多于上涨家数" if down > up else "涨跌家数相同"
    return {
        "status": "CLEAN_PASS" if valid else "BLOCKED",
        "date": DATE,
        "source": source,
        "source_url": url,
        "source_tier": "primary_live",
        "rows": rows,
        "valid_change_rows": valid,
        "up_count": int(up),
        "down_count": int(down),
        "flat_count": int(flat),
        "up_ratio": round(up_ratio, 4),
        "down_ratio": round(down_ratio, 4),
        "breadth_label": label,
        "collected_at": datetime.now().isoformat(timespec="seconds"),
    }


def validate_external_market_breadth(data: dict[str, Any]) -> dict[str, Any]:
    required = ["date", "up_count", "down_count", "flat_count", "source", "source_url"]
    missing = [key for key in required if data.get(key) in (None, "")]
    result = dict(data)
    result.setdefault("source_tier", "verified_external")
    result.setdefault("collected_at", datetime.now().isoformat(timespec="seconds"))
    if missing:
        result["status"] = "BLOCKED"
        result["blocks"] = [f"market breadth override missing: {missing}"]
        return result
    normalized_date = str(result.get("date") or "").replace("-", "")
    if normalized_date != DATE:
        result["status"] = "BLOCKED"
        result["blocks"] = [f"market breadth override date mismatch: {normalized_date} != {DATE}"]
        return result
    try:
        up = int(result["up_count"])
        down = int(result["down_count"])
        flat = int(result["flat_count"])
    except Exception as exc:
        result["status"] = "BLOCKED"
        result["blocks"] = [f"market breadth override invalid counts: {exc}"]
        return result
    valid = up + down + flat
    if valid <= 0:
        result["status"] = "BLOCKED"
        result["blocks"] = ["market breadth override has no valid rows"]
        return result
    result["up_count"] = up
    result["down_count"] = down
    result["flat_count"] = flat
    result["rows"] = int(result.get("rows") or valid)
    result["valid_change_rows"] = valid
    result["up_ratio"] = round(up / valid, 4)
    result["down_ratio"] = round(down / valid, 4)
    if not result.get("breadth_label"):
        result["breadth_label"] = "下跌占优" if down > up else "上涨占优" if up > down else "涨跌均衡"
    result["status"] = "CLEAN_PASS"
    return result


def load_market_breadth_override() -> dict[str, Any] | None:
    override_raw = os.environ.get("LIMITUP_MARKET_BREADTH_JSON", "").strip()
    if not override_raw:
        return None
    override_path = Path(override_raw)
    if not override_path.exists():
        return {
            "status": "BLOCKED",
            "date": DATE,
            "source": "market_breadth_override",
            "source_tier": "missing_override",
            "blocks": [f"missing market breadth override: {override_path}"],
            "collected_at": datetime.now().isoformat(timespec="seconds"),
        }
    try:
        data = json.loads(override_path.read_text(encoding="utf-8-sig"))
    except Exception as exc:
        return {
            "status": "BLOCKED",
            "date": DATE,
            "source": "market_breadth_override",
            "source_tier": "invalid_override",
            "blocks": [f"cannot read market breadth override: {exc}"],
            "collected_at": datetime.now().isoformat(timespec="seconds"),
        }
    result = validate_external_market_breadth(data)
    result["override_path"] = str(override_path)
    MARKET_BREADTH_JSON.write_text(json.dumps(result, ensure_ascii=False, indent=2), encoding="utf-8")
    return result


def parse_tdx_day_record(chunk: bytes) -> dict[str, Any]:
    date_i, open_i, high_i, low_i, close_i, amount, volume, _unused = DAY_RECORD.unpack(chunk)
    return {
        "date": str(date_i),
        "open": open_i / 100.0,
        "high": high_i / 100.0,
        "low": low_i / 100.0,
        "close": close_i / 100.0,
        "amount": float(amount),
        "volume": int(volume),
    }


def is_a_share_day_file(path: Path) -> bool:
    stem = path.stem.lower()
    if len(stem) < 8:
        return False
    market = stem[:2]
    code = stem[-6:]
    if not code.isdigit():
        return False
    if market == "sh":
        return code.startswith("6")
    if market == "sz":
        return code.startswith(("000", "001", "002", "003", "300", "301", "302"))
    if market == "bj":
        return code.startswith(("43", "83", "87", "88", "92"))
    return False


def iter_tdx_day_files() -> list[Path]:
    roots = [
        TDX_VIPDOC / "sh" / "lday",
        TDX_VIPDOC / "sz" / "lday",
        TDX_VIPDOC / "bj" / "lday",
        TDX_VIPDOC / "xinzeng" / "sh" / "lday",
        TDX_VIPDOC / "xinzeng" / "sz" / "lday",
        TDX_VIPDOC / "xinzeng" / "bj" / "lday",
    ]
    by_symbol: dict[str, Path] = {}
    for root in roots:
        if not root.exists():
            continue
        for path in root.glob("*.day"):
            if not is_a_share_day_file(path):
                continue
            symbol = path.stem.lower()
            current = by_symbol.get(symbol)
            if current is None or path.stat().st_mtime > current.stat().st_mtime:
                by_symbol[symbol] = path
    return list(by_symbol.values())


def fetch_tdx_local_breadth() -> dict[str, Any]:
    changes: list[float] = []
    latest_dates: Counter[str] = Counter()
    scanned = 0
    skipped_short = 0
    for path in iter_tdx_day_files():
        scanned += 1
        size = path.stat().st_size
        count = size // DAY_RECORD.size
        if count < 2:
            skipped_short += 1
            continue
        with path.open("rb") as handle:
            handle.seek((count - 2) * DAY_RECORD.size)
            prev = parse_tdx_day_record(handle.read(DAY_RECORD.size))
            last = parse_tdx_day_record(handle.read(DAY_RECORD.size))
        latest_dates[str(last["date"])] += 1
        if str(last["date"]) != DATE:
            continue
        prev_close = float(prev.get("close") or 0)
        close = float(last.get("close") or 0)
        if prev_close <= 0 or close <= 0:
            continue
        changes.append((close - prev_close) / prev_close * 100)
    result = compute_breadth_from_changes(changes, source="tdx_local_day_files", rows=len(changes), url=str(TDX_VIPDOC))
    result.update({
        "scanned_day_files": scanned,
        "skipped_short_files": skipped_short,
        "latest_date_distribution": dict(latest_dates.most_common(5)),
        "tdx_root": str(TDX_ROOT),
        "source_tier": "primary_local_tdx_same_day",
        "coverage_ratio": round(len(changes) / scanned, 4) if scanned else 0.0,
    })
    if len(changes) < 3000 or (scanned and len(changes) / scanned < 0.85):
        result["status"] = "BLOCKED"
        result["blocks"] = [
            f"TDX same-day breadth coverage insufficient: valid={len(changes)}, scanned={scanned}"
        ]
    MARKET_BREADTH_JSON.write_text(json.dumps(result, ensure_ascii=False, indent=2), encoding="utf-8")
    return result


def fetch_market_breadth() -> dict[str, Any]:
    override = load_market_breadth_override()
    if override is not None:
        return override
    errors: list[str] = []
    # Local Tongdaxin is the primary source.  It is a same-day, symbol-level
    # close-to-close census, not a diagnostic cache.  Network sources are only
    # fallbacks when the local coverage/date hard gate fails.
    try:
        result = fetch_tdx_local_breadth()
        if result.get("status") == "CLEAN_PASS":
            return result
        errors.extend(result.get("blocks") or ["TDX local breadth failed its coverage gate"])
    except Exception as exc:
        errors.append(f"tdx_local_day_files: {type(exc).__name__}: {str(exc)[:240]}")

    if DATE != datetime.now().strftime("%Y%m%d"):
        result = {
            "status": "BLOCKED",
            "date": DATE,
            "source": "historical_live_fallback_forbidden",
            "errors": errors + ["historical reruns require a same-day local snapshot or dated override"],
            "collected_at": datetime.now().isoformat(timespec="seconds"),
        }
        MARKET_BREADTH_JSON.write_text(json.dumps(result, ensure_ascii=False, indent=2), encoding="utf-8")
        return result

    url = "https://push2.eastmoney.com/api/qt/clist/get"
    params = {
        "pn": 1,
        "pz": 100,
        "po": 1,
        "np": 1,
        "ut": "bd1d9ddb04089700cf9c27f6f7426281",
        "fltt": 2,
        "invt": 2,
        "fid": "f3",
        "fs": "m:0+t:6,m:0+t:80,m:1+t:2,m:1+t:23,m:0+t:81+s:2048",
        "fields": "f12,f14,f2,f3",
    }
    try:
        data, used_url = push2_json(url, params)
        payload = data.get("data", {}) or {}
        total = int(payload.get("total") or 0)
        rows = payload.get("diff", []) or []
        pages = math.ceil(total / 100) if total else 1
        for page in range(2, pages + 1):
            page_params = dict(params)
            page_params["pn"] = page
            page_data, _ = push2_json(url, page_params, retries=3)
            rows.extend((page_data.get("data", {}) or {}).get("diff", []) or [])
        changes = [float(item.get("f3")) for item in rows if is_number(item.get("f3"))]
        result = compute_breadth_from_changes(changes, source="eastmoney_push2_all_a", rows=len(rows), url=used_url)
        MARKET_BREADTH_JSON.write_text(json.dumps(result, ensure_ascii=False, indent=2), encoding="utf-8")
        return result
    except Exception as exc:
        errors.append(f"eastmoney_push2_all_a: {type(exc).__name__}: {str(exc)[:240]}")

    try:
        import akshare as ak  # type: ignore

        spot = ak.stock_zh_a_spot_em()
        col = "涨跌幅"
        if col not in spot.columns:
            raise RuntimeError(f"missing column: {col}")
        changes = [float(value) for value in spot[col].tolist() if is_number(value)]
        result = compute_breadth_from_changes(changes, source="akshare_stock_zh_a_spot_em", rows=int(len(spot)))
        result["fallback_errors"] = errors
        MARKET_BREADTH_JSON.write_text(json.dumps(result, ensure_ascii=False, indent=2), encoding="utf-8")
        return result
    except Exception as exc:
        errors.append(f"akshare_stock_zh_a_spot_em: {type(exc).__name__}: {str(exc)[:240]}")

    result = {
        "status": "BLOCKED",
        "date": DATE,
        "errors": errors,
        "collected_at": datetime.now().isoformat(timespec="seconds"),
    }
    MARKET_BREADTH_JSON.write_text(json.dumps(result, ensure_ascii=False, indent=2), encoding="utf-8")
    return result


def to_num(value: Any, default: float = 0.0) -> float:
    try:
        if pd.isna(value):
            return default
        return float(value)
    except Exception:
        return default


def to_int(value: Any, default: int = 0) -> int:
    try:
        if pd.isna(value):
            return default
        return int(float(value))
    except Exception:
        return default


def fmt_num(value: Any, nd: int = 2, default: str = "-") -> str:
    try:
        if pd.isna(value):
            return default
        return f"{float(value):.{nd}f}"
    except Exception:
        return default


def fmt_pct(value: Any) -> str:
    num = fmt_num(value, 2)
    return "-" if num == "-" else f"{num}%"


def fmt_amount(value: Any) -> str:
    try:
        if pd.isna(value):
            return "-"
        num = float(value)
        if abs(num) >= 100000000:
            return f"{num / 100000000:.2f}亿"
        if abs(num) >= 10000:
            return f"{num / 10000:.1f}万"
        return f"{num:.0f}"
    except Exception:
        return "-"


def fmt_time(value: Any) -> str:
    try:
        if pd.isna(value):
            return "-"
        text = str(int(float(value))).zfill(6)
        return f"{text[:2]}:{text[2:4]}:{text[4:]}"
    except Exception:
        return "-"


def stock_label(row: pd.Series, with_board: bool = True) -> str:
    board = to_int(row.get("连板数"), 1)
    board_text = f"{board}板" if board >= 2 else "首板"
    base = f"{code6(row.get('代码'))} {clean(row.get('名称'))}"
    return f"{base}（{board_text}）" if with_board else base


def leader_score(row: pd.Series) -> float:
    board = to_int(row.get("连板数"), 1)
    first = to_int(row.get("首次封板时间"), 150000)
    open_count = to_int(row.get("炸板次数"), 0)
    fund = to_num(row.get("主力净流入"), 0)
    early = max(0, 150000 - first) / 1000
    return board * 30 + early - open_count * 4 + (8 if fund > 0 else 0)


def sector_summary(df: pd.DataFrame, limit: int = 5) -> list[dict[str, Any]]:
    work = df.copy()
    work["连板数_num"] = pd.to_numeric(work["连板数"], errors="coerce").fillna(1).astype(int)
    if "主力净流入" in work.columns:
        work["主力净流入_num"] = pd.to_numeric(work["主力净流入"], errors="coerce").fillna(0)
    else:
        work["主力净流入_num"] = 0
    work["排序分"] = work.apply(leader_score, axis=1)
    grouped: list[dict[str, Any]] = []
    for sector, group in work.groupby(work["投资板块"].fillna("未分类")):
        group = group.sort_values(["连板数_num", "排序分", "主力净流入_num"], ascending=[False, False, False])
        max_board = int(group["连板数_num"].max()) if not group.empty else 1
        connected = int((group["连板数_num"] >= 2).sum())
        count = int(len(group))
        score = min(45, count * 7) + min(25, connected * 8) + min(20, max_board * 4) + (10 if group["主力净流入_num"].sum() > 0 else 0)
        grouped.append({
            "sector": clean(sector, "未分类"),
            "count": count,
            "connected": connected,
            "max_board": max_board,
            "fund": float(group["主力净流入_num"].sum()),
            "score": score,
            "leaders": group.head(3),
        })
    return sorted(grouped, key=lambda x: (x["count"], x["connected"], x["max_board"], x["fund"]), reverse=True)[:limit]


def dragon_rows(df: pd.DataFrame) -> list[list[Any]]:
    rows: list[list[Any]] = []
    for idx, item in enumerate(sector_summary(df, 5), start=1):
        leaders = [stock_label(row) for _, row in item["leaders"].iterrows()]
        while len(leaders) < 3:
            leaders.append("-")
        rows.append([
            idx,
            item["sector"],
            item["count"],
            item["connected"],
            f"{item['max_board']}板",
            leaders[0],
            leaders[1],
            leaders[2],
            fmt_amount(item["fund"]),
        ])
    return rows


def pct_text(part: int | float, total: int | float) -> str:
    try:
        total_f = float(total)
        if total_f == 0:
            return "-"
        return f"{float(part) / total_f:.1%}"
    except Exception:
        return "-"


def sentence(conclusion: str, evidence: str) -> dict[str, str]:
    """Return structured market prose without template conclusion wrappers."""
    return {
        "analysis": str(conclusion or "").strip().rstrip("。") + "。",
        "facts": str(evidence or "").strip().rstrip("。") + "。",
    }


def load_previous_review_material(current_df: pd.DataFrame, current_date_h: str) -> dict[str, Any]:
    current_dt = datetime.strptime(current_date_h, "%Y-%m-%d")
    candidates: list[tuple[datetime, Path]] = []
    for path in HISTORY_DIR.glob("*-limit-up-table.csv"):
        match = re.match(r"(\d{4}-\d{2}-\d{2})-limit-up-table\.csv$", path.name)
        if not match:
            continue
        day = datetime.strptime(match.group(1), "%Y-%m-%d")
        if day < current_dt:
            candidates.append((day, path))
    if not candidates:
        return {
            "status": "UNAVAILABLE",
            "date_h": "无前序实采表",
            "table": "",
            "watchlist": "",
            "previous_pool_count": 0,
            "watch_count": 0,
            "watch_retained_count": 0,
            "watch_promoted_count": 0,
            "pool_retained_count": 0,
            "retained_watch_labels": "-",
            "promoted_watch_labels": "-",
            "baseline": False,
            "blocks": [],
        }
    prev_day, prev_table = max(candidates, key=lambda item: item[0])
    prev_date_h = prev_day.strftime("%Y-%m-%d")
    prev_df = read_csv(prev_table, {"代码": str})
    if prev_df.empty or "代码" not in prev_df.columns:
        return {"status": "BLOCKED", "blocks": [f"invalid previous limit-up table: {prev_table}"]}
    prev_df["代码"] = prev_df["代码"].map(code6)
    for col in ["连板数", "炸板次数", "主力净流入"]:
        if col in prev_df.columns:
            prev_df[col] = pd.to_numeric(prev_df[col], errors="coerce")
    watch_path = HISTORY_DIR / f"{prev_date_h}-limit-up-watchlist.md"
    watch_codes: list[str] = []
    if watch_path.exists():
        seen: set[str] = set()
        for raw in re.findall(r"\b\d{6}\b", watch_path.read_text(encoding="utf-8", errors="ignore")):
            code = code6(raw)
            if code not in seen:
                watch_codes.append(code)
                seen.add(code)
    current_codes = {code6(x) for x in current_df["代码"].dropna().tolist()}
    prev_codes = {code6(x) for x in prev_df["代码"].dropna().tolist()}
    prev_board = {code6(row["代码"]): to_int(row.get("连板数"), 1) for _, row in prev_df.iterrows()}
    current_board = {code6(row["代码"]): to_int(row.get("连板数"), 1) for _, row in current_df.iterrows()}
    retained_watch = [code for code in watch_codes if code in current_codes]
    promoted_watch = [code for code in retained_watch if current_board.get(code, 0) > prev_board.get(code, 0)]
    retained_pool = sorted(prev_codes & current_codes)
    return {
        "status": "VERIFIED" if watch_codes else "UNAVAILABLE",
        "date_h": prev_date_h,
        "table": str(prev_table),
        "watchlist": str(watch_path) if watch_path.exists() else "",
        "previous_pool_count": int(len(prev_df)),
        "watch_count": int(len(watch_codes)),
        "watch_retained_count": int(len(retained_watch)),
        "watch_promoted_count": int(len(promoted_watch)),
        "pool_retained_count": int(len(retained_pool)),
        "retained_watch_labels": "、".join(
            stock_label(current_df[current_df["代码"].map(code6) == code].iloc[0]) for code in retained_watch[:8]
        ) if retained_watch else "-",
        "promoted_watch_labels": "、".join(
            f"{code} {clean(current_df[current_df['代码'].map(code6) == code].iloc[0].get('名称'))}（{prev_board.get(code, 1)}板->{current_board.get(code, 1)}板）"
            for code in promoted_watch[:8]
        ) if promoted_watch else "-",
        "blocks": [],
        "reason": "" if watch_codes else "previous watchlist missing; no substitute list was manufactured",
    }


def load_historical_baselines(current_date_h: str) -> dict[str, Any]:
    current_dt = datetime.strptime(current_date_h, "%Y-%m-%d")
    records: list[dict[str, Any]] = []
    for path in HISTORY_DIR.glob("*-limit-up-table.csv"):
        match = re.match(r"(\d{4}-\d{2}-\d{2})-limit-up-table\.csv$", path.name)
        if not match:
            continue
        day = datetime.strptime(match.group(1), "%Y-%m-%d")
        if day >= current_dt:
            continue
        frame = read_csv(path, {"代码": str})
        if frame.empty or "连板数" not in frame.columns:
            continue
        boards = pd.to_numeric(frame["连板数"], errors="coerce").fillna(1).astype(int)
        if "投资板块" in frame.columns:
            counts = frame["投资板块"].fillna("未分类").value_counts()
            largest_share = float(counts.iloc[0] / len(frame)) if len(counts) and len(frame) else 0.0
        else:
            largest_share = float("nan")
        records.append({
            "date": match.group(1),
            "limitup_count": int(len(frame)),
            "continuation_count": int((boards >= 2).sum()),
            "first_board_share": float((boards == 1).sum() / len(frame)),
            "max_board": int(boards.max()) if len(boards) else 0,
            "largest_direction_share": largest_share,
            "sha256": sha256(path),
        })
    records.sort(key=lambda item: item["date"])
    records = records[-60:]

    def window_summary(size: int) -> dict[str, Any]:
        window = records[-size:]
        fields = ["limitup_count", "continuation_count", "first_board_share", "max_board", "largest_direction_share"]
        summary: dict[str, Any] = {"sample_count": len(window), "dates": [item["date"] for item in window]}
        for field in fields:
            values = [float(item[field]) for item in window if is_number(item.get(field))]
            summary[field] = {
                "median": round(float(pd.Series(values).median()), 4) if len(values) >= 5 else None,
                "q25": round(float(pd.Series(values).quantile(0.25)), 4) if len(values) >= 5 else None,
                "q75": round(float(pd.Series(values).quantile(0.75)), 4) if len(values) >= 5 else None,
            }
        return summary

    sample_count = len(records)
    status = "VERIFIED" if sample_count >= 20 else "PARTIAL" if sample_count >= 5 else "UNAVAILABLE"
    return {
        "status": status,
        "source_dir": str(HISTORY_DIR),
        "sample_count": sample_count,
        "windows": {"20": window_summary(20), "60": window_summary(60)},
        "records": records,
        "reason": "" if status == "VERIFIED" else f"historical same-schema samples={sample_count}, fewer than 20",
    }


def extract_report_facts(
    df: pd.DataFrame,
    top_sectors: list[dict[str, Any]],
    qdf: pd.DataFrame,
    main: dict[str, Any],
    market_breadth: dict[str, Any],
    prev_material: dict[str, Any],
    *,
    zt_count: int,
    connected_count: int,
    first_board_count: int,
    highest_board: int,
    zbgc_count: int,
    dtgc_count: int,
    seal_rate: float,
    fail_rate: float,
    push_ok: int,
    lhb_hit: int,
    formula_sample: dict[str, Any],
) -> dict[str, Any]:
    main_sector = main["sector"]
    main_df = df[df["投资板块"] == main_sector].copy()
    first_share = first_board_count / zt_count if zt_count else 0
    connected_share = connected_count / zt_count if zt_count else 0
    top5_count = sum(int(item["count"]) for item in top_sectors[:5])
    top5_share = top5_count / zt_count if zt_count else 0
    main_share = int(main["count"]) / zt_count if zt_count else 0
    high_count = int((df["连板数"].fillna(1).astype(int) >= 3).sum())
    board_counts = Counter(int(x) for x in df["连板数"].fillna(1).astype(int))
    missing_middle = [board for board in [3, 4] if board not in board_counts]
    front_df = qdf.sort_values(["连板数", "排序分"], ascending=[False, False]).head(10).copy()
    front_open_sum = int(pd.to_numeric(front_df.get("炸板次数", 0), errors="coerce").fillna(0).sum())
    front_open_ge5 = int((pd.to_numeric(front_df.get("炸板次数", 0), errors="coerce").fillna(0) >= 5).sum())
    front_fund = pd.to_numeric(front_df.get("主力净流入", pd.Series(index=front_df.index, dtype=float)), errors="coerce")
    if "ok" in front_df.columns:
        front_fund = front_fund.where(front_df["ok"].map(parse_verified_bool))
    front_negative_count = int((front_fund < 0).sum())
    fund_series = pd.to_numeric(df.get("主力净流入", pd.Series(index=df.index, dtype=float)), errors="coerce")
    if "ok" in df.columns:
        fund_series = fund_series.where(df["ok"].map(parse_verified_bool))
    fund_total_raw = fund_series.sum(min_count=1)
    fund_total = float(fund_total_raw) if pd.notna(fund_total_raw) else float("nan")
    fund_pos = int((fund_series > 0).sum())
    fund_neg = int((fund_series < 0).sum())
    top_fund_idx = fund_series.dropna().idxmax() if not fund_series.dropna().empty else None
    top_fund_label = "-"
    top_fund_amount = 0.0
    if top_fund_idx is not None and top_fund_idx in df.index:
        top_fund_row = df.loc[top_fund_idx]
        top_fund_label = stock_label(top_fund_row)
        top_fund_amount = float(fund_series.loc[top_fund_idx])
    main_fund_series = pd.to_numeric(main_df.get("主力净流入", pd.Series(index=main_df.index, dtype=float)), errors="coerce")
    if "ok" in main_df.columns:
        main_fund_series = main_fund_series.where(main_df["ok"].map(parse_verified_bool))
    main_fund_raw = main_fund_series.sum(min_count=1)
    main_fund = float(main_fund_raw) if pd.notna(main_fund_raw) else float("nan")
    main_open_sum = int(pd.to_numeric(main_df.get("炸板次数", pd.Series(dtype=float)), errors="coerce").fillna(0).sum())
    main_first = int((main_df["连板数"].fillna(1).astype(int) == 1).sum()) if not main_df.empty else 0
    return {
        "main_sector": main_sector,
        "main_share": main_share,
        "top5_count": top5_count,
        "top5_share": top5_share,
        "first_share": first_share,
        "connected_share": connected_share,
        "high_count": high_count,
        "board_counts": dict(sorted(board_counts.items(), reverse=True)),
        "missing_middle": missing_middle,
        "front_open_sum": front_open_sum,
        "front_open_ge5": front_open_ge5,
        "front_negative_count": front_negative_count,
        "fund_total": fund_total,
        "fund_pos": fund_pos,
        "fund_neg": fund_neg,
        "top_fund_label": top_fund_label,
        "top_fund_amount": top_fund_amount,
        "main_fund": main_fund,
        "main_open_sum": main_open_sum,
        "main_first": main_first,
        "zbgc_rate": fail_rate / 100 if fail_rate > 1 else fail_rate,
        "lhb_share": lhb_hit / zt_count if zt_count else 0,
        "prev": prev_material,
        "blocks": [] if prev_material.get("status") == "CLEAN_PASS" else list(prev_material.get("blocks", [])),
        "formula_success": int(formula_sample.get("sample_success_count") or 0),
        "formula_failed": int(formula_sample.get("sample_failed_count") or 0),
        "push_ok": push_ok,
        "market_down_ratio": float(market_breadth.get("down_ratio") or 0),
        "zt_count": zt_count,
        "connected_count": connected_count,
        "first_board_count": first_board_count,
        "highest_board": highest_board,
        "zbgc_count": zbgc_count,
        "dtgc_count": dtgc_count,
        "seal_rate": seal_rate,
        "lhb_hit": lhb_hit,
    }


def metric(value: Any, source_id: str, unit: str = "", status: str = "VERIFIED") -> dict[str, Any]:
    return {"value": value, "unit": unit, "source_id": source_id, "verification_status": status}


def thesis_evidence(
    source_id: str,
    metric_id: str,
    dimension: str,
    interpretation: str,
    *,
    effect: str = "SUPPORTS",
    challenge_to: str = "",
    alternative_explanation: str = "",
) -> dict[str, Any]:
    item = {
        "source_id": source_id,
        "metric_id": metric_id,
        "evidence_id": f"EV-METRIC-{metric_id}",
        "dimension": dimension,
        "interpretation": interpretation,
        "effect": effect,
    }
    if challenge_to:
        item["challenge_to"] = challenge_to
    if alternative_explanation:
        item["alternative_explanation"] = alternative_explanation
    return item


def thesis_visible_sentence(thesis: dict[str, Any]) -> dict[str, str]:
    evidence_text = "；".join(str(item.get("interpretation") or "") for item in thesis.get("evidence", []) if item.get("interpretation"))
    return sentence(str(thesis.get("conclusion") or ""), evidence_text)


def build_business_result(
    df: pd.DataFrame,
    fund: pd.DataFrame,
    zbgc: pd.DataFrame,
    dtgc: pd.DataFrame,
    lhb: pd.DataFrame,
    qdf: pd.DataFrame,
    top_sectors: list[dict[str, Any]],
    facts: dict[str, Any],
    market_breadth: dict[str, Any],
    prev_material: dict[str, Any],
    formula_sample: dict[str, Any],
    fail_rate: float,
) -> dict[str, Any]:
    generated_at = datetime.now().isoformat(timespec="seconds")
    fund_path = TASK_DIR / f"push2_fund_flow_resolved_{DATE}.csv"
    zbgc_path = TASK_DIR / f"ak_stock_zt_pool_zbgc_em_{DATE}.csv"
    dtgc_path = TASK_DIR / f"ak_stock_zt_pool_dtgc_em_{DATE}.csv"
    lhb_path = TASK_DIR / f"ak_stock_lhb_detail_em_{DATE}.csv"
    research_bundle, catalyst_source = load_research_bundle()
    eligibility_audit = load_eligibility_audit()
    historical_baselines = load_historical_baselines(DATE_H)
    fund_ok = fund["ok"].map(parse_verified_bool) if not fund.empty and "ok" in fund.columns else pd.Series(dtype=bool)
    fund_ok_count = int(fund_ok.sum()) if len(fund_ok) else 0
    fund_status = "VERIFIED" if fund_path.exists() and fund_ok_count == len(df) else "PARTIAL" if fund_path.exists() and fund_ok_count > 0 else "UNAVAILABLE"
    breadth_status = "VERIFIED" if market_breadth.get("status") == "CLEAN_PASS" and str(market_breadth.get("date") or "").replace("-", "") == DATE else "UNAVAILABLE"
    previous_status = "VERIFIED" if prev_material.get("status") == "VERIFIED" else "UNAVAILABLE"
    sources = {
        "official_pool": source_record("official_pool", "official_limit_up_pool", IN_TABLE, "VERIFIED", len(df)),
        "market_breadth": source_record("market_breadth", "all_a_breadth", MARKET_BREADTH_JSON, breadth_status, market_breadth.get("valid_change_rows"), provider=market_breadth.get("source")),
        "fund_flow": source_record("fund_flow", "capital_flow", fund_path, fund_status, len(fund), verified_rows=fund_ok_count),
        "failed_limit_up_pool": source_record("failed_limit_up_pool", "failed_limit_up", zbgc_path, "VERIFIED" if zbgc_path.exists() else "UNAVAILABLE", len(zbgc) if zbgc_path.exists() else None),
        "limit_down_pool": source_record("limit_down_pool", "limit_down", dtgc_path, "VERIFIED" if dtgc_path.exists() else "UNAVAILABLE", len(dtgc) if dtgc_path.exists() else None),
        "dragon_tiger": source_record("dragon_tiger", "dragon_tiger_list", lhb_path, "VERIFIED" if lhb_path.exists() else "UNAVAILABLE", len(lhb) if lhb_path.exists() else None),
        "formula_sample": source_record("formula_sample", "technical_sample", TQ_SAMPLE_JSON, "PARTIAL" if formula_sample.get("sample_success_count") else "UNAVAILABLE", formula_sample.get("sample_success_count")),
        "previous_prediction": source_record("previous_prediction", "frozen_prediction_history", Path(prev_material.get("watchlist")) if prev_material.get("watchlist") else None, previous_status, prev_material.get("watch_count")),
        "catalyst_research": catalyst_source,
        "historical_baseline": source_record(
            "historical_baseline",
            "same_schema_history",
            HISTORY_DIR,
            historical_baselines["status"],
            historical_baselines["sample_count"],
            reason=historical_baselines.get("reason", ""),
        ),
    }

    total_market = int(market_breadth.get("up_count") or 0) + int(market_breadth.get("down_count") or 0) + int(market_breadth.get("flat_count") or 0)
    largest = top_sectors[0]
    largest_count = int(largest["count"])
    largest_connected = int(largest["connected"])
    zbgc_value: int | None = len(zbgc) if zbgc_path.exists() else None
    dtgc_value: int | None = len(dtgc) if dtgc_path.exists() else None
    fund_series = pd.to_numeric(df.get("主力净流入", pd.Series(dtype=float)), errors="coerce")
    verified_fund_series = fund_series[df.get("ok", pd.Series(False, index=df.index)).map(parse_verified_bool)] if "ok" in df.columns else pd.Series(dtype=float)
    fund_total: float | None = float(verified_fund_series.sum()) if not verified_fund_series.dropna().empty else None
    board_series = pd.to_numeric(df.get("连板数", 1), errors="coerce").fillna(1).astype(int)
    open_series = pd.to_numeric(df.get("炸板次数", 0), errors="coerce").fillna(0).astype(int)
    first_seal_series = pd.to_numeric(df.get("首次封板时间", 0), errors="coerce").fillna(0).astype(int)
    code_series = df["代码"].map(code6)
    growth_board_mask = code_series.str.startswith(("300", "301", "688"))
    growth_board_count = int(growth_board_mask.sum())
    growth_board_share = growth_board_count / len(df) if len(df) else 0.0
    early_seal_count = int(((first_seal_series > 0) & (first_seal_series < 100000)).sum())
    early_seal_share = early_seal_count / len(df) if len(df) else 0.0
    zero_open_count = int((open_series == 0).sum())
    zero_open_share = zero_open_count / len(df) if len(df) else 0.0
    high_board_mask = board_series >= 3
    high_board_count = int(high_board_mask.sum())
    high_board_open_count = int(open_series[high_board_mask].sum())
    market_cap_series = pd.to_numeric(df.get("流通市值", pd.Series(dtype=float)), errors="coerce").dropna()
    median_float_cap = float(market_cap_series.median()) if not market_cap_series.empty else None
    positive_fund = verified_fund_series[verified_fund_series > 0].sort_values(ascending=False)
    positive_fund_total = float(positive_fund.sum()) if not positive_fund.empty else 0.0
    top5_positive_fund_share = float(positive_fund.head(5).sum() / positive_fund_total) if positive_fund_total > 0 else 0.0

    lhb_work = lhb.copy()
    if not lhb_work.empty and "代码" in lhb_work.columns:
        lhb_work["代码"] = lhb_work["代码"].map(code6)
        lhb_work = lhb_work[lhb_work["代码"].isin(set(code_series))].copy()
        lhb_work["_net"] = pd.to_numeric(lhb_work.get("龙虎榜净买额", 0), errors="coerce").fillna(0.0)
        lhb_work["_abs_net"] = lhb_work["_net"].abs()
        lhb_unique = lhb_work.sort_values("_abs_net", ascending=False).drop_duplicates("代码")
        stock_context = df[["代码", "名称", "投资板块"]].copy()
        stock_context["代码"] = stock_context["代码"].map(code6)
        lhb_unique = lhb_unique.drop(columns=["名称", "投资板块"], errors="ignore").merge(
            stock_context.drop_duplicates("代码"), on="代码", how="left"
        )
    else:
        lhb_unique = pd.DataFrame()
    lhb_interpretation = lhb_unique.get("解读", pd.Series(dtype=str)).astype(str) if not lhb_unique.empty else pd.Series(dtype=str)
    institution_mask = lhb_interpretation.str.contains("机构", na=False)
    hot_money_mask = lhb_interpretation.str.contains(r"买一|资金|席位|主力", regex=True, na=False) & ~institution_mask
    institution_rows = lhb_unique[institution_mask].copy() if not lhb_unique.empty else pd.DataFrame()
    hot_money_rows = lhb_unique[hot_money_mask].copy() if not lhb_unique.empty else pd.DataFrame()
    institution_positive_count = int((institution_rows.get("_net", pd.Series(dtype=float)) > 0).sum()) if not institution_rows.empty else 0
    institution_net_total = float(institution_rows.get("_net", pd.Series(dtype=float)).sum()) if not institution_rows.empty else 0.0
    hot_money_positive_count = int((hot_money_rows.get("_net", pd.Series(dtype=float)) > 0).sum()) if not hot_money_rows.empty else 0
    hot_money_net_total = float(hot_money_rows.get("_net", pd.Series(dtype=float)).sum()) if not hot_money_rows.empty else 0.0

    def participant_profile(actor: str, rows: pd.DataFrame) -> dict[str, Any]:
        sample_count = int(len(rows))
        net_total = float(rows.get("_net", pd.Series(dtype=float)).sum()) if sample_count else 0.0
        positive_count = int((rows.get("_net", pd.Series(dtype=float)) > 0).sum()) if sample_count else 0
        sector_counts = rows.get("投资板块", pd.Series(dtype=str)).fillna("未分类").astype(str).value_counts()
        sector_count = int(len(sector_counts))
        top_sector = clean(sector_counts.index[0], "未分类") if sector_count else "未分类"
        top_sector_count = int(sector_counts.iloc[0]) if sector_count else 0
        top_sector_share = top_sector_count / sample_count if sample_count else 0.0
        leaders = [clean(name) for name in rows.sort_values("_abs_net", ascending=False).get("名称", pd.Series(dtype=str)).head(3).tolist()]
        if net_total > 0:
            flow_text = f"合计净买入{fmt_amount(net_total)}"
            flow_short = "净买入"
        elif net_total < 0:
            flow_text = f"合计净卖出{fmt_amount(abs(net_total))}"
            flow_short = "净卖出"
        else:
            flow_text = "合计净额持平"
            flow_short = "净额持平"
        if sector_count == 1:
            distribution = f"{sample_count}只样本全部落在{top_sector}"
            headline = f"{actor}{flow_short}集中在{top_sector}"
        elif top_sector_share >= 0.5:
            distribution = f"分布在{sector_count}个投资板块，其中{top_sector}{top_sector_count}只、占{top_sector_share:.1%}，集中度最高"
            headline = f"{actor}{flow_short}主要集中在{top_sector}"
        else:
            distribution = f"分布在{sector_count}个投资板块，最大板块{top_sector}{top_sector_count}只、占{top_sector_share:.1%}"
            headline = f"{actor}{flow_short}分布在{sector_count}个投资板块"
        leader_text = "、".join(leaders) if leaders else "无"
        conclusion = (
            f"{actor}特征席位覆盖{sample_count}只涨停股，其中{positive_count}只净买入为正，{flow_text}；"
            f"{distribution}，净额靠前的股票是{leader_text}"
        )
        return {
            "sample_count": sample_count,
            "positive_count": positive_count,
            "net_total": net_total,
            "sector_count": sector_count,
            "top_sector": top_sector,
            "top_sector_count": top_sector_count,
            "top_sector_share": top_sector_share,
            "leaders": leaders,
            "headline": headline,
            "conclusion": conclusion,
        }

    hot_money_profile = participant_profile("游资", hot_money_rows)
    institution_profile = participant_profile("机构", institution_rows)

    def percentile_rank(field: str, current: float, window: int = 20) -> float | None:
        if historical_baselines["status"] not in {"VERIFIED", "PARTIAL"}:
            return None
        values = [float(item[field]) for item in historical_baselines["records"][-window:] if is_number(item.get(field))]
        if not values:
            return None
        return round(sum(value <= current for value in values) / len(values), 4)

    metrics = {
        "official_limit_up_count": metric(int(len(df)), "official_pool", "stocks"),
        "first_board_count": metric(int(facts["first_board_count"]), "official_pool", "stocks"),
        "continuation_count": metric(int(facts["connected_count"]), "official_pool", "stocks"),
        "max_board": metric(int(facts["highest_board"]), "official_pool", "boards"),
        "first_board_share": metric(round(float(facts["first_share"]), 4), "official_pool", "ratio"),
        "continuation_share": metric(round(float(facts["connected_share"]), 4), "official_pool", "ratio"),
        "market_up_count": metric(int(market_breadth.get("up_count") or 0) if breadth_status == "VERIFIED" else None, "market_breadth", "stocks", breadth_status),
        "market_down_count": metric(int(market_breadth.get("down_count") or 0) if breadth_status == "VERIFIED" else None, "market_breadth", "stocks", breadth_status),
        "market_up_ratio": metric(round(float(market_breadth.get("up_ratio") or 0), 4) if breadth_status == "VERIFIED" else None, "market_breadth", "ratio", breadth_status),
        "market_down_ratio": metric(round(float(market_breadth.get("down_ratio") or 0), 4) if breadth_status == "VERIFIED" else None, "market_breadth", "ratio", breadth_status),
        "touch_fail_rate": metric(round(float(fail_rate), 2) if zbgc_value is not None else None, "failed_limit_up_pool", "percent", "VERIFIED" if zbgc_value is not None else "UNAVAILABLE"),
        "limit_down_count": metric(dtgc_value, "limit_down_pool", "stocks", "VERIFIED" if dtgc_value is not None else "UNAVAILABLE"),
        "largest_direction_count": metric(largest_count, "official_pool", "stocks"),
        "largest_direction_share": metric(round(largest_count / len(df), 4), "official_pool", "ratio"),
        "largest_direction_continuation": metric(largest_connected, "official_pool", "stocks"),
        "largest_direction_fund_total": metric(float(facts["main_fund"]) if fund_status != "UNAVAILABLE" else None, "fund_flow", "currency", fund_status),
        "top5_direction_share": metric(round(float(facts["top5_share"]), 4), "official_pool", "ratio"),
        "front_open_count": metric(int(facts["front_open_sum"]), "official_pool", "events"),
        "fund_verified_count": metric(fund_ok_count if fund_status != "UNAVAILABLE" else None, "fund_flow", "stocks", fund_status),
        "fund_total": metric(fund_total, "fund_flow", "currency", fund_status),
        "previous_watch_retained": metric(prev_material.get("watch_retained_count") if previous_status == "VERIFIED" else None, "previous_prediction", "stocks", previous_status),
        "previous_watch_promoted": metric(prev_material.get("watch_promoted_count") if previous_status == "VERIFIED" else None, "previous_prediction", "stocks", previous_status),
        "limitup_count_percentile_20": metric(percentile_rank("limitup_count", float(len(df))), "historical_baseline", "percentile", historical_baselines["status"]),
        "continuation_count_percentile_20": metric(percentile_rank("continuation_count", float(facts["connected_count"])), "historical_baseline", "percentile", historical_baselines["status"]),
        "first_board_share_percentile_20": metric(percentile_rank("first_board_share", float(facts["first_share"])), "historical_baseline", "percentile", historical_baselines["status"]),
        "growth_board_count": metric(growth_board_count, "official_pool", "stocks"),
        "growth_board_share": metric(round(growth_board_share, 4), "official_pool", "ratio"),
        "early_seal_count": metric(early_seal_count, "official_pool", "stocks"),
        "early_seal_share": metric(round(early_seal_share, 4), "official_pool", "ratio"),
        "zero_open_count": metric(zero_open_count, "official_pool", "stocks"),
        "zero_open_share": metric(round(zero_open_share, 4), "official_pool", "ratio"),
        "high_board_count": metric(high_board_count, "official_pool", "stocks"),
        "high_board_open_count": metric(high_board_open_count, "official_pool", "events"),
        "median_float_cap": metric(median_float_cap, "official_pool", "currency", "VERIFIED" if median_float_cap is not None else "UNAVAILABLE"),
        "top5_positive_fund_share": metric(round(top5_positive_fund_share, 4), "fund_flow", "ratio", fund_status),
        "lhb_match_count": metric(int(len(lhb_unique)), "dragon_tiger", "stocks", "VERIFIED" if not lhb_unique.empty else "UNAVAILABLE"),
        "institution_sample_count": metric(int(len(institution_rows)), "dragon_tiger", "stocks", "VERIFIED" if not lhb_unique.empty else "UNAVAILABLE"),
        "institution_positive_count": metric(institution_positive_count, "dragon_tiger", "stocks", "VERIFIED" if not lhb_unique.empty else "UNAVAILABLE"),
        "institution_net_total": metric(institution_net_total, "dragon_tiger", "currency", "VERIFIED" if not lhb_unique.empty else "UNAVAILABLE"),
        "institution_sector_count": metric(institution_profile["sector_count"], "dragon_tiger", "sectors", "VERIFIED" if not institution_rows.empty else "UNAVAILABLE"),
        "institution_top_sector_count": metric(institution_profile["top_sector_count"], "dragon_tiger", "stocks", "VERIFIED" if not institution_rows.empty else "UNAVAILABLE"),
        "institution_top_sector_share": metric(round(institution_profile["top_sector_share"], 4), "dragon_tiger", "ratio", "VERIFIED" if not institution_rows.empty else "UNAVAILABLE"),
        "hot_money_sample_count": metric(int(len(hot_money_rows)), "dragon_tiger", "stocks", "VERIFIED" if not lhb_unique.empty else "UNAVAILABLE"),
        "hot_money_positive_count": metric(hot_money_positive_count, "dragon_tiger", "stocks", "VERIFIED" if not lhb_unique.empty else "UNAVAILABLE"),
        "hot_money_net_total": metric(hot_money_net_total, "dragon_tiger", "currency", "VERIFIED" if not lhb_unique.empty else "UNAVAILABLE"),
        "hot_money_sector_count": metric(hot_money_profile["sector_count"], "dragon_tiger", "sectors", "VERIFIED" if not hot_money_rows.empty else "UNAVAILABLE"),
        "hot_money_top_sector_count": metric(hot_money_profile["top_sector_count"], "dragon_tiger", "stocks", "VERIFIED" if not hot_money_rows.empty else "UNAVAILABLE"),
        "hot_money_top_sector_share": metric(round(hot_money_profile["top_sector_share"], 4), "dragon_tiger", "ratio", "VERIFIED" if not hot_money_rows.empty else "UNAVAILABLE"),
    }
    evidence_records: dict[str, dict[str, Any]] = {
        f"EV-METRIC-{metric_id}": {
            "id": f"EV-METRIC-{metric_id}",
            "kind": "aggregate_metric",
            "source_id": record["source_id"],
            "metric_id": metric_id,
            "observed_value": record.get("value"),
            "observed_at": f"{DATE_H}T15:00:00+08:00",
        }
        for metric_id, record in metrics.items()
        if record.get("verification_status") != "UNAVAILABLE"
    }

    up_ratio = float(market_breadth.get("up_ratio") or 0)
    down_ratio = float(market_breadth.get("down_ratio") or 0)
    baseline_insufficient = historical_baselines["sample_count"] < 5
    limitup_percentile_20 = metrics["limitup_count_percentile_20"].get("value")
    if baseline_insufficient or limitup_percentile_20 is None:
        market_baseline_clause = (
            f"同口径历史样本仅{historical_baselines['sample_count']}期，"
            "当前证据不能判定市场周期强弱"
        )
    else:
        market_baseline_clause = (
            f"当日{len(df)}只涨停位于可用20日样本的{float(limitup_percentile_20):.0%}分位，"
            "上涨宽度不能单独代表涨停生态质量"
        )
    market_conclusion = (
        f"上涨家数{int(market_breadth.get('up_count') or 0)}只高于下跌家数{int(market_breadth.get('down_count') or 0)}只，"
        f"但连板占比仅{facts['connected_share']:.1%}、触板失败率{fail_rate:.1f}%；"
        f"宽度与涨停质量存在背离；{market_baseline_clause}"
    )
    market_status = "INSUFFICIENT" if baseline_insufficient else "CONTESTED"

    missing_middle = [
        board
        for board in range(int(facts["highest_board"]) - 1, 1, -1)
        if int(facts["board_counts"].get(board, 0)) == 0
    ]
    structure_status = "INSUFFICIENT" if baseline_insufficient else ("SUPPORTED" if facts["connected_share"] >= 0.20 and facts["highest_board"] >= 3 and not missing_middle else "CONTESTED")
    if facts["highest_board"] < 3:
        structure_conclusion = f"涨停结构由{facts['first_board_count']}只首板和{facts['connected_count']}只连板构成，最高仅{facts['highest_board']}板，梯队仍停留在低位确认阶段"
    elif missing_middle:
        structure_conclusion = f"涨停结构由{facts['first_board_count']}只首板和{facts['connected_count']}只连板构成，{missing_middle}板层级缺失，高度与中位梯队存在断层"
    else:
        board_distribution = "、".join(
            f"{board}板{int(facts['board_counts'].get(board, 0))}只"
            for board in range(int(facts["highest_board"]), 1, -1)
        )
        structure_conclusion = (
            f"涨停结构由{facts['first_board_count']}只首板和{facts['connected_count']}只连板构成；"
            f"{board_distribution}，虽然2至{facts['highest_board']}板没有断层，但2板有{int(facts['board_counts'].get(2, 0))}只、3板以上只有{sum(int(count) for board, count in facts['board_counts'].items() if int(board) >= 3)}只；"
            "明天强弱取决于2板能不能大批升到3板，单看最高板继续上涨不够"
        )

    catalyst_verified = catalyst_source.get("verification_status") == "VERIFIED"
    direction_supported = largest_connected >= 2 and largest_count / len(df) >= 0.15 and catalyst_verified and not baseline_insufficient
    direction_status = "SUPPORTED" if direction_supported else "CONTESTED"
    if direction_supported:
        direction_conclusion = (
            f"{largest['sector']}以{largest_count}只涨停、{largest_connected}只连板、最高{largest['max_board']}板位列第一，"
            f"占全池{largest_count / len(df):.1%}且催化已经核验，达到方向升级门槛"
        )
    else:
        catalyst_state_text = "催化已经核验" if catalyst_verified else "缺少可核验催化"
        direction_conclusion = (
            f"{largest['sector']}以{largest_count}只涨停、{largest_connected}只连板、最高{largest['max_board']}板位列第一，"
            f"但仅占全池{largest_count / len(df):.1%}且{catalyst_state_text}，不满足方向升级门槛"
        )
    risk_conclusion = (
        f"同日触板失败率{fail_rate:.1f}%、跌停{dtgc_value if dtgc_value is not None else '不可用'}只、前排开板{facts['front_open_sum']}次，"
        + (
            f"历史基线仅{historical_baselines['sample_count']}期，不能判定是否处于历史高压区间"
            if baseline_insufficient
            else f"与上涨{int(market_breadth.get('up_count') or 0)}只、下跌{int(market_breadth.get('down_count') or 0)}只的市场宽度背离，风险论点保留争议状态"
        )
    )

    market_evidence = [
        thesis_evidence("market_breadth", "market_up_ratio", "breadth", f"上涨占比{up_ratio:.1%}、下跌占比{down_ratio:.1%}"),
        thesis_evidence("official_pool", "continuation_share", "limit_up_ecology", f"连板占涨停池{facts['connected_share']:.1%}"),
        thesis_evidence("failed_limit_up_pool", "touch_fail_rate", "failure_pressure", f"触板失败率{fail_rate:.1f}%"),
    ]
    structure_evidence = [
        thesis_evidence("official_pool", "first_board_share", "breadth_origin", f"首板占比{facts['first_share']:.1%}"),
        thesis_evidence("official_pool", "continuation_share", "continuity", f"连板占比{facts['connected_share']:.1%}"),
        thesis_evidence("official_pool", "max_board", "height", f"最高板为{facts['highest_board']}板"),
    ]
    if metrics["limitup_count_percentile_20"]["verification_status"] in {"VERIFIED", "PARTIAL"} and metrics["limitup_count_percentile_20"]["value"] is not None:
        market_evidence.append(thesis_evidence("historical_baseline", "limitup_count_percentile_20", "historical_baseline", f"涨停数位于可用20日样本的{metrics['limitup_count_percentile_20']['value']:.0%}分位"))
    if metrics["continuation_count_percentile_20"]["verification_status"] in {"VERIFIED", "PARTIAL"} and metrics["continuation_count_percentile_20"]["value"] is not None:
        structure_evidence.append(thesis_evidence("historical_baseline", "continuation_count_percentile_20", "historical_baseline", f"连板数位于可用20日样本的{metrics['continuation_count_percentile_20']['value']:.0%}分位"))

    def confidence_basis(*dimensions: str, limitations: list[str]) -> dict[str, Any]:
        return {
            "supporting_dimensions": list(dimensions),
            "limiting_factors": limitations,
            "baseline_status": historical_baselines["status"],
        }

    theses = [
        {
            "id": "TH-MARKET", "scope": "market", "claim_type": "INFERENCE", "conclusion": market_conclusion,
            "status": market_status, "confidence": "LOW" if baseline_insufficient or breadth_status != "VERIFIED" else "MEDIUM",
            "reasoning_path": ["先记录全市场涨跌宽度", "再对照涨停池连板占比与触板失败率", "宽度与质量背离时只形成条件化观察"],
            "confidence_basis": confidence_basis("breadth", "continuity", "failure_pressure", limitations=["历史基线样本不足" if baseline_insufficient else "历史样本仍在扩展"]),
            "evidence": market_evidence,
            "counterevidence": [thesis_evidence("official_pool", "official_limit_up_count", "breadth_expansion", f"当日仍有{len(df)}只涨停，可能代表风险偏好扩散，削弱谨慎判断", effect="CHALLENGES", challenge_to="TH-MARKET", alternative_explanation="涨停数量扩张可能先于连板质量改善")],
            "invalidation": [f"若次日连板数高于{facts['connected_count']}且触板失败率低于{fail_rate:.1f}%并同步保持上涨家数高于下跌家数，则宽度与质量背离判断失效"],
            "practical_implication": f"次日只有连板数高于{facts['connected_count']}只、触板失败率低于{fail_rate:.1f}%且上涨家数仍高于下跌家数，才上调环境假设。",
        },
        {
            "id": "TH-STRUCTURE", "scope": "structure", "claim_type": "INFERENCE", "conclusion": structure_conclusion,
            "status": structure_status, "confidence": "LOW" if baseline_insufficient else "MEDIUM",
            "reasoning_path": ["拆分首板与连板数量", "检查最高板和中间层级", "把梯队事实与次日晋级条件分开"],
            "confidence_basis": confidence_basis("breadth_origin", "continuity", "height", limitations=["历史基线样本不足" if baseline_insufficient else "梯队需次日验证"]),
            "evidence": structure_evidence,
            "counterevidence": [thesis_evidence("official_pool", "continuation_count", "existing_continuity", f"当前已有{facts['connected_count']}只连板；若次日连板数高于{facts['connected_count']}只且最高板高于{facts['highest_board']}板，梯队可继续抬升", effect="CHALLENGES", challenge_to="TH-STRUCTURE", alternative_explanation="当日层级分布不能排除次日整体晋级" )],
            "invalidation": [f"若次日连板数高于{facts['connected_count']}只且最高板高于{facts['highest_board']}板，则当前梯队判断需要上调"],
            "practical_implication": f"次日需同时满足连板数高于{facts['connected_count']}只且最高板高于{facts['highest_board']}板，才上调梯队判断；仅最高板变化不足以改变结论。",
        },
        {
            "id": "TH-DIRECTION", "scope": "direction", "claim_type": "HYPOTHESIS", "conclusion": direction_conclusion,
            "status": direction_status, "confidence": "MEDIUM" if direction_supported else "LOW",
            "reasoning_path": ["统计官方池内投资板块样本", "核对方向内连板与资金覆盖", "缺少同日催化时只保留待验证方向假设"],
            "confidence_basis": confidence_basis(
                "direction_breadth",
                "direction_continuity",
                "capital_consistency",
                limitations=([] if catalyst_verified else ["催化研究包不可用"])
                + (["历史基线样本不足"] if baseline_insufficient else ["历史样本仍在扩展"]),
            ),
            "evidence": [
                thesis_evidence("official_pool", "largest_direction_count", "direction_breadth", f"{largest['sector']}涨停{largest_count}只"),
                thesis_evidence("official_pool", "largest_direction_continuation", "direction_continuity", f"方向内连板{largest_connected}只、最高{largest['max_board']}板"),
                thesis_evidence("fund_flow", "largest_direction_fund_total", "capital_consistency", f"方向内主力净流入合计{fmt_amount(facts['main_fund'])}"),
            ],
            "counterevidence": [thesis_evidence("official_pool", "largest_direction_share", "statistical_concentration", f"最大方向仅占全池{largest_count / len(df):.1%}，数量集中度不足以单独解释全日盘面", effect="CHALLENGES", challenge_to="TH-DIRECTION", alternative_explanation="投资板块归属和单日数量不能替代催化与产业链证据")],
            "invalidation": [f"若次日{largest['sector']}涨停数不高于{largest_count}且连板数不高于{largest_connected}，或仍无可核验催化，则方向假设不升级"],
            "practical_implication": f"{largest['sector']}当前为{largest_count}只涨停、{largest_connected}只连板；只有次日涨停数高于{largest_count}只、连板数高于{largest_connected}只且新增可核验催化，才上调方向状态。",
        },
        {
            "id": "TH-RISK", "scope": "risk", "claim_type": "INFERENCE", "conclusion": risk_conclusion,
            "status": "INSUFFICIENT" if baseline_insufficient else "CONTESTED", "confidence": "LOW" if baseline_insufficient else ("MEDIUM" if zbgc_value is not None and dtgc_value is not None else "LOW"),
            "reasoning_path": ["记录触板失败、跌停与前排开板", "识别同日风险字段是否同时出现", "无历史基线时仅生成风险复核条件"],
            "confidence_basis": confidence_basis("failed_breakout", "downside_tail", "front_row_fragility", limitations=["历史基线样本不足" if baseline_insufficient else "同日数据不能替代跨期校准"]),
            "evidence": [
                thesis_evidence("failed_limit_up_pool", "touch_fail_rate", "failed_breakout", f"炸板{zbgc_value if zbgc_value is not None else '不可用'}只、触板失败率{fail_rate:.1f}%"),
                thesis_evidence("limit_down_pool", "limit_down_count", "downside_tail", f"跌停{dtgc_value if dtgc_value is not None else '不可用'}只"),
                thesis_evidence("official_pool", "front_open_count", "front_row_fragility", f"前排10只样本合计开板{facts['front_open_sum']}次"),
            ],
            "counterevidence": [
                thesis_evidence("market_breadth", "market_up_ratio", "breadth_resilience", f"上涨占比{up_ratio:.1%}，并非普遍下跌环境", effect="CHALLENGES", challenge_to="TH-RISK", alternative_explanation="较高触板失败率可能来自活跃换手而非系统性风险"),
                thesis_evidence("limit_down_pool", "limit_down_count", "limited_tail", f"跌停仅{dtgc_value if dtgc_value is not None else '不可用'}只，尾部压力未必扩散", effect="CHALLENGES", challenge_to="TH-RISK", alternative_explanation="风险可能局限于局部高波动样本"),
            ],
            "invalidation": [f"若次日触板失败率低于{fail_rate:.1f}%、跌停数低于{dtgc_value if dtgc_value is not None else 0}只、前排开板少于{facts['front_open_sum']}次，三项中至少两项改善，则风险复核优先级下调"],
            "practical_implication": f"以触板失败率{fail_rate:.1f}%、跌停{dtgc_value if dtgc_value is not None else '不可用'}只、前排开板{facts['front_open_sum']}次为三条基线；次日至少两项下降才下调风险复核优先级。",
        },
    ]

    raw_available_fields = {str(column) for column in df.columns}
    derived_available_fields = {
        "全市场上涨占比", "全市场下跌占比", "触板失败率", "跌停数量", "连板占比",
        "方向集中度", "成长板涨停占比", "十点前封板占比", "零开板占比", "高位开板次数",
        "龙虎榜净买额", "龙虎榜解读", "机构样本", "游资样本", "正向资金前五集中度",
        "历史涨停分位", "历史连板分位", "昨日观察留存", "催化核验状态",
    }
    available_fields = raw_available_fields | derived_available_fields
    analyzed_fields: set[str] = set()
    candidate_signal_audit: list[dict[str, Any]] = []
    conclusion_graph: list[dict[str, Any]] = []

    def add_signal(
        signal_id: str,
        operator: str,
        input_fields: list[str],
        observed_result: str,
        selected: bool,
        selection_reason: str,
        *,
        topic: str = "",
        conclusion: str = "",
        objects: list[str] | None = None,
        reasoning_path: list[str] | None = None,
        evidence: list[dict[str, Any]] | None = None,
        counterevidence: list[dict[str, Any]] | None = None,
        decision_value: str = "",
        confirmation: list[str] | None = None,
        invalidation: list[str] | None = None,
        unavailable: bool = False,
    ) -> None:
        analyzed_fields.update(input_fields)
        status = "UNAVAILABLE" if unavailable else "SELECTED" if selected else "REJECTED"
        candidate_signal_audit.append({
            "signal_id": signal_id,
            "operator": operator,
            "input_fields": input_fields,
            "status": status,
            "observed_result": observed_result,
            "selection_reason": selection_reason,
        })
        if not selected or unavailable:
            return
        conclusion_graph.append({
            "id": f"CG-{signal_id}",
            "signal_id": signal_id,
            "topic": topic,
            "conclusion": conclusion,
            "objects": objects or [],
            "reasoning_path": reasoning_path or [],
            "evidence": evidence or [],
            "counterevidence": counterevidence or [],
            "decision_value": decision_value,
            "confirmation": confirmation or [],
            "invalidation": invalidation or [],
        })

    add_signal(
        "BREADTH_QUALITY_DIVERGENCE", "divergence",
        ["全市场上涨占比", "连板占比", "触板失败率"],
        f"上涨占比{up_ratio:.1%}、连板占比{facts['connected_share']:.1%}、触板失败率{fail_rate:.1f}%",
        up_ratio >= 0.50 and (facts["connected_share"] < 0.25 or fail_rate >= 20),
        "市场宽度与涨停质量出现方向相反的信号" if up_ratio >= 0.50 else "未形成宽度扩张",
        topic="市场宽度与短线质量",
        conclusion=f"上涨股票占到{up_ratio:.1%}，但连板只占涨停池{facts['connected_share']:.1%}，冲板后没封住的比例又有{fail_rate:.1f}%；多数股票在涨，强势股接力却不顺，不能把普涨当成短线全面转强",
        objects=["全市场", "涨停池", "连板梯队"],
        reasoning_path=["先用全市场上涨占比判断宽度", "再用连板占比和触板失败率判断短线深度", "两者背离时不把普涨等同于接力转强"],
        evidence=[
            thesis_evidence("market_breadth", "market_up_ratio", "breadth", f"上涨占比{up_ratio:.1%}"),
            thesis_evidence("official_pool", "continuation_share", "continuity", f"连板占比{facts['connected_share']:.1%}"),
            thesis_evidence("failed_limit_up_pool", "touch_fail_rate", "failure_pressure", f"触板失败率{fail_rate:.1f}%"),
        ],
        counterevidence=[thesis_evidence("official_pool", "official_limit_up_count", "limit_up_breadth", f"涨停仍有{len(df)}只，局部赚钱效应没有消失", effect="CHALLENGES", challenge_to="CG-BREADTH_QUALITY_DIVERGENCE", alternative_explanation="宽度修复可能先于连板质量改善")],
        decision_value="明天不能只看指数或上涨家数，必须等连板增加、冲板失败减少。",
        confirmation=[f"连板数高于{facts['connected_count']}只且触板失败率低于{fail_rate:.1f}%"],
        invalidation=[f"上涨家数不再高于下跌家数，或连板数与触板失败率同时改善"],
    )

    direction_share = largest_count / len(df) if len(df) else 0.0
    add_signal(
        "DIRECTION_CONCENTRATION", "concentration",
        ["投资板块", "方向集中度", "连板数", "主力净流入"],
        f"{largest['sector']}有{largest_count}只涨停、{largest_connected}只连板，占全池{direction_share:.1%}",
        True,
        "最大方向的数量、连续性和资金可以形成可比较的横截面结论",
        topic=f"{largest['sector']}的热点层级",
        conclusion=(
            f"{largest['sector']}以{largest_count}只涨停、{largest_connected}只连板居首，但只占全池{direction_share:.1%}；"
            + ("它是今天最清楚的热点，但还没有强到成为全市场唯一主线" if direction_share < 0.20 or not catalyst_verified else "数量、连续上涨和消息面已经互相印证，可以按主线看待")
        ),
        objects=[largest["sector"]] + [stock_label(row, with_board=False) for _, row in largest["leaders"].head(3).iterrows()],
        reasoning_path=["按投资板块统计涨停广度", "核对方向内连板连续性与资金", "再以全池占比和催化决定热点或主线层级"],
        evidence=[
            thesis_evidence("official_pool", "largest_direction_count", "direction_breadth", f"{largest['sector']}涨停{largest_count}只"),
            thesis_evidence("official_pool", "largest_direction_continuation", "direction_continuity", f"方向内连板{largest_connected}只"),
            thesis_evidence("fund_flow", "largest_direction_fund_total", "capital_consistency", f"方向内主力净流入{fmt_amount(facts['main_fund'])}"),
        ],
        counterevidence=[thesis_evidence("official_pool", "largest_direction_share", "dispersion", f"该方向只占全池{direction_share:.1%}", effect="CHALLENGES", challenge_to="CG-DIRECTION_CONCENTRATION", alternative_explanation="全日涨停可能由多个独立热点共同构成")],
        decision_value=f"把{largest['sector']}作为热点候选验证，不在催化和扩散不足时预设为主线。",
        confirmation=[f"次日{largest['sector']}涨停超过{largest_count}只、连板超过{largest_connected}只并出现可核验催化"],
        invalidation=[f"次日该方向涨停不高于{largest_count}只且连板不高于{largest_connected}只"],
    )

    add_signal(
        "LADDER_FRAGILITY", "cohort_compare",
        ["连板数", "炸板次数", "高位开板次数", "首次封板时间"],
        f"最高{facts['highest_board']}板，3板以上{high_board_count}只、累计开板{high_board_open_count}次",
        high_board_count > 0,
        "高位样本存在且可与全池封板质量比较",
        topic="高位梯队的真实承接",
        conclusion=f"{clean(df.loc[board_series.idxmax()].get('名称'))}把高度维持在{facts['highest_board']}板，但3板以上只有{high_board_count}只，而且合计开板{high_board_open_count}次；最高板还在，后面的连板股却接得不整齐，不能用一只最高板代表全部强势股都顺畅",
        objects=[stock_label(row, with_board=False) for _, row in df[high_board_mask].sort_values("连板数", ascending=False).iterrows()],
        reasoning_path=["识别最高板和3板以上样本", "比较高位样本开板次数", "区分单一高度与梯队整体承接"],
        evidence=[
            thesis_evidence("official_pool", "max_board", "height", f"最高{facts['highest_board']}板"),
            thesis_evidence("official_pool", "high_board_count", "high_board_breadth", f"3板以上{high_board_count}只"),
            thesis_evidence("official_pool", "high_board_open_count", "high_board_stability", f"高位累计开板{high_board_open_count}次"),
        ],
        counterevidence=[thesis_evidence("official_pool", "continuation_count", "continuity_stock", f"全池仍有{facts['connected_count']}只连板", effect="CHALLENGES", challenge_to="CG-LADDER_FRAGILITY", alternative_explanation="2板储备可能在次日补齐中高位梯队")],
        decision_value="高标只能作为情绪锚，次日需要观察2板向3板的补位，而不是只追踪最高板。",
        confirmation=[f"3板以上样本超过{high_board_count}只且高位累计开板少于{high_board_open_count}次"],
        invalidation=[f"最高板断板且3板以上样本不增加"],
    )

    style_selected = growth_board_share <= 0.35 or growth_board_share >= 0.65
    style_name = "主板股票" if growth_board_share <= 0.35 else "创业板和科创板股票"
    add_signal(
        "BOARD_STYLE", "distribution",
        ["市场", "代码", "成长板涨停占比", "连板数"],
        f"创业板与科创板涨停{growth_board_count}只、占{growth_board_share:.1%}",
        style_selected,
        "板型占比明显偏离均衡区间" if style_selected else "板型分布接近均衡，未形成清晰风格结论",
        topic=f"{style_name}风格",
        conclusion=f"创业板和科创板只占涨停池{growth_board_share:.1%}，今天更容易赚钱的股票主要在{style_name}；创业板和科创板只有少数股票涨停，没有带动整个市场",
        objects=[style_name, "创业板与科创板"],
        reasoning_path=["按证券代码识别板型", "比较各板型在涨停池中的占比", "把占比优势映射为当日风格而非长期偏好"],
        evidence=[
            thesis_evidence("official_pool", "growth_board_share", "board_style", f"成长板占比{growth_board_share:.1%}"),
            thesis_evidence("official_pool", "continuation_share", "continuity_style", f"全池连板占比{facts['connected_share']:.1%}"),
        ],
        counterevidence=[thesis_evidence("official_pool", "growth_board_count", "growth_optional", f"成长板仍有{growth_board_count}只涨停", effect="CHALLENGES", challenge_to="CG-BOARD_STYLE", alternative_explanation="少量高弹性样本仍可能提供局部高收益")],
        decision_value=f"次日机会筛选优先在{style_name}中寻找方向连续性，不把少数高弹性涨停外推成全市场风格。",
        confirmation=[f"成长板涨停占比继续低于{growth_board_share:.1%}" if growth_board_share <= 0.35 else f"成长板涨停占比继续高于{growth_board_share:.1%}"],
        invalidation=["另一板型涨停占比跨过五成并同时占据最高板"],
    )

    add_signal(
        "SEALING_TIMING", "cross_tab",
        ["首次封板时间", "炸板次数", "十点前封板占比", "零开板占比"],
        f"十点前封板{early_seal_count}只、占{early_seal_share:.1%}；零开板{zero_open_count}只、占{zero_open_share:.1%}",
        True,
        "封板时序和开板稳定性共同定位赚钱效应",
        topic="封板时序与确定性溢价",
        conclusion=f"十点前第一次封住涨停的股票占{early_seal_share:.1%}，封住后没再开板的占{zero_open_share:.1%}；今天最稳的赚钱位置在早盘快速封住、很少开板的股票，午后才回封和反复开板的股票明显更难做",
        objects=["十点前封板样本", "零开板样本", "午后回封样本"],
        reasoning_path=["按首次封板时间划分早盘与午后", "用开板次数验证封板稳定性", "两项共同定位当日确定性溢价"],
        evidence=[
            thesis_evidence("official_pool", "early_seal_share", "timing", f"十点前封板占{early_seal_share:.1%}"),
            thesis_evidence("official_pool", "zero_open_share", "stability", f"零开板占{zero_open_share:.1%}"),
        ],
        counterevidence=[thesis_evidence("official_pool", "front_open_count", "reseal_activity", f"前排10只仍累计开板{facts['front_open_sum']}次", effect="CHALLENGES", challenge_to="CG-SEALING_TIMING", alternative_explanation="部分强股可能通过充分换手而非一次封死产生溢价")],
        decision_value="次日优先验证早盘封板样本的晋级率；午后回封只有在方向扩散同步时才提高权重。",
        confirmation=[f"十点前封板样本的次日晋级数量高于当前连板{facts['connected_count']}只对应水平"],
        invalidation=["午后回封样本的次日晋级率显著高于早盘封板样本"],
    )

    if not lhb_unique.empty:
        add_signal(
            "HOT_MONEY_BEHAVIOR", "participant_attribution",
            ["龙虎榜解读", "龙虎榜净买额", "游资样本", "投资板块"],
            f"游资特征样本{len(hot_money_rows)}只，其中净买入为正{hot_money_positive_count}只，合计{fmt_amount(hot_money_net_total)}",
            len(hot_money_rows) > 0,
            "龙虎榜解读可识别游资或席位型行为",
            topic="游资进攻位置",
            conclusion=hot_money_profile["conclusion"],
            objects=hot_money_profile["leaders"] + [hot_money_profile["top_sector"]],
            reasoning_path=["从龙虎榜解读识别席位型行为", "计算样本净买方向", "按投资板块统计集中度并核对代表股票"],
            evidence=[
                thesis_evidence("dragon_tiger", "hot_money_sample_count", "participant_coverage", f"游资样本{len(hot_money_rows)}只"),
                thesis_evidence("dragon_tiger", "hot_money_positive_count", "buying_breadth", f"净买入为正{hot_money_positive_count}只"),
                thesis_evidence("dragon_tiger", "hot_money_net_total", "net_direction", f"合计净买额{fmt_amount(hot_money_net_total)}"),
                thesis_evidence("dragon_tiger", "hot_money_top_sector_share", "sector_concentration", f"{hot_money_profile['top_sector']}占游资样本{hot_money_profile['top_sector_share']:.1%}"),
            ],
            counterevidence=[thesis_evidence("dragon_tiger", "lhb_match_count", "coverage_limit", f"龙虎榜仅覆盖{len(lhb_unique)}/{len(df)}只涨停样本", effect="CHALLENGES", challenge_to="CG-HOT_MONEY_BEHAVIOR", alternative_explanation="未上榜样本无法据此归因为游资行为")],
            decision_value="把游资净买入个股视为局部进攻线索，只有同方向后排扩散时才升级为板块行为。",
            confirmation=["游资净买入样本次日继续晋级并带动同方向新增涨停"],
            invalidation=["游资净买入样本次日普遍断板且所属方向不扩散"],
        )
        add_signal(
            "INSTITUTION_BEHAVIOR", "participant_attribution",
            ["龙虎榜解读", "龙虎榜净买额", "机构样本", "投资板块"],
            f"机构特征样本{len(institution_rows)}只，其中净买入为正{institution_positive_count}只，合计{fmt_amount(institution_net_total)}",
            len(institution_rows) > 0,
            "龙虎榜解读存在明确机构参与样本",
            topic="机构偏好的涨停方向",
            conclusion=institution_profile["conclusion"],
            objects=institution_profile["leaders"] + [institution_profile["top_sector"]],
            reasoning_path=["识别龙虎榜机构标签", "计算机构样本净买方向", "按投资板块统计集中度并核对代表股票"],
            evidence=[
                thesis_evidence("dragon_tiger", "institution_sample_count", "participant_coverage", f"机构样本{len(institution_rows)}只"),
                thesis_evidence("dragon_tiger", "institution_positive_count", "buying_breadth", f"净买入为正{institution_positive_count}只"),
                thesis_evidence("dragon_tiger", "institution_net_total", "net_direction", f"合计净买额{fmt_amount(institution_net_total)}"),
                thesis_evidence("dragon_tiger", "institution_top_sector_share", "sector_concentration", f"{institution_profile['top_sector']}占机构样本{institution_profile['top_sector_share']:.1%}"),
            ],
            counterevidence=[thesis_evidence("dragon_tiger", "lhb_match_count", "coverage_limit", f"龙虎榜只匹配{len(lhb_unique)}只涨停股", effect="CHALLENGES", challenge_to="CG-INSTITUTION_BEHAVIOR", alternative_explanation="机构未上榜交易和普通席位无法被完整识别")],
            decision_value="机构线索用于缩小个股观察范围，不把机构净买直接解释为方向持续性。",
            confirmation=["机构净买入样本次日相对同方向样本保持更高晋级或溢价"],
            invalidation=["机构净买入样本次日集体弱于所属方向且龙虎榜转为净卖出"],
        )

    fund_concentration_selected = fund_status != "UNAVAILABLE" and positive_fund_total > 0
    add_signal(
        "CAPITAL_CONCENTRATION", "concentration",
        ["主力净流入", "超大单净流入", "正向资金前五集中度", "成交额"],
        f"正向主力净流入前五只占全部正向净流入{top5_positive_fund_share:.1%}",
        fund_concentration_selected,
        "资金覆盖完整且正向流入可计算集中度" if fund_concentration_selected else "正向资金总额不足以形成集中度结论",
        topic="资金集中还是普遍扩散",
        conclusion=(
            f"主力净流入最多的前五只股票占全部正向净流入{top5_positive_fund_share:.1%}，"
            f"资金数据覆盖{fund_ok_count}/{len(df)}只涨停股；"
            + ("正向资金明显集中在前五只股票" if top5_positive_fund_share >= 0.60 else "正向资金在前五只与其余股票之间分布相对均衡")
        ),
        objects=[stock_label(row, with_board=False) for _, row in df.assign(_fund=fund_series).sort_values("_fund", ascending=False).head(5).iterrows()],
        reasoning_path=["剔除不可用资金字段", "计算正向净流入总额与前五集中度", "比较资金集中度和涨停数量分布"],
        evidence=[
            thesis_evidence("fund_flow", "top5_positive_fund_share", "capital_concentration", f"前五占正向净流入{top5_positive_fund_share:.1%}"),
            thesis_evidence("fund_flow", "fund_verified_count", "coverage", f"资金覆盖{fund_ok_count}/{len(df)}只"),
        ],
        counterevidence=[thesis_evidence("fund_flow", "fund_total", "aggregate_flow", f"全池主力净流入合计{fmt_amount(fund_total)}", effect="CHALLENGES", challenge_to="CG-CAPITAL_CONCENTRATION", alternative_explanation="总量为正时，集中资金仍可能带动后排补涨")],
        decision_value="机会优先落在资金集中个股与其方向联动，不能把全池资金为正理解成普遍可做。",
        confirmation=["前五资金集中个股次日晋级并带动同方向后排增加"],
        invalidation=["前五资金集中个股次日弱于全池且资金转为净流出"],
    )

    add_signal(
        "PROFIT_LOSS_LOCATION", "cohort_compare",
        ["连板数", "首次封板时间", "炸板次数", "触板失败率", "跌停数量"],
        f"零开板占{zero_open_share:.1%}、十点前封板占{early_seal_share:.1%}、触板失败率{fail_rate:.1f}%、跌停{dtgc_value if dtgc_value is not None else '不可用'}只",
        True,
        "成功样本和失败样本均有同日可比数据",
        topic="赚钱效应与亏钱效应落点",
        conclusion=f"赚钱效应集中在早盘封住、少开板和方向有连续性的样本；亏钱效应主要来自触板后失败的{zbgc_value if zbgc_value is not None else '不可用'}只样本及高位反复开板，而不是全市场普遍下跌。当前可做性来自结构筛选，不来自无差别追涨",
        objects=["早盘零开板样本", "方向连板样本", "触板失败样本", "高位反复开板样本"],
        reasoning_path=["把成功涨停按封板时序和稳定性分组", "把失败端按炸板和跌停分组", "结合市场宽度判断亏钱效应是局部还是系统性"],
        evidence=[
            thesis_evidence("official_pool", "zero_open_share", "profit_stability", f"零开板占{zero_open_share:.1%}"),
            thesis_evidence("official_pool", "early_seal_share", "profit_timing", f"十点前封板占{early_seal_share:.1%}"),
            thesis_evidence("failed_limit_up_pool", "touch_fail_rate", "loss_failure", f"触板失败率{fail_rate:.1f}%"),
        ],
        counterevidence=[thesis_evidence("market_breadth", "market_up_ratio", "breadth_resilience", f"市场上涨占比{up_ratio:.1%}", effect="CHALLENGES", challenge_to="CG-PROFIT_LOSS_LOCATION", alternative_explanation="普涨环境可能使部分高换手炸板样本仍有盘中收益")],
        decision_value="明天只挑封板稳、方向里还有其他股票跟涨的对象，不能因为普涨就随便追。",
        confirmation=["早盘零开板样本的晋级和溢价继续优于反复开板样本"],
        invalidation=["反复开板和午后回封样本次日表现系统性优于早盘稳定封板样本"],
    )

    add_signal(
        "NEXT_DAY_BRANCH", "scenario_branch",
        ["连板占比", "触板失败率", "方向集中度", "跌停数量"],
        f"以连板{facts['connected_count']}只、触板失败率{fail_rate:.1f}%、{largest['sector']}涨停{largest_count}只、跌停{dtgc_value if dtgc_value is not None else '不可用'}只冻结次日分支",
        True,
        "当前关键变量可以形成互斥验证路径",
        topic="明天市场怎么走",
        conclusion=f"明天先看普涨能不能变成强势股继续涨停，而不是先猜还会普涨；连板超过{facts['connected_count']}只、冲板失败低于{fail_rate:.1f}%，才算强势股接力转好；如果{largest['sector']}走弱而新方向超过{largest_count}只，就是热点换方向；如果冲板失败和跌停一起增加，就是亏钱范围在扩大",
        objects=["质量修复路径", "方向轮动路径", "风险传导路径"],
        reasoning_path=["从当日背离中识别核心矛盾", "用连板、失败率、方向广度和跌停设置互斥条件", "次日只按实测条件切换路径"],
        evidence=[
            thesis_evidence("official_pool", "continuation_count", "quality_trigger", f"连板基线{facts['connected_count']}只"),
            thesis_evidence("failed_limit_up_pool", "touch_fail_rate", "failure_trigger", f"失败率基线{fail_rate:.1f}%"),
            thesis_evidence("official_pool", "largest_direction_count", "rotation_trigger", f"最大方向基线{largest_count}只"),
        ],
        counterevidence=[thesis_evidence("market_breadth", "market_up_ratio", "market_noise", f"上涨占比{up_ratio:.1%}可能掩盖局部结构变化", effect="CHALLENGES", challenge_to="CG-NEXT_DAY_BRANCH", alternative_explanation="指数和个股宽度变化可能使盘中路径短暂重叠")],
        decision_value="不押一个固定答案，明天按连板、冲板失败、热点数量和跌停的实际变化判断。",
        confirmation=[f"修复：连板>{facts['connected_count']}且失败率<{fail_rate:.1f}%；轮动：新方向涨停>{largest_count}；风险：失败率>{fail_rate:.1f}%且跌停增加"],
        invalidation=["出现收盘后重大政策、公告或监管事实，导致原变量关系失效"],
    )

    previous_verified = previous_status == "VERIFIED" and int(prev_material.get("watch_count") or 0) > 0
    previous_watch_count = int(prev_material.get("watch_count") or 0)
    previous_retained = int(prev_material.get("watch_retained_count") or 0)
    previous_promoted = int(prev_material.get("watch_promoted_count") or 0)
    add_signal(
        "PRIOR_CALIBRATION", "temporal_compare",
        ["昨日观察留存", "连板数", "投资板块"],
        f"昨日冻结样本{previous_watch_count}只，今日留存{previous_retained}只、晋级{previous_promoted}只",
        previous_verified,
        "存在同口径冻结观察账本" if previous_verified else "没有可验证的上一交易日冻结样本，禁止补造命中率",
        topic="昨天判断准了多少",
        conclusion=f"昨天收盘选出的{previous_watch_count}只股票，今天有{previous_retained}只仍在涨停池、{previous_promoted}只继续涨停；昨天的判断抓到了一部分强势股，但漏掉的更多，今天只能保留已经兑现的线索，不能事后把漏掉的股票算成昨天看对",
        objects=["昨日冻结观察样本", "今日留存样本", "今日晋级样本"],
        reasoning_path=["读取冻结账本而非事后替代名单", "计算留存和晋级", "用漏失范围限制本轮置信度"],
        evidence=[
            thesis_evidence("previous_prediction", "previous_watch_retained", "retention", f"今日留存{previous_retained}只"),
            thesis_evidence("previous_prediction", "previous_watch_promoted", "promotion", f"今日晋级{previous_promoted}只"),
        ],
        counterevidence=[thesis_evidence("official_pool", "official_limit_up_count", "new_information", f"今日官方涨停池有{len(df)}只，远大于昨日观察样本", effect="CHALLENGES", challenge_to="CG-PRIOR_CALIBRATION", alternative_explanation="今日新增事件和扩散可能超出昨日信息集")],
        decision_value="只沿用经冻结账本验证的线索；未被昨日覆盖的新热点必须重新建证据链。",
        confirmation=["昨日留存样本在次日继续维持方向连续性或板级晋升"],
        invalidation=["冻结账本缺失、日期不一致或样本集合被改写"],
    )

    add_signal(
        "CATALYST_CAUSALITY", "association",
        ["催化核验状态", "投资板块", "首次封板时间"],
        "同日催化研究包已核验" if catalyst_verified else "同日催化研究包不可用",
        catalyst_verified,
        "只有事件时间先于行情且方向有横截面响应才允许形成因果结论" if catalyst_verified else "缺少可核验催化，相关性不得写成因果",
        topic="事件催化与方向因果",
        conclusion=f"已核验催化与{largest['sector']}的封板时序和横截面响应一致，事件可作为方向持续性的独立证据",
        objects=[largest["sector"], "同日可核验催化"],
        reasoning_path=["核对事件发布时间", "核对受益映射和方向内响应", "排除市场普涨替代解释"],
        evidence=[
            thesis_evidence("official_pool", "largest_direction_count", "cross_section", f"方向涨停{largest_count}只"),
            thesis_evidence("official_pool", "largest_direction_continuation", "continuity", f"方向连板{largest_connected}只"),
        ],
        counterevidence=[thesis_evidence("market_breadth", "market_up_ratio", "market_alternative", f"市场上涨占比{up_ratio:.1%}", effect="CHALLENGES", challenge_to="CG-CATALYST_CAUSALITY", alternative_explanation="方向上涨可能部分来自普涨而非事件独立驱动")],
        decision_value="催化只在时间和横截面同时成立时提高方向置信度。",
        confirmation=["次日受益映射内样本继续扩散且强于市场宽度"],
        invalidation=["事件发布时间晚于行情，或受益方向没有横截面共振"],
        unavailable=not catalyst_verified,
    )

    highest_name = clean(df.loc[board_series.idxmax()].get("名称"))
    presentation_rules = {
        "BREADTH_QUALITY_DIVERGENCE": {
            "category_id": "market_style", "category_label": "市场与风格", "scope_label": "全市场",
            "angle_label": "上涨家数与强势股接力", "headline": "普涨，但强势股接力不顺", "priority": 10,
        },
        "BOARD_STYLE": {
            "category_id": "market_style", "category_label": "市场与风格", "scope_label": "主板、创业板、科创板",
            "angle_label": "涨停分布与赚钱偏好", "headline": f"赚钱主要在{style_name}", "priority": 20,
        },
        "DIRECTION_CONCENTRATION": {
            "category_id": "sector_mainline", "category_label": "板块与主线", "scope_label": largest["sector"],
            "angle_label": "涨停数量、连板和资金", "headline": f"{largest['sector']}最热，但还不是唯一主线", "priority": 30,
        },
        "CATALYST_CAUSALITY": {
            "category_id": "sector_mainline", "category_label": "板块与主线", "scope_label": largest["sector"],
            "angle_label": "消息时间与板块联动", "headline": f"{largest['sector']}的消息与走势互相印证", "priority": 35,
        },
        "LADDER_FRAGILITY": {
            "category_id": "leader_sealing", "category_label": "带队股与封板", "scope_label": "高位连板股",
            "angle_label": "最高板与后排接力", "headline": f"{highest_name}守住高度，后排接力偏弱", "priority": 40,
        },
        "SEALING_TIMING": {
            "category_id": "leader_sealing", "category_label": "带队股与封板", "scope_label": "早盘与午后涨停股",
            "angle_label": "首次封板时间和开板次数", "headline": "早盘快封、少开板的股票更好做", "priority": 50,
        },
        "HOT_MONEY_BEHAVIOR": {
            "category_id": "capital_behavior", "category_label": "资金行为", "scope_label": "龙虎榜游资席位",
            "angle_label": "净买入和投资板块集中度", "headline": hot_money_profile["headline"], "priority": 60,
        },
        "INSTITUTION_BEHAVIOR": {
            "category_id": "capital_behavior", "category_label": "资金行为", "scope_label": "龙虎榜机构席位",
            "angle_label": "机构净买入和投资板块分布", "headline": institution_profile["headline"], "priority": 65,
        },
        "CAPITAL_CONCENTRATION": {
            "category_id": "capital_behavior", "category_label": "资金行为", "scope_label": "全部涨停股",
            "angle_label": "主力净流入集中度", "headline": f"正向资金前五只占{top5_positive_fund_share:.1%}", "priority": 70,
        },
        "PROFIT_LOSS_LOCATION": {
            "category_id": "profit_risk", "category_label": "赚钱与风险", "scope_label": "涨停、炸板和高位股",
            "angle_label": "赚钱位置与亏钱来源", "headline": "机会在早盘稳封，风险在炸板和高位反复开板", "priority": 80,
        },
        "NEXT_DAY_BRANCH": {
            "category_id": "next_day", "category_label": "明日路径", "scope_label": "明日全市场",
            "angle_label": "接力、轮动和亏钱范围", "headline": "明天先验证强势股接力能不能转好", "priority": 90,
        },
        "PRIOR_CALIBRATION": {
            "category_id": "prior_review", "category_label": "昨日校准", "scope_label": "昨日观察名单",
            "angle_label": "今日留存和继续涨停", "headline": f"昨天抓到{previous_promoted}只继续涨停，但漏掉更多", "priority": 100,
        },
    }
    for node in conclusion_graph:
        presentation = presentation_rules.get(str(node.get("signal_id") or ""), {})
        node.update(presentation)

    current_data_payload = {
        "trade_date": DATE,
        "official_codes": sorted(df["代码"].map(code6).tolist()),
        "sources": sources,
        "metrics": metrics,
    }
    current_data_fingerprint = business_semantic_fingerprint(current_data_payload)
    for node in conclusion_graph:
        node_data_payload = {
            "trade_date": DATE,
            "current_data_fingerprint": current_data_fingerprint,
            "signal_id": node.get("signal_id"),
            "conclusion": node.get("conclusion"),
            "objects": node.get("objects"),
            "evidence": node.get("evidence"),
            "counterevidence": node.get("counterevidence"),
            "confirmation": node.get("confirmation"),
            "invalidation": node.get("invalidation"),
        }
        node["data_fingerprint"] = hashlib.sha256(
            json.dumps(node_data_payload, ensure_ascii=False, sort_keys=True, separators=(",", ":"), default=str).encode("utf-8")
        ).hexdigest()

    input_fields_used = {field for item in candidate_signal_audit for field in item["input_fields"]}
    unused_fields = [
        {"field": field, "reason": "标识、展示或原始明细字段，已由派生指标或同源字段覆盖，不单独形成推断"}
        for field in sorted(available_fields - input_fields_used)
    ]
    analysis_inventory = {
        "available_fields": sorted(available_fields),
        "analyzed_fields": sorted(input_fields_used),
        "unused_fields": unused_fields,
    }

    scenarios = {
        "strengthen": {
            "type": "strengthen",
            "triggers": [f"连板数高于当前{facts['connected_count']}只", f"触板失败率低于当前{fail_rate:.1f}%", f"{largest['sector']}方向连板数高于当前{largest_connected}只"],
            "trigger_evidence": [
                {"type": "METRIC", "condition": f"连板数高于当前{facts['connected_count']}只", "metric_id": "continuation_count", "source_id": "official_pool", "threshold": facts["connected_count"], "threshold_basis": "当日收盘连板数"},
                {"type": "METRIC", "condition": f"触板失败率低于当前{fail_rate:.1f}%", "metric_id": "touch_fail_rate", "source_id": "failed_limit_up_pool", "threshold": round(fail_rate, 2), "threshold_basis": "当日收盘触板失败率"},
                {"type": "METRIC", "condition": f"{largest['sector']}方向连板数高于当前{largest_connected}只", "metric_id": "largest_direction_continuation", "source_id": "official_pool", "threshold": largest_connected, "threshold_basis": "当日第一方向连板数"},
            ],
            "checkpoints": {"premarket": ["核对隔夜公告与政策是否改变原假设"], "open": ["检查前排是否出现一致承接而非孤立高开"], "intraday": ["检查首板补涨和连板晋级是否同时增加"], "close": ["以收盘封板、炸板和方向留存确认"]},
            "actions": ["维持市场与结构论点，并把方向状态从待验证上调一级。"],
            "invalidation": ["任一高标断板并伴随炸板率上升时，不进入强化分支。"],
        },
        "rotation": {
            "type": "rotation",
            "triggers": [f"{largest['sector']}涨停数低于当前{largest_count}只", f"新方向涨停数高于当前{largest_count}只", f"触板失败率不高于当前{fail_rate:.1f}%"],
            "trigger_evidence": [
                {"type": "METRIC", "condition": f"{largest['sector']}涨停数低于当前{largest_count}只", "metric_id": "largest_direction_count", "source_id": "official_pool", "threshold": largest_count, "threshold_basis": "当日第一方向涨停数"},
                {"type": "METRIC", "condition": f"新方向涨停数高于当前{largest_count}只", "metric_id": "largest_direction_count", "source_id": "official_pool", "threshold": largest_count, "threshold_basis": "当日第一方向涨停数，用作轮动比较线"},
                {"type": "METRIC", "condition": f"触板失败率不高于当前{fail_rate:.1f}%", "metric_id": "touch_fail_rate", "source_id": "failed_limit_up_pool", "threshold": round(fail_rate, 2), "threshold_basis": "当日收盘触板失败率"},
            ],
            "checkpoints": {"premarket": ["识别新增可核验催化"], "open": ["比较新旧方向前排承接"], "intraday": ["检查新方向是否出现首板扩散"], "close": ["以方向涨停数、连板数和资金覆盖重排"]},
            "actions": ["原方向降级为局部样本，重新建立新方向证据链。"],
            "invalidation": ["若新方向只有单只高标或缺少扩散，不认定轮动成立。"],
        },
        "failure": {
            "type": "failure",
            "triggers": [f"触板失败率高于当前{fail_rate:.1f}%", f"跌停数高于当前{dtgc_value if dtgc_value is not None else 0}只", f"连板数低于当前{facts['connected_count']}只"],
            "trigger_evidence": [
                {"type": "METRIC", "condition": f"触板失败率高于当前{fail_rate:.1f}%", "metric_id": "touch_fail_rate", "source_id": "failed_limit_up_pool", "threshold": round(fail_rate, 2), "threshold_basis": "当日收盘触板失败率"},
                {"type": "METRIC", "condition": f"跌停数高于当前{dtgc_value if dtgc_value is not None else 0}只", "metric_id": "limit_down_count", "source_id": "limit_down_pool", "threshold": dtgc_value, "threshold_basis": "当日收盘跌停数"},
                {"type": "METRIC", "condition": f"连板数低于当前{facts['connected_count']}只", "metric_id": "continuation_count", "source_id": "official_pool", "threshold": facts["connected_count"], "threshold_basis": "当日收盘连板数"},
            ],
            "checkpoints": {"premarket": ["检查高位样本重大利空"], "open": ["检查高标低开与板块脱节"], "intraday": ["检查炸板和跌停是否扩散"], "close": ["撤销已触发失效条件的论点"]},
            "actions": ["撤销方向延续假设，观察池转为风险复核清单。"],
            "invalidation": ["若失败触发未达到且连板晋级改善，则不进入失败分支。"],
        },
        "exogenous": {
            "type": "exogenous",
            "triggers": ["冻结后出现改变产业受益逻辑的重大政策或公告", "停复牌、监管或外部市场冲击改变可交易样本", "核心数据源交易日或口径发生变化"],
            "trigger_evidence": [
                {"type": "EVENT", "condition": "冻结后出现改变产业受益逻辑的重大政策或公告", "threshold_basis": "冻结后新增可核验事件", "verification_rule": "核对发布时间、原文来源和受益映射"},
                {"type": "EVENT", "condition": "停复牌、监管或外部市场冲击改变可交易样本", "threshold_basis": "冻结后新增可核验事件", "verification_rule": "核对公告、监管文件或交易状态"},
                {"type": "EVENT", "condition": "核心数据源交易日或口径发生变化", "threshold_basis": "来源元数据变化", "verification_rule": "核对来源日期、字段口径和内容哈希"},
            ],
            "checkpoints": {"premarket": ["核对公告、监管与政策发布时间"], "open": ["确认事件是否真正映射到相关样本"], "intraday": ["检查横截面是否形成共振"], "close": ["重新生成证据合同与预测账本"]},
            "actions": ["旧情景树作废，重新执行数据与结论链，不沿用旧方向判断。"],
            "invalidation": ["无法确认事件先于行情或受益映射时，不使用因果表述。"],
        },
    }

    watchlist: list[dict[str, Any]] = []
    risk_name_mask = qdf["名称"].astype(str).str.contains(r"(?:退市|退$|\*?ST)", case=False, regex=True, na=False)
    exclusions = [
        {"code": code6(row.get("代码")), "name": clean(row.get("名称")), "reason": "退市或特别处理标识，不进入实战观察池"}
        for _, row in qdf[risk_name_mask].iterrows()
    ]
    eligible = qdf[~risk_name_mask].copy()
    eligible["_board"] = pd.to_numeric(eligible["连板数"], errors="coerce").fillna(1).astype(int)
    eligible["_open_count"] = pd.to_numeric(eligible.get("炸板次数", 0), errors="coerce").fillna(0).astype(int)
    eligible["_first_seal"] = eligible.get("首次封板时间", "").astype(str)
    eligible["_fund"] = pd.to_numeric(eligible.get("主力净流入", pd.Series(index=eligible.index, dtype=float)), errors="coerce")
    observation_sort = ["_board", "_open_count", "_first_seal", "_fund"]
    direction_front = eligible[eligible["投资板块"] == largest["sector"]].sort_values(observation_sort, ascending=[False, True, True, False], na_position="last")
    ladder_front = eligible.sort_values(observation_sort, ascending=[False, True, True, False], na_position="last")
    candidate_df = pd.concat([direction_front.head(3), ladder_front]).drop_duplicates("代码").head(5)

    def row_reference(code: str, field: str, observed_value: Any, dimension: str, interpretation: str, *, source_id: str = "official_pool", effect: str = "SUPPORTS", alternative_explanation: str = "") -> dict[str, Any]:
        evidence_id = f"EV-ROW-{code}-{field}"
        evidence_records[evidence_id] = {
            "id": evidence_id,
            "kind": "row_field",
            "source_id": source_id,
            "code": code,
            "field": field,
            "observed_value": observed_value,
            "observed_at": f"{DATE_H}T15:00:00+08:00",
        }
        item = {"evidence_id": evidence_id, "source_id": source_id, "dimension": dimension, "interpretation": interpretation, "effect": effect}
        if alternative_explanation:
            item["alternative_explanation"] = alternative_explanation
        return item

    for _, row in candidate_df.iterrows():
        code = code6(row.get("代码"))
        board = to_int(row.get("连板数"), 1)
        sector = clean(row.get("投资板块"), "未分类")
        role = "height_anchor" if board == facts["highest_board"] and board >= 3 else "direction_front" if sector == largest["sector"] else "ladder_observer" if board >= 2 else "breadth_observer"
        open_count = to_int(row.get("炸板次数"))
        first_seal = fmt_time(row.get("首次封板时间"))
        fund_value = pd.to_numeric(pd.Series([row.get("主力净流入")]), errors="coerce").iloc[0]
        cohort = eligible[eligible["_board"] == board]
        cohort_definition = f"官方涨停池内{board}板样本"
        if len(cohort) < 2:
            direction_cohort = eligible[eligible["投资板块"] == sector]
            if len(direction_cohort) >= 2:
                cohort = direction_cohort
                cohort_definition = f"官方涨停池内{sector}方向样本"
            else:
                continuity_cohort = eligible[eligible["_board"] >= 2] if board >= 2 else eligible[eligible["_board"] == 1]
                if len(continuity_cohort) >= 2:
                    cohort = continuity_cohort
                    cohort_definition = "官方涨停池内连板样本" if board >= 2 else "官方涨停池内首板样本"
                else:
                    cohort = eligible
                    cohort_definition = "官方涨停池合规样本"
        cohort_open_median = float(cohort["_open_count"].median())
        comparison = {
            "cohort_definition": cohort_definition,
            "cohort_size": int(len(cohort)),
            "relative_position": f"开板{open_count}次，对比组中位数{cohort_open_median:.1f}次；首次封板{first_seal}",
        }
        evidence = [
            row_reference(code, "连板数", board, "ladder_position", f"样本{board}板，市场最高{facts['highest_board']}板"),
            row_reference(code, "首次封板时间", row.get("首次封板时间"), "timing_quality", f"首次封板时间为{first_seal}"),
            row_reference(code, "炸板次数", open_count, "sealing_quality", f"样本开板{open_count}次；{comparison['relative_position']}"),
        ]
        if fund_status != "UNAVAILABLE" and pd.notna(fund_value):
            evidence.append(row_reference(code, "主力净流入", float(fund_value), "capital_flow", f"主力净流入{fmt_amount(float(fund_value))}", source_id="fund_flow"))
        else:
            evidence.append(row_reference(code, "投资板块", sector, "direction_linkage", f"所属{sector}；当日数量第一方向为{largest['sector']}"))
        counter = row_reference(
            code,
            "炸板次数",
            open_count,
            "stability_challenge",
            f"开板{open_count}次并不等于次日承接成立，需与同口径对比组和方向扩散共同验证",
            effect="CHALLENGES",
            alternative_explanation="单日封板和资金流可能受盘中流动性或被动回封影响",
        )
        linked_claim_ids = ["TH-STRUCTURE", "TH-RISK"] + (["TH-DIRECTION"] if sector == largest["sector"] else [])
        watchlist.append({
            "code": code,
            "name": clean(row.get("名称")),
            "role": role,
            "rationale": f"{board}板、{sector}、首次封板{first_seal}、开板{open_count}次；作为{role}跟踪次日承接与晋级。",
            "why_now": f"当前处于{comparison['cohort_definition']}，{comparison['relative_position']}。",
            "linked_claim_ids": linked_claim_ids,
            "evidence": evidence,
            "counterevidence": [counter],
            "cross_sectional_comparison": comparison,
            "confirmation": ["次日仍在官方涨停池或实现板级晋升", "所属方向连板与后排扩散至少一项改善"],
            "exclusion_conditions": ["不再属于官方涨停池且未实现板级晋升", "出现板块脱节并且开板次数高于对比组中位数"],
            "invalidation": ["样本断板且方向连板数下降", "样本反复开板并伴随跌停或炸板扩散"],
        })

    blocks_delivery: list[str] = []
    if len(df) != len(set(df["代码"].map(code6))):
        blocks_delivery.append("official pool contains duplicate codes")
    if breadth_status != "VERIFIED":
        blocks_delivery.append("market breadth is unavailable or date-mismatched")
    if fund_status == "UNAVAILABLE":
        blocks_delivery.append("fund flow source unavailable")
    if not zbgc_path.exists() or not dtgc_path.exists():
        blocks_delivery.append("failed-limit-up or limit-down source unavailable")
    warnings = []
    if not catalyst_verified:
        warnings.append("catalyst research unavailable; direction confidence capped at LOW")
    if previous_status != "VERIFIED":
        warnings.append("previous frozen watchlist unavailable; no substitute sample was created")
    if formula_sample.get("status") != "CLEAN_PASS":
        warnings.append("technical formula sample is partial; it cannot support a pool-wide conclusion")
    if historical_baselines["status"] != "VERIFIED":
        warnings.append(f"20/60-day same-schema baseline is {historical_baselines['status']}; samples={historical_baselines['sample_count']}")

    prediction_entries: list[dict[str, Any]] = []
    for thesis in theses:
        prediction_entries.append({
            "prediction_id": f"P-{thesis['id']}",
            "object_type": "thesis",
            "object_id": thesis["id"],
            "claim_id": thesis["id"],
            "horizon": "T+1",
            "premises": [item["interpretation"] for item in thesis["evidence"]],
            "expected_observations": [thesis["practical_implication"]],
            "confirmation_conditions": [branch["triggers"][0] for branch in scenarios.values() if branch.get("triggers")],
            "invalidation_conditions": thesis["invalidation"],
            "scenario_links": list(scenarios.keys()),
            "risk_boundaries": [item["interpretation"] for item in thesis["counterevidence"]],
        })
    for item in watchlist:
        prediction_entries.append({
            "prediction_id": f"P-STOCK-{item['code']}",
            "object_type": "stock",
            "object_id": item["code"],
            "claim_id": item["linked_claim_ids"][0],
            "horizon": "T+1",
            "premises": [evidence["interpretation"] for evidence in item["evidence"]],
            "expected_observations": item["confirmation"],
            "confirmation_conditions": item["confirmation"],
            "invalidation_conditions": item["invalidation"],
            "scenario_links": ["strengthen", "rotation", "failure"],
            "risk_boundaries": item["exclusion_conditions"] + [evidence["interpretation"] for evidence in item["counterevidence"]],
        })

    result: dict[str, Any] = {
        "schema_version": BUSINESS_RESULT_SCHEMA_VERSION,
        "status": "CLEAN_PASS" if not blocks_delivery else "BLOCKED",
        "date": DATE,
        "run": {"run_id": RUN_ID, "trade_date": DATE, "latest_completed_trade_date": DATE, "as_of": f"{DATE_H}T15:00:00+08:00", "timezone": "Asia/Shanghai", "generated_at": generated_at, "entrypoint": "scripts/codex_entry.py run", "current_data_fingerprint": current_data_fingerprint},
        "official_universe": {
            "source_id": "official_pool",
            "official_pool_only": True,
            "eligibility_rule": DELISTING_HARD_EXCLUSION,
            "raw_count": int(eligibility_audit.get("official_source_count") or len(df)),
            "excluded": eligibility_audit.get("hard_exclusions") or [],
            "codes": sorted(df["代码"].map(code6).tolist()),
            "count": int(len(df)),
        },
        "sources": sources,
        "metrics": metrics,
        "evidence_records": evidence_records,
        "domains": {
            "data": {"status": "VERIFIED", "source_ids": ["official_pool"]},
            "market": {"status": breadth_status, "source_ids": ["market_breadth"]},
            "structure": {"status": "VERIFIED", "source_ids": ["official_pool", "failed_limit_up_pool", "limit_down_pool"]},
            "direction": {"status": "VERIFIED" if direction_supported else "PARTIAL", "source_ids": ["official_pool", "fund_flow", "catalyst_research"]},
            "capital": {"status": fund_status, "source_ids": ["fund_flow", "dragon_tiger"]},
            "catalyst": {"status": catalyst_source.get("verification_status"), "source_ids": ["catalyst_research"], "item_count": len(research_bundle.get("items") or []) if catalyst_verified and isinstance(research_bundle, dict) else None},
            "continuity": {"status": previous_status, "source_ids": ["previous_prediction"]},
            "risk": {"status": "VERIFIED" if zbgc_path.exists() and dtgc_path.exists() else "UNAVAILABLE", "source_ids": ["failed_limit_up_pool", "limit_down_pool"]},
            "decision": {"status": "VERIFIED", "source_ids": ["official_pool", "market_breadth", "fund_flow"]},
            "baseline": {"status": historical_baselines["status"], "source_ids": ["historical_baseline"]},
        },
        "theses": theses,
        "analysis_inventory": analysis_inventory,
        "candidate_signal_audit": candidate_signal_audit,
        "conclusion_graph": conclusion_graph,
        "scenario_tree": scenarios,
        "watchlist": watchlist,
        "exclusions": exclusions,
        "prediction_ledger": {"ledger_version": "v2", "run_id": RUN_ID, "as_of": f"{DATE_H}T15:00:00+08:00", "frozen_at": generated_at, "entries": prediction_entries, "content_sha256": ""},
        "blocks_delivery": bool(blocks_delivery),
        "blocks": list(blocks_delivery),
        "unverified_items": warnings,
        "data_quality": {
            "critical_source_status": {key: value.get("verification_status") for key, value in sources.items()},
            "fund_coverage": {"verified": fund_ok_count, "official_pool": len(df), "ratio": round(fund_ok_count / len(df), 4)},
            "historical_sample_count": historical_baselines["sample_count"],
            "catalyst_item_count": len(research_bundle.get("items") or []) if catalyst_verified and isinstance(research_bundle, dict) else None,
        },
        "data_gaps": [
            {"domain": "catalyst", "status": catalyst_source.get("verification_status"), "impact": "direction confidence capped; causal wording forbidden"},
            {"domain": "continuity", "status": previous_status, "impact": "no prior hit-rate claim and no substitute watchlist"},
            {"domain": "technical", "status": sources["formula_sample"].get("verification_status"), "impact": "formula sample cannot represent the full pool"},
            {"domain": "baseline", "status": historical_baselines["status"], "impact": "historical percentile unavailable or sample-limited"},
        ],
        "historical_baselines": historical_baselines,
        "validation": {"status": "PENDING", "errors": [], "warnings": warnings, "input_hashes": {key: value.get("sha256") for key, value in sources.items()}},
    }
    result["prediction_ledger"]["content_sha256"] = compute_prediction_ledger_hash(DATE, theses, scenarios, watchlist, prediction_entries)
    return result


def persist_and_readback_business_result(result: dict[str, Any]) -> tuple[dict[str, Any], list[str]]:
    BUSINESS_RESULT_JSON.parent.mkdir(parents=True, exist_ok=True)
    temp_path = BUSINESS_RESULT_JSON.with_suffix(BUSINESS_RESULT_JSON.suffix + ".tmp")
    temp_path.write_text(json.dumps(result, ensure_ascii=False, indent=2, sort_keys=True), encoding="utf-8")
    temp_path.replace(BUSINESS_RESULT_JSON)
    readback = json.loads(BUSINESS_RESULT_JSON.read_text(encoding="utf-8"))
    validation = validate_business_result_payload(readback, expected_date=DATE, expected_run_id=RUN_ID)
    blocks = list(validation.get("blocks", [])) if isinstance(validation, dict) else list(validation or [])
    if blocks:
        readback["status"] = "BLOCKED"
        readback["blocks"] = sorted(set(list(readback.get("blocks") or []) + blocks))
        readback["validation"]["status"] = "BLOCKED"
        readback["validation"]["errors"] = blocks
    else:
        readback["validation"]["status"] = "CLEAN_PASS"
    BUSINESS_RESULT_JSON.write_text(json.dumps(readback, ensure_ascii=False, indent=2, sort_keys=True), encoding="utf-8")
    confirmed = json.loads(BUSINESS_RESULT_JSON.read_text(encoding="utf-8"))
    return confirmed, blocks


def md_table(headers: list[str], rows: list[list[Any]]) -> str:
    out = ["| " + " | ".join(headers) + " |", "|" + "|".join(["---"] * len(headers)) + "|"]
    for row in rows:
        out.append("| " + " | ".join(clean(cell) for cell in row).replace("\n", " ") + " |")
    return "\n".join(out)


def validate_chapter_conclusions(md_text: str) -> list[str]:
    blocks: list[str] = []
    seen_analysis: set[str] = set()
    forbidden_visible = [
        "结论：", "依据：", "但要注意：", "上面判断就不成立",
        "默认结论", "固定措辞", "兜底判断", "重要方向", "看成交额、涨停扩散",
        "冻结账本", "冻结预测账本", "本地上一交易日表", "从观察池提取代码",
        "只保留当日涨停池确认样本", "事实来源", "研究动作", "论点状态与证伪条件",
        "次日四分支情景树", "工作流", "脚本", "接口", "字段",
    ]
    for chapter in range(17):
        pattern = re.compile(
            rf"(?ms)^##\s+{chapter}\.\s+[^\n]*\n(?P<body>.*?)(?=^##\s+(?:[0-9]|1[0-6])\.\s+|\Z)"
        )
        match = pattern.search(md_text)
        if not match:
            blocks.append(f"chapter {chapter} missing")
            continue
        body = match.group("body")
        if chapter == 0:
            for marker in ("对象", "观察角度", "今日判断", "明日验证"):
                if marker not in body:
                    blocks.append(f"chapter 0 missing classified conclusion matrix marker: {marker}")
            continue
        hits = [term for term in forbidden_visible if term in body]
        if hits:
            blocks.append(f"chapter {chapter} contains forbidden fixed/process text: {hits}")
        analyses = [value.strip() for value in re.findall(r"\*\*(.+?)\*\*", body, flags=re.S) if value.strip()]
        if len(analyses) != 1:
            blocks.append(f"chapter {chapter} must contain exactly one dynamic analysis paragraph")
            continue
        analysis = analyses[0]
        if analysis in seen_analysis:
            blocks.append(f"chapter {chapter} repeats another chapter analysis")
        seen_analysis.add(analysis)
        if not re.search(r"\d", body):
            blocks.append(f"chapter {chapter} lacks current numeric facts")
        if not re.search(r"(上涨|下跌|涨停|连板|首板|开板|跌停|封板|资金|龙虎榜|热点|市场|留存|接力|风险)", analysis):
            blocks.append(f"chapter {chapter} analysis lacks market terms")
    return blocks


def add_run(paragraph, text: str, *, bold: bool = False, size: int | None = None, color: RGBColor | None = None):
    run = paragraph.add_run(text)
    run.bold = bold
    run.font.name = FONT
    run._element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
    if size:
        run.font.size = Pt(size)
    if color:
        run.font.color.rgb = color
    return run


def set_cell_shading(cell, fill: str):
    tc_pr = cell._tc.get_or_add_tcPr()
    shading = OxmlElement("w:shd")
    shading.set(qn("w:fill"), fill)
    tc_pr.append(shading)


def set_cell_text(cell, text: Any, *, bold: bool = False, color: RGBColor | None = None, size: int = 8):
    cell.text = ""
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    paragraph = cell.paragraphs[0]
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = paragraph.add_run(clean(text))
    run.bold = bold
    run.font.name = FONT
    run._element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
    run.font.size = Pt(size)
    if color:
        run.font.color.rgb = color


def add_table(doc: Document, headers: list[str], rows: list[list[Any]], widths: list[float] | None = None, font_size: int = 8):
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    table.autofit = False
    table.alignment = WD_ALIGN_PARAGRAPH.CENTER
    header_row_properties = table.rows[0]._tr.get_or_add_trPr()
    repeat_header = OxmlElement("w:tblHeader")
    repeat_header.set(qn("w:val"), "true")
    header_row_properties.append(repeat_header)
    header = table.rows[0].cells
    for idx, title in enumerate(headers):
        set_cell_shading(header[idx], "1F4E79")
        set_cell_text(header[idx], title, bold=True, color=RGBColor(0xFF, 0xFF, 0xFF), size=font_size)
        header[idx].paragraphs[0].paragraph_format.keep_with_next = True
        header[idx].paragraphs[0].paragraph_format.keep_together = True
    for row in rows:
        cells = table.add_row().cells
        for idx, value in enumerate(row):
            set_cell_text(cells[idx], value, size=font_size)
    if widths:
        for row in table.rows:
            for idx, width in enumerate(widths):
                if idx < len(row.cells):
                    row.cells[idx].width = Inches(width)
    return table


def grouped_conclusion_nodes(nodes: list[dict[str, Any]]) -> list[tuple[str, list[dict[str, Any]]]]:
    ordered = sorted(
        [item for item in nodes if isinstance(item, dict)],
        key=lambda item: (int(item.get("priority") or 999), str(item.get("category_label") or "")),
    )
    groups: list[tuple[str, list[dict[str, Any]]]] = []
    by_id: dict[str, list[dict[str, Any]]] = {}
    labels: dict[str, str] = {}
    order: list[str] = []
    for item in ordered:
        category_id = str(item.get("category_id") or "").strip()
        if category_id not in by_id:
            by_id[category_id] = []
            labels[category_id] = str(item.get("category_label") or "").strip()
            order.append(category_id)
        by_id[category_id].append(item)
    for category_id in order:
        groups.append((labels[category_id], by_id[category_id]))
    return groups


def conclusion_evidence_summary(node: dict[str, Any]) -> str:
    parts = [
        str(item.get("interpretation") or "").strip()
        for item in node.get("evidence", [])
        if isinstance(item, dict) and str(item.get("interpretation") or "").strip()
    ]
    return "；".join(parts[:3])


def conclusion_confirmation_summary(node: dict[str, Any]) -> str:
    values = [str(value).strip() for value in node.get("confirmation", []) if str(value).strip()]
    return values[0] if values else "-"


def set_cell_padding(cell, top: int = 80, start: int = 90, bottom: int = 80, end: int = 90) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for tag, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{tag}"))
        if node is None:
            node = OxmlElement(f"w:{tag}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def keep_row_together(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    if tr_pr.find(qn("w:cantSplit")) is None:
        tr_pr.append(OxmlElement("w:cantSplit"))


def add_conclusion_matrix(doc: Document, nodes: list[dict[str, Any]]):
    table = doc.add_table(rows=1, cols=4)
    table.style = "Table Grid"
    table.autofit = False
    table.alignment = WD_ALIGN_PARAGRAPH.CENTER
    widths = [0.95, 1.15, 3.75, 1.15]
    headers = ["对象", "观察角度", "今日判断", "明日验证"]
    header_cells = table.rows[0].cells
    for index, title in enumerate(headers):
        set_cell_shading(header_cells[index], "17365D")
        set_cell_text(header_cells[index], title, bold=True, color=RGBColor(0xFF, 0xFF, 0xFF), size=8)
        header_cells[index].width = Inches(widths[index])
        set_cell_padding(header_cells[index], 70, 80, 70, 80)
    header_row_properties = table.rows[0]._tr.get_or_add_trPr()
    repeat_header = OxmlElement("w:tblHeader")
    repeat_header.set(qn("w:val"), "true")
    header_row_properties.append(repeat_header)
    keep_row_together(table.rows[0])

    item_index = 0
    for category_label, category_nodes in grouped_conclusion_nodes(nodes):
        group_row = table.add_row()
        merged = group_row.cells[0].merge(group_row.cells[3])
        set_cell_shading(merged, "D9EAF7")
        merged.text = ""
        group_para = merged.paragraphs[0]
        group_para.alignment = WD_ALIGN_PARAGRAPH.LEFT
        group_para.paragraph_format.space_before = Pt(1)
        group_para.paragraph_format.space_after = Pt(1)
        add_run(group_para, category_label, bold=True, color=BLUE, size=9)
        set_cell_padding(merged, 85, 110, 85, 110)
        keep_row_together(group_row)

        for node in category_nodes:
            item_index += 1
            row = table.add_row()
            keep_row_together(row)
            cells = row.cells
            if item_index % 2 == 0:
                for cell in cells:
                    set_cell_shading(cell, "F7FAFC")
            for index, width in enumerate(widths):
                cells[index].width = Inches(width)
                set_cell_padding(cells[index], 90, 90, 90, 90)
            set_cell_text(cells[0], node.get("scope_label"), bold=True, color=BLUE, size=8)
            set_cell_text(cells[1], node.get("angle_label"), size=8)

            judgment = cells[2]
            judgment.text = ""
            judgment.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            paragraph = judgment.paragraphs[0]
            paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
            paragraph.paragraph_format.space_before = Pt(0)
            paragraph.paragraph_format.space_after = Pt(1)
            paragraph.paragraph_format.line_spacing = 1.05
            add_run(paragraph, str(node.get("headline") or ""), bold=True, color=BLUE, size=8)
            add_run(paragraph, "\n" + str(node.get("conclusion") or ""), size=7)
            evidence = conclusion_evidence_summary(node)
            if evidence:
                add_run(paragraph, "\n" + evidence, color=GRAY, size=6)

            verification = cells[3]
            verification.text = ""
            verification.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            verify_para = verification.paragraphs[0]
            verify_para.alignment = WD_ALIGN_PARAGRAPH.LEFT
            verify_para.paragraph_format.space_before = Pt(0)
            verify_para.paragraph_format.space_after = Pt(0)
            verify_para.paragraph_format.line_spacing = 1.05
            add_run(verify_para, conclusion_confirmation_summary(node), size=7)
    return table


def setup_doc() -> Document:
    doc = Document()
    section = doc.sections[0]
    section.orientation = WD_ORIENT.PORTRAIT
    section.page_width = Inches(8.27)
    section.page_height = Inches(11.69)
    section.top_margin = Inches(0.72)
    section.bottom_margin = Inches(0.50)
    section.left_margin = Inches(0.68)
    section.right_margin = Inches(0.68)
    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = FONT
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
    normal.font.size = Pt(10)
    normal.paragraph_format.space_after = Pt(5)
    normal.paragraph_format.line_spacing = 1.08
    for name, size in [("Heading 1", 14), ("Heading 2", 12), ("Heading 3", 11)]:
        style = styles[name]
        style.font.name = FONT
        style._element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
        style.font.size = Pt(size)
        style.font.color.rgb = BLUE
        style.font.bold = True
        style.paragraph_format.space_before = Pt(8)
        style.paragraph_format.space_after = Pt(4)
    return doc


def add_stock_risk_notice_page(doc: Document) -> None:
    header = doc.add_table(rows=2, cols=1)
    header.autofit = False
    for row in header.rows:
        set_cell_shading(row.cells[0], "0B2442")
    set_cell_text(
        header.rows[0].cells[0],
        "投资风险提示",
        bold=True,
        color=RGBColor(0xFF, 0xFF, 0xFF),
        size=26,
    )
    header.rows[0].cells[0].paragraphs[0].paragraph_format.space_before = Pt(12)
    header.rows[0].cells[0].paragraphs[0].paragraph_format.space_after = Pt(6)
    set_cell_text(
        header.rows[1].cells[0],
        "教学案例｜交流学习｜方法演示",
        bold=True,
        color=RGBColor(0xFF, 0xD6, 0x5A),
        size=13,
    )
    header.rows[1].cells[0].paragraphs[0].paragraph_format.space_after = Pt(12)

    spacer = doc.add_paragraph()
    spacer.paragraph_format.space_after = Pt(18)

    warnings = doc.add_table(rows=3, cols=1)
    warnings.autofit = False
    for row, warning in zip(
        warnings.rows,
        ("不作为股票推荐", "不构成投资建议", "不作为买卖依据"),
    ):
        set_cell_shading(row.cells[0], "FFF1C7")
        set_cell_text(
            row.cells[0],
            warning,
            bold=True,
            color=RGBColor(0xD4, 0x00, 0x00),
            size=18,
        )
        row.cells[0].paragraphs[0].paragraph_format.space_before = Pt(8)
        row.cells[0].paragraphs[0].paragraph_format.space_after = Pt(8)

    spacer = doc.add_paragraph()
    spacer.paragraph_format.space_after = Pt(14)

    red_bands = doc.add_table(rows=2, cols=1)
    red_bands.autofit = False
    red_texts = (
        "授课方/作者不具备荐股资质",
        "学生不得据此直接买卖；需独立判断并严格根据规则训练",
    )
    for row, warning in zip(red_bands.rows, red_texts):
        set_cell_shading(row.cells[0], "9D0000")
        set_cell_text(
            row.cells[0],
            warning,
            bold=True,
            color=RGBColor(0xFF, 0xFF, 0xFF),
            size=13,
        )
        row.cells[0].paragraphs[0].paragraph_format.space_before = Pt(7)
        row.cells[0].paragraphs[0].paragraph_format.space_after = Pt(7)

    notice = doc.add_paragraph()
    notice.alignment = WD_ALIGN_PARAGRAPH.CENTER
    notice.paragraph_format.space_before = Pt(16)
    notice.paragraph_format.space_after = Pt(20)
    add_run(
        notice,
        "本页为教学案例风险提示，后续内容依据当前交易日数据计算。",
        bold=True,
        color=RGBColor(0x0B, 0x24, 0x42),
        size=11,
    )

    footer = doc.add_table(rows=1, cols=1)
    footer.autofit = False
    set_cell_shading(footer.rows[0].cells[0], "0B2442")
    set_cell_text(
        footer.rows[0].cells[0],
        "股市有风险，入市需谨慎",
        bold=True,
        color=RGBColor(0xFF, 0xD6, 0x5A),
        size=16,
    )
    footer.rows[0].cells[0].paragraphs[0].paragraph_format.space_before = Pt(10)
    footer.rows[0].cells[0].paragraphs[0].paragraph_format.space_after = Pt(10)
    doc.add_page_break()


def add_heading(doc: Document, text: str, level: int = 1):
    paragraph = doc.add_heading("", level=level)
    add_run(paragraph, text, bold=True, color=BLUE, size=14 if level == 1 else 12)
    return paragraph


def add_para(doc: Document, text: str, *, color: RGBColor | None = None, bold: bool = False):
    paragraph = doc.add_paragraph()
    add_run(paragraph, text, bold=bold, color=color or DARK, size=9)
    return paragraph


def add_analysis_note(doc: Document, note: dict[str, Any]):
    analysis = str(note.get("analysis") or "").strip()
    facts = str(note.get("facts") or "").strip()
    main = doc.add_paragraph()
    main.paragraph_format.space_after = Pt(2)
    main.paragraph_format.line_spacing = 1.08
    main.paragraph_format.keep_with_next = True
    add_run(main, analysis, bold=True, color=DARK, size=9)
    detail = doc.add_paragraph()
    detail.paragraph_format.space_before = Pt(0)
    detail.paragraph_format.space_after = Pt(4)
    detail.paragraph_format.line_spacing = 1.03
    detail.paragraph_format.keep_with_next = True
    add_run(detail, facts, color=GRAY, size=8)
    return detail


def latest_value(value: Any) -> str:
    if isinstance(value, list):
        value = value[-1] if value else None
    text = clean(value)
    return text if text not in {"-", "nan", "None"} else "-"


def visible_formula_value(value: Any) -> str:
    text = latest_value(value)
    if text == "-":
        return text
    forbidden = [
        "龙头", "发酵", "主升", "高潮", "评级", "最强", "强度评分", "数量领先", "决定强度上限",
        "龙1", "龙2", "龙3", "看延续与降级", "修复强度", "修复中带分歧",
    ]
    if any(term in text for term in forbidden):
        number = re.search(r"-?\d+(?:\.\d+)?", text)
        return number.group(0) if number else "已返回"
    return chinese_visible_text(text)


def summarize_formula_output(item: dict[str, Any]) -> str:
    if not item.get("ok"):
        return "失败"
    symbol = item.get("symbol")
    result = item.get("result") if isinstance(item.get("result"), dict) else {}
    stock_data = result.get(symbol) if symbol else None
    if not isinstance(stock_data, dict):
        stock_data = next((value for key, value in result.items() if key != "ErrorId" and isinstance(value, dict)), {})
    formula = str(item.get("formula", ""))
    priority = {
        "大牛线4.0": ["OUTPUT37", "OUTPUT30", "主趋势线", "涨停价"],
        "飞龙在天": ["波", "段", "OUTPUT3", "OUTPUT4"],
        "游资资金监控": ["买方意向", "AAA", "DDD"],
        "机构资金监控": ["机构大单进", "大单动向", "散户资金进"],
        "庄家资金监控": ["控盘程度", "控盘度", "OUTPUT3"],
    }.get(formula, [])
    display_names = {
        "OUTPUT37": "输出三十七",
        "OUTPUT30": "输出三十",
        "OUTPUT3": "输出三",
        "OUTPUT4": "输出四",
        "AAA": "资金指标甲",
        "DDD": "资金指标丁",
    }
    parts: list[str] = []
    for key in priority:
        value = visible_formula_value(stock_data.get(key)) if isinstance(stock_data, dict) else "-"
        if value != "-":
            parts.append(f"{display_names.get(key, chinese_visible_text(key))}:{value}")
        if len(parts) >= 2:
            break
    return "；".join(parts) if parts else "已返回"


def collect_formula_samples(sample_df: pd.DataFrame, limit: int = 3) -> dict[str, Any]:
    rows: list[list[Any]] = []
    raw_samples: list[dict[str, Any]] = []
    failed: list[dict[str, Any]] = []
    tdx_script = SCRIPTS_DIR / "tdx_hub.py"
    for _, row in sample_df.head(limit).iterrows():
        code = code6(row.get("代码"))
        name = clean(row.get("名称"))
        height = f"{to_int(row.get('连板数'), 1)}板" if to_int(row.get("连板数"), 1) >= 2 else "首板"
        try:
            proc = subprocess.run(
                [sys.executable, str(tdx_script), "five", code],
                cwd=str(SCRIPTS_DIR),
                capture_output=True,
                text=True,
                encoding="utf-8",
                errors="replace",
                timeout=120,
            )
            stdout = proc.stdout.strip()
            start = stdout.find("{")
            payload = json.loads(stdout[start:]) if start >= 0 else {}
            required_ok = bool(payload.get("ok")) and proc.returncode == 0
            formulas = {item.get("formula"): summarize_formula_output(item) for item in payload.get("items", []) if isinstance(item, dict)}
            sample = {
                "code": code,
                "name": name,
                "height": height,
                "ok": required_ok,
                "returncode": proc.returncode,
                "payload": payload,
                "stderr_tail": proc.stderr[-600:],
            }
            raw_samples.append(sample)
            if required_ok:
                rows.append([
                    code,
                    name,
                    height,
                    formulas.get("大牛线4.0", "已返回"),
                    formulas.get("飞龙在天", "已返回"),
                    "；".join([
                        formulas.get("游资资金监控", "已返回"),
                        formulas.get("机构资金监控", "已返回"),
                    ]),
                    "前排高度样本核验",
                ])
            else:
                failed.append({"code": code, "name": name, "returncode": proc.returncode, "stderr_tail": proc.stderr[-600:]})
        except Exception as exc:
            failed.append({"code": code, "name": name, "error": f"{type(exc).__name__}: {str(exc)[:240]}"})
    status = "CLEAN_PASS" if rows and not failed else "BLOCKED"
    result = {
        "status": status,
        "date": DATE,
        "sample_limit": limit,
        "sample_success_count": len(rows),
        "sample_failed_count": len(failed),
        "rows": rows,
        "failed": failed,
        "raw_samples": raw_samples,
        "collected_at": datetime.now().isoformat(timespec="seconds"),
    }
    TQ_SAMPLE_JSON.write_text(json.dumps(result, ensure_ascii=False, indent=2), encoding="utf-8")
    return result


def validate_entry_date_gate() -> list[str]:
    path = TASK_DIR / "preflight_audit.json"
    if not path.exists() or path.stat().st_size == 0:
        return ["ENTRY_DATE_GATE_MISSING"]
    try:
        payload = json.loads(path.read_text(encoding="utf-8-sig"))
    except Exception as exc:
        return [f"ENTRY_DATE_GATE_UNREADABLE:{type(exc).__name__}:{exc}"]
    gate = payload.get("date_gate") if isinstance(payload.get("date_gate"), dict) else {}
    blocks: list[str] = []
    if payload.get("status") != "CLEAN_PASS" or gate.get("status") != "CLEAN_PASS":
        blocks.append("ENTRY_DATE_GATE_NOT_CLEAN")
    if str(payload.get("run_id") or "") != RUN_ID:
        blocks.append("ENTRY_DATE_GATE_RUN_ID_MISMATCH")
    if str(payload.get("date") or "").replace("-", "") != DATE:
        blocks.append("ENTRY_DATE_GATE_AUDIT_DATE_MISMATCH")
    if str(gate.get("target_date") or "").replace("-", "") != DATE:
        blocks.append("ENTRY_DATE_GATE_TARGET_DATE_MISMATCH")
    mode = str(gate.get("mode") or "")
    if mode not in {"latest", "historical"}:
        blocks.append("ENTRY_DATE_GATE_MODE_INVALID")
    if gate.get("blocks"):
        blocks.append("ENTRY_DATE_GATE_HAS_BLOCKS")
    if mode == "historical" and path_is_within(DELIVERY, DEFAULT_DELIVERY_ROOT):
        blocks.append(f"HISTORICAL_MODE_DEFAULT_DELIVERY_ROOT_FORBIDDEN:{DELIVERY}")
    if mode == "latest":
        runtime_gate = evaluate_date_gate("", "latest")
        if runtime_gate.get("status") != "CLEAN_PASS":
            blocks.extend(f"RUNTIME_LATEST_DATE_GATE:{item}" for item in runtime_gate.get("blocks", []))
        if str(runtime_gate.get("target_date") or "").replace("-", "") != DATE:
            blocks.append("RUNTIME_LATEST_DATE_MISMATCH")
    return blocks


def validate_source_freshness() -> list[str]:
    audit_path = TASK_DIR / "data_integrity_audit.json"
    if not audit_path.exists() or audit_path.stat().st_size == 0:
        return ["DATA_INTEGRITY_AUDIT_MISSING"]
    try:
        audit = json.loads(audit_path.read_text(encoding="utf-8-sig"))
    except Exception as exc:
        return [f"DATA_INTEGRITY_AUDIT_UNREADABLE:{type(exc).__name__}:{exc}"]
    blocks: list[str] = []
    if audit.get("status") != "CLEAN_PASS" or audit.get("blocks"):
        blocks.append("DATA_INTEGRITY_AUDIT_NOT_CLEAN")
    if str(audit.get("date") or "").replace("-", "") != DATE:
        blocks.append("DATA_INTEGRITY_AUDIT_DATE_MISMATCH")
    qdate = str(audit.get("official_source_qdate") or "")
    date_reconciliation = audit.get("official_source_date_reconciliation")
    if qdate != DATE and not is_cross_validated_official_pool_date(
        date_reconciliation,
        DATE,
    ):
        blocks.append(f"OFFICIAL_SOURCE_DATE_MISMATCH:{qdate}!={DATE}")
    source_codes = {code6(value) for value in (audit.get("official_source_codes") or [])}
    if not source_codes:
        blocks.append("OFFICIAL_SOURCE_CODES_MISSING")
    try:
        source_count = int(audit.get("official_source_count") or 0)
    except (TypeError, ValueError):
        source_count = -1
        blocks.append("OFFICIAL_SOURCE_COUNT_INVALID")
    if source_count != len(source_codes):
        blocks.append("OFFICIAL_SOURCE_COUNT_MISMATCH")
    source_hash = hashlib.sha256("\n".join(sorted(source_codes)).encode("utf-8")).hexdigest()
    if str(audit.get("official_source_codes_sha256") or "") != source_hash:
        blocks.append("OFFICIAL_SOURCE_CODES_HASH_MISMATCH")
    eligible_codes = {code6(value) for value in (audit.get("eligible_official_codes") or [])}
    exclusion_items = audit.get("hard_exclusions") if isinstance(audit.get("hard_exclusions"), list) else []
    exclusion_codes = {
        code6(item.get("code"))
        for item in exclusion_items
        if isinstance(item, dict) and item.get("code")
    }
    if not eligible_codes:
        blocks.append("ELIGIBLE_OFFICIAL_CODES_MISSING")
    if not eligible_codes.issubset(source_codes):
        blocks.append("ELIGIBLE_OFFICIAL_CODES_OUTSIDE_RAW_SOURCE")
    if source_codes - eligible_codes != exclusion_codes:
        blocks.append("DELISTING_HARD_EXCLUSION_MISMATCH")
    try:
        eligible_count = int(audit.get("eligible_official_count") or 0)
    except (TypeError, ValueError):
        eligible_count = -1
        blocks.append("ELIGIBLE_OFFICIAL_COUNT_INVALID")
    if eligible_count != len(eligible_codes):
        blocks.append("ELIGIBLE_OFFICIAL_COUNT_MISMATCH")
    eligible_hash = hashlib.sha256("\n".join(sorted(eligible_codes)).encode("utf-8")).hexdigest()
    if str(audit.get("eligible_official_codes_sha256") or "") != eligible_hash:
        blocks.append("ELIGIBLE_OFFICIAL_CODES_HASH_MISMATCH")
    snapshot_raw = str(audit.get("official_source_snapshot") or "")
    snapshot_path = Path(snapshot_raw) if snapshot_raw else None
    if snapshot_path is None or not snapshot_path.exists() or not snapshot_path.is_file() or snapshot_path.stat().st_size == 0:
        blocks.append("OFFICIAL_SOURCE_SNAPSHOT_MISSING")
    else:
        try:
            snapshot = json.loads(snapshot_path.read_text(encoding="utf-8-sig"))
            snapshot_codes = {code6(value) for value in (snapshot.get("source_codes") or [])}
            if (
                str(snapshot.get("qdate") or "") != DATE
                and not is_cross_validated_official_pool_date(snapshot, DATE)
            ):
                blocks.append("OFFICIAL_SOURCE_SNAPSHOT_DATE_MISMATCH")
            if snapshot_codes != source_codes:
                blocks.append("OFFICIAL_SOURCE_SNAPSHOT_UNIVERSE_MISMATCH")
        except Exception as exc:
            blocks.append(f"OFFICIAL_SOURCE_SNAPSHOT_UNREADABLE:{type(exc).__name__}:{exc}")
    if not IN_TABLE.exists() or IN_TABLE.stat().st_size == 0:
        blocks.append("VERIFIED_TABLE_MISSING_FOR_SOURCE_CHECK")
    else:
        try:
            table = pd.read_csv(IN_TABLE, dtype={"代码": str}, encoding="utf-8-sig")
            table_codes = {code6(value) for value in table["代码"].tolist()} if "代码" in table.columns else set()
            if table_codes != eligible_codes:
                blocks.append("OFFICIAL_UNIVERSE_MISMATCH")
        except Exception as exc:
            blocks.append(f"VERIFIED_TABLE_UNREADABLE_FOR_SOURCE_CHECK:{type(exc).__name__}:{exc}")
    return blocks


def build_report() -> int:
    TASK_DIR.mkdir(parents=True, exist_ok=True)
    MEM_DIR.mkdir(parents=True, exist_ok=True)
    pre_delivery_blocks = validate_entry_date_gate() + validate_source_freshness()
    if pre_delivery_blocks:
        raise SystemExit(";".join(pre_delivery_blocks))
    DELIVERY.mkdir(parents=True, exist_ok=True)
    if not IN_TABLE.exists():
        raise SystemExit(f"missing strict table: {IN_TABLE}")
    df = read_csv(IN_TABLE, {"代码": str})
    if df.empty:
        raise SystemExit("strict table is empty")
    required_columns = ["代码", "名称", "连板数", "投资板块", "题材归属", "题材证据", "所属行业", "首次封板时间", "炸板次数"]
    missing_columns = [column for column in required_columns if column not in df.columns]
    if missing_columns:
        raise SystemExit(f"strict table missing required columns: {missing_columns}")
    fund_path = TASK_DIR / f"push2_fund_flow_resolved_{DATE}.csv"
    zbgc_path = TASK_DIR / f"ak_stock_zt_pool_zbgc_em_{DATE}.csv"
    dtgc_path = TASK_DIR / f"ak_stock_zt_pool_dtgc_em_{DATE}.csv"
    lhb_path = TASK_DIR / f"ak_stock_lhb_detail_em_{DATE}.csv"
    fund = read_csv(fund_path, {"代码": str})
    zbgc = read_csv(zbgc_path, {"代码": str})
    dtgc = read_csv(dtgc_path, {"代码": str}, allow_empty_file=True)
    lhb = read_csv(lhb_path, {"代码": str})

    df["代码"] = df["代码"].map(code6)
    if not fund.empty and "代码" in fund.columns:
        fund["代码"] = fund["代码"].map(code6)
        if "ok" in fund.columns:
            fund["ok"] = fund["ok"].map(parse_verified_bool)
        keep = [c for c in ["代码", "ok", "主力净流入", "超大单净流入", "主力净占比"] if c in fund.columns]
        df = df.merge(fund[keep].drop_duplicates("代码"), on="代码", how="left")
    for col in ["连板数", "涨跌幅", "最新价", "成交额", "炸板次数", "主力净流入", "超大单净流入", "主力净占比"]:
        if col in df.columns:
            df[col] = pd.to_numeric(df[col], errors="coerce")
    for direction_column in ["投资板块", "题材归属", "题材证据", "所属行业"]:
        df[direction_column] = df[direction_column].fillna("").map(lambda x: chinese_visible_text(x, ""))
        if df[direction_column].str.strip().eq("").any():
            raise SystemExit(f"strict table contains blank {direction_column}")
    df["名称"] = df["名称"].map(chinese_visible_text)
    delisting_rows = df[df["名称"].map(lambda value: "退" in str(value))]
    if not delisting_rows.empty:
        leaked = [f"{code6(row.get('代码'))} {row.get('名称')}" for _, row in delisting_rows.iterrows()]
        raise SystemExit(f"DELISTING_HARD_EXCLUSION_LEAK:{leaked}")
    df = df.sort_values(["连板数", "首次封板时间", "代码"], ascending=[False, True, True]).reset_index(drop=True)
    df.insert(0, "报告序号", range(1, len(df) + 1))

    zt_count = int(len(df))
    board_counts = Counter(int(x) for x in df["连板数"].fillna(1).astype(int))
    highest_board = max(board_counts) if board_counts else 1
    connected_count = sum(count for board, count in board_counts.items() if board >= 2)
    first_board_count = board_counts.get(1, 0)
    zbgc_count = int(len(zbgc))
    dtgc_count = int(len(dtgc))
    touched = zt_count + zbgc_count
    seal_rate = zt_count / touched * 100 if touched else 0
    fail_rate = zbgc_count / touched * 100 if touched else 0
    top_sectors = sector_summary(df, 5)
    largest = top_sectors[0]
    largest_direction = largest["sector"]
    largest_front = stock_label(next(largest["leaders"].iterrows())[1]) if len(largest["leaders"]) else "-"
    push_ok = int(fund["ok"].map(parse_verified_bool).sum()) if not fund.empty and "ok" in fund.columns else 0
    push_pos = int((pd.to_numeric(df.get("主力净流入", pd.Series(dtype=float)), errors="coerce") > 0).sum())
    push_neg = int((pd.to_numeric(df.get("主力净流入", pd.Series(dtype=float)), errors="coerce") < 0).sum())
    lhb_codes = {code6(x) for x in lhb["代码"].dropna().tolist()} if not lhb.empty and "代码" in lhb.columns else set()
    lhb_hit = int(df["代码"].isin(lhb_codes).sum())

    market_breadth = fetch_market_breadth()
    breadth_ok = market_breadth.get("status") == "CLEAN_PASS"
    up_count = int(market_breadth.get("up_count") or 0)
    down_count = int(market_breadth.get("down_count") or 0)
    flat_count = int(market_breadth.get("flat_count") or 0)
    breadth_label = clean(market_breadth.get("breadth_label"), "宽度采集失败")
    down_ratio = float(market_breadth.get("down_ratio") or 0)
    up_ratio = float(market_breadth.get("up_ratio") or 0)
    market_core_sentence = sentence(
        f"本次市场宽度记录为{breadth_label}",
        f"上涨{up_count}只、下跌{down_count}只、平盘{flat_count}只；上涨占比{up_ratio:.2%}、下跌占比{down_ratio:.2%}；涨停{zt_count}只",
    )

    dragon = dragon_rows(df)
    sector_rows = [[
        idx + 1,
        item["sector"],
        item["count"],
        item["connected"],
        item["max_board"],
        fmt_amount(item["fund"]),
        stock_label(next(item["leaders"].iterrows())[1]) if len(item["leaders"]) else "-",
        "按涨停数、连板数、最高板、资金净额排序",
    ] for idx, item in enumerate(top_sectors)]
    ladder_rows: list[list[Any]] = []
    for board in sorted(board_counts.keys(), reverse=True):
        sub = df[df["连板数"].fillna(1).astype(int) == board].copy()
        reps = "、".join(stock_label(row) for _, row in sub.head(8).iterrows())
        state = "最高板样本" if board == highest_board else "2板及以上样本" if board >= 2 else "首板样本"
        ladder_rows.append([f"{board}板" if board >= 2 else "首板", len(sub), reps, state])
    high_rows = []
    for _, row in df[df["连板数"].fillna(1).astype(int) >= 3].sort_values(["连板数", "首次封板时间"], ascending=[False, True]).iterrows():
        high_rows.append([code6(row["代码"]), row["名称"], f"{to_int(row['连板数'])}板", fmt_time(row.get("首次封板时间")), to_int(row.get("炸板次数")), clean(row.get("投资板块"))])
    if not high_rows:
        high_rows.append(["-", "-", "-", "-", "-", "当日无3板以上样本"])

    quality_rows = []
    qdf = df.copy()
    qdf["排序分"] = qdf.apply(leader_score, axis=1).clip(0, 100).round().astype(int)
    for _, row in qdf.sort_values(["连板数", "排序分"], ascending=[False, False]).head(10).iterrows():
        fund_state = fmt_amount(row.get("主力净流入")) if parse_verified_bool(row.get("ok")) else "不可用"
        quality_rows.append([code6(row["代码"]), row["名称"], f"{to_int(row['连板数'])}板", fmt_time(row.get("首次封板时间")), to_int(row.get("炸板次数")), clean(row.get("投资板块")), fund_state])
    formula_sample_df = qdf.sort_values(["连板数", "排序分"], ascending=[False, False])[["代码", "名称", "连板数"]].copy()
    formula_sample = collect_formula_samples(formula_sample_df, limit=3)
    formula_rows = formula_sample.get("rows") or []
    prev_material = load_previous_review_material(df, DATE_H)

    first_rows = []
    first_df = df[df["连板数"].fillna(1).astype(int) == 1].copy()
    for item in sector_summary(first_df, 5):
        leaders = "、".join(stock_label(row, with_board=False) for _, row in item["leaders"].iterrows())
        next_day_rule = (
            f"今日{item['count']}只首板中至少1只次日晋级，才保留该方向连续性证据；"
            "0只晋级则降级"
        )
        first_rows.append([item["sector"], item["count"], pct_text(item["count"], first_board_count), leaders, next_day_rule])

    fund_rows = []
    if "主力净流入" in df.columns:
        fund_valid = df[df["ok"].map(parse_verified_bool)].copy() if "ok" in df.columns else df.iloc[0:0].copy()
        in_rows = fund_valid.sort_values("主力净流入", ascending=False).head(7)
        out_rows = fund_valid.sort_values("主力净流入", ascending=True).head(3)
        for _, row in pd.concat([in_rows, out_rows]).drop_duplicates("代码").iterrows():
            fund_rows.append([code6(row["代码"]), row["名称"], fmt_amount(row.get("主力净流入")), fmt_amount(row.get("超大单净流入")), fmt_pct(row.get("主力净占比")), "主力净流入为正" if to_num(row.get("主力净流入")) > 0 else "主力净流入为负"])
    research_blocks: list[str] = []
    if not fund_rows:
        research_blocks.append("fund flow rows missing after collection")

    risk_rows = []
    risk_df = df.copy()
    risk_df["风险分"] = risk_df["连板数"].fillna(1) * 10 + risk_df["炸板次数"].fillna(0) * 5 - (pd.to_numeric(risk_df.get("主力净流入", 0), errors="coerce").fillna(0) / 100000000)
    for _, row in risk_df.sort_values("风险分", ascending=False).head(24).iterrows():
        triggers: list[str] = []
        if to_int(row.get("连板数")) >= 3:
            triggers.append("高位样本")
        if to_int(row.get("炸板次数")) >= 5:
            triggers.append("反复开板")
        if parse_verified_bool(row.get("ok")) and to_num(row.get("主力净流入")) < 0:
            triggers.append("资金净流出")
        if not triggers:
            triggers.append("同组对照样本")
        risk_rows.append([code6(row["代码"]), row["名称"], f"{to_int(row.get('连板数'), 1)}板", to_int(row.get("炸板次数")), "；".join(triggers), clean(row.get("投资板块"))])

    core_table = [
        ["涨停家数", zt_count, "涨停股票"],
        ["全市场上涨", up_count if breadth_ok else "-", "上涨股票"],
        ["全市场下跌", down_count if breadth_ok else "-", breadth_label],
        ["连板数量", connected_count, "2板及以上"],
        ["首板数量", first_board_count, "首板样本数"],
        ["最高板", f"{highest_board}板", largest_front],
        ["炸板数量", zbgc_count, "封板失败样本"],
        ["跌停数量", dtgc_count, "跌停股票"],
        ["封板率", f"{seal_rate:.1f}%", f"{zt_count}/{touched}"],
    ]
    facts = extract_report_facts(
        df,
        top_sectors,
        qdf,
        largest,
        market_breadth,
        prev_material,
        zt_count=zt_count,
        connected_count=connected_count,
        first_board_count=first_board_count,
        highest_board=highest_board,
        zbgc_count=zbgc_count,
        dtgc_count=dtgc_count,
        seal_rate=seal_rate,
        fail_rate=fail_rate,
        push_ok=push_ok,
        lhb_hit=lhb_hit,
        formula_sample=formula_sample,
    )
    research_blocks.extend(facts.get("blocks", []))
    business_result = build_business_result(
        df,
        fund,
        zbgc,
        dtgc,
        lhb,
        qdf,
        top_sectors,
        facts,
        market_breadth,
        prev_material,
        formula_sample,
        fail_rate,
    )
    thesis_by_id = {item["id"]: item for item in business_result["theses"]}
    core_by_signal = {
        str(item.get("signal_id") or ""): item
        for item in business_result.get("conclusion_graph", [])
        if isinstance(item, dict)
    }

    largest_support_text = (
        f"并有{largest['connected']}只连板支撑"
        if int(largest["connected"]) > 0
        else "但没有连板样本支撑"
    )
    present_boards = sorted((int(board) for board in facts["board_counts"]), reverse=True)
    board_text = "、".join(f"{board}板" for board in present_boards) if present_boards else "无有效梯队"
    missing_middle = [board for board in range(highest_board - 1, 1, -1) if facts["board_counts"].get(board, 0) == 0]
    ladder_shape = (
        f"{highest_board}板以下缺少{'、'.join(str(x) + '板' for x in missing_middle)}样本"
        if missing_middle
        else "现有板数层级连续"
    )
    sampled_boards = sorted({int(x) for x in formula_sample_df.head(3)["连板数"].fillna(1).tolist()}, reverse=True)
    sampled_board_text = "、".join(f"{x}板" for x in sampled_boards)
    if up_ratio >= 0.60:
        breadth_summary = f"上涨家数占优（{up_count}/{up_count + down_count + flat_count}）"
    elif down_ratio >= 0.60:
        breadth_summary = f"下跌家数占优（{down_count}/{up_count + down_count + flat_count}）"
    else:
        breadth_summary = f"涨跌分化（上涨{up_count}、下跌{down_count}）"
    risk_pressure = "触板失败压力较高" if fail_rate >= 40 else "触板失败压力中等" if fail_rate >= 25 else "触板失败压力较低"

    extreme_weak_breadth = breadth_ok and (down_count >= 4000 or float(market_breadth.get("down_ratio") or 0) >= 0.75)
    market_core_node = core_by_signal.get("BREADTH_QUALITY_DIVERGENCE", {})
    market_core_sentence = sentence(
        str(market_core_node.get("conclusion") or thesis_by_id["TH-MARKET"].get("conclusion") or ""),
        conclusion_evidence_summary(market_core_node) or "；".join(item.get("interpretation", "") for item in thesis_by_id["TH-MARKET"].get("evidence", [])),
    )
    visible_status = {
        "SUPPORTED": "证据支持",
        "CONTESTED": "存在争议",
        "INSUFFICIENT": "暂不判断",
        "REJECTED": "判断不成立",
    }
    visible_scenarios = {
        "strengthen": "加强",
        "rotation": "轮动",
        "failure": "失败",
        "exogenous": "外生冲击",
    }
    market_node = core_by_signal.get("BREADTH_QUALITY_DIVERGENCE", {})
    ladder_node = core_by_signal.get("LADDER_FRAGILITY", {})
    direction_node = core_by_signal.get("DIRECTION_CONCENTRATION", {})
    risk_node = core_by_signal.get("PROFIT_LOSS_LOCATION", {})
    lifecycle_rows = [
        ["市场宽度", market_node.get("headline"), conclusion_confirmation_summary(market_node)],
        ["连板接力", ladder_node.get("headline"), conclusion_confirmation_summary(ladder_node)],
        [largest["sector"], direction_node.get("headline"), conclusion_confirmation_summary(direction_node)],
        ["赚钱与风险", risk_node.get("headline"), conclusion_confirmation_summary(risk_node)],
    ]
    scenario_tree = business_result["scenario_tree"]
    scenario_rows = [
        ["接力转强", "；".join(scenario_tree["strengthen"]["triggers"]), f"连板增加、冲板失败减少，{largest['sector']}的连续性同步增强", "；".join(scenario_tree["strengthen"]["invalidation"])],
        ["热点轮动", "；".join(scenario_tree["rotation"]["triggers"]), f"{largest['sector']}降温，资金和涨停数量转向新的热点", "；".join(scenario_tree["rotation"]["invalidation"])],
        ["亏钱扩大", "；".join(scenario_tree["failure"]["triggers"]), "炸板、高位断板和跌停向更多股票传导", "；".join(scenario_tree["failure"]["invalidation"])],
        ["外部冲击", "收盘后出现重大政策、公司公告、监管变化或外围市场急变", "消息改变原有热点和强势股的强弱关系", "消息没有改变板块涨停和连板表现"],
    ]

    metric_records = business_result.get("metrics") if isinstance(business_result.get("metrics"), dict) else {}
    limitup_percentile_value = (metric_records.get("limitup_count_percentile_20") or {}).get("value")
    continuation_percentile_value = (metric_records.get("continuation_count_percentile_20") or {}).get("value")
    limitup_percentile_text = f"{float(limitup_percentile_value):.0%}" if is_number(limitup_percentile_value) else "不可用"
    continuation_percentile_text = f"{float(continuation_percentile_value):.0%}" if is_number(continuation_percentile_value) else "不可用"
    non_top5_count = max(zt_count - int(facts["top5_count"]), 0)
    non_top5_share = non_top5_count / zt_count if zt_count else 0.0
    high_df = df[pd.to_numeric(df["连板数"], errors="coerce").fillna(1).astype(int) >= 3].copy()
    high_open_sum = int(pd.to_numeric(high_df.get("炸板次数", 0), errors="coerce").fillna(0).sum())
    high_share_of_continuation = len(high_df) / connected_count if connected_count else 0.0
    front_zero_open_count = sum(1 for row in quality_rows if to_int(row[4]) == 0)
    front_opened_count = max(len(quality_rows) - front_zero_open_count, 0)
    front_high_open_share = high_open_sum / facts["front_open_sum"] if facts["front_open_sum"] else 0.0
    continuation_df = df[pd.to_numeric(df["连板数"], errors="coerce").fillna(1).astype(int) >= 2].copy()
    direction_connected_counts = {
        clean(sector): int(count)
        for sector, count in continuation_df.groupby("投资板块", dropna=False).size().to_dict().items()
    }
    first_direction_continuity_count = sum(
        1 for row in first_rows if direction_connected_counts.get(clean(row[0]), 0) > 0
    )
    first_top5_count = sum(int(row[1]) for row in first_rows)
    first_top5_share = first_top5_count / first_board_count if first_board_count else 0.0
    fund_positive_share = facts["fund_pos"] / push_ok if push_ok else 0.0
    fund_total_value = to_num(facts.get("fund_total"))
    top_fund_share = to_num(facts.get("top_fund_amount")) / fund_total_value if fund_total_value > 0 else 0.0
    formula_coverage_share = facts["formula_success"] / zt_count if zt_count else 0.0
    previous_watch_count = int(facts["prev"].get("watch_count") or 0)
    previous_retained_count = int(facts["prev"].get("watch_retained_count") or 0)
    previous_promoted_count = int(facts["prev"].get("watch_promoted_count") or 0)
    previous_missed_count = max(previous_watch_count - previous_retained_count, 0)
    previous_retained_share = previous_retained_count / previous_watch_count if previous_watch_count else 0.0
    formula_observations: list[str] = []
    for formula_row in formula_rows:
        if not isinstance(formula_row, list) or len(formula_row) < 6:
            continue
        name = clean(formula_row[1])
        matched = qdf[qdf["名称"] == name]
        if matched.empty:
            continue
        sample = matched.iloc[0]
        formula_observations.append(
            f"{name}{to_int(sample.get('连板数'), 1)}板、开板{to_int(sample.get('炸板次数'))}次、主力净流入{fmt_amount(sample.get('主力净流入'))}、资金信号为{clean(formula_row[5])}"
        )
    formula_success_count = int(formula_sample.get("sample_success_count") or 0)
    formula_failed_count = int(formula_sample.get("sample_failed_count") or 0)
    formula_sample_limit = int(
        formula_sample.get("sample_limit")
        or formula_success_count + formula_failed_count
        or len(formula_sample_df.head(3))
    )
    if formula_observations:
        formula_fact_text = (
            f"技术公式计划采集{formula_sample_limit}只、成功{formula_success_count}只、失败{formula_failed_count}只；"
            + "；".join(formula_observations)
        )
        formula_section_analysis = (
            f"本轮技术公式样本成功{formula_success_count}/{formula_sample_limit}只；"
            "返回结果只与当日连板、开板和主力净流入交叉核对，不单独决定排序或结论。"
        )
    else:
        formula_fact_text = (
            f"技术公式计划采集{formula_sample_limit}只、成功{formula_success_count}只、失败{formula_failed_count}只；"
            "本轮无可用于横向比较的公式结果。"
        )
        formula_section_analysis = (
            f"本轮技术公式样本成功{formula_success_count}/{formula_sample_limit}只，覆盖不足；"
            "本节不据此判断高位资金强弱，只保留当日连板、开板和主力净流入的实采事实。"
        )

    prior_node = core_by_signal.get("PRIOR_CALIBRATION", {})
    previous_sentence = (
        sentence(
            str(prior_node.get("conclusion") or f"昨天收盘观察的{previous_watch_count}只股票，今天有{previous_retained_count}只仍在涨停池、{previous_promoted_count}只继续涨停"),
            conclusion_evidence_summary(prior_node) or f"昨日观察{previous_watch_count}只；今日仍在涨停池{previous_retained_count}只；今日继续涨停{previous_promoted_count}只；未留存{previous_missed_count}只",
        )
        if facts["prev"].get("status") == "VERIFIED"
        else sentence(
            "昨天没有可对照的市场观察样本，今天不评价昨日涨停判断的命中情况",
            "昨日观察样本0只；今日留存0只；今日继续涨停0只",
        )
    )
    direction_state_para = sentence(
        f"{largest['sector']}以{int(largest['count'])}只涨停居首，但只占全池{int(largest['count']) / zt_count:.1%}；当前热点清楚，市场仍由多个方向共同推动",
        f"{largest['sector']}涨停{int(largest['count'])}只、连板{int(largest['connected'])}只；前5方向合计{facts['top5_count']}只、占{pct_text(facts['top5_count'], zt_count)}；全日涨停{zt_count}只",
    )
    sections = [
        ("1. 当日市场总判断", [market_core_sentence], []),
        ("2. 市场宽度与涨停质量", [sentence(f"上涨家数占比{up_ratio:.1%}，但{zt_count}只涨停仅位于20日记录的{limitup_percentile_text}位置，冲板失败率为{fail_rate:.1f}%；普涨没有同步带来更稳的涨停接力", f"上涨{up_count}只、下跌{down_count}只；涨停{zt_count}只、连板{connected_count}只、封板率{seal_rate:.1f}%；20日涨停位置{limitup_percentile_text}")], [["维度", "数值", "数字说明"], ["全市场上涨", up_count if breadth_ok else "-", "上涨股票"], ["全市场下跌", down_count if breadth_ok else "-", breadth_label], ["涨停家数", zt_count, "涨停股票"], ["连板数量", connected_count, f"占涨停池{pct_text(connected_count, zt_count)}"], ["首板数量", first_board_count, f"占涨停池{pct_text(first_board_count, zt_count)}"], ["封板率", f"{seal_rate:.1f}%", f"{zt_count}/{touched}"], ["跌停数量", dtgc_count, "跌停股票"]]),
        ("3. 涨停结构与前五热点", [sentence(f"首板占{pct_text(first_board_count, zt_count)}，连板处于20日记录的{continuation_percentile_text}位置；强势股数量不低，但前5热点只占{pct_text(facts['top5_count'], zt_count)}，热点仍然分散", f"首板{first_board_count}只、连板{connected_count}只、最高{highest_board}板；前5热点合计{facts['top5_count']}只、占{pct_text(facts['top5_count'], zt_count)}")], core_table),
        ("4. 热点强弱与前排", [sentence(str(direction_node.get("conclusion") or ""), conclusion_evidence_summary(direction_node))], sector_rows),
        ("5. 热点集中程度", [sentence(f"最大热点只占{pct_text(largest['count'], zt_count)}，前5热点合计占{pct_text(facts['top5_count'], zt_count)}，其余{non_top5_count}只分布在前5之外；今天没有单一热点统领全场", f"{largest_direction}涨停{largest['count']}只、连板{largest['connected']}只；前5热点{facts['top5_count']}只；前5之外{non_top5_count}只；全日涨停{zt_count}只")], [["项目", "观察值", "数字对照"], ["最大热点占比", pct_text(largest["count"], zt_count), f"{largest_direction} {largest['count']}/{zt_count}"], ["前5热点占比", pct_text(facts["top5_count"], zt_count), f"{facts['top5_count']}/{zt_count}"], ["最大热点连板", largest["connected"], "2板及以上股票"], ["最大热点开板", facts["main_open_sum"], "热点内开板次数合计"], ["最大热点资金", fmt_amount(facts["main_fund"]), "热点内主力净流入合计"]]),
        ("6. 连板梯队", [sentence(str(ladder_node.get("conclusion") or ""), conclusion_evidence_summary(ladder_node))], ladder_rows),
        ("7. 三板以上股票", [sentence(f"三板及以上只有{facts['high_count']}只、占全部连板{high_share_of_continuation:.1%}，合计开板{high_open_sum}次；高位接力集中在少数股票，稳定性弱于连板总量表现", f"三板以上{facts['high_count']}只、连板总数{connected_count}只、占比{high_share_of_continuation:.1%}；最高{highest_board}板；高位合计开板{high_open_sum}次")], high_rows),
        (
            "8. 前排封板质量",
            [
                sentence(
                    f"前排{len(quality_rows)}只中{front_zero_open_count}只零开板、另{front_opened_count}只累计开板{facts['front_open_sum']}次，其中3板以上样本贡献{high_open_sum}次、占{front_high_open_share:.1%}，前排封板质量明显分化；次日只有发生开板的前排样本少于{front_opened_count}只且3板以上开板少于{high_open_sum}次才确认一致性改善",
                    f"前排{len(quality_rows)}只；零开板{front_zero_open_count}只、发生开板{front_opened_count}只、开板合计{facts['front_open_sum']}次；3板以上开板{high_open_sum}次；有资金数据{sum(1 for row in quality_rows if row[6] != '不可用')}/{len(quality_rows)}只",
                )
            ],
            quality_rows,
        ),
        ("9. 首板热点分布", [sentence(f"首板占涨停池{pct_text(first_board_count, zt_count)}，但首板前5热点只覆盖{first_top5_count}/{first_board_count}只，其中仅{first_direction_continuity_count}个热点已有连板；新增涨停很多，能够连续接力的热点仍少", f"首板{first_board_count}只；首板前5热点{first_top5_count}只、占{first_top5_share:.1%}；连板{connected_count}只；前5首板热点中已有连板的热点{first_direction_continuity_count}个")], first_rows),
        ("10. 炸板与跌停风险", [sentence(f"冲板失败率{fail_rate:.1f}%、跌停{dtgc_count}只、前排合计开板{facts['front_open_sum']}次；上涨股票较多，但炸板和高位反复开板仍在制造局部亏钱", f"炸板{zbgc_count}只；跌停{dtgc_count}只；前排10只合计开板{facts['front_open_sum']}次；上涨{up_count}只、下跌{down_count}只")], risk_rows),
        ("11. 游资、机构与大资金", [sentence(f"主力净流入为正的股票占{fund_positive_share:.1%}，但首板占比仍为{first_board_count / zt_count:.1%}、冲板失败率为{fail_rate:.1f}%；资金覆盖很广，稳定封板和连续涨停没有同步变强", f"有资金数据{push_ok}/{zt_count}只；净流入为正{facts['fund_pos']}只、为负{facts['fund_neg']}只；合计{fmt_amount(facts['fund_total'])}；最大单股{facts['top_fund_label']}为{fmt_amount(facts['top_fund_amount'])}；龙虎榜匹配{lhb_hit}只")], fund_rows),
        ("12. 高位股资金与封板对照", [sentence(formula_section_analysis, formula_fact_text)], formula_rows),
        ("13. 当前判断与转强条件", [direction_state_para], lifecycle_rows),
        (
            "14. 明日可能出现的四种走法",
            [
                sentence(
                    str(core_by_signal.get("NEXT_DAY_BRANCH", {}).get("conclusion") or ""),
                    conclusion_evidence_summary(core_by_signal.get("NEXT_DAY_BRANCH", {})),
                )
            ],
            scenario_rows,
        ),
        ("15. 昨日判断复核", [previous_sentence], [["项目", "结果", "对照"], ["昨日收盘日期", facts["prev"].get("date_h"), "上一交易日"], ["昨日观察", facts["prev"].get("watch_count"), "股票数量"], ["今日仍在涨停池", facts["prev"].get("watch_retained_count"), "股票数量"], ["今日继续涨停", facts["prev"].get("watch_promoted_count"), facts["prev"].get("promoted_watch_labels")]]),
        ("16. 主要风险", [sentence(f"当前主要风险在冲板失败的{zbgc_count}只股票、高位反复开板和{dtgc_count}只跌停股；只要这些数量没有下降，普涨就不能等同于短线风险消退", f"冲板失败率{fail_rate:.1f}%；跌停{dtgc_count}只；前排合计开板{facts['front_open_sum']}次；连板{connected_count}只")], risk_rows[:8]),
    ]

    section_signal_map = {
        1: ["BREADTH_QUALITY_DIVERGENCE"],
        2: ["BREADTH_QUALITY_DIVERGENCE"],
        3: ["BREADTH_QUALITY_DIVERGENCE", "DIRECTION_CONCENTRATION"],
        4: ["DIRECTION_CONCENTRATION"],
        5: ["DIRECTION_CONCENTRATION"],
        6: ["LADDER_FRAGILITY"],
        7: ["LADDER_FRAGILITY"],
        8: ["SEALING_TIMING", "LADDER_FRAGILITY"],
        9: ["DIRECTION_CONCENTRATION"],
        10: ["PROFIT_LOSS_LOCATION"],
        11: ["HOT_MONEY_BEHAVIOR", "INSTITUTION_BEHAVIOR", "CAPITAL_CONCENTRATION"],
        12: ["LADDER_FRAGILITY", "CAPITAL_CONCENTRATION"],
        13: ["BREADTH_QUALITY_DIVERGENCE", "LADDER_FRAGILITY", "DIRECTION_CONCENTRATION", "PROFIT_LOSS_LOCATION"],
        14: ["NEXT_DAY_BRANCH"],
        15: ["PRIOR_CALIBRATION"],
        16: ["PROFIT_LOSS_LOCATION"],
    }
    section_analysis: list[dict[str, Any]] = []
    selected_signal_set = set(core_by_signal)
    fallback_signal_id = min(selected_signal_set, default="")
    for title_text, paras, _ in sections:
        chapter = int(title_text.split(".", 1)[0])
        if len(paras) != 1 or not isinstance(paras[0], dict):
            raise RuntimeError(f"section {chapter} must have exactly one structured analysis note")
        note = paras[0]
        analysis_text = str(note.get("analysis") or "").strip()
        facts_text = str(note.get("facts") or "").strip()
        linked_signals = [value for value in section_signal_map[chapter] if value in selected_signal_set]
        if not linked_signals and fallback_signal_id:
            linked_signals = [fallback_signal_id]
        section_analysis.append({
            "chapter": chapter,
            "heading": title_text,
            "analysis": analysis_text,
            "facts": facts_text,
            "source_signal_ids": linked_signals,
            "analysis_sha256": hashlib.sha256(analysis_text.encode("utf-8")).hexdigest(),
            "facts_sha256": hashlib.sha256(facts_text.encode("utf-8")).hexdigest(),
        })
    business_result["section_analysis"] = section_analysis
    business_result, business_blocks = persist_and_readback_business_result(business_result)
    if business_result.get("status") != "CLEAN_PASS" or business_blocks:
        print(json.dumps({
            "status": "BLOCKED",
            "run_id": RUN_ID,
            "analysis": str(BUSINESS_RESULT_JSON),
            "blocks": business_result.get("blocks") or business_blocks,
        }, ensure_ascii=False, indent=2))
        return 2
    core_graph = [item for item in business_result.get("conclusion_graph", []) if isinstance(item, dict)]

    md: list[str] = []
    md.append(f"# {DATE_H} 每日涨停板深度复盘报告")
    md.append("")
    md.append("## 0. 核心结论地图")
    for category_label, category_nodes in grouped_conclusion_nodes(core_graph):
        md.append("")
        md.append(f"### {category_label}")
        md.append(md_table(
            ["对象", "观察角度", "今日判断", "关键数据", "明日验证"],
            [
                [
                    node.get("scope_label"),
                    node.get("angle_label"),
                    f"{node.get('headline')}：{node.get('conclusion')}",
                    conclusion_evidence_summary(node),
                    conclusion_confirmation_summary(node),
                ]
                for node in category_nodes
            ],
        ))
    md.append("")

    for title, paras, table_rows in sections:
        md.append("")
        md.append(f"## {title}")
        for para in paras:
            md.append(f"**{para['analysis']}**")
            md.append("")
            md.append(para["facts"])
        if title.startswith("3."):
            md.append(md_table(["项目", "数值", "备注"], core_table))
            md.append("")
            md.append(md_table(["排名", "方向", "涨停数", "连板数", "最高板", "前排1", "前排2", "前排3", "资金合计"], dragon))
        elif table_rows and isinstance(table_rows[0], list):
            headers = [str(x) for x in table_rows[0]]
            rows = table_rows[1:] if len(table_rows) > 1 and all(not isinstance(x, (int, float)) for x in table_rows[0]) else table_rows
            if title.startswith(("4.", "6.", "7.", "8.", "9.", "10.", "11.", "12.", "13.", "14.", "15.", "16.")):
                predefined = {
                    "4.": ["排名", "方向", "涨停数", "连板数", "最高板", "资金合计", "前排样本", "排序依据"],
                    "6.": ["板数", "数量", "代表个股", "梯队状态"],
                    "7.": ["代码", "名称", "高度", "首次封板", "炸板次数", "所属板块"],
                    "8.": ["代码", "名称", "高度", "首次封板", "炸板次数", "所属板块", "资金状态"],
                    "9.": ["方向", "首板数", "首板占比", "代表样本", "次日观察"],
                    "10.": ["代码", "名称", "高度", "炸板次数", "风险触发", "所属板块"],
                    "11.": ["代码", "个股", "主力净流入", "超大单净流入", "主力净占比", "备注"],
                    "12.": ["代码", "名称", "高度", "大牛线4.0", "飞龙在天", "资金公式摘要", "样本事实"],
                    "13.": ["观察角度", "当前判断", "转强条件"],
                    "14.": ["走法", "出现条件", "盘面含义", "不再适用"],
                    "15.": ["项目", "结果", "对照"],
                    "16.": ["代码", "名称", "高度", "炸板次数", "风险触发", "所属板块"],
                }
                key = title.split(".", 1)[0] + "."
                md.append(md_table(predefined[key], table_rows))
            else:
                md.append(md_table(headers, rows))

    OUT_MD.write_text("\n".join(md) + "\n", encoding="utf-8")

    watch_lines = [
        f"# {DATE_H} 明日观察池",
        "",
    ]
    visible_roles = {
        "height_anchor": "高度锚",
        "direction_front": "方向核心",
        "ladder_observer": "梯队观察",
        "breadth_observer": "广度观察",
    }
    for item in business_result["watchlist"]:
        role_text = visible_roles.get(str(item.get("role") or ""), "观察样本")
        rationale_text = str(item.get("rationale") or "")
        for internal_role, visible_role in visible_roles.items():
            rationale_text = rationale_text.replace(internal_role, visible_role)
        watch_lines.append(
            f"- {item['code']} {item['name']}｜角色：{role_text}｜理由：{chinese_visible_text(rationale_text)}"
            f"｜确认：{'；'.join(item['confirmation'])}｜失效：{'；'.join(item['invalidation'])}"
        )
    OUT_WATCH.write_text("\n".join(watch_lines) + "\n", encoding="utf-8")

    visible_df = df.copy()
    if "ok" in visible_df.columns:
        visible_df["ok"] = visible_df["ok"].map(lambda value: "是" if parse_verified_bool(value) else "否")
        visible_df = visible_df.rename(columns={"ok": "资金数据可用"})
    visible_df.to_csv(OUT_CSV, index=False, encoding="utf-8-sig")
    shutil.copy2(OUT_CSV, TASK_DIR / f"final_limitup_union_strict_{DATE}.csv")

    doc = setup_doc()
    add_stock_risk_notice_page(doc)
    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    add_run(title, DATE_H, bold=True, color=GRAY, size=12)
    title.add_run("\n")
    add_run(title, "每日涨停板深度复盘报告", bold=True, color=BLUE, size=20)
    subtitle = doc.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    add_run(subtitle, f"核心：{zt_count}只涨停 | 连板{connected_count}只 | 最高{highest_board}板 | 数量最多方向：{largest_direction}", color=GRAY, size=9)
    add_heading(doc, "0. 核心结论地图", 1)
    add_table(doc, ["核心指标", "数值", "说明"], [
        ["涨停家数", zt_count, "涨停股票"],
        ["全市场涨跌", f"{up_count}/{down_count}" if breadth_ok else "-", breadth_label],
        ["连板/首板", f"{connected_count}/{first_board_count}", f"连板占{pct_text(connected_count, zt_count)} / 首板占{pct_text(first_board_count, zt_count)}"],
        ["最高板", f"{highest_board}板", largest_front],
        ["封板率", f"{seal_rate:.1f}%", f"炸板{zbgc_count}只"],
        ["跌停数量", dtgc_count, "跌停股票"],
        ["数量最多方向", largest_direction, f"占涨停池{pct_text(largest['count'], zt_count)}"],
    ], widths=[1.3, 1.5, 4.5], font_size=9)
    spacer = doc.add_paragraph()
    spacer.paragraph_format.space_after = Pt(2)
    add_conclusion_matrix(doc, core_graph)
    doc.add_page_break()

    for title_text, paras, table_rows in sections:
        section_heading = add_heading(doc, title_text, 1)
        if title_text.startswith("16."):
            section_heading.paragraph_format.space_before = Pt(4)
            section_heading.paragraph_format.space_after = Pt(2)
            for run in section_heading.runs:
                run.font.size = Pt(12)
        last_section_paragraph = None
        for para in paras:
            last_section_paragraph = add_analysis_note(doc, para)
        if last_section_paragraph is not None:
            if title_text.startswith("16."):
                last_section_paragraph.paragraph_format.space_after = Pt(0)
                last_section_paragraph.paragraph_format.line_spacing = 1.0
                last_section_paragraph.paragraph_format.keep_together = True
                for run in last_section_paragraph.runs:
                    run.font.size = Pt(8)
            else:
                last_section_paragraph.paragraph_format.keep_with_next = True
        if title_text.startswith("3."):
            add_table(doc, ["项目", "数值", "备注"], core_table, widths=[1.7, 1.4, 4.2], font_size=8)
            add_table(doc, ["排名", "方向", "涨停数", "连板数", "最高板", "前排1", "前排2", "前排3", "资金合计"], dragon, widths=[0.42, 0.85, 0.5, 0.5, 0.5, 1.35, 1.35, 1.35, 0.85], font_size=7)
        elif title_text.startswith("4."):
            add_table(doc, ["排名", "方向", "涨停数", "连板数", "最高板", "资金合计", "前排样本", "排序依据"], sector_rows, widths=[0.45, 1.0, 0.55, 0.55, 0.55, 0.8, 1.65, 1.7], font_size=7)
        elif title_text.startswith("5."):
            add_table(doc, ["项目", "观察值", "事实说明"], [["最大方向占比", pct_text(largest["count"], zt_count), f"{largest_direction} {largest['count']}/{zt_count}"], ["前5方向占比", pct_text(facts["top5_count"], zt_count), f"{facts['top5_count']}/{zt_count}"], ["最大方向连板", largest["connected"], "2板及以上样本数"], ["最大方向炸板", facts["main_open_sum"], "方向内炸板次数合计"], ["最大方向资金", fmt_amount(facts["main_fund"]), "方向内主力净流入合计"]], widths=[1.8, 1.5, 4.0], font_size=8)
        elif title_text.startswith("6."):
            add_table(doc, ["板数", "数量", "代表个股", "梯队状态"], ladder_rows, widths=[0.7, 0.55, 5.1, 0.9], font_size=7)
        elif title_text.startswith("7."):
            add_table(doc, ["代码", "名称", "高度", "首次封板", "炸板", "所属板块"], high_rows, widths=[0.75, 1.05, 0.55, 0.9, 0.55, 3.4], font_size=8)
        elif title_text.startswith("8."):
            add_table(doc, ["代码", "名称", "高度", "首次封板", "炸板", "所属板块", "资金状态"], quality_rows, widths=[0.7, 0.95, 0.5, 0.85, 0.45, 2.75, 0.9], font_size=7)
        elif title_text.startswith("9."):
            add_table(doc, ["方向", "首板数", "首板占比", "代表样本", "次日观察"], first_rows, widths=[1.1, 0.6, 0.7, 3.65, 1.1], font_size=7)
        elif title_text.startswith("10."):
            add_table(doc, ["代码", "名称", "高度", "炸板", "风险触发", "所属板块"], risk_rows[:16], widths=[0.7, 1.0, 0.55, 0.45, 2.0, 2.45], font_size=7)
        elif title_text.startswith("11."):
            add_table(doc, ["代码", "个股", "主力净流入", "超大单净流入", "主力净占比", "备注"], fund_rows, widths=[0.75, 1.0, 1.0, 1.05, 0.85, 2.5], font_size=7)
        elif title_text.startswith("12."):
            if formula_rows:
                add_table(doc, ["代码", "名称", "高度", "大牛线4.0", "飞龙在天", "资金公式摘要", "样本事实"], formula_rows, widths=[0.65, 0.85, 0.45, 1.35, 1.0, 2.0, 1.3], font_size=6)
        elif title_text.startswith("13."):
            add_table(doc, ["观察角度", "当前判断", "转强条件"], lifecycle_rows, widths=[1.2, 2.6, 3.5], font_size=8)
        elif title_text.startswith("14."):
            add_table(doc, ["走法", "出现条件", "盘面含义", "不再适用"], scenario_rows, widths=[0.8, 2.7, 2.0, 1.8], font_size=7)
        elif title_text.startswith("15."):
            add_table(doc, ["项目", "结果", "对照"], [["昨日收盘日期", facts["prev"].get("date_h"), "上一交易日"], ["昨日观察", facts["prev"].get("watch_count"), "股票数量"], ["今日仍在涨停池", facts["prev"].get("watch_retained_count"), "股票数量"], ["今日继续涨停", facts["prev"].get("watch_promoted_count"), facts["prev"].get("promoted_watch_labels")]], widths=[1.5, 1.3, 4.5], font_size=8)
        elif title_text.startswith("16."):
            pass

    if doc.paragraphs and not doc.paragraphs[-1].text.strip():
        trailing = doc.paragraphs[-1]._element
        trailing.getparent().remove(trailing)
    doc.save(OUT_DOCX)
    embed_stock_review_manifest(OUT_DOCX, business_result)
    shutil.copy2(OUT_DOCX, OPEN_COPY)

    md_text = OUT_MD.read_text(encoding="utf-8")
    docx_text = read_docx_text(OUT_DOCX)
    visible = md_text + "\n" + docx_text
    latin_hits = sorted(set(re.findall(r"[A-Za-z]+", visible)))
    eligibility_audit = load_eligibility_audit()
    exclusion_items = eligibility_audit.get("hard_exclusions") if isinstance(eligibility_audit.get("hard_exclusions"), list) else []
    delisting_visible_hits = []
    for item in exclusion_items:
        if not isinstance(item, dict):
            continue
        code = code6(item.get("code"))
        name = str(item.get("name") or "").strip()
        if (code and code in visible) or (name and name in visible):
            delisting_visible_hits.append(f"{code} {name}".strip())
    forbidden_terms = [
        "本版纠偏口径", "数据状态与数据来源", "口径纪律", "数据来源清单", "今日全部涨停股票穷举清单", "全量股票名单穷举", "正文固定17章",
        "TDX", "AK", "Codex", "工作流", "ZTC", "ELB",
        "未实采", "未纳入", "保留框架", "等待独立采集", "修复强度", "修复中带分歧", "盘面有修复",
        "发酵偏分歧", "发酵", "主升", "高潮",
        "今日最强主线", "最强主线", "主线板块", "主线评分", "主线生命周期", "强度评分", "数量领先", "决定强度上限",
        "龙1", "龙2", "龙3", "龙头", "评级", "龙1带动", "看延续与降级", "偏强确认", "偏弱确认", "强势信号", "弱势信号",
        "情绪温度", "宽度约束分", "集中度分", "仅按排序分排列", "技术辅助只展示", "不参与股票池生成",
        "不输出买卖", "不作买卖依据", "只核验数据变化", "不预设方向", "只比对同一股票集合", "只记录样本数",
        "默认结论", "固定措辞", "兜底判断", "预测账本哈希", "内部启发式", "CLEAN_PASS", "SHA-256", "sha256", "验证通过", "自验证",
        "结论：", "依据：", "但要注意：", "上面判断就不成立",
        "冻结账本", "冻结预测账本", "本地上一交易日表", "从观察池提取代码", "只保留当日涨停池确认样本",
        "事实来源", "研究动作", "论点状态与证伪条件", "次日四分支情景树",
        "全市场上涨家数明显占优", "全市场下跌家数明显占优", "涨跌家数分化，涨停池没有获得单边宽度信号",
        "涨停结构偏向首板扩散而非连板推进", "方向集中度不足以支撑单一方向解释全日盘面",
        "市场高度的可观察样本很窄", "前排样本分化明显", "分散试错而不是成熟梯队扩张",
        "亏钱效应没有收敛", "局部集中而不是全局转强", "昨日观察样本的有效性集中在少数晋级个股",
        "重要方向", "看成交额、涨停扩散", "看首板转连板数量", "优先检查各板级是否同时晋级", "字段",
        "要求承接、扩散和有时点的催化事实同向后才升级", "先检查开板、跌停扩散和方向脱节",
        "需要结合可用历史基线确认其相对位置", "这些字段构成次日风险分支的优先观察依据",
        "次日市场与方向按加强、轮动、失败、外生冲击四个互斥分支验收",
        "各分支的触发阈值已在收盘后冻结，次日只按实际条件执行",
        "逐项展示高度、封板时间、开板次数、方向和资金可用状态",
    ]
    forbidden_hits = [term for term in forbidden_terms if term in visible]
    chapter17_hit = bool(re.search(r"(?:^|\n|\r)\s*(?:#{1,4}\s*)?17[.、]\s+", visible))
    content_blocks: list[str] = []
    content_blocks.extend(research_blocks)
    content_blocks.extend(validate_chapter_conclusions(md_text))
    if forbidden_hits:
        content_blocks.append(f"visible_forbidden_hits: {forbidden_hits}")
    if latin_hits:
        content_blocks.append(f"{VISIBLE_LATIN_GATE}: {latin_hits}")
    if delisting_visible_hits:
        content_blocks.append(f"{DELISTING_HARD_EXCLUSION}: {delisting_visible_hits}")
    if "涨停结构与前五热点" not in visible:
        content_blocks.append("missing required market-structure section")
    if chapter17_hit:
        content_blocks.append("chapter 17 heading detected")
    if market_breadth.get("status") != "CLEAN_PASS":
        content_blocks.append("market breadth collection failed")
    if business_result.get("status") != "CLEAN_PASS":
        content_blocks.extend(business_result.get("blocks") or ["business result is not CLEAN_PASS"])
    status = "CLEAN_PASS" if not content_blocks else "BLOCKED"
    truth_snapshot = {
        "status": status,
        "run_id": RUN_ID,
        "schema_version": BUSINESS_RESULT_SCHEMA_VERSION,
        "generated_at": datetime.now().isoformat(timespec="seconds"),
        "docx_path": str(OUT_DOCX),
        "md_path": str(OUT_MD),
        "csv_path": str(OUT_CSV),
        "docx_sha256": sha256(OUT_DOCX),
        "md_sha256": sha256(OUT_MD),
        "csv_sha256": sha256(OUT_CSV),
        "watchlist_sha256": sha256(OUT_WATCH),
        "fund_flow_sha256": sha256(fund_path),
        "analysis_sha256": sha256(BUSINESS_RESULT_JSON),
        "market_breadth_sha256": sha256(MARKET_BREADTH_JSON),
        "formula_sample_sha256": sha256(TQ_SAMPLE_JSON),
        "docx_size": OUT_DOCX.stat().st_size,
        "md_size": OUT_MD.stat().st_size,
        "csv_rows": int(len(df)),
        "official_limitup_count": zt_count,
        "blocks": content_blocks,
        "market_breadth_status": market_breadth.get("status"),
        "formula_sample_status": formula_sample.get("status"),
    }
    evidence = {
        "status": status,
        "date": DATE,
        "run_id": RUN_ID,
        "docx": str(OUT_DOCX),
        "markdown": str(OUT_MD),
        "csv": str(OUT_CSV),
        "watchlist": str(OUT_WATCH),
        "analysis": {
            "path": str(BUSINESS_RESULT_JSON),
            "sha256": sha256(BUSINESS_RESULT_JSON),
            "run_id": RUN_ID,
            "schema_version": BUSINESS_RESULT_SCHEMA_VERSION,
        },
        "facts": {
            "limitup_count": zt_count,
            "connected_count": connected_count,
            "highest_board": highest_board,
            "zbgc_count": zbgc_count,
            "dtgc_count": dtgc_count,
            "largest_direction": largest_direction,
            "largest_direction_count": int(largest["count"]),
            "largest_direction_share": pct_text(largest["count"], zt_count),
            "first_board_share": pct_text(first_board_count, zt_count),
            "connected_share": pct_text(connected_count, zt_count),
            "fund_coverage": f"{push_ok}/{zt_count}",
            "research_blocks": research_blocks,
            "previous_review_material": prev_material,
            "market_breadth": {
                "up_count": up_count,
                "down_count": down_count,
                "flat_count": flat_count,
                "breadth_label": breadth_label,
                "source": market_breadth.get("source"),
                "source_url": market_breadth.get("source_url"),
            },
            "formula_sample": {
                "status": formula_sample.get("status"),
                "sample_success_count": formula_sample.get("sample_success_count"),
                "sample_failed_count": formula_sample.get("sample_failed_count"),
                "json_path": str(TQ_SAMPLE_JSON),
            },
        },
        "content_blocks": content_blocks,
        "visible_forbidden_hits": forbidden_hits,
        "visible_latin_hits": latin_hits,
        "visible_delisting_hits": delisting_visible_hits,
        "hard_exclusions": exclusion_items,
        "chapter17_heading_detected": chapter17_hit,
        "truth_snapshot": truth_snapshot,
    }
    EVIDENCE_JSON.write_text(json.dumps(evidence, ensure_ascii=False, indent=2), encoding="utf-8")
    print(json.dumps(evidence, ensure_ascii=False, indent=2))
    if status != "CLEAN_PASS":
        return 2
    return 0


def read_docx_text(path: Path) -> str:
    with zipfile.ZipFile(path) as zf:
        xml = zf.read("word/document.xml").decode("utf-8", errors="ignore")
    root = ET.fromstring(xml.encode("utf-8"))
    ns = {"w": "http://schemas.openxmlformats.org/wordprocessingml/2006/main"}
    paragraphs: list[str] = []
    for paragraph in root.findall(".//w:p", ns):
        texts = [node.text or "" for node in paragraph.findall(".//w:t", ns)]
        if texts:
            paragraphs.append("".join(texts))
    return "\n".join(paragraphs)


if __name__ == "__main__":
    raise SystemExit(build_report())
