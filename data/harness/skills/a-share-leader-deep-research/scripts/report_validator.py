#!/usr/bin/env python3
from __future__ import annotations
# ONESTOCK_EMBEDDED_LOCAL_IMPORT_BOOTSTRAP
import sys as _onestock_embedded_sys
from pathlib import Path as _OneStockEmbeddedPath
_onestock_embedded_dir = str(_OneStockEmbeddedPath(__file__).resolve().parent)
if _onestock_embedded_dir not in _onestock_embedded_sys.path:
    _onestock_embedded_sys.path.insert(0, _onestock_embedded_dir)

import argparse
import hashlib
import json
import math
import re
from pathlib import Path
from zipfile import ZipFile

from docx import Document
from PIL import Image, ImageStat


def sha256_path(path: Path) -> str:
    return hashlib.sha256(path.read_bytes()).hexdigest()


def effective_style_size(document: Document, style_name: str) -> float | None:
    style = document.styles[style_name]
    return style.font.size.pt if style.font.size else None


def page_geometry_cm(section) -> dict[str, float]:
    return {
        "width_cm": section.page_width.cm,
        "height_cm": section.page_height.cm,
        "top_margin_cm": section.top_margin.cm,
        "bottom_margin_cm": section.bottom_margin.cm,
        "left_margin_cm": section.left_margin.cm,
        "right_margin_cm": section.right_margin.cm,
    }


def extract_visible_text(document: Document) -> tuple[list[str], list[str], str, str]:
    visible_paragraphs = [
        paragraph.text.strip()
        for paragraph in document.paragraphs
        if paragraph.text.strip()
    ]
    visible_table_cells = [
        cell.text.strip()
        for table in document.tables
        for row in table.rows
        for cell in row.cells
        if cell.text.strip()
    ]
    visible_text = "\n".join(visible_paragraphs + visible_table_cells)
    normalized_text = re.sub(r"\s+", "", visible_text)
    return visible_paragraphs, visible_table_cells, visible_text, normalized_text


def check_content_identity(
    document: Document,
    payload: dict,
    contract: dict,
    visible_text: str | None = None,
    normalized_text: str | None = None,
) -> tuple[list[str], dict]:
    errors: list[str] = []
    gate = contract.get("content_identity_gate")
    if not isinstance(gate, dict):
        return ["content_identity_contract_missing"], {
            "status": "BLOCKED",
            "contract_present": False,
        }

    if visible_text is None or normalized_text is None:
        _, _, visible_text, normalized_text = extract_visible_text(document)

    core_title = (document.core_properties.title or "").strip()
    core_subject = (document.core_properties.subject or "").strip()
    normalized_title = re.sub(r"\s+", "", core_title)
    normalized_subject = re.sub(r"\s+", "", core_subject)
    normalized_metadata = normalized_title + normalized_subject

    for term in gate["required_core_title_terms"]:
        if re.sub(r"\s+", "", term) not in normalized_title:
            errors.append(f"content_identity_title_term_missing:{term}")
    for term in gate["required_core_subject_terms"]:
        if re.sub(r"\s+", "", term) not in normalized_subject:
            errors.append(f"content_identity_subject_term_missing:{term}")
    for term in gate["forbidden_core_metadata_terms"]:
        if re.sub(r"\s+", "", term) in normalized_metadata:
            errors.append(f"content_identity_core_metadata_forbidden:{term}")
    for term in gate["forbidden_visible_substitute_terms"]:
        if re.sub(r"\s+", "", term) in normalized_text:
            errors.append(f"content_identity_visible_substitute:{term}")
    for term in gate.get("forbidden_visible_process_terms", []):
        if re.sub(r"\s+", "", term) in normalized_text:
            errors.append(f"content_identity_visible_process_pollution:{term}")
    for term in gate.get("forbidden_causal_overclaim_terms", []):
        if re.sub(r"\s+", "", term) in normalized_text:
            errors.append(f"content_identity_causal_overclaim:{term}")
    statistics_only_patterns = (
        r"上涨逻辑\s*.+共有\d+只涨停[，,；;].*连板\d+只.*零开板\d+只",
        r"核心逻辑\s*.+共有\d+只涨停[，,；;].*连板\d+只.*零开板\d+只",
    )
    for pattern in statistics_only_patterns:
        if re.search(pattern, re.sub(r"\s+", "", visible_text)):
            errors.append("content_identity_statistics_only_logic")
            break
    for phrase in (
        "居板块内涨停地位首位。排序依据直接采用已授权龙头研究回执",
        "不以高度锚、机制锚或容量锚替代地位榜",
    ):
        if re.sub(r"\s+", "", phrase) in normalized_text:
            errors.append(f"content_identity_shallow_position_conclusion:{phrase}")

    minimum_characters = int(gate["minimum_non_whitespace_characters"])
    visible_character_count = len(normalized_text)
    if visible_character_count < minimum_characters:
        errors.append(
            "content_identity_text_too_short:"
            f"{visible_character_count}<{minimum_characters}"
        )

    payload_stock_codes = {
        str(item.get("stock_code") or "").strip()
        for item in payload.get("stocks") or []
        if re.fullmatch(r"\d{6}", str(item.get("stock_code") or "").strip())
    }
    visible_six_digit_codes = set(re.findall(r"(?<!\d)\d{6}(?!\d)", visible_text))
    matched_payload_codes = payload_stock_codes & visible_six_digit_codes
    minimum_ratio = float(gate["minimum_payload_stock_code_coverage_ratio"])
    minimum_codes = int(gate["minimum_unique_payload_stock_codes"])
    required_code_count = 0
    coverage_ratio = 0.0
    if payload_stock_codes:
        required_code_count = min(
            len(payload_stock_codes),
            max(
                minimum_codes,
                math.ceil(len(payload_stock_codes) * minimum_ratio),
            ),
        )
        coverage_ratio = len(matched_payload_codes) / len(payload_stock_codes)
        if len(matched_payload_codes) < required_code_count:
            errors.append(
                "content_identity_stock_code_coverage_low:"
                f"{len(matched_payload_codes)}<{required_code_count}"
            )
    else:
        errors.append("content_identity_payload_stock_codes_missing")

    detail = {
        "status": "PASS" if not errors else "BLOCKED",
        "contract_present": True,
        "core_title": core_title,
        "core_subject": core_subject,
        "visible_character_count": visible_character_count,
        "minimum_visible_character_count": minimum_characters,
        "payload_stock_code_count": len(payload_stock_codes),
        "visible_six_digit_code_count": len(visible_six_digit_codes),
        "matched_payload_stock_code_count": len(matched_payload_codes),
        "required_payload_stock_code_count": required_code_count,
        "payload_stock_code_coverage_ratio": round(coverage_ratio, 4),
    }
    return errors, detail



def portable_mean_luminance(value: float) -> float:
    return round(float(value), 0)


def check_render(render_dir: Path, contract: dict) -> tuple[list[str], dict]:
    errors: list[str] = []
    page_files = sorted(
        render_dir.glob("page-*.png"),
        key=lambda path: int(path.stem.split("-")[-1]),
    )
    render_contract = contract["render"]
    if not page_files:
        return ["render_pages_missing"], {"page_count": 0, "pages": []}
    page_count = len(page_files)
    if not render_contract["minimum_pages"] <= page_count <= render_contract["maximum_pages"]:
        errors.append(f"render_page_count_out_of_range:{page_count}")
    page_rows = []
    for page in page_files:
        with Image.open(page) as image:
            gray = image.convert("L")
            if gray.width < 1000 or gray.height < 1000:
                errors.append(f"render_resolution_too_small:{page.name}:{gray.width}x{gray.height}")
            histogram = gray.histogram()
            dark_pixels = sum(histogram[:245])
            ratio = dark_pixels / float(gray.width * gray.height)
            mean = ImageStat.Stat(gray).mean[0]
            if ratio < render_contract["minimum_ink_ratio"]:
                errors.append(f"render_page_nearly_blank:{page.name}:{ratio:.5f}")
            if ratio > render_contract["maximum_ink_ratio"]:
                errors.append(f"render_page_overdense:{page.name}:{ratio:.5f}")
            page_rows.append(
                {
                    "page": page.name,
                    "width": gray.width,
                    "height": gray.height,
                    "ink_ratio": round(ratio, 5),
                    "mean_luminance": portable_mean_luminance(mean),
                }
            )
    final_page_minimum = float(
        render_contract.get(
            "minimum_final_page_ink_ratio",
            render_contract["minimum_ink_ratio"],
        )
    )
    final_ratio = float(page_rows[-1]["ink_ratio"])
    if page_count > 1 and final_ratio < final_page_minimum:
        errors.append(
            f"render_final_page_orphaned:{page_rows[-1]['page']}:{final_ratio:.5f}"
        )
    return errors, {"page_count": page_count, "pages": page_rows}


def validate(
    docx_path: Path,
    payload_path: Path,
    manifest_path: Path,
    template_contract_path: Path,
    render_dir: Path | None,
) -> dict:
    errors: list[str] = []
    payload = json.loads(payload_path.read_text(encoding="utf-8-sig"))
    manifest = json.loads(manifest_path.read_text(encoding="utf-8-sig"))
    contract = json.loads(template_contract_path.read_text(encoding="utf-8-sig"))
    payload_sha = sha256_path(payload_path)
    contract_sha = sha256_path(template_contract_path)
    if manifest.get("result_sha256") != payload_sha:
        errors.append("manifest_payload_hash_mismatch")
    if manifest.get("template_contract_sha256") != contract_sha:
        errors.append("manifest_template_contract_hash_mismatch")

    try:
        document = Document(docx_path)
    except Exception as exc:
        return {
            "status": "BLOCKED",
            "errors": [f"docx_open_failed:{type(exc).__name__}:{exc}"],
        }

    (
        visible_paragraphs,
        visible_table_cells,
        visible_text,
        normalized_text,
    ) = extract_visible_text(document)
    identity_errors, identity_detail = check_content_identity(
        document,
        payload,
        contract,
        visible_text,
        normalized_text,
    )
    errors.extend(identity_errors)
    section_contract = contract["required_section_order"]
    positions = []
    for required in section_contract:
        position = normalized_text.find(re.sub(r"\s+", "", required))
        positions.append(position)
        if position < 0:
            errors.append(f"required_section_missing:{required}")
    if all(position >= 0 for position in positions) and positions != sorted(positions):
        errors.append("required_section_order_invalid")
    for forbidden in contract["forbidden_visible_sections"]:
        if forbidden in visible_text:
            errors.append(f"forbidden_section_present:{forbidden}")

    structure = contract["structure"]
    table_count = len(document.tables)
    if not structure["minimum_tables"] <= table_count <= structure["maximum_tables"]:
        errors.append(f"table_count_out_of_range:{table_count}")
    one_by_one = sum(
        1 for table in document.tables
        if len(table.rows) == 1 and len(table.columns) == 1
    )
    if one_by_one > structure["maximum_one_by_one_tables"]:
        errors.append(f"decorative_one_by_one_tables_present:{one_by_one}")

    section = document.sections[0]
    actual_geometry = page_geometry_cm(section)
    for field, expected in contract["page"].items():
        if field in {"size", "orientation"}:
            continue
        if abs(actual_geometry[field] - float(expected)) > 0.08:
            errors.append(
                f"page_geometry_mismatch:{field}:{actual_geometry[field]:.2f}!={expected}"
            )

    typography = contract["typography"]
    style_expectations = {
        "Normal": typography["normal_pt"],
        "Title": typography["title_pt"],
        "Heading 1": typography["heading_1_pt"],
        "Heading 2": typography["heading_2_pt"],
        "Heading 3": typography["heading_3_pt"],
    }
    for style_name, expected in style_expectations.items():
        actual = effective_style_size(document, style_name)
        if actual is None or abs(actual - float(expected)) > 0.1:
            errors.append(f"style_size_mismatch:{style_name}:{actual}!={expected}")
    normal_font = document.styles["Normal"].font.name
    if normal_font != typography["font_family"]:
        errors.append(f"normal_font_mismatch:{normal_font}!={typography['font_family']}")

    undersized_table_runs = []
    minimum_table_pt = float(typography["minimum_table_pt"])
    for table_index, table in enumerate(document.tables):
        for row_index, row in enumerate(table.rows):
            for cell_index, cell in enumerate(row.cells):
                for paragraph in cell.paragraphs:
                    for run in paragraph.runs:
                        if not run.text.strip():
                            continue
                        size = run.font.size.pt if run.font.size else None
                        if size is not None and size + 0.01 < minimum_table_pt:
                            undersized_table_runs.append(
                                f"{table_index}:{row_index}:{cell_index}:{size}"
                            )
    if undersized_table_runs:
        errors.append(
            "table_font_below_minimum:" + ",".join(undersized_table_runs[:20])
        )

    comments = document.core_properties.comments or ""
    required_markers = {
        "data_evidence_sha256": payload_sha,
        "template_contract_sha256": contract_sha,
        "STOCK_DATA_TRADE_DATE": payload["trade_date"],
    }
    for marker, expected in required_markers.items():
        if f"{marker}={expected}" not in comments:
            errors.append(f"machine_marker_missing_or_wrong:{marker}")

    for item in payload["concept_summary"]:
        if item["concept"] not in visible_text:
            errors.append(f"concept_missing:{item['concept']}")
        count_text = f"{item['verified_limit_up_count']}只"
        if count_text not in visible_text:
            errors.append(f"concept_count_missing:{item['concept']}:{count_text}")
        model = item.get("logic_model") or {}
        position_conclusion = str(model.get("position_conclusion") or "")
        if position_conclusion and position_conclusion not in visible_text:
            errors.append(f"position_conclusion_missing:{item['concept']}")
        for row in item.get("position_top") or []:
            for field in ("relative_comparison", "weakness", "invalidation"):
                value = str(row.get(field) or "")
                if value and value not in visible_text:
                    errors.append(
                        f"position_reason_field_missing:{item['concept']}:"
                        f"{row.get('stock_code')}:{field}"
                    )
    for item in payload["report_claims"]:
        if item["text"] not in visible_text:
            errors.append(f"report_claim_missing:{item['claim_id']}")
    for item in payload["morning_limit_ups"]:
        if item["stock_code"] not in visible_text:
            errors.append(f"morning_stock_missing:{item['stock_code']}")
    for item in payload["lhb_net_buy_ge_100m"]:
        if item["stock_code"] not in visible_text:
            errors.append(f"lhb_stock_missing:{item['stock_code']}")

    with ZipFile(docx_path) as archive:
        document_xml = archive.read("word/document.xml").decode("utf-8")
    business_position = document_xml.find("A股每日涨停板")
    if business_position < 0:
        errors.append("business_cover_missing")

    render_detail = {"page_count": None, "pages": []}
    if render_dir is None:
        errors.append("render_dir_required")
    else:
        render_errors, render_detail = check_render(render_dir, contract)
        errors.extend(render_errors)

    return {
        "status": "PASS" if not errors else "BLOCKED",
        "errors": errors,
        "docx": str(docx_path),
        "docx_sha256": sha256_path(docx_path),
        "payload_sha256": payload_sha,
        "template_contract_sha256": contract_sha,
        "paragraph_count": len(document.paragraphs),
        "table_count": table_count,
        "one_by_one_table_count": one_by_one,
        "content_identity_gate": identity_detail,
        "render": render_detail,
    }


def main() -> int:
    parser = argparse.ArgumentParser(description="龙头深度研究 Word 模板与交付验收")
    parser.add_argument("--docx", required=True, type=Path)
    parser.add_argument("--input", required=True, type=Path)
    parser.add_argument("--manifest", required=True, type=Path)
    parser.add_argument("--template-contract", required=True, type=Path)
    parser.add_argument("--render-dir", required=True, type=Path)
    args = parser.parse_args()
    result = validate(
        args.docx.resolve(),
        args.input.resolve(),
        args.manifest.resolve(),
        args.template_contract.resolve(),
        args.render_dir.resolve(),
    )
    print(json.dumps(result, ensure_ascii=False, indent=2))
    return 0 if result["status"] == "PASS" else 2


if __name__ == "__main__":
    raise SystemExit(main())
