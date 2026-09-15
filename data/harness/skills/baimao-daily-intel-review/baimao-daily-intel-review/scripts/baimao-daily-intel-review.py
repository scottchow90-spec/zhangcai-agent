#!/usr/bin/env python3
# -*- coding: utf-8 -*-
from __future__ import annotations
# ONESTOCK_EMBEDDED_LOCAL_IMPORT_BOOTSTRAP
import sys as _onestock_embedded_sys
from pathlib import Path as _OneStockEmbeddedPath
_onestock_embedded_dir = str(_OneStockEmbeddedPath(__file__).resolve().parent)
if _onestock_embedded_dir not in _onestock_embedded_sys.path:
    _onestock_embedded_sys.path.insert(0, _onestock_embedded_dir)

import argparse
import datetime as dt
import hashlib
import importlib.util
import json
import math
import os
import re
import subprocess
import sys
from pathlib import Path
from typing import Any

os.environ.setdefault("PYTHONIOENCODING", "utf-8")
os.environ.setdefault("PYTHONUTF8", "1")
try:
    sys.stdout.reconfigure(encoding="utf-8")
    sys.stderr.reconfigure(encoding="utf-8")
except Exception:
    pass

WORKSPACE = Path(r"D:\C盘转移\日志\codex")
SKILL_DIR = Path(__file__).resolve().parents[1]
SCRIPTS_DIR = SKILL_DIR / "scripts"
TEMPLATES_DIR = SKILL_DIR / "templates"
REPORTS_DIR = SKILL_DIR / "reports"
GENERATED_DIR = REPORTS_DIR / "generated"
EVIDENCE_DIR = REPORTS_DIR / "evidence"
QUALITY_DIR = REPORTS_DIR / "quality"
LOCK = WORKSPACE / "hooks" / "skill_workflow_lock.py"
SUBSTANTIVE_GATE = WORKSPACE / "hooks" / "stock_workflow_substantive_gate.py"
TDX_HUB = WORKSPACE / "skills" / "tdx-local-hub" / "scripts" / "tdx_hub.py"
PYTHON_EXE = sys.executable
SKILL_NAME_DISPLAY = "白猫老师每日资讯及复盘报告"
LOCKED_PRODUCTION_SCRIPT = "baimao-daily-intel-review_closure_gate.py"

MODE_TO_TITLE = {
    "morning": "每日早盘资讯",
    "midday": "每日午盘异动推送",
    "closing": "每日复盘报告",
}
MODE_TO_TEMPLATE = {
    "morning": "morning_brief.md",
    "midday": "midday_push.md",
    "closing": "closing_review.md",
}
REQUIRED_INDEX_NAMES = ["上证", "深证", "创业", "科创", "北证"]
REQUIRED_BUSINESS_SECTIONS = [
    "## 事件池与主题聚类",
    "## 市场反馈验证",
    "## 资金与风格行为",
    "## 风险分层",
    "## 次日观察计划",
]


def now_local() -> dt.datetime:
    return dt.datetime.now().astimezone()


def today() -> str:
    return now_local().strftime("%Y-%m-%d")


def emit(payload: dict[str, Any], code: int = 0) -> int:
    print(json.dumps(payload, ensure_ascii=False, indent=2, default=str))
    return code


def run_capture(cmd: list[str], timeout: int = 120, cwd: Path | None = None) -> dict[str, Any]:
    try:
        r = subprocess.run(
            cmd,
            cwd=str(cwd or SKILL_DIR),
            capture_output=True,
            text=True,
            encoding="utf-8",
            errors="ignore",
            timeout=timeout,
        )
        return {
            "ok": r.returncode == 0,
            "exit": r.returncode,
            "cmd": cmd,
            "stdout": r.stdout,
            "stdout_tail": r.stdout[-1800:],
            "stderr_tail": r.stderr[-800:],
        }
    except Exception as exc:
        return {"ok": False, "cmd": cmd, "error": f"{type(exc).__name__}: {exc}"}


def run_json(cmd: list[str], timeout: int = 60) -> dict[str, Any]:
    result = run_capture(cmd, timeout=timeout, cwd=WORKSPACE)
    try:
        payload = json.loads(result.get("stdout") or "{}")
    except Exception:
        payload = {"ok": False, "parse_error": "stdout is not json", "stdout_tail": result.get("stdout_tail", "")}
    payload["_cmd"] = cmd
    payload["_exit"] = result.get("exit")
    if result.get("stderr_tail"):
        payload["_stderr_tail"] = result["stderr_tail"]
    return payload


def load_json(path: Path) -> dict[str, Any]:
    return json.loads(path.read_text(encoding="utf-8-sig"))


def write_json(path: Path, payload: dict[str, Any]) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(json.dumps(payload, ensure_ascii=False, indent=2, default=str), encoding="utf-8")


def norm_text(value: Any) -> str:
    return str(value if value is not None else "").replace("\r", " ").replace("\n", " ").strip()


def clean_cell(value: Any) -> str:
    text = norm_text(value)
    return text.replace("|", "/")


def pct(value: Any) -> float:
    try:
        return float(value)
    except Exception:
        return 0.0


def fmt_pct(value: Any) -> str:
    return f"{pct(value):+.2f}%"


def clamp(value: float, low: float, high: float) -> float:
    return max(low, min(high, value))


def markdown_cells(line: str) -> list[str]:
    return [cell.strip() for cell in line.strip().strip("|").split("|")]


def add_table(lines: list[str], headers: list[str], rows: list[list[Any]]) -> None:
    lines.append("| " + " | ".join(headers) + " |")
    lines.append("|" + "|".join(["---"] * len(headers)) + "|")
    for row in rows:
        padded = [clean_cell(x) for x in row] + [""] * max(0, len(headers) - len(row))
        lines.append("| " + " | ".join(padded[: len(headers)]) + " |")
    lines.append("")


def write_docx_from_markdown(md_text: str, docx_path: Path, title: str) -> None:
    try:
        from docx import Document
        from docx.shared import Pt
    except Exception as exc:
        raise RuntimeError(f"python-docx unavailable: {exc}") from exc

    document = Document()
    styles = document.styles
    styles["Normal"].font.name = "Microsoft YaHei"
    styles["Normal"].font.size = Pt(10.5)
    document.add_heading(title, level=0)
    lines = md_text.splitlines()
    i = 0
    while i < len(lines):
        line = lines[i].rstrip()
        if not line:
            i += 1
            continue
        if line.startswith("<!--"):
            i += 1
            continue
        if line.startswith("# "):
            document.add_heading(line[2:].strip(), level=1)
            i += 1
            continue
        if line.startswith("## "):
            document.add_heading(line[3:].strip(), level=2)
            i += 1
            continue
        if line.startswith("### "):
            document.add_heading(line[4:].strip(), level=3)
            i += 1
            continue
        if line.startswith("> "):
            document.add_paragraph(line[2:].strip())
            i += 1
            continue
        if line.startswith("|") and "|" in line[1:]:
            table_lines: list[str] = []
            while i < len(lines) and lines[i].strip().startswith("|"):
                table_lines.append(lines[i].strip())
                i += 1
            rows = [markdown_cells(row) for row in table_lines if "---" not in row]
            if rows:
                table = document.add_table(rows=len(rows), cols=max(len(row) for row in rows))
                table.style = "Table Grid"
                for r_idx, row in enumerate(rows):
                    for c_idx, cell in enumerate(row):
                        table.cell(r_idx, c_idx).text = cell
            continue
        document.add_paragraph(line)
        i += 1
    docx_path.parent.mkdir(parents=True, exist_ok=True)
    document.save(str(docx_path))


def resolve_mode_auto(moment: dt.datetime) -> str:
    if moment.weekday() >= 5:
        return "morning"
    if 11 <= moment.hour < 14:
        return "midday"
    if moment.hour >= 15:
        return "closing"
    return "morning"


def source_map(evidence: dict[str, Any]) -> dict[str, dict[str, Any]]:
    return {str(item.get("id")): item for item in evidence.get("sources", []) if item.get("id")}


def source_label(evidence: dict[str, Any], source_id: Any) -> str:
    sid = str(source_id or "")
    item = source_map(evidence).get(sid, {})
    name = item.get("name") or sid
    collected = item.get("collected_at") or item.get("time") or ""
    return f"{name} {collected}".strip()


def ensure_source(blocks: list[str], evidence: dict[str, Any], obj: dict[str, Any], path: str) -> None:
    sid = str(obj.get("source") or "")
    if not sid:
        blocks.append(f"{path}.source 缺失")
    elif sid not in source_map(evidence):
        blocks.append(f"{path}.source={sid} 未在 sources 中登记")


def require_list(blocks: list[str], evidence: dict[str, Any], key: str, min_count: int) -> list[dict[str, Any]]:
    rows = evidence.get(key, [])
    if not isinstance(rows, list) or len(rows) < min_count:
        blocks.append(f"{key} 数量不足，至少需要 {min_count}")
        return []
    return [row for row in rows if isinstance(row, dict)]


def require_keys(blocks: list[str], obj: dict[str, Any], keys: list[str], path: str) -> None:
    for key in keys:
        value = obj.get(key)
        if value is None or value == "":
            blocks.append(f"{path}.{key} 缺失")


def validate_indices(blocks: list[str], evidence: dict[str, Any]) -> None:
    rows = require_list(blocks, evidence, "indices", 5)
    names = {norm_text(row.get("name")) for row in rows}
    for name in REQUIRED_INDEX_NAMES:
        if name not in names:
            blocks.append(f"indices 缺少 {name}")
    for idx, row in enumerate(rows):
        require_keys(blocks, row, ["name", "pct", "reason", "source"], f"indices[{idx}]")
        ensure_source(blocks, evidence, row, f"indices[{idx}]")


def split_sectors(evidence: dict[str, Any]) -> tuple[list[dict[str, Any]], list[dict[str, Any]]]:
    rows = [row for row in evidence.get("sectors", []) if isinstance(row, dict)]
    sorted_rows = sorted(rows, key=lambda item: pct(item.get("pct")), reverse=True)
    up = [row for row in sorted_rows if pct(row.get("pct")) >= 0][:3]
    down = [row for row in reversed(sorted_rows) if pct(row.get("pct")) < 0][:3]
    return up, down


def validate_common_market(blocks: list[str], evidence: dict[str, Any]) -> None:
    validate_indices(blocks, evidence)
    breadth = evidence.get("breadth")
    if not isinstance(breadth, dict):
        blocks.append("breadth 缺失")
    else:
        require_keys(blocks, breadth, ["up", "down", "source"], "breadth")
        ensure_source(blocks, evidence, breadth, "breadth")
    sectors = require_list(blocks, evidence, "sectors", 6)
    up, down = split_sectors(evidence)
    if len(up) < 3:
        blocks.append("sectors 涨幅前三不足")
    if len(down) < 3:
        blocks.append("sectors 跌幅前三不足")
    for idx, row in enumerate(sectors):
        require_keys(blocks, row, ["name", "pct", "reason", "source"], f"sectors[{idx}]")
        ensure_source(blocks, evidence, row, f"sectors[{idx}]")


def validate_enhanced_modules(blocks: list[str], evidence: dict[str, Any]) -> None:
    headline = evidence.get("headline_conclusion")
    if not isinstance(headline, dict):
        blocks.append("headline_conclusion 缺失")
    else:
        require_keys(
            blocks,
            headline,
            ["market_direction", "strongest_theme", "second_theme", "major_risk", "capital_behavior", "next_session_observation", "source"],
            "headline_conclusion",
        )
        ensure_source(blocks, evidence, headline, "headline_conclusion")

    events = require_list(blocks, evidence, "event_pool", 3)
    for idx, row in enumerate(events):
        require_keys(blocks, row, ["time", "title", "category", "source", "heat", "keywords", "summary", "impact", "confidence"], f"event_pool[{idx}]")
        ensure_source(blocks, evidence, row, f"event_pool[{idx}]")
        keywords = row.get("keywords")
        if not isinstance(keywords, list) or not keywords:
            blocks.append(f"event_pool[{idx}].keywords 缺失")

    clusters = require_list(blocks, evidence, "theme_clusters", 2)
    for idx, row in enumerate(clusters):
        require_keys(
            blocks,
            row,
            ["name", "score", "conclusion", "evidence", "logic", "market_feedback", "capital_feedback", "watch_targets", "next_observation", "invalidation", "source"],
            f"theme_clusters[{idx}]",
        )
        ensure_source(blocks, evidence, row, f"theme_clusters[{idx}]")
        try:
            score = float(row.get("score"))
        except (TypeError, ValueError):
            blocks.append(f"theme_clusters[{idx}].score 必须是数值")
        else:
            if not math.isfinite(score) or not 0 <= score <= 100:
                blocks.append(f"theme_clusters[{idx}].score 必须在0到100之间")
        targets = row.get("watch_targets")
        if not isinstance(targets, list) or not targets:
            blocks.append(f"theme_clusters[{idx}].watch_targets 缺失")

    feedback = evidence.get("market_feedback")
    if not isinstance(feedback, dict):
        blocks.append("market_feedback 缺失")
    else:
        require_keys(blocks, feedback, ["price_breadth", "sector_diffusion", "limit_up_structure", "volume_state", "risk_appetite", "source"], "market_feedback")
        ensure_source(blocks, evidence, feedback, "market_feedback")

    capital = evidence.get("capital_behavior")
    if not isinstance(capital, dict):
        blocks.append("capital_behavior 缺失")
    else:
        require_keys(blocks, capital, ["main_style", "active_side", "defensive_side", "large_vs_micro", "source"], "capital_behavior")
        ensure_source(blocks, evidence, capital, "capital_behavior")

    risks = require_list(blocks, evidence, "risk_items", 2)
    for idx, row in enumerate(risks):
        require_keys(blocks, row, ["risk", "level", "trigger", "invalidation", "source"], f"risk_items[{idx}]")
        ensure_source(blocks, evidence, row, f"risk_items[{idx}]")

    plans = require_list(blocks, evidence, "next_session_plan", 2)
    for idx, row in enumerate(plans):
        require_keys(blocks, row, ["direction", "observe", "trigger", "invalidation"], f"next_session_plan[{idx}]")


def validate_evidence(mode: str, evidence: dict[str, Any], allow_fixture: bool = False) -> list[str]:
    blocks: list[str] = []
    if evidence.get("workflow") not in ("baimao-daily-intel-review", SKILL_DIR.name):
        blocks.append("workflow 不匹配 baimao-daily-intel-review")
    if evidence.get("mode") != mode:
        blocks.append(f"mode 不匹配，期望 {mode}")
    if evidence.get("is_test_fixture") and not allow_fixture:
        blocks.append("检测到 SELFTEST_FIXTURE，正式运行禁止使用")
    if not evidence.get("trade_date"):
        blocks.append("trade_date 缺失")
    if not isinstance(evidence.get("sources"), list) or not evidence.get("sources"):
        blocks.append("sources 缺失")
    validate_enhanced_modules(blocks, evidence)
    if mode in ("midday", "closing"):
        validate_common_market(blocks, evidence)
        style = evidence.get("style")
        if not isinstance(style, dict):
            blocks.append("style 缺失")
        else:
            require_keys(blocks, style, ["hs300_pct", "microcap_pct", "source"], "style")
            ensure_source(blocks, evidence, style, "style")
    if mode == "morning":
        events = require_list(blocks, evidence, "after_close_events", 1)
        if len(events) > 2:
            blocks.append("after_close_events 超过 2 条，需要筛选成 1-2 条")
        for idx, row in enumerate(events):
            require_keys(blocks, row, ["title", "summary", "impact", "source"], f"after_close_events[{idx}]")
            ensure_source(blocks, evidence, row, f"after_close_events[{idx}]")
        regulatory = require_list(blocks, evidence, "regulatory_events", 1)
        for idx, row in enumerate(regulatory):
            require_keys(blocks, row, ["type", "subject", "event", "risk", "source"], f"regulatory_events[{idx}]")
            ensure_source(blocks, evidence, row, f"regulatory_events[{idx}]")
        overseas = evidence.get("overseas")
        if not isinstance(overseas, dict):
            blocks.append("overseas 缺失")
        else:
            require_keys(blocks, overseas, ["indices", "source"], "overseas")
            ensure_source(blocks, evidence, overseas, "overseas")
        calendar = evidence.get("calendar")
        if not isinstance(calendar, dict):
            blocks.append("calendar 缺失")
        else:
            require_keys(blocks, calendar, ["items", "source"], "calendar")
            ensure_source(blocks, evidence, calendar, "calendar")
        viewpoints = evidence.get("viewpoints", [])
        if not isinstance(viewpoints, list) or not (1 <= len(viewpoints) <= 5):
            blocks.append("viewpoints 需要 1-5 条")
    if mode == "midday":
        news = evidence.get("news", [])
        if news:
            for idx, row in enumerate(news):
                if isinstance(row, dict):
                    ensure_source(blocks, evidence, row, f"news[{idx}]")
    if mode == "closing":
        news = require_list(blocks, evidence, "news", 2)
        if len(news) > 3:
            blocks.append("news 超过 3 条，需要筛选成 2-3 条")
        for idx, row in enumerate(news):
            require_keys(blocks, row, ["time", "title", "heat", "summary", "impact", "source"], f"news[{idx}]")
            ensure_source(blocks, evidence, row, f"news[{idx}]")
        summary = norm_text(evidence.get("summary"))
        if not summary:
            blocks.append("summary 缺失")
        elif len(summary) > 90:
            blocks.append("summary 超过 90 字，需要压缩到约 50 字")
        directions = evidence.get("watch_directions", [])
        if not isinstance(directions, list) or not (2 <= len(directions) <= 3):
            blocks.append("watch_directions 需要 2-3 个方向")
    return blocks


def index_score(value: Any) -> int:
    return int(round(clamp(50 + pct(value) * 12, 0, 100)))


def index_lamp(value: Any) -> str:
    v = pct(value)
    if v >= 0.3:
        return "绿灯"
    if v <= -0.3:
        return "红灯"
    return "黄灯"


def sentiment_state(breadth: dict[str, Any]) -> dict[str, Any]:
    up = int(float(breadth.get("up", 0) or 0))
    down = int(float(breadth.get("down", 0) or 0))
    total = max(up + down, 1)
    ratio = up / total
    if ratio >= 0.65:
        lamp = "绿灯"
    elif ratio < 0.35:
        lamp = "红灯"
    else:
        lamp = "黄灯"
    return {"up": up, "down": down, "ratio": ratio, "lamp": lamp}


def top_indices(evidence: dict[str, Any]) -> tuple[dict[str, Any], dict[str, Any]]:
    rows = [row for row in evidence.get("indices", []) if isinstance(row, dict)]
    if not rows:
        return {}, {}
    sorted_rows = sorted(rows, key=lambda item: pct(item.get("pct")), reverse=True)
    return sorted_rows[0], sorted_rows[-1]


def style_text(style: dict[str, Any]) -> str:
    hs300 = pct(style.get("hs300_pct"))
    micro = pct(style.get("microcap_pct"))
    diff = hs300 - micro
    if diff >= 0.5:
        return "偏大盘股"
    if diff <= -0.5:
        return "偏微盘股"
    return "大小盘均衡"


def hot_news(news: list[dict[str, Any]], limit: int) -> list[dict[str, Any]]:
    return sorted(news, key=lambda item: float(item.get("heat", 0) or 0), reverse=True)[:limit]


def evidence_source_table(evidence: dict[str, Any]) -> list[list[Any]]:
    rows = []
    for item in evidence.get("sources", []):
        rows.append([
            item.get("id"),
            item.get("name"),
            item.get("type"),
            item.get("collected_at") or item.get("time"),
            item.get("path") or item.get("url") or item.get("note"),
        ])
    return rows


def join_items(value: Any) -> str:
    if isinstance(value, list):
        return "、".join(clean_cell(item) for item in value)
    return clean_cell(value)


def headline_rows(evidence: dict[str, Any]) -> list[list[Any]]:
    row = evidence.get("headline_conclusion", {}) if isinstance(evidence.get("headline_conclusion"), dict) else {}
    src = source_label(evidence, row.get("source"))
    return [
        ["市场方向", row.get("market_direction"), src],
        ["最强主题", row.get("strongest_theme"), src],
        ["次强主题", row.get("second_theme"), src],
        ["主要风险", row.get("major_risk"), src],
        ["资金行为", row.get("capital_behavior"), src],
        ["下一窗口观察", row.get("next_session_observation"), src],
    ]


def event_pool_rows(evidence: dict[str, Any]) -> list[list[Any]]:
    rows: list[list[Any]] = []
    events = [row for row in evidence.get("event_pool", []) if isinstance(row, dict)]
    events = sorted(events, key=lambda item: float(item.get("heat", 0) or 0), reverse=True)
    for idx, row in enumerate(events, 1):
        rows.append([
            idx,
            row.get("time"),
            row.get("category"),
            row.get("title"),
            row.get("heat"),
            join_items(row.get("keywords")),
            row.get("impact"),
            row.get("confidence"),
            source_label(evidence, row.get("source")),
        ])
    return rows


def theme_cluster_rows(evidence: dict[str, Any]) -> list[list[Any]]:
    rows: list[list[Any]] = []
    clusters = [row for row in evidence.get("theme_clusters", []) if isinstance(row, dict)]
    clusters = sorted(clusters, key=lambda item: float(item.get("score", 0) or 0), reverse=True)
    for idx, row in enumerate(clusters, 1):
        rows.append([
            idx,
            row.get("name"),
            row.get("score"),
            row.get("conclusion"),
            row.get("evidence"),
            row.get("market_feedback"),
            row.get("capital_feedback"),
            join_items(row.get("watch_targets")),
            row.get("next_observation"),
            row.get("invalidation"),
            source_label(evidence, row.get("source")),
        ])
    return rows


def market_feedback_rows(evidence: dict[str, Any]) -> list[list[Any]]:
    row = evidence.get("market_feedback", {}) if isinstance(evidence.get("market_feedback"), dict) else {}
    src = source_label(evidence, row.get("source"))
    return [
        ["价格宽度", row.get("price_breadth"), src],
        ["板块扩散", row.get("sector_diffusion"), src],
        ["涨停结构", row.get("limit_up_structure"), src],
        ["成交状态", row.get("volume_state"), src],
        ["风险偏好", row.get("risk_appetite"), src],
    ]


def capital_behavior_rows(evidence: dict[str, Any]) -> list[list[Any]]:
    row = evidence.get("capital_behavior", {}) if isinstance(evidence.get("capital_behavior"), dict) else {}
    src = source_label(evidence, row.get("source"))
    return [
        ["主导风格", row.get("main_style"), src],
        ["进攻方向", row.get("active_side"), src],
        ["防御方向", row.get("defensive_side"), src],
        ["大盘/微盘", row.get("large_vs_micro"), src],
    ]


def risk_item_rows(evidence: dict[str, Any]) -> list[list[Any]]:
    return [
        [idx + 1, row.get("risk"), row.get("level"), row.get("trigger"), row.get("invalidation"), source_label(evidence, row.get("source"))]
        for idx, row in enumerate(evidence.get("risk_items", []))
        if isinstance(row, dict)
    ]


def next_session_plan_rows(evidence: dict[str, Any]) -> list[list[Any]]:
    return [
        [idx + 1, row.get("direction"), row.get("observe"), row.get("trigger"), row.get("invalidation")]
        for idx, row in enumerate(evidence.get("next_session_plan", []))
        if isinstance(row, dict)
    ]


def append_enhanced_modules(lines: list[str], evidence: dict[str, Any]) -> None:
    add_table(lines, ["证据ID", "来源", "类型", "采集时间", "路径/链接/说明"], evidence_source_table(evidence))
    lines.append("## 事件池与主题聚类")
    add_table(lines, ["序号", "时间", "类别", "事件", "热度", "关键词", "市场影响", "置信度", "来源"], event_pool_rows(evidence))
    add_table(lines, ["排名", "主题", "主题分", "结论", "事实证据", "盘面验证", "资金行为", "观察标的", "下一观察", "失效条件", "来源"], theme_cluster_rows(evidence))
    lines.append("## 市场反馈验证")
    add_table(lines, ["维度", "反馈", "来源"], market_feedback_rows(evidence))
    lines.append("## 资金与风格行为")
    add_table(lines, ["维度", "判断", "来源"], capital_behavior_rows(evidence))
    lines.append("## 风险分层")
    add_table(lines, ["序号", "风险", "级别", "触发条件", "失效条件", "来源"], risk_item_rows(evidence))
    lines.append("## 次日观察计划")
    add_table(lines, ["序号", "方向", "观察点", "触发条件", "失效条件"], next_session_plan_rows(evidence))


def render_morning(evidence: dict[str, Any]) -> str:
    lines: list[str] = []
    lines.append(f"# {SKILL_NAME_DISPLAY} - 每日早盘资讯")
    lines.append("")
    lines.append(f"> 交易日/自然日：{evidence.get('trade_date')}；生成时间：{now_local().isoformat(timespec='seconds')}")
    if evidence.get("is_test_fixture"):
        lines.append("> SELFTEST_FIXTURE：仅用于流程验收，不代表真实市场。")
    lines.append("")
    lines.append("## 核心结论")
    add_table(lines, ["维度", "结论", "来源"], headline_rows(evidence))
    lines.append("### 盘前观点")
    for i, view in enumerate(evidence.get("viewpoints", []), 1):
        lines.append(f"{i}. {clean_cell(view)}")
    lines.append("")
    append_enhanced_modules(lines, evidence)
    lines.append("## 原始需求模块")
    add_table(
        lines,
        ["序号", "昨日盘后大事", "热度", "影响", "来源"],
        [[i + 1, row.get("title"), row.get("repeat_count", row.get("heat", "")), row.get("impact"), source_label(evidence, row.get("source"))] for i, row in enumerate(evidence.get("after_close_events", []))],
    )
    add_table(
        lines,
        ["序号", "类型", "主体", "事件", "风险含义", "来源"],
        [[i + 1, row.get("type"), row.get("subject"), row.get("event"), row.get("risk"), source_label(evidence, row.get("source"))] for i, row in enumerate(evidence.get("regulatory_events", []))],
    )
    overseas = evidence.get("overseas", {})
    overseas_rows = []
    for row in overseas.get("indices", []):
        trigger = "触发" if abs(pct(row.get("pct"))) >= 1.0 else "未触发"
        overseas_rows.append([row.get("name"), fmt_pct(row.get("pct")), trigger, row.get("reason"), source_label(evidence, overseas.get("source"))])
    for row in overseas.get("sectors", []):
        trigger = "触发" if abs(pct(row.get("pct"))) >= 1.0 else "未触发"
        overseas_rows.append([row.get("name"), fmt_pct(row.get("pct")), trigger, row.get("reason"), source_label(evidence, overseas.get("source"))])
    add_table(lines, ["指数/板块", "涨跌幅", "1%阈值", "原因", "来源"], overseas_rows)
    calendar = evidence.get("calendar", {})
    add_table(
        lines,
        ["类型", "事项", "影响范围", "来源"],
        [[row.get("type"), row.get("event"), row.get("impact"), source_label(evidence, calendar.get("source"))] for row in calendar.get("items", [])],
    )
    lines.append("## 风险与失效")
    lines.append("盘前结论只覆盖已采集证据窗口；若开盘后指数、成交额、涨跌家数或监管消息发生明显变化，早盘观点自动降级。")
    lines.append("")
    return "\n".join(lines)


def render_midday(evidence: dict[str, Any]) -> str:
    sentiment = sentiment_state(evidence.get("breadth", {}))
    up_sectors, down_sectors = split_sectors(evidence)
    style = evidence.get("style", {})
    lines: list[str] = []
    lines.append(f"# {SKILL_NAME_DISPLAY} - 每日午盘异动推送")
    lines.append("")
    lines.append(f"> 交易日：{evidence.get('trade_date')}；生成时间：{now_local().isoformat(timespec='seconds')}")
    if evidence.get("is_test_fixture"):
        lines.append("> SELFTEST_FIXTURE：仅用于流程验收，不代表真实市场。")
    lines.append("")
    lines.append("## 核心结论")
    lines.append(f"- 市场情绪：{sentiment['lamp']}，上涨占比 {sentiment['ratio']:.1%}。")
    lines.append(f"- 风格偏好：{style_text(style)}，沪深300 {fmt_pct(style.get('hs300_pct'))}，微盘指数 {fmt_pct(style.get('microcap_pct'))}。")
    if up_sectors:
        lines.append(f"- 领涨主线：{up_sectors[0].get('name')}；拖累方向：{down_sectors[0].get('name') if down_sectors else ''}。")
    lines.append("")
    add_table(lines, ["维度", "结论", "来源"], headline_rows(evidence))
    append_enhanced_modules(lines, evidence)
    lines.append("## 原始需求模块")
    add_table(
        lines,
        ["指数", "涨跌幅", "阈值处理", "原因", "评分", "来源"],
        [[row.get("name"), fmt_pct(row.get("pct")), threshold_note(row), row.get("reason"), index_score(row.get("pct")), source_label(evidence, row.get("source"))] for row in evidence.get("indices", [])],
    )
    add_table(lines, ["上涨家数", "下跌家数", "上涨占比", "情绪灯", "来源"], [[sentiment["up"], sentiment["down"], f"{sentiment['ratio']:.1%}", sentiment["lamp"], source_label(evidence, evidence.get("breadth", {}).get("source"))]])
    add_table(lines, ["排名", "方向", "板块", "涨跌幅", "异动原因", "来源"], sector_rows(evidence, up_sectors, down_sectors))
    add_table(lines, ["维度", "表现", "结论", "来源"], [["沪深300", fmt_pct(style.get("hs300_pct")), style_text(style), source_label(evidence, style.get("source"))], ["微盘指数", fmt_pct(style.get("microcap_pct")), style_text(style), source_label(evidence, style.get("source"))]])
    lines.append("## 风险与失效")
    lines.append("午盘推送是 12:00 附近窗口判断；若午后成交额、涨跌家数、领涨板块或指数方向反转，风格判断自动降级。")
    lines.append("")
    return "\n".join(lines)


def threshold_note(row: dict[str, Any]) -> str:
    name = norm_text(row.get("name"))
    value = abs(pct(row.get("pct")))
    if name == "上证" and value >= 1:
        return "上证超 1%，必须解释"
    if name in ("创业", "科创") and value >= 2:
        return f"{name}超 2%，必须解释"
    return "未触发强阈值"


def sector_rows(evidence: dict[str, Any], up_sectors: list[dict[str, Any]], down_sectors: list[dict[str, Any]]) -> list[list[Any]]:
    rows: list[list[Any]] = []
    for i, row in enumerate(up_sectors, 1):
        rows.append([i, "领涨", row.get("name"), fmt_pct(row.get("pct")), row.get("reason"), source_label(evidence, row.get("source"))])
    for i, row in enumerate(down_sectors, 1):
        rows.append([i, "领跌", row.get("name"), fmt_pct(row.get("pct")), row.get("reason"), source_label(evidence, row.get("source"))])
    return rows


def render_closing(evidence: dict[str, Any]) -> str:
    best, worst = top_indices(evidence)
    sentiment = sentiment_state(evidence.get("breadth", {}))
    up_sectors, down_sectors = split_sectors(evidence)
    style = evidence.get("style", {})
    selected_news = hot_news([row for row in evidence.get("news", []) if isinstance(row, dict)], 3)
    breadth = evidence.get("breadth", {})
    lines: list[str] = []
    lines.append(f"# {SKILL_NAME_DISPLAY} - 每日复盘报告")
    lines.append("")
    lines.append(f"> 交易日：{evidence.get('trade_date')}；生成时间：{now_local().isoformat(timespec='seconds')}")
    if evidence.get("is_test_fixture"):
        lines.append("> SELFTEST_FIXTURE：仅用于流程验收，不代表真实市场。")
    lines.append("")
    lines.append("## 核心结论")
    lines.append(f"- 指数强弱：上涨最多为 {best.get('name')} {fmt_pct(best.get('pct'))}，下跌最多为 {worst.get('name')} {fmt_pct(worst.get('pct'))}。")
    lines.append(f"- 市场情绪：{sentiment['lamp']}，上涨占比 {sentiment['ratio']:.1%}，涨停 {breadth.get('limit_up')}，跌停 {breadth.get('limit_down')}，连板 {breadth.get('lianban')}。")
    lines.append(f"- 板块风格：领涨 {up_sectors[0].get('name') if up_sectors else ''}，领跌 {down_sectors[0].get('name') if down_sectors else ''}，大小盘风格为 {style_text(style)}。")
    lines.append(f"- 50字总结：{clean_cell(evidence.get('summary'))}")
    lines.append("")
    add_table(lines, ["维度", "结论", "来源"], headline_rows(evidence))
    append_enhanced_modules(lines, evidence)
    lines.append("## 原始需求模块")
    add_table(
        lines,
        ["指数", "涨跌幅", "红绿灯", "评分", "原因", "来源"],
        [[row.get("name"), fmt_pct(row.get("pct")), index_lamp(row.get("pct")), index_score(row.get("pct")), row.get("reason"), source_label(evidence, row.get("source"))] for row in evidence.get("indices", [])],
    )
    ratio_text = f"{sentiment['up']}:{sentiment['down']}"
    add_table(lines, ["上涨家数", "下跌家数", "涨跌比", "涨停", "跌停", "连板", "情绪灯", "来源"], [[sentiment["up"], sentiment["down"], ratio_text, breadth.get("limit_up"), breadth.get("limit_down"), breadth.get("lianban"), sentiment["lamp"], source_label(evidence, breadth.get("source"))]])
    add_table(lines, ["排名", "方向", "板块", "涨跌幅", "异动原因", "来源"], sector_rows(evidence, up_sectors, down_sectors))
    add_table(
        lines,
        ["序号", "时间", "事件", "热度", "市场影响", "来源"],
        [[i + 1, row.get("time"), row.get("title"), row.get("heat"), row.get("impact"), source_label(evidence, row.get("source"))] for i, row in enumerate(selected_news)],
    )
    add_table(lines, ["维度", "表现", "结论", "来源"], [["沪深300", fmt_pct(style.get("hs300_pct")), style_text(style), source_label(evidence, style.get("source"))], ["微盘指数", fmt_pct(style.get("microcap_pct")), style_text(style), source_label(evidence, style.get("source"))]])
    add_table(lines, ["未来关注方向", "触发条件/观察点"], [[row.get("name") if isinstance(row, dict) else row, row.get("trigger") if isinstance(row, dict) else "看证据更新"] for row in evidence.get("watch_directions", [])])
    lines.append("## 风险与失效")
    lines.append("若次日指数低开后市场宽度转弱、领涨板块兑现、监管或公告风险发酵，今日复盘结论降级；所有方向仅作研究观察，不构成交易指令。")
    lines.append("")
    return "\n".join(lines)


def render_report(mode: str, evidence: dict[str, Any]) -> str:
    if mode == "morning":
        return render_morning(evidence)
    if mode == "midday":
        return render_midday(evidence)
    if mode == "closing":
        return render_closing(evidence)
    raise ValueError(mode)


def quality_gate(payload: dict[str, Any], allow_fixture: bool = False) -> dict[str, Any]:
    checks: list[dict[str, Any]] = []

    def add(name: str, ok: bool, detail: Any = "") -> None:
        checks.append({"name": name, "ok": bool(ok), "detail": detail})

    artifacts = payload.get("artifacts", {})
    md_path = Path(artifacts.get("markdown", ""))
    docx_path = Path(artifacts.get("word", ""))
    json_path = Path(artifacts.get("json", ""))
    add("status_clean_pass", payload.get("status") == "CLEAN_PASS", payload.get("status"))
    add("markdown_exists", md_path.exists() and md_path.stat().st_size > 1000, str(md_path))
    add("word_exists", docx_path.exists() and docx_path.suffix.lower() == ".docx" and docx_path.stat().st_size > 5000, str(docx_path))
    add("json_exists", json_path.exists() and json_path.stat().st_size > 500, str(json_path))
    add("source_count_ok", int(payload.get("source_count", 0)) >= 1, payload.get("source_count"))
    add("not_fixture_for_formal", allow_fixture or not payload.get("is_test_fixture"), payload.get("is_test_fixture"))
    if md_path.exists():
        text = md_path.read_text(encoding="utf-8-sig", errors="ignore")
        add("no_data_required_in_final_markdown", "DATA_REQUIRED" not in text, "DATA_REQUIRED scan")
        add("has_core_conclusion", "## 核心结论" in text, "核心结论")
        add("has_source_table", "证据ID" in text and "来源" in text, "source table")
        present_sections = [section for section in REQUIRED_BUSINESS_SECTIONS if section in text]
        add("has_enhanced_business_modules", len(present_sections) == len(REQUIRED_BUSINESS_SECTIONS), present_sections)
        add("has_theme_scoring_and_feedback", "主题分" in text and "盘面验证" in text and "资金行为" in text, "theme score/feedback")
    if docx_path.exists():
        try:
            from docx import Document

            doc = Document(str(docx_path))
            add("word_readable", len(doc.paragraphs) >= 6 and len(doc.tables) >= 7, {"paragraphs": len(doc.paragraphs), "tables": len(doc.tables)})
        except Exception as exc:
            add("word_readable", False, f"{type(exc).__name__}: {exc}")
    ok = all(item["ok"] for item in checks)
    return {"status": "CLEAN_PASS" if ok else "BLOCKED", "checks": checks, "blocks": [item for item in checks if not item["ok"]]}


def build_artifact_paths(mode: str, trade_date: str, output_dir: Path | None = None) -> dict[str, Path]:
    stamp = now_local().strftime("%Y%m%d_%H%M%S")
    day_dir = output_dir or (GENERATED_DIR / trade_date)
    day_dir.mkdir(parents=True, exist_ok=True)
    stem = f"{mode}_{stamp}"
    return {"markdown": day_dir / f"{stem}.md", "word": day_dir / f"{stem}.docx", "json": day_dir / f"{stem}.json"}


def run_workflow(mode: str, evidence: dict[str, Any], allow_fixture: bool = False, output_dir: Path | None = None) -> tuple[dict[str, Any], int]:
    blocks = validate_evidence(mode, evidence, allow_fixture=allow_fixture)
    trade_date = evidence.get("trade_date") or today()
    if blocks:
        blocked_path = (QUALITY_DIR / trade_date / f"{mode}_blocked_{now_local().strftime('%Y%m%d_%H%M%S')}.json")
        payload = {
            "skill": SKILL_DIR.name,
            "display_name": SKILL_NAME_DISPLAY,
            "mode": mode,
            "trade_date": trade_date,
            "status": "BLOCKED",
            "blocks": blocks,
            "evidence_file_required": True,
        }
        write_json(blocked_path, payload)
        payload["artifact"] = str(blocked_path)
        return payload, 2

    paths = build_artifact_paths(mode, trade_date, output_dir=output_dir)
    md_text = render_report(mode, evidence)
    paths["markdown"].write_text(md_text, encoding="utf-8")
    write_docx_from_markdown(md_text, paths["word"], f"{SKILL_NAME_DISPLAY} - {MODE_TO_TITLE[mode]}")
    payload = {
        "skill": SKILL_DIR.name,
        "display_name": SKILL_NAME_DISPLAY,
        "mode": mode,
        "trade_date": trade_date,
        "generated_at": now_local().isoformat(timespec="seconds"),
        "status": "CLEAN_PASS",
        "is_test_fixture": bool(evidence.get("is_test_fixture")),
        "source_count": len(evidence.get("sources", [])),
        "artifacts": {key: str(path) for key, path in paths.items()},
        "evidence_hash": hashlib.sha256(json.dumps(evidence, ensure_ascii=False, sort_keys=True, default=str).encode("utf-8")).hexdigest(),
    }
    write_json(paths["json"], payload)
    gate = quality_gate(payload, allow_fixture=allow_fixture)
    payload["quality_gate"] = gate
    if gate["status"] != "CLEAN_PASS":
        payload["status"] = "BLOCKED"
    write_json(paths["json"], payload)
    return payload, 0 if payload["status"] == "CLEAN_PASS" else 2


def fixture_payload(mode: str, trade_date: str) -> dict[str, Any]:
    sources = [
        {"id": "S_MARKET", "name": "SELFTEST_FIXTURE行情包", "type": "selftest_market", "collected_at": f"{trade_date} 15:05", "note": "仅用于流程验收"},
        {"id": "S_NEWS", "name": "SELFTEST_FIXTURE新闻包", "type": "selftest_news", "collected_at": f"{trade_date} 15:10", "note": "仅用于流程验收"},
        {"id": "S_CAL", "name": "SELFTEST_FIXTURE日历包", "type": "selftest_calendar", "collected_at": f"{trade_date} 08:00", "note": "仅用于流程验收"},
        {"id": "S_US", "name": "SELFTEST_FIXTURE外盘包", "type": "selftest_overseas", "collected_at": f"{trade_date} 07:30", "note": "仅用于流程验收"},
    ]
    indices = [
        {"name": "上证", "pct": 0.62, "reason": "权重稳定且情绪修复", "source": "S_MARKET"},
        {"name": "深证", "pct": 0.91, "reason": "成长方向走强", "source": "S_MARKET"},
        {"name": "创业", "pct": 1.35, "reason": "新能源与科技回暖", "source": "S_MARKET"},
        {"name": "科创", "pct": 2.21, "reason": "硬科技板块拉升并触发解释阈值", "source": "S_MARKET"},
        {"name": "北证", "pct": -0.28, "reason": "小市值分化", "source": "S_MARKET"},
    ]
    sectors = [
        {"name": "机器人", "pct": 3.20, "reason": "政策预期与成交放大", "source": "S_MARKET"},
        {"name": "半导体", "pct": 2.60, "reason": "科创权重共振", "source": "S_MARKET"},
        {"name": "电力设备", "pct": 1.85, "reason": "成长风格回补", "source": "S_MARKET"},
        {"name": "煤炭", "pct": -1.10, "reason": "防御方向回落", "source": "S_MARKET"},
        {"name": "银行", "pct": -0.82, "reason": "权重承压", "source": "S_MARKET"},
        {"name": "旅游", "pct": -0.65, "reason": "消费分歧", "source": "S_MARKET"},
    ]
    base = {
        "workflow": SKILL_DIR.name,
        "mode": mode,
        "trade_date": trade_date,
        "is_test_fixture": True,
        "sources": sources,
        "indices": indices,
        "breadth": {"up": 3450, "down": 1620, "limit_up": 58, "limit_down": 7, "lianban": 13, "source": "S_MARKET"},
        "sectors": sectors,
        "style": {"hs300_pct": 0.44, "microcap_pct": -0.31, "source": "S_MARKET"},
        "headline_conclusion": {
            "market_direction": "偏强修复，成长方向占优",
            "strongest_theme": "机器人",
            "second_theme": "半导体设备",
            "major_risk": "高位题材分歧与防御板块回落",
            "capital_behavior": "资金从防御切向成长，沪深300强于微盘",
            "next_session_observation": "观察机器人和半导体是否继续扩散并获得成交承接",
            "source": "S_MARKET",
        },
        "event_pool": [
            {"time": "09:40", "title": "机器人产业链出现集中催化", "category": "产业事件", "source": "S_NEWS", "heat": 91, "keywords": ["机器人", "减速器", "控制器"], "summary": "产业政策与订单预期共同抬升热度", "impact": "带动机器人链条成为当日强方向", "confidence": "高"},
            {"time": "11:10", "title": "半导体设备方向成交放大", "category": "盘面反馈", "source": "S_NEWS", "heat": 84, "keywords": ["半导体", "设备", "科创"], "summary": "科创权重和设备链同步走强", "impact": "推动科创指数强于其他指数", "confidence": "高"},
            {"time": "14:05", "title": "防御板块午后回落", "category": "风格切换", "source": "S_NEWS", "heat": 72, "keywords": ["煤炭", "银行", "高股息"], "summary": "高股息方向回落，资金转向成长", "impact": "强化成长修复结论", "confidence": "中"},
            {"time": "15:00", "title": "涨停结构维持活跃", "category": "市场结构", "source": "S_MARKET", "heat": 78, "keywords": ["涨停", "连板", "市场宽度"], "summary": "涨停和连板维持一定数量", "impact": "支持情绪偏暖但仍需识别高位分歧", "confidence": "中"},
        ],
        "theme_clusters": [
            {
                "name": "机器人",
                "score": 86,
                "conclusion": "事件和盘面共振最强，是当日最强主题",
                "evidence": "产业事件热度最高，机器人板块涨幅居前",
                "logic": "政策预期、订单想象和成交放大共同推动主题扩散",
                "market_feedback": "板块涨幅居首且带动相关方向扩散",
                "capital_feedback": "成长资金主动承接，防御方向回落",
                "watch_targets": ["机器人", "减速器", "控制器"],
                "next_observation": "观察是否继续放量并维持板块前排",
                "invalidation": "板块跌出涨幅前列或高位股集中回落",
                "source": "S_MARKET",
            },
            {
                "name": "半导体设备",
                "score": 78,
                "conclusion": "科创强势带动设备链，属于次强主题",
                "evidence": "半导体设备事件热度较高，科创指数触发强阈值",
                "logic": "指数强势与设备链成交放大形成共振",
                "market_feedback": "科创指数涨幅突出，半导体板块居前",
                "capital_feedback": "资金偏向硬科技成长方向",
                "watch_targets": ["半导体", "设备", "科创权重"],
                "next_observation": "观察科创指数能否维持强势并扩散到细分设备",
                "invalidation": "科创指数回落且设备链成交收缩",
                "source": "S_NEWS",
            },
        ],
        "market_feedback": {
            "price_breadth": "上涨家数高于下跌家数，市场宽度偏暖",
            "sector_diffusion": "机器人、半导体、电力设备领涨，煤炭、银行、旅游领跌",
            "limit_up_structure": "涨停 58 家、跌停 7 家、连板 13 家，短线情绪仍活跃",
            "volume_state": "成长方向成交放大，防御方向承接减弱",
            "risk_appetite": "风险偏好回升但高位题材存在分歧",
            "source": "S_MARKET",
        },
        "capital_behavior": {
            "main_style": "成长修复",
            "active_side": "机器人、半导体、电力设备",
            "defensive_side": "煤炭、银行等防御方向回落",
            "large_vs_micro": "沪深300强于微盘，资金更偏指数与成长权重",
            "source": "S_MARKET",
        },
        "risk_items": [
            {"risk": "高位题材兑现", "level": "中", "trigger": "领涨板块冲高后跌出前排或炸板增多", "invalidation": "主题继续放量并维持前排扩散", "source": "S_MARKET"},
            {"risk": "防御方向拖累指数", "level": "中", "trigger": "银行、煤炭持续下行并拖累权重", "invalidation": "权重企稳且成长板块继续承接", "source": "S_MARKET"},
        ],
        "next_session_plan": [
            {"direction": "机器人", "observe": "板块排名、成交和核心分支扩散", "trigger": "继续位于涨幅前列且成交放大", "invalidation": "核心分支回落并跌出强势区"},
            {"direction": "半导体设备", "observe": "科创指数与设备链联动", "trigger": "科创指数维持强势且设备链扩散", "invalidation": "科创指数走弱并出现缩量"},
            {"direction": "市场宽度", "observe": "上涨占比和涨停连板结构", "trigger": "上涨占比维持偏高且跌停不扩散", "invalidation": "下跌家数反超并高位分歧扩大"},
        ],
        "news": [
            {"time": "09:40", "title": "机器人产业链出现集中催化", "heat": 91, "summary": "产业政策与订单预期共同抬升热度", "impact": "带动机器人、减速器和控制器方向", "source": "S_NEWS"},
            {"time": "11:10", "title": "半导体设备方向成交放大", "heat": 84, "summary": "科创权重和设备链共振", "impact": "推动科创指数走强", "source": "S_NEWS"},
            {"time": "14:05", "title": "防御板块午后回落", "heat": 72, "summary": "资金从高股息转向成长", "impact": "煤炭、银行承压", "source": "S_NEWS"},
        ],
        "summary": "成长方向修复，科创最强，市场宽度偏暖；关注机器人、半导体和电力设备持续性。",
        "watch_directions": [
            {"name": "机器人", "trigger": "继续放量并保持板块涨幅前列"},
            {"name": "半导体", "trigger": "科创指数维持强势且设备链扩散"},
            {"name": "电力设备", "trigger": "成交额改善并出现主线共振"},
        ],
    }
    if mode == "morning":
        base.update(
            {
                "after_close_events": [
                    {"title": "政策催化集中指向先进制造", "summary": "多源头条重复出现先进制造主题", "repeat_count": 4, "impact": "提升机器人与半导体关注度", "source": "S_NEWS"},
                    {"title": "海外科技股带动风险偏好", "summary": "外盘科技方向强于大盘", "repeat_count": 3, "impact": "利于成长风格开盘情绪", "source": "S_US"},
                ],
                "regulatory_events": [
                    {"type": "公告监管", "subject": "高位题材股", "event": "异动公告增多", "risk": "追高风险上升", "source": "S_NEWS"},
                ],
                "overseas": {
                    "source": "S_US",
                    "indices": [{"name": "纳指", "pct": 1.18, "reason": "大型科技股走强"}],
                    "sectors": [{"name": "半导体ETF", "pct": 1.42, "reason": "AI链条修复"}],
                },
                "calendar": {
                    "source": "S_CAL",
                    "items": [
                        {"type": "新股打新", "event": "新股A申购", "impact": "关注资金分流"},
                        {"type": "新股上市", "event": "新股B上市", "impact": "观察次新情绪"},
                    ],
                },
                "viewpoints": [
                    "盘前关注成长方向延续性，机器人和半导体需要成交确认。",
                    "若外盘科技映射兑现不足，早盘冲高方向要降低追涨预期。",
                    "监管异动增多时，高位题材股以风险识别优先。",
                ],
            }
        )
    return base


def build_draft_report(mode: str, trade_date: str, source_note: str = "") -> dict[str, Any]:
    template_path = TEMPLATES_DIR / MODE_TO_TEMPLATE[mode]
    stamp = now_local().strftime("%Y%m%d_%H%M%S")
    day_dir = GENERATED_DIR / trade_date
    day_dir.mkdir(parents=True, exist_ok=True)
    md_path = day_dir / f"{mode}_data_required_{stamp}.md"
    json_path = day_dir / f"{mode}_data_required_{stamp}.json"
    docx_path = day_dir / f"{mode}_data_required_{stamp}.docx"
    template = template_path.read_text(encoding="utf-8-sig", errors="ignore")
    md_text = "\n".join([
        f"# {SKILL_NAME_DISPLAY} - {MODE_TO_TITLE[mode]} 数据准备单",
        "",
        "> 这不是正式市场报告。正式报告必须运行 run --evidence 并通过 quality-gate。",
        f"> source_note: {source_note}" if source_note else "",
        "",
        template,
    ])
    md_path.write_text(md_text, encoding="utf-8")
    write_docx_from_markdown(md_text, docx_path, f"{SKILL_NAME_DISPLAY} - 数据准备单")
    payload = {
        "skill": SKILL_DIR.name,
        "display_name": SKILL_NAME_DISPLAY,
        "mode": mode,
        "trade_date": trade_date,
        "status": "DATA_REQUIRED",
        "message": "正式报告请使用 run --evidence <证据包.json>",
        "artifacts": {"markdown": str(md_path), "word": str(docx_path), "json": str(json_path), "template": str(template_path)},
    }
    write_json(json_path, payload)
    return payload


def cmd_info(_: argparse.Namespace) -> int:
    return emit(
        {
            "ok": True,
            "skill": SKILL_DIR.name,
            "display_name": SKILL_NAME_DISPLAY,
            "skill_dir": str(SKILL_DIR),
            "entry": str(Path(__file__).resolve()),
            "locked_execution": str(SCRIPTS_DIR / LOCKED_PRODUCTION_SCRIPT),
            "workflow": str(SKILL_DIR / "references" / "workflow.md"),
            "evidence_schema": str(SKILL_DIR / "references" / "evidence-schema.json"),
            "business_blueprint": str(SKILL_DIR / "references" / "business-blueprint.md"),
            "data_source_matrix": str(SKILL_DIR / "references" / "data-source-matrix.json"),
            "scoring_model": str(SKILL_DIR / "references" / "scoring-model.json"),
            "tdx_hub": str(TDX_HUB),
            "primary_commands": ["collect-local", "run --evidence", "quality-gate", "auto"],
        }
    )


def cmd_list(_: argparse.Namespace) -> int:
    return emit(
        {
            "ok": True,
            "modes": [
                {"mode": "morning", "schedule": "周一到周日每日早上，具体小时由调度传入", "delivery": "docx/md/json"},
                {"mode": "midday", "schedule": "周一到周五 12:00", "delivery": "docx/md/json"},
                {"mode": "closing", "schedule": "周一到周五 16:00", "delivery": "docx/md/json"},
            ],
            "commands": {
                "collect-local": "采集本地TDX状态、板块、新闻缓存，生成证据草稿",
                "run": "读取证据包并生成正式Word报告",
                "quality-gate": "检查报告质量和Word可读性",
                "generate": "仅生成DATA_REQUIRED准备单",
            },
        }
    )


def cmd_collect_local(args: argparse.Namespace) -> int:
    mode = args.mode if args.mode != "auto" else resolve_mode_auto(now_local())
    trade_date = args.date or today()
    status = run_json([PYTHON_EXE, str(TDX_HUB), "status"], timeout=45) if TDX_HUB.exists() else {"ok": False, "error": "tdx_hub missing"}
    blocks = run_json([PYTHON_EXE, str(TDX_HUB), "blocks"], timeout=45) if TDX_HUB.exists() else {"ok": False}
    news = run_json([PYTHON_EXE, str(TDX_HUB), "news", "--limit", "10", "--preview-chars", "180"], timeout=45) if TDX_HUB.exists() else {"ok": False}
    payload = {
        "workflow": SKILL_DIR.name,
        "mode": mode,
        "trade_date": trade_date,
        "status": "DATA_REQUIRED",
        "collected_at": now_local().isoformat(timespec="seconds"),
        "sources": [
            {"id": "TDX_STATUS", "name": "tdx-local-hub status", "type": "local_inventory", "collected_at": now_local().isoformat(timespec="seconds"), "path": str(TDX_HUB)},
            {"id": "TDX_BLOCKS", "name": "tdx-local-hub blocks", "type": "local_block_inventory", "collected_at": now_local().isoformat(timespec="seconds"), "path": str(TDX_HUB)},
            {"id": "TDX_NEWS_CACHE", "name": "tdx-local-hub news cache", "type": "local_news_cache", "collected_at": now_local().isoformat(timespec="seconds"), "path": str(TDX_HUB)},
        ],
        "local_inventory": {"status": status, "blocks": blocks, "news": news},
        "next_required": [
            "补充五大指数、涨跌家数、板块涨跌幅、新闻热度、公告监管、外盘或交易日历字段",
            "补充 headline_conclusion、event_pool、theme_clusters、market_feedback、capital_behavior、risk_items、next_session_plan",
            "字段格式按 references/evidence-schema.json",
            "补齐后执行 run --mode <mode> --evidence <本文件>",
        ],
    }
    out = Path(args.output) if args.output else (EVIDENCE_DIR / trade_date / f"{mode}_local_draft_{now_local().strftime('%Y%m%d_%H%M%S')}.json")
    write_json(out, payload)
    payload["artifact"] = str(out)
    return emit(payload)


def cmd_fixture(args: argparse.Namespace) -> int:
    mode = args.mode
    trade_date = args.date or today()
    payload = fixture_payload(mode, trade_date)
    out = Path(args.output) if args.output else (EVIDENCE_DIR / trade_date / f"{mode}_selftest_fixture.json")
    write_json(out, payload)
    payload["artifact"] = str(out)
    return emit(payload)


def cmd_run(args: argparse.Namespace) -> int:
    evidence_path = Path(args.evidence)
    evidence = load_json(evidence_path)
    mode = args.mode or evidence.get("mode") or "closing"
    payload, code = run_workflow(mode, evidence, allow_fixture=args.allow_fixture, output_dir=Path(args.output_dir) if args.output_dir else None)
    payload["evidence_path"] = str(evidence_path)
    return emit(payload, code)


def cmd_quality_gate(args: argparse.Namespace) -> int:
    report_path = Path(args.report_json)
    payload = load_json(report_path)
    result = quality_gate(payload, allow_fixture=args.allow_fixture)
    result["report_json"] = str(report_path)
    out = QUALITY_DIR / (payload.get("trade_date") or today()) / f"quality_{now_local().strftime('%Y%m%d_%H%M%S')}.json"
    write_json(out, result)
    result["artifact"] = str(out)
    return emit(result, 0 if result["status"] == "CLEAN_PASS" else 2)


def cmd_generate(args: argparse.Namespace) -> int:
    mode = args.mode if args.mode != "auto" else resolve_mode_auto(now_local())
    payload = build_draft_report(mode, args.date or today(), args.source_note or "")
    return emit(payload)


def cmd_selftest(_: argparse.Namespace) -> int:
    steps: list[dict[str, Any]] = []
    if LOCK.exists():
        r = run_capture([PYTHON_EXE, str(LOCK), "check", "--skill", SKILL_DIR.name], timeout=120)
        r["step"] = "skill_workflow_lock_check"
        r["ok"] = r.get("ok") and '"status": "CLEAN_PASS"' in r.get("stdout", "")
        steps.append(r)
    else:
        steps.append({"step": "skill_workflow_lock_check", "ok": False, "error": f"missing {LOCK}"})
    if SUBSTANTIVE_GATE.exists():
        r = run_capture([PYTHON_EXE, str(SUBSTANTIVE_GATE), "check", "--skill", SKILL_DIR.name], timeout=180)
        r["step"] = "stock_workflow_substantive_gate"
        r["ok"] = r.get("ok") and '"status": "CLEAN_PASS"' in r.get("stdout", "")
        steps.append(r)
    else:
        steps.append({"step": "stock_workflow_substantive_gate", "ok": False, "error": f"missing {SUBSTANTIVE_GATE}"})
    steps.append({"step": "python_docx_available", "ok": importlib.util.find_spec("docx") is not None})
    fixture = fixture_payload("closing", "2026-07-01")
    payload, code = run_workflow("closing", fixture, allow_fixture=True, output_dir=REPORTS_DIR / "selftest")
    steps.append({"step": "fixture_run_closing_report", "ok": code == 0 and payload.get("status") == "CLEAN_PASS", "artifacts": payload.get("artifacts")})
    ok = all(step.get("ok") for step in steps)
    return emit({"skill": SKILL_DIR.name, "mode": "selftest", "status": "CLEAN_PASS" if ok else "BLOCKED", "steps": steps, "blocks": [s for s in steps if not s.get("ok")]}, 0 if ok else 1)


def cmd_auto(_: argparse.Namespace) -> int:
    steps: list[dict[str, Any]] = []
    selftest = run_capture([PYTHON_EXE, str(Path(__file__).resolve()), "selftest"], timeout=300)
    selftest["step"] = "selftest"
    selftest["ok"] = selftest.get("ok") and int(selftest.get("exit", 1)) == 0
    steps.append(selftest)
    target = SCRIPTS_DIR / LOCKED_PRODUCTION_SCRIPT
    if target.exists():
        r = run_capture([PYTHON_EXE, str(target)], timeout=300)
        r["step"] = "locked_execution"
        r["script"] = target.name
        steps.append(r)
    else:
        steps.append({"step": "locked_execution", "ok": False, "error": f"missing {target}"})
    ok = all(step.get("ok") for step in steps)
    return emit(
        {
            "skill": SKILL_DIR.name,
            "display_name": SKILL_NAME_DISPLAY,
            "mode": "auto",
            "executed_at": now_local().isoformat(timespec="seconds"),
            "status": "CLEAN_PASS" if ok else "BLOCKED",
            "all_ok": ok,
            "steps": steps,
        },
        0 if ok else 1,
    )


def cmd_hash(_: argparse.Namespace) -> int:
    files = [
        SKILL_DIR / "SKILL.md",
        SKILL_DIR / "references" / "workflow.md",
        SKILL_DIR / "references" / "source-requirement.md",
        SKILL_DIR / "references" / "evidence-schema.json",
        SKILL_DIR / "references" / "business-blueprint.md",
        SKILL_DIR / "references" / "data-source-matrix.json",
        SKILL_DIR / "references" / "scoring-model.json",
        TEMPLATES_DIR / "morning_brief.md",
        TEMPLATES_DIR / "midday_push.md",
        TEMPLATES_DIR / "closing_review.md",
        Path(__file__).resolve(),
        SCRIPTS_DIR / LOCKED_PRODUCTION_SCRIPT,
    ]
    digest = hashlib.sha256()
    for path in files:
        digest.update(path.read_bytes())
    return emit({"ok": True, "skill": SKILL_DIR.name, "sha256": digest.hexdigest(), "files": [str(p) for p in files]})


def build_parser() -> argparse.ArgumentParser:
    parser = argparse.ArgumentParser(description=f"{SKILL_NAME_DISPLAY} Codex execution entry")
    sub = parser.add_subparsers(dest="cmd")
    sub.add_parser("info")
    sub.add_parser("list")
    sub.add_parser("selftest")
    sub.add_parser("auto")
    sub.add_parser("hash")

    collect = sub.add_parser("collect-local")
    collect.add_argument("--mode", choices=["auto", "morning", "midday", "closing"], default="auto")
    collect.add_argument("--date", default="")
    collect.add_argument("--output", default="")

    fixture = sub.add_parser("fixture")
    fixture.add_argument("--mode", choices=["morning", "midday", "closing"], default="closing")
    fixture.add_argument("--date", default="")
    fixture.add_argument("--output", default="")

    run_p = sub.add_parser("run")
    run_p.add_argument("--mode", choices=["morning", "midday", "closing"], default="")
    run_p.add_argument("--evidence", required=True)
    run_p.add_argument("--output-dir", default="")
    run_p.add_argument("--allow-fixture", action="store_true")

    q = sub.add_parser("quality-gate")
    q.add_argument("--report-json", required=True)
    q.add_argument("--allow-fixture", action="store_true")

    gen = sub.add_parser("generate")
    gen.add_argument("--mode", choices=["auto", "morning", "midday", "closing"], default="auto")
    gen.add_argument("--date", default="")
    gen.add_argument("--source-note", default="")
    return parser


def main() -> int:
    parser = build_parser()
    if len(sys.argv) == 1:
        return cmd_auto(None)
    args = parser.parse_args()
    if args.cmd == "info":
        return cmd_info(args)
    if args.cmd == "list":
        return cmd_list(args)
    if args.cmd == "collect-local":
        return cmd_collect_local(args)
    if args.cmd == "fixture":
        return cmd_fixture(args)
    if args.cmd == "run":
        return cmd_run(args)
    if args.cmd == "quality-gate":
        return cmd_quality_gate(args)
    if args.cmd == "generate":
        return cmd_generate(args)
    if args.cmd == "selftest":
        return cmd_selftest(args)
    if args.cmd == "auto":
        return cmd_auto(args)
    if args.cmd == "hash":
        return cmd_hash(args)
    parser.print_help()
    return 0


if __name__ == "__main__" and __import__("os").environ.get("ONESTOCK_STOCK_CANONICAL_CHILD") != "1":
    print("canonical_stock_legacy_entry_direct_execution_blocked", file=__import__("sys").stderr)
    raise SystemExit(2)


if __name__ == "__main__":
    raise SystemExit(main())
