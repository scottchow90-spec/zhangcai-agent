#!/usr/bin/env python3
# -*- coding: utf-8 -*-
from __future__ import annotations
# ONESTOCK_EMBEDDED_LOCAL_IMPORT_BOOTSTRAP
import sys as _onestock_embedded_sys
from pathlib import Path as _OneStockEmbeddedPath
_onestock_embedded_dir = str(_OneStockEmbeddedPath(__file__).resolve().parent)
if _onestock_embedded_dir not in _onestock_embedded_sys.path:
    _onestock_embedded_sys.path.insert(0, _onestock_embedded_dir)

import argparse
import html
import json
import subprocess
import sys
import zipfile
from pathlib import Path
from typing import Any

try:
    sys.stdout.reconfigure(encoding="utf-8")
    sys.stderr.reconfigure(encoding="utf-8")
except Exception:
    pass

SKILL_DIR = Path(__file__).resolve().parents[1]
REPORTS = SKILL_DIR / "reports" / "comprehensive"
PYTHON_EXE = sys.executable

KNOWN_COMPANY: dict[str, dict[str, Any]] = {
    "301372": {
        "name": "科净源",
        "business": "水处理产品、水环境综合治理方案、项目运营服务",
        "industry": "环保水处理/水环境治理",
        "finance": [
            ["2026Q1收入", "约7025.97万元，同比+8.07%", "收入有增长，但增速不高。"],
            ["2026Q1归母净利润", "约573.81万元，同比-48.94%", "利润端明显承压。"],
            ["2026Q1扣非净利润", "约511.45万元，同比约+5.23%", "扣非略增，但不足以改变谨慎判断。"],
            ["2026Q1经营现金流", "约-2063.75万元，同比下降约113.01%", "现金流质量偏弱，回款压力需跟踪。"],
            ["2025收入", "约2.76亿元，同比+44.43%", "业务规模修复。"],
            ["2025归母净利润", "约-3770万元", "全年仍亏损，盈利稳定性不足。"],
        ],
        "announcements": "近期以股东会、董事会换届和高管聘任类公告为主，未构成直接业绩强催化。",
        "sources": "巨潮资讯2025年年度报告/年度报告摘要、公开披露的2026年一季报信息、公开行业资讯交叉核验。",
    }
}


def normalize_code(symbol: str) -> str:
    digits = "".join(ch for ch in symbol if ch.isdigit())
    if len(digits) < 6:
        raise ValueError(f"invalid symbol: {symbol}")
    return digits[-6:]


def parse_json(text: str) -> Any:
    clean = (text or "").lstrip("\ufeff").strip()
    if not clean:
        raise ValueError("empty json")
    clean = clean[clean.find("{"):]
    return json.loads(clean)


def run_score(symbol: str) -> dict[str, Any]:
    result = subprocess.run(
        [PYTHON_EXE, str(SKILL_DIR / "scripts" / "baimao_stock_score.py"), "score", symbol],
        cwd=str(SKILL_DIR),
        capture_output=True,
        text=True,
        encoding="utf-8",
        errors="ignore",
        timeout=180,
    )
    payload = parse_json(result.stdout)
    payload["_exit"] = result.returncode
    return payload


def comprehensive(symbol: str) -> dict[str, Any]:
    code = normalize_code(symbol)
    score = run_score(symbol)
    if not score.get("ok"):
        return {
            "ok": False,
            "status": score.get("status", "BLOCKED"),
            "symbol": score.get("symbol", symbol),
            "data_gate": score.get("data_gate", {}),
            "error": score.get("error", "score blocked"),
            "technical": score,
            "conclusion": "数据闸门未通过，禁止输出综合结论。",
        }
    info = KNOWN_COMPANY.get(code, {
        "name": score.get("symbol"),
        "business": "未配置本地业务画像",
        "industry": "未配置本地行业画像",
        "finance": [["财务信息", "未配置结构化财务快照", "综合结论降级，仅保留技术评分。"]],
        "announcements": "未配置公告摘要，需接入iFinD/交易所公告后补全。",
        "sources": "本地K线与白猫老师评分体系；基本面画像未配置。",
    })
    levels = score["levels"]
    tech_state = "技术面弱势未修复" if score["score"] < 40 else "技术面可观察但需复核"
    final_rating = "排除 / 不纳入候选池" if score["score"] < 40 else "候选 / 基本面复核后再评估"
    conclusion = (
        f"{info['name']}当前综合评级为{final_rating}。"
        f"白猫老师技术评分{score['score']}/100，{score['rating']}；{tech_state}。"
        "在财务质量、现金流、公告催化和行业属性纳入后，只有技术与基本面同时改善才可复评。"
    )
    return {
        "ok": True,
        "status": "CLEAN_PASS",
        "symbol": score["symbol"],
        "name": info["name"],
        "trade_date": score["trade_date"],
        "latest": score["latest"],
        "technical_score": score["score"],
        "technical_rating": score["rating"],
        "final_rating": final_rating,
        "conclusion": conclusion,
        "business": info["business"],
        "industry": info["industry"],
        "finance": info["finance"],
        "announcements": info["announcements"],
        "sources": info["sources"],
        "technical": score,
        "risk_summary": score.get("risks", []),
        "recheck_conditions": [
            f"放量站回触发线{levels['trigger']}，且平均成本线由弱转强。",
            "躲猫猫、白猫渡劫、白猫队长、白猫RSI至少形成多模块共振。",
            "后续季度利润率修复、经营现金流改善。",
            "出现高质量订单、中标、回款改善或行业政策催化。",
        ],
        "invalidation": f"跌破{levels['invalidation']}或继续受长期EMA压制时维持排除。",
    }


def _xml(text: Any) -> str:
    return html.escape(str(text), quote=True)


def write_minimal_docx(out: Path, title: str, sections: list[tuple[str, list[list[Any]]]]) -> None:
    def p(text: Any, bold: bool = False) -> str:
        tag1 = "<w:b/>" if bold else ""
        return f"<w:p><w:r><w:rPr>{tag1}<w:rFonts w:eastAsia=\"Microsoft YaHei\"/></w:rPr><w:t>{_xml(text)}</w:t></w:r></w:p>"

    def table_xml(headers: list[str], rows: list[list[Any]]) -> str:
        all_rows = [headers] + rows
        out_rows = []
        for row in all_rows:
            cells = []
            for value in row:
                cells.append(f"<w:tc><w:tcPr><w:tcW w:w=\"2400\" w:type=\"dxa\"/></w:tcPr>{p(value)}</w:tc>")
            out_rows.append("<w:tr>" + "".join(cells) + "</w:tr>")
        return "<w:tbl><w:tblPr><w:tblW w:w=\"0\" w:type=\"auto\"/><w:tblBorders><w:top w:val=\"single\" w:sz=\"4\"/><w:left w:val=\"single\" w:sz=\"4\"/><w:bottom w:val=\"single\" w:sz=\"4\"/><w:right w:val=\"single\" w:sz=\"4\"/><w:insideH w:val=\"single\" w:sz=\"4\"/><w:insideV w:val=\"single\" w:sz=\"4\"/></w:tblBorders></w:tblPr>" + "".join(out_rows) + "</w:tbl>"

    body = [p(title, True)]
    for heading, rows in sections:
        body.append(p(heading, True))
        body.append(table_xml(["项目", "内容"], rows))
    document_xml = (
        '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
        '<w:document xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main">'
        "<w:body>"
        + "".join(body)
        + '<w:sectPr><w:pgSz w:w="11906" w:h="16838"/><w:pgMar w:top="900" w:right="900" w:bottom="900" w:left="900"/></w:sectPr>'
        "</w:body></w:document>"
    )
    with zipfile.ZipFile(out, "w", zipfile.ZIP_DEFLATED) as z:
        z.writestr("[Content_Types].xml", '<?xml version="1.0" encoding="UTF-8"?><Types xmlns="http://schemas.openxmlformats.org/package/2006/content-types"><Default Extension="rels" ContentType="application/vnd.openxmlformats-package.relationships+xml"/><Default Extension="xml" ContentType="application/xml"/><Override PartName="/word/document.xml" ContentType="application/vnd.openxmlformats-officedocument.wordprocessingml.document.main+xml"/></Types>')
        z.writestr("_rels/.rels", '<?xml version="1.0" encoding="UTF-8"?><Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships"><Relationship Id="rId1" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/officeDocument" Target="word/document.xml"/></Relationships>')
        z.writestr("word/document.xml", document_xml)


def write_docx_with_python_docx(payload: dict[str, Any], out: Path) -> None:
    from docx import Document
    from docx.enum.table import WD_TABLE_ALIGNMENT
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Cm, Pt, RGBColor

    def east(run):
        run.font.name = "微软雅黑"
        run._element.rPr.rFonts.set(qn("w:eastAsia"), "微软雅黑")

    def shade(cell, fill):
        tc_pr = cell._tc.get_or_add_tcPr()
        shd = OxmlElement("w:shd")
        shd.set(qn("w:fill"), fill)
        tc_pr.append(shd)

    def table(doc: Document, headers: list[str], rows: list[list[Any]]) -> None:
        t = doc.add_table(rows=1, cols=len(headers))
        t.style = "Table Grid"
        t.alignment = WD_TABLE_ALIGNMENT.CENTER
        for i, head in enumerate(headers):
            shade(t.rows[0].cells[i], "17324D")
            p = t.rows[0].cells[i].paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            r = p.add_run(str(head))
            r.bold = True
            r.font.color.rgb = RGBColor(255, 255, 255)
            east(r)
        for row in rows:
            cells = t.add_row().cells
            for i, val in enumerate(row):
                cells[i].text = str(val)
        doc.add_paragraph()

    doc = Document()
    sec = doc.sections[0]
    sec.top_margin = Cm(1.5)
    sec.bottom_margin = Cm(1.5)
    sec.left_margin = Cm(1.5)
    sec.right_margin = Cm(1.5)
    for name, size, color in [("Normal", 10.2, "1F2933"), ("Title", 22, "17324D"), ("Heading 1", 15, "17324D")]:
        st = doc.styles[name]
        st.font.name = "微软雅黑"
        st._element.rPr.rFonts.set(qn("w:eastAsia"), "微软雅黑")
        st.font.size = Pt(size)
        st.font.color.rgb = RGBColor.from_string(color)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run(f"{payload['name']}综合分析报告")
    r.bold = True
    r.font.size = Pt(22)
    r.font.color.rgb = RGBColor.from_string("17324D")
    east(r)

    table(doc, ["项目", "结论"], [
        ["综合评级", payload["final_rating"]],
        ["核心结论", payload["conclusion"]],
        ["技术评分", f"{payload['technical_score']}/100，{payload['technical_rating']}"],
        ["最新K线日期", payload["trade_date"]],
        ["行业/业务", f"{payload['industry']}；{payload['business']}"],
    ])

    doc.add_heading("一、技术体系", level=1)
    table(doc, ["模块", "得分", "事实", "风险"], [
        [m["name"], f"{m['score']}/{m['max_score']}", "；".join(m.get("facts", [])[:4]), "；".join(m.get("risks", [])) or "无"]
        for m in payload["technical"]["modules"]
    ])

    doc.add_heading("二、财务与现金流", level=1)
    table(doc, ["项目", "数据/事实", "结论"], payload["finance"])

    doc.add_heading("三、行业公告与复评条件", level=1)
    table(doc, ["维度", "内容"], [
        ["公告事件", payload["announcements"]],
        ["风险", "；".join(payload["risk_summary"]) or "六公式层面未出现额外风险项"],
        ["复评条件", "；".join(payload["recheck_conditions"])],
        ["失效条件", payload["invalidation"]],
        ["来源", payload["sources"]],
    ])

    doc.save(out)


def write_docx(payload: dict[str, Any]) -> Path:
    REPORTS.mkdir(parents=True, exist_ok=True)
    out = REPORTS / f"{payload['symbol'].replace('.', '')}_{payload['trade_date']}_综合分析.docx"
    try:
        write_docx_with_python_docx(payload, out)
    except Exception:
        sections = [
            ("核心结论", [
                ["综合评级", payload["final_rating"]],
                ["核心结论", payload["conclusion"]],
                ["技术评分", f"{payload['technical_score']}/100，{payload['technical_rating']}"],
                ["最新K线日期", payload["trade_date"]],
                ["行业/业务", f"{payload['industry']}；{payload['business']}"],
            ]),
            ("财务与现金流", payload["finance"]),
            ("公告风险与复评", [
                ["公告事件", payload["announcements"]],
                ["风险", "；".join(payload["risk_summary"]) or "六公式层面未出现额外风险项"],
                ["复评条件", "；".join(payload["recheck_conditions"])],
                ["失效条件", payload["invalidation"]],
                ["来源", payload["sources"]],
            ]),
        ]
        write_minimal_docx(out, f"{payload['name']}综合分析报告", sections)
    return out


def cmd_analyze(args: argparse.Namespace) -> int:
    payload = comprehensive(args.symbol)
    print(json.dumps(payload, ensure_ascii=False, indent=2))
    return 0 if payload.get("ok") else 2


def cmd_report(args: argparse.Namespace) -> int:
    payload = comprehensive(args.symbol)
    if not payload.get("ok"):
        print(json.dumps(payload, ensure_ascii=False, indent=2))
        return 2
    out = write_docx(payload)
    result = {"ok": True, "status": "CLEAN_PASS", "symbol": payload["symbol"], "docx": str(out), "rating": payload["final_rating"]}
    print(json.dumps(result, ensure_ascii=False, indent=2))
    return 0


def main() -> int:
    parser = argparse.ArgumentParser(description="白猫老师评分体系综合分析")
    sub = parser.add_subparsers(dest="cmd", required=True)
    a = sub.add_parser("analyze")
    a.add_argument("symbol")
    a.set_defaults(func=cmd_analyze)
    r = sub.add_parser("report")
    r.add_argument("symbol")
    r.set_defaults(func=cmd_report)
    args = parser.parse_args()
    return args.func(args)


if __name__ == "__main__":
    raise SystemExit(main())
