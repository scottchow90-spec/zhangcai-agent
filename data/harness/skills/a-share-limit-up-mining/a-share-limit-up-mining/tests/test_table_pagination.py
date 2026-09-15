from __future__ import annotations
# ONESTOCK_EMBEDDED_LOCAL_IMPORT_BOOTSTRAP
import sys as _onestock_embedded_sys
from pathlib import Path as _OneStockEmbeddedPath
_onestock_embedded_dir = str(_OneStockEmbeddedPath(__file__).resolve().parent)
if _onestock_embedded_dir not in _onestock_embedded_sys.path:
    _onestock_embedded_sys.path.insert(0, _onestock_embedded_dir)

import importlib.util
from pathlib import Path

from docx import Document
from docx.oxml.ns import qn


SKILL_ROOT = Path(__file__).resolve().parents[1]


def _load_build_docx():
    script = SKILL_ROOT / "scripts" / "build_docx.py"
    spec = importlib.util.spec_from_file_location("limit_up_build_docx", script)
    assert spec is not None and spec.loader is not None
    module = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(module)
    return module


def test_table_header_is_kept_with_first_body_row() -> None:
    module = _load_build_docx()
    table = Document().add_table(rows=2, cols=2)

    module.keep_table_header_with_first_body_row(table)

    for cell in table.rows[0].cells:
        for paragraph in cell.paragraphs:
            assert paragraph._p.get_or_add_pPr().find(qn("w:keepNext")) is not None
    first_body_properties = table.rows[1]._tr.get_or_add_trPr()
    assert first_body_properties.find(qn("w:cantSplit")) is not None


def test_header_body_pagination_is_limited_to_stock_detail_tables() -> None:
    module = _load_build_docx()

    selected = [
        index
        for index in range(19)
        if module.requires_header_body_pagination(index)
    ]

    assert selected == [8, 12, 13, 14, 15, 16]


def test_compact_result_table_reduces_vertical_footprint() -> None:
    module = _load_build_docx()
    table = Document().add_table(rows=2, cols=2)

    module.apply_compact_result_table_style(table)

    for row in table.rows:
        for cell in row.cells:
            margins = cell._tc.get_or_add_tcPr().find(qn("w:tcMar"))
            assert margins is not None
            assert int(margins.find(qn("w:top")).get(qn("w:w"))) <= 30
            assert int(margins.find(qn("w:bottom")).get(qn("w:w"))) <= 30
            for paragraph in cell.paragraphs:
                assert paragraph.paragraph_format.space_before.pt == 0
                assert paragraph.paragraph_format.space_after.pt == 0
                assert paragraph.paragraph_format.line_spacing == 1.0
                assert paragraph.runs[0].font.size.pt <= 8.5
