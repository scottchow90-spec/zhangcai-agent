from __future__ import annotations
# ONESTOCK_EMBEDDED_LOCAL_IMPORT_BOOTSTRAP
import sys as _onestock_embedded_sys
from pathlib import Path as _OneStockEmbeddedPath
_onestock_embedded_dir = str(_OneStockEmbeddedPath(__file__).resolve().parent)
if _onestock_embedded_dir not in _onestock_embedded_sys.path:
    _onestock_embedded_sys.path.insert(0, _onestock_embedded_dir)

import csv
import importlib.util
import json
from datetime import datetime
from pathlib import Path
from zoneinfo import ZoneInfo

from docx import Document


SOURCE = Path(__file__).resolve().parents[1] / "scripts" / "run_a_share_15d.py"
AUDIT_SOURCE = Path(__file__).resolve().parents[1] / "scripts" / "audit_engine.py"

ENABLED_FACTORS = [
    *( (name, weight, "爆发力") for name, weight in [
        ("强势收盘强度", 4.0), ("前高突破强度", 4.0), ("短周期收益加速度", 4.0),
        ("健康放量质量", 4.0), ("成交额放大强度", 3.0), ("大牛线攻击信号", 4.0),
        ("飞龙启动信号", 5.0), ("游资点火强度", 4.0), ("机构净流入强度", 3.0),
    ]),
    *( (name, weight, "持续性") for name, weight in [
        ("多周期相对强度", 5.0), ("均线多头结构", 4.0), ("趋势斜率", 4.0),
        ("趋势效率", 4.0), ("飞龙波段阶段", 4.0), ("趋势乖离健康度", 3.0),
        ("量价一致性", 3.0), ("资金合力持续性", 4.0), ("催化延续证据", 4.0),
    ]),
    *( (name, weight, "市场协同") for name, weight in [
        ("市场广度", 4.0), ("涨停生态", 4.0), ("主线题材共振", 7.0), ("全市场相对强度", 5.0),
    ]),
    *( (name, weight, "可交易性") for name, weight in [
        ("流动性分位", 3.0), ("日内承接强度", 3.0), ("流动性稳定度", 2.0), ("波动可控性", 2.0),
    ]),
]
FACTOR_ROLES = {
    **{name: "candidate_ranking" for name, _weight, _axis in ENABLED_FACTORS},
    "强势收盘强度": "pool_confirmation",
    "催化延续证据": "candidate_evidence",
    "市场广度": "market_regime",
    "涨停生态": "market_regime",
}


def load_module():
    spec = importlib.util.spec_from_file_location("a_share_15d_sparse_test", SOURCE)
    assert spec is not None and spec.loader is not None
    module = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(module)
    return module


def load_audit_module():
    spec = importlib.util.spec_from_file_location("a_share_15d_audit_test", AUDIT_SOURCE)
    assert spec is not None and spec.loader is not None
    module = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(module)
    return module


def no_candidate_payload() -> dict:
    return {
        "status": "CLEAN_PASS",
        "selection_status": "NO_CANDIDATES",
        "target_count": 5,
        "result_count": 0,
        "trade_date": "20260818",
        "generated_at": "2026-08-18T09:00:00+08:00",
        "source": {
            "template_sha256": "EC71F04176CB13FBF56AC73D60EE6701CCC71AB877C457D365E130A1795993F5",
        },
        "candidate_pool": {
            "mode": "current_limit_up",
            "name": "当日真实涨停池",
            "member_count": 0,
        },
        "macro": {
            "market_count": 5000,
            "advancers": 2000,
            "decliners": 2800,
            "flat": 200,
            "total_amount_yi": 12000,
            "scan_coverage": {
                "eligible_file_count": 5000,
                "current_row_count": 5000,
                "stale_or_short_count": 0,
                "read_error_count": 0,
                "markets": ["BJ", "SH", "SZ"],
                "complete": True,
            },
        },
        "raw_limit_up_count": 0,
        "pool_candidate_count": 0,
        "preliminary_count": 0,
        "scored_count": 0,
        "excluded_count": 0,
        "ranking_count": 0,
        "ranking": [],
        "top5": [],
        "excluded": [],
        "factor_registry": {
            "enabled": [
                {
                    "name": name,
                    "axis": axis,
                    "role": FACTOR_ROLES[name],
                    "ranking_relevant": FACTOR_ROLES[name] in {"candidate_ranking", "candidate_evidence"},
                    "weight": weight,
                    "status": "ENABLED_REPRODUCIBLE",
                }
                for name, weight, axis in ENABLED_FACTORS
            ],
            "pending_data": [{"name": f"待接入{i}"} for i in range(10)],
            "rejected": [{"name": f"拒绝{i}"} for i in range(5)],
            "enabled_count": 26,
            "pending_count": 10,
            "rejected_count": 5,
        },
        "factor_diagnostics": {
            "sample_count": 0,
            "effective_factor_count": 0,
            "ranking_eligible_factor_count": 23,
            "ranking_effective_factor_count": 0,
            "context_factor_count": 3,
            "dead_factors": [],
            "saturated_factors": [],
            "ranking_dead_factors": [],
            "ranking_saturated_factors": [],
        },
        "evaluation_contract": {
            "decision_time": "trade_date_after_close",
            "entry_assumption": "next_trade_day_open",
            "forward_horizons": [1, 3, 5],
            "walk_forward_required": True,
            "lookahead_policy": "features_and_evidence_timestamp_at_or_before_decision_time",
        },
        "predictive_validation": {
            "status": "UNVERIFIED_REQUIRES_POINT_IN_TIME_FORWARD_LABELS",
        },
        "validation": {
            "top5_count": 0,
            "complete_scan": True,
            "result_count_matches": True,
            "result_count_lte_target": True,
            "scarcity_disclosed": True,
            "all_top5_hard_gates": True,
            "all_top5_26_factors": True,
            "all_top5_15_risks": True,
            "all_top5_7_hard_exclusions": True,
            "all_top5_required_formulas": True,
            "all_top5_segment_le_80": True,
            "all_top5_score_math": True,
            "ranking_consistent": True,
            "score_contract_v6": True,
            "ranking_count_matches": True,
            "all_ranked_hard_gates": True,
            "all_ranked_26_factors": True,
            "all_ranked_15_risks": True,
            "all_ranked_7_hard_exclusions": True,
            "all_ranked_required_formulas": True,
            "all_ranked_segment_le_80": True,
            "all_ranked_score_math": True,
            "full_ranking_consistent": True,
            "top5_is_ranking_prefix": True,
            "pool_coverage_complete": True,
            "market_scope_complete": True,
            "current_trade_date": True,
        },
    }


def test_runtime_home_is_derived_from_relocated_skill_directory(tmp_path: Path):
    module = load_module()
    relocated_skill = tmp_path / "Home" / "skills" / "a-share-15d-selection"

    assert module.runtime_home(relocated_skill) == tmp_path / "Home"


def test_execute_persists_truthful_no_candidate_result(monkeypatch, tmp_path: Path):
    module = load_module()
    started = datetime(2026, 8, 18, 9, 0, tzinfo=ZoneInfo("Asia/Shanghai"))
    macro = {
        "market_count": 5000,
        "advancers": 2000,
        "decliners": 2800,
        "flat": 200,
        "total_amount_yi": 12000,
        "limit_up_count": 0,
        "scan_coverage": {
            "eligible_file_count": 5000,
            "current_row_count": 5000,
            "stale_or_short_count": 0,
            "read_error_count": 0,
            "markets": ["BJ", "SH", "SZ"],
            "complete": True,
        },
    }
    monkeypatch.setattr(module, "now_cn", lambda: started)
    monkeypatch.setattr(module, "load_names", lambda: {})
    monkeypatch.setattr(module, "scan_market", lambda names: ("20260818", [], macro))
    monkeypatch.setattr(module, "expected_completed_trade_date", lambda value: "20260818")
    monkeypatch.setattr(module, "collect_local_text", lambda: [])
    monkeypatch.setattr(
        module,
        "render_report_from_json",
        lambda json_path, docx_path: docx_path.write_bytes(b"docx") or docx_path,
    )
    monkeypatch.setattr(
        module,
        "validate_artifacts",
        lambda json_path, csv_path, docx_path, output_path: {
            "status": "CLEAN_PASS",
            "errors": [],
        },
    )

    result = module.execute(tmp_path)

    assert result["status"] == "CLEAN_PASS"
    assert result["selection_status"] == "NO_CANDIDATES"
    assert result["target_count"] == 5
    assert result["result_count"] == 0
    assert result["top5"] == []
    assert result["validation"]["complete_scan"] is True
    assert result["validation"]["scarcity_disclosed"] is True
    assert len(list(tmp_path.glob("*.json"))) == 1
    assert len(list(tmp_path.glob("*.csv"))) == 1
    assert (tmp_path / "a-share-15d-selection-result.json").is_file()
    assert (tmp_path / "a-share-15d-selection-result.csv").is_file()
    assert (tmp_path / "a-share-15d-selection-result.docx").is_file()


def test_render_report_discloses_no_candidate_result(tmp_path: Path):
    module = load_module()
    json_path = tmp_path / "empty.json"
    docx_path = tmp_path / "empty.docx"
    json_path.write_text(
        json.dumps(no_candidate_payload(), ensure_ascii=False),
        encoding="utf-8",
    )

    module.render_report_from_json(json_path, docx_path)

    assert docx_path.is_file()
    text = "\n".join(paragraph.text for paragraph in Document(docx_path).paragraphs)
    assert "本轮无候选" in text


def test_audit_accepts_complete_truthful_no_candidate_result():
    audit = load_audit_module()
    assert audit.classify_engine_result(no_candidate_payload()) == "CLEAN_PASS"


def write_no_candidate_artifacts(tmp_path: Path) -> tuple[Path, Path, Path]:
    json_path = tmp_path / "a-share-15d-selection-result.json"
    csv_path = tmp_path / "a-share-15d-selection-result.csv"
    docx_path = tmp_path / "a-share-15d-selection-result.docx"
    json_path.write_text(
        json.dumps(no_candidate_payload(), ensure_ascii=False, indent=2),
        encoding="utf-8",
    )
    with csv_path.open("w", newline="", encoding="utf-8-sig") as handle:
        writer = csv.writer(handle)
        writer.writerow([
            "排名", "代码", "名称", "现价", "涨幅", "正向得分",
            "风险扣分", "最终得分", "飞龙波", "飞龙段", "主趋势线", "失效条件",
        ])
    document = Document()
    document.add_paragraph("交易日 2026-08-18｜本轮无候选")
    document.save(docx_path)
    return json_path, csv_path, docx_path


def test_artifact_validator_accepts_consistent_json_csv_docx(tmp_path: Path):
    audit = load_audit_module()
    json_path, csv_path, docx_path = write_no_candidate_artifacts(tmp_path)
    output_path = tmp_path / "artifact_validation.json"

    result = audit.validate_artifacts(json_path, csv_path, docx_path, output_path)

    assert result["status"] == "CLEAN_PASS"
    assert result["errors"] == []
    assert output_path.is_file()
    assert result["artifacts"]["json"]["size"] == json_path.stat().st_size
    assert len(result["artifacts"]["docx"]["sha256"]) == 64


def test_artifact_validator_blocks_tampered_csv_and_docx(tmp_path: Path):
    audit = load_audit_module()
    json_path, csv_path, docx_path = write_no_candidate_artifacts(tmp_path)
    with csv_path.open("a", newline="", encoding="utf-8-sig") as handle:
        csv.writer(handle).writerow([1, "000001", "伪造候选"])
    docx_path.write_bytes(b"not-a-docx")

    result = audit.validate_artifacts(
        json_path,
        csv_path,
        docx_path,
        tmp_path / "artifact_validation.json",
    )

    assert result["status"] == "BLOCKED"
    assert "csv_result_count_mismatch" in result["errors"]
    assert any(error.startswith("docx_invalid:") for error in result["errors"])
